import copy
import datetime as dt
import hashlib
import importlib.util
import json
import os
import re
import subprocess
import sys
from pathlib import Path

import jsonschema
import openpyxl
import pytest

MODULE_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = MODULE_ROOT.parents[1]
SPEC = importlib.util.spec_from_file_location(
    "boiling_module_rollup_tests",
    MODULE_ROOT / "src" / "boiling_module.py",
)
assert SPEC and SPEC.loader
boiling_module = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(boiling_module)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def make_sample(path, *, month=1, batch_no="MONTH-001", product_usage="家具"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "水煮实验"
    sheet.append(
        ["批号", batch_no, "产品名称", "脱敏竹材", "规格型号", "100×100×10"]
    )
    sheet.append(
        [
            "压制日期",
            dt.date(2026, month, 3),
            "班长",
            "脱敏班长",
            "产品用途",
            product_usage,
        ]
    )
    sheet.append(
        [
            "检测日期",
            dt.date(2026, month, 16),
            "浸胶日期",
            dt.date(2026, month, 10),
            "订单号",
            "SANITIZED-ORDER",
        ]
    )
    sheet.cell(7, 1, "试件")
    sheet.cell(7, 2, "长度(mm)")
    sheet.cell(7, 3, "宽度(mm)")
    sheet.cell(7, 4, "煮前厚度(mm)")
    sheet.cell(7, 5, "重量(g)")
    sheet.cell(7, 6, "密度(g/cm³)")
    sheet.cell(7, 7, "煮后厚度(mm)")
    sheet.cell(7, 8, "膨胀率")
    sheet.cell(8, 1, "端1-1")
    sheet.cell(8, 2, 100)
    sheet.cell(8, 3, 100)
    sheet.cell(9, 5, 140)
    for row in range(8, 14):
        sheet.cell(row, 4, 10.0)
        sheet.cell(row, 7, 10.5)
    sheet.cell(14, 6, 1.4)
    sheet.cell(14, 8, 0.05)
    sheet.cell(
        17,
        1,
        "检测结论：密度检测合格，2h膨胀率检测合格，浸渍剥离检测合格",
    )
    workbook.save(path)


def make_job(sample, output, *, phase="draft", **extra):
    job = {
        "schema_version": "1.0",
        "job_id": "ROLLUP-TEST-001",
        "module_id": "boiling",
        "phase": phase,
        "run_mode": "single",
        "input_files": [str(sample)],
        "output_dir": str(output),
        "module_config_file": None,
        "ai_enabled": False,
        "allowed_knowledge_scopes": [],
        "parent_job_id": None,
        "confirmation_file": None,
        "service_response_file": None,
        "task_name": "水煮月度测试",
        "period_year": 2026,
        "period_month": 1,
        "learning_enabled": True,
    }
    job.update(extra)
    return job


def write_json(path, value):
    Path(path).write_text(
        json.dumps(value, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def run_job(job, path):
    write_json(path, job)
    assert boiling_module.run_job(path) == 0
    return json.loads(
        (Path(job["output_dir"]) / "result.json").read_text(encoding="utf-8")
    )


def finalized_snapshot(tmp_path):
    sample = tmp_path / "month.xlsx"
    make_sample(sample)
    draft_output = tmp_path / "draft"
    draft = run_job(
        make_job(sample, draft_output),
        tmp_path / "draft-job.json",
    )
    conclusion_confirmation = next(
        item
        for item in draft["confirmations"]
        if item["confirmation_id"].startswith("C-CONCLUSION-RULE-")
    )
    ai_request = next(
        item for item in draft["service_requests"] if item["service"] == "ai"
    )
    confirmation_file = tmp_path / "confirmations.json"
    write_json(
        confirmation_file,
        {
            "schema_version": "1.0",
            "task_id": draft["job_id"],
            "module_id": "boiling",
            "answers": [
                {
                    "confirmation_id": conclusion_confirmation["confirmation_id"],
                    "value": conclusion_confirmation["suggested_value"],
                    "apply_to_batch": False,
                }
            ],
        },
    )
    ai_items = []
    for candidate in ai_request["payload"]["candidates"]:
        ai_items.append(
            {
                "candidate_id": candidate["candidate_id"],
                "risk_summary": "按月度证据复核。",
                "possible_causes": ["需要核对工艺记录。"],
                "check_items": ["核对对应试件记录。"],
                "retest_recommendations": ["必要时复测。"],
                "priority": "medium",
                "evidence_refs": candidate["evidence_refs"],
            }
        )
    service_file = tmp_path / "service.json"
    write_json(
        service_file,
        {
            "schema_version": "1.0",
            "task_id": draft["job_id"],
            "module_id": "boiling",
            "responses": [
                {
                    "request_id": ai_request["request_id"],
                    "service": "ai",
                    "status": "ok",
                    "text": json.dumps(
                        {
                            "schema_version": "1.0",
                            "overall_summary": "既有 AI 月度建议。",
                            "items": ai_items,
                            "limitations": ["不改变确定性业务输出。"],
                        },
                        ensure_ascii=False,
                    ),
                }
            ],
        },
    )
    final_output = tmp_path / "final"
    final = run_job(
        make_job(
            sample,
            final_output,
            phase="finalize",
            job_id="ROLLUP-FINAL-001",
            parent_job_id=draft["job_id"],
            confirmation_file=str(confirmation_file),
            service_response_file=str(service_file),
        ),
        tmp_path / "final-job.json",
    )
    assert final["monthly_snapshot"] is not None
    assert all(item["type"] != "monthly_snapshot" for item in final["outputs"])
    snapshot_path = Path(final["monthly_snapshot"]["path"])
    return final, snapshot_path


def snapshot_variant(snapshot_path, target, *, month, batch_no, year=2026):
    snapshot = json.loads(Path(snapshot_path).read_text(encoding="utf-8"))
    snapshot["year"] = year
    snapshot["month"] = month
    snapshot["payload"]["records"][0]["test_date"] = f"{year:04d}-{month:02d}-16"
    snapshot["payload"]["records"][0]["batch_no"] = batch_no
    snapshot["payload"]["records"][0]["crack_inspection_confirmed"] = True
    snapshot["payload"]["records"][0]["crack_inspection_status"] = (
        "已检查且符合对应等级"
    )
    write_json(target, snapshot)
    return {
        "snapshot_id": f"SNAPSHOT-{year:04d}-{month:02d}-{batch_no[-3:]}",
        "job_id": f"ROLLUP-JOB-{year:04d}-{month:02d}",
        "module_id": "boiling",
        "year": year,
        "month": month,
        "path": str(Path(target).resolve()),
        "size": Path(target).stat().st_size,
        "sha256": digest(target),
    }


def rollup_request(root, output, *, snapshots, year=2026, target_month=3):
    return {
        "schema_version": "1.0",
        "transaction_id": "ROLLUP-TRANSACTION-001",
        "module_id": "boiling",
        "module_version": "1.3.1",
        "year": year,
        "target_month": target_month,
        "snapshot_root": str(Path(root).resolve()),
        "output_dir": str(Path(output).resolve()),
        "snapshots": snapshots,
        "configuration_identity": {
            "version": boiling_module.load_rules()["rule_version"],
            "sha256": digest(boiling_module.DEFAULT_RULES_FILE),
        },
        "offline_required": True,
    }


def test_finalize_writes_confirmed_monthly_snapshot_with_existing_ai(tmp_path):
    final, snapshot_path = finalized_snapshot(tmp_path)
    assert final["status"] == "success"
    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    assert snapshot["module_version"] == "1.3.1"
    assert (snapshot["year"], snapshot["month"]) == (2026, 1)
    record = snapshot["payload"]["records"][0]
    assert record["final_confirmation_decisions"][0]["value"] == "不合格"
    assert snapshot["payload"]["ai_review"]["status"] == "ok"
    serialized = json.dumps(snapshot, ensure_ascii=False)
    assert not re.search(r"(?i)[a-z]:[\\/]", serialized)


def test_finalize_preserves_standalone_slash_business_placeholders(tmp_path):
    sample = tmp_path / "slash-placeholders.xlsx"
    make_sample(sample, month=4, batch_no="SLASH-APR")
    workbook = openpyxl.load_workbook(sample)
    try:
        sheet = workbook["水煮实验"]
        for coordinate in ("B2", "D2", "D3", "F3"):
            sheet[coordinate] = "/"
        workbook.save(sample)
    finally:
        workbook.close()

    draft_output = tmp_path / "slash-draft"
    draft = run_job(
        make_job(
            sample,
            draft_output,
            job_id="SLASH-DRAFT-001",
            period_month=4,
        ),
        tmp_path / "slash-draft-job.json",
    )
    confirmation_file = tmp_path / "slash-confirmations.json"
    write_json(
        confirmation_file,
        {
            "schema_version": "1.0",
            "task_id": draft["job_id"],
            "module_id": "boiling",
            "answers": [
                {
                    "confirmation_id": item["confirmation_id"],
                    "value": (
                        item["options"][0]
                        if item["options"]
                        else item["suggested_value"]
                    ),
                    "apply_to_batch": False,
                    "remember_mapping": False,
                }
                for item in draft["confirmations"]
            ],
        },
    )
    final_output = tmp_path / "slash-final"
    final = run_job(
        make_job(
            sample,
            final_output,
            phase="finalize",
            job_id="SLASH-FINAL-001",
            parent_job_id=draft["job_id"],
            confirmation_file=str(confirmation_file),
            period_month=4,
        ),
        tmp_path / "slash-final-job.json",
    )

    assert final["status"] == "success"
    assert final["monthly_snapshot"] is not None
    assert all(item["type"] != "monthly_snapshot" for item in final["outputs"])
    snapshot = json.loads(
        Path(final["monthly_snapshot"]["path"]).read_text(encoding="utf-8")
    )
    record = snapshot["payload"]["records"][0]
    assert {
        record["pressing_date"],
        record["team_leader"],
        record["impregnation_date"],
        record["order_info"],
    } == {"/"}


@pytest.mark.parametrize(
    "unsafe_path",
    [
        r"C:\customer\boiling.xlsx",
        r"\\server\share\boiling.xlsx",
        "/home/customer/boiling.xlsx",
    ],
    ids=["windows", "unc", "posix"],
)
def test_monthly_snapshot_rejects_true_absolute_paths(unsafe_path):
    payload = {
        "rule_version": "test-rule-v1",
        "records": [{"order_info": unsafe_path}],
        "excluded_records": [],
        "anomalies": [],
        "confirmation_answers": [],
        "learning_descriptions": [],
        "ai_review": None,
        "knowledge_review": None,
    }

    with pytest.raises(
        boiling_module.RollupContractError,
        match="月度快照 payload 不得包含绝对路径",
    ):
        boiling_module.build_snapshot(
            module_version="1.3.1",
            year=2026,
            month=4,
            payload=payload,
        )


def test_snapshot_uses_final_human_density_and_expansion_values(tmp_path):
    sample = tmp_path / "conflicts.xlsx"
    make_sample(sample)
    workbook = openpyxl.load_workbook(sample)
    try:
        sheet = workbook["水煮实验"]
        sheet.cell(14, 6, 1.2)
        sheet.cell(14, 8, 0.06)
        workbook.save(sample)
    finally:
        workbook.close()
    draft_output = tmp_path / "draft-conflicts"
    draft = run_job(
        make_job(sample, draft_output),
        tmp_path / "draft-conflicts-job.json",
    )
    conflict_prompts = [
        item
        for item in draft["confirmations"]
        if item["confirmation_id"].startswith("C-CONFLICT-")
    ]
    assert len(conflict_prompts) == 2
    confirmation_file = tmp_path / "conflict-confirmations.json"
    answers = []
    for prompt in conflict_prompts:
        answers.append(
            {
                "confirmation_id": prompt["confirmation_id"],
                "value": prompt["options"][0],
                "apply_to_batch": False,
            }
        )
    write_json(
        confirmation_file,
        {
            "schema_version": "1.0",
            "task_id": draft["job_id"],
            "module_id": "boiling",
            "answers": answers,
        },
    )
    final_output = tmp_path / "final-conflicts"
    final = run_job(
        make_job(
            sample,
            final_output,
            phase="finalize",
            job_id="ROLLUP-FINAL-VALUES",
            parent_job_id=draft["job_id"],
            confirmation_file=str(confirmation_file),
        ),
        tmp_path / "final-conflicts-job.json",
    )
    assert final["monthly_snapshot"] is not None
    snapshot_path = Path(final["monthly_snapshot"]["path"])
    record = json.loads(snapshot_path.read_text(encoding="utf-8"))["payload"][
        "records"
    ][0]
    assert record["samples"][0]["density_g_cm3"] == pytest.approx(1.2)
    assert record["samples"][0]["expansion_ratio"] == pytest.approx(0.06)
    assert record["average_density_g_cm3"] == pytest.approx(1.2)
    assert record["average_expansion_ratio"] == pytest.approx(0.06)
    assert {
        item["choice"] for item in record["confirmed_calculation_values"]
    } == {"source"}


def test_finalize_with_unanswered_blocking_confirmation_fails_without_snapshot(
    tmp_path,
):
    sample = tmp_path / "unanswered.xlsx"
    make_sample(sample, product_usage="")
    draft_output = tmp_path / "draft-unanswered"
    draft = run_job(
        make_job(sample, draft_output),
        tmp_path / "draft-unanswered-job.json",
    )
    assert draft["confirmations"]
    confirmation_file = tmp_path / "empty-confirmations.json"
    write_json(
        confirmation_file,
        {
            "schema_version": "1.0",
            "task_id": draft["job_id"],
            "module_id": "boiling",
            "answers": [],
        },
    )
    final_output = tmp_path / "final-unanswered"
    final_job = make_job(
        sample,
        final_output,
        phase="finalize",
        job_id="ROLLUP-FINAL-NO-SNAPSHOT",
        parent_job_id=draft["job_id"],
        confirmation_file=str(confirmation_file),
    )
    final_job_path = tmp_path / "final-unanswered-job.json"
    write_json(final_job_path, final_job)
    assert boiling_module.run_job(final_job_path) == 3
    final = json.loads((final_output / "result.json").read_text(encoding="utf-8"))
    assert final["status"] == "failed"
    assert any(
        error.startswith("BOILING-UNRESOLVED-BLOCKING-CONFIRMATION-001:")
        for error in final["errors"]
    )
    assert final["monthly_snapshot"] is None
    assert all(item["type"] != "monthly_snapshot" for item in final["outputs"])
    assert not (final_output / "boiling_monthly_snapshot.json").exists()


def test_source_bat_draft_finalize_snapshot_and_rollup(tmp_path):
    sample = tmp_path / "bat-month.xlsx"
    make_sample(sample, month=4, batch_no="BAT-APR")
    environment = os.environ.copy()
    environment["BAMBOO_PYTHON"] = sys.executable
    launcher = MODULE_ROOT / "run_module.bat"

    draft_output = tmp_path / "bat-draft"
    draft_job = make_job(
        sample,
        draft_output,
        job_id="BAT-DRAFT-001",
        period_month=4,
        task_name="水煮 BAT 2026-04",
    )
    draft_job_path = tmp_path / "bat-draft-job.json"
    write_json(draft_job_path, draft_job)
    draft_run = subprocess.run(
        [str(launcher), "run", "--job", str(draft_job_path)],
        cwd=REPO_ROOT,
        env=environment,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    assert draft_run.returncode == 0, draft_run.stderr
    draft = json.loads((draft_output / "result.json").read_text(encoding="utf-8"))
    assert draft["status"] == "needs_confirmation"
    draft_ids = {item["confirmation_id"] for item in draft["confirmations"]}

    confirmation_file = tmp_path / "bat-confirmations.json"
    write_json(
        confirmation_file,
        {
            "schema_version": "1.0",
            "task_id": draft["job_id"],
            "module_id": "boiling",
            "answers": [
                {
                    "confirmation_id": item["confirmation_id"],
                    "value": (
                        item["options"][0]
                        if item["options"]
                        else item["suggested_value"]
                    ),
                    "apply_to_batch": False,
                    "remember_mapping": False,
                }
                for item in draft["confirmations"]
            ],
        },
    )
    final_output = tmp_path / "bat-final"
    final_job = make_job(
        sample,
        final_output,
        phase="finalize",
        job_id="BAT-FINAL-001",
        parent_job_id=draft["job_id"],
        confirmation_file=str(confirmation_file),
        period_month=4,
        task_name="水煮 BAT 2026-04",
    )
    final_job_path = tmp_path / "bat-final-job.json"
    write_json(final_job_path, final_job)
    final_run = subprocess.run(
        [str(launcher), "run", "--job", str(final_job_path)],
        cwd=REPO_ROOT,
        env=environment,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    assert final_run.returncode == 0, final_run.stderr
    final = json.loads((final_output / "result.json").read_text(encoding="utf-8"))
    assert final["status"] == "success"
    assert final["confirmations"] == []
    assert final["monthly_snapshot"] is not None
    assert (
        final["monthly_snapshot"]["year"],
        final["monthly_snapshot"]["month"],
    ) == (2026, 4)
    assert all(item["type"] != "monthly_snapshot" for item in final["outputs"])
    snapshot = json.loads(
        Path(final["monthly_snapshot"]["path"]).read_text(encoding="utf-8")
    )
    decision_ids = {
        decision["confirmation_id"]
        for record in snapshot["payload"]["records"]
        for decision in record["final_confirmation_decisions"]
    }
    assert decision_ids <= draft_ids

    snapshot_root = tmp_path / "bat-snapshots"
    snapshot_root.mkdir()
    snapshot_path = snapshot_root / "2026-04.json"
    snapshot_path.write_bytes(Path(final["monthly_snapshot"]["path"]).read_bytes())
    snapshot_ref = {
        "snapshot_id": "BAT-SNAPSHOT-2026-04",
        "job_id": final["job_id"],
        "module_id": "boiling",
        "year": 2026,
        "month": 4,
        "path": str(snapshot_path.resolve()),
        "size": snapshot_path.stat().st_size,
        "sha256": digest(snapshot_path),
    }
    rollup_output = tmp_path / "bat-rollup"
    request_path = tmp_path / "bat-rollup-request.json"
    response_path = rollup_output / "result.json"
    write_json(
        request_path,
        rollup_request(
            snapshot_root,
            rollup_output,
            snapshots=[snapshot_ref],
            target_month=4,
        ),
    )
    rollup_run = subprocess.run(
        [
            str(launcher),
            "rollup",
            "--request",
            str(request_path),
            "--response",
            str(response_path),
        ],
        cwd=REPO_ROOT,
        env=environment,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    assert rollup_run.returncode == 0, rollup_run.stderr
    rollup = json.loads(response_path.read_text(encoding="utf-8"))
    assert rollup["status"] == "success"
    assert rollup["used_months"] == [4]
    assert rollup["network_call_count"] == 0


def test_rollup_allows_month_gap_and_never_builds_ai_request(tmp_path, monkeypatch):
    _, source_snapshot = finalized_snapshot(tmp_path)
    snapshot_root = tmp_path / "snapshots"
    snapshot_root.mkdir()
    january = snapshot_variant(
        source_snapshot,
        snapshot_root / "01.json",
        month=1,
        batch_no="JAN-NEW",
    )
    march = snapshot_variant(
        source_snapshot,
        snapshot_root / "03.json",
        month=3,
        batch_no="MAR-NEW",
    )
    output = tmp_path / "rollup"
    request = rollup_request(
        snapshot_root,
        output,
        snapshots=[january, march],
    )
    request_path = tmp_path / "rollup-request.json"
    response_path = output / "rollup-response.json"
    write_json(request_path, request)
    monkeypatch.setattr(
        boiling_module,
        "build_ai_service_request",
        lambda *_args, **_kwargs: pytest.fail("累计重算不得构造 AI 请求"),
    )

    assert boiling_module.run_monthly_rollup(request_path, response_path) == 0
    response = json.loads(response_path.read_text(encoding="utf-8"))
    expected = json.loads(
        (
            MODULE_ROOT
            / "samples"
            / "expected"
            / "monthly_rollup_gap.expected.json"
        ).read_text(encoding="utf-8")
    )
    assert response["status"] == expected["expected_status"]
    assert response["used_months"] == expected["expected_used_months"]
    assert response["network_call_count"] == expected["expected_network_call_count"]
    assert response["record_count"] == expected["expected_record_count"]
    assert [
        (item["snapshot_id"], item["month"], item["sha256"])
        for item in response["snapshot_hashes"]
    ] == [
        (item["snapshot_id"], item["month"], item["sha256"])
        for item in (january, march)
    ]
    assert [item["type"] for item in response["outputs"]] == expected[
        "expected_output_types"
    ]
    assert all(Path(item["path"]).parent == output for item in response["outputs"])
    assert all(digest(item["path"]) == item["sha256"] for item in response["outputs"])
    assert all("snapshot" not in item["type"] for item in response["outputs"])
    request_schema = json.loads(
        (MODULE_ROOT / "schemas" / "monthly-rollup-request.schema.json").read_text(
            encoding="utf-8"
        )
    )
    response_schema = json.loads(
        (MODULE_ROOT / "schemas" / "monthly-rollup-result.schema.json").read_text(
            encoding="utf-8"
        )
    )
    jsonschema.Draft202012Validator(request_schema).validate(request)
    jsonschema.Draft202012Validator(response_schema).validate(response)
    central_request_schema = json.loads(
        (REPO_ROOT / "contracts" / "schemas" / "monthly-rollup-request.schema.json").read_text(
            encoding="utf-8"
        )
    )
    central_result_schema = json.loads(
        (REPO_ROOT / "contracts" / "schemas" / "monthly-rollup-result.schema.json").read_text(
            encoding="utf-8"
        )
    )
    jsonschema.Draft202012Validator(central_request_schema).validate(request)
    jsonschema.Draft202012Validator(central_result_schema).validate(response)
    workbook = openpyxl.load_workbook(output / "boiling_summary.xlsx", data_only=True)
    try:
        assert (
            workbook["水煮检测汇总"].max_row - 1
            == expected["expected_record_count"]
        )
        assert workbook["多月趋势"].max_row == 3
    finally:
        workbook.close()
    report = (output / "boiling_analysis.md").read_text(encoding="utf-8")
    assert "既有 AI 月度建议" in report
    assert "网络调用次数为 0" in report


def test_early_month_candidate_replaces_existing_snapshot(tmp_path):
    _, source_snapshot = finalized_snapshot(tmp_path)
    snapshot_root = tmp_path / "snapshots"
    snapshot_root.mkdir()
    snapshot_variant(
        source_snapshot,
        snapshot_root / "01-old.json",
        month=1,
        batch_no="JAN-OLD",
    )
    january_new = snapshot_variant(
        source_snapshot,
        snapshot_root / "01-new.json",
        month=1,
        batch_no="JAN-REPLACED",
    )
    march = snapshot_variant(
        source_snapshot,
        snapshot_root / "03.json",
        month=3,
        batch_no="MAR-KEEP",
    )
    output = tmp_path / "rollup"
    request = rollup_request(
        snapshot_root,
        output,
        snapshots=[january_new, march],
    )
    request_path = tmp_path / "request.json"
    response_path = output / "response.json"
    write_json(request_path, request)

    assert boiling_module.run_monthly_rollup(request_path, response_path) == 0
    records = json.loads(
        (output / "boiling_records.json").read_text(encoding="utf-8")
    )["records"]
    assert [record["batch_no"] for record in records] == ["JAN-REPLACED", "MAR-KEEP"]
    assert all(record["crack_inspection_confirmed"] is True for record in records)
    assert all(
        record["crack_inspection_status"] == "已检查且符合对应等级"
        for record in records
    )


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        ("out_of_order", "月份乱序"),
        ("duplicate", "重复月份"),
        ("hash", "哈希不一致"),
        ("cross_year", "跨年快照"),
        ("size", "大小不一致"),
        ("other_module", "module_id 必须为 boiling"),
        ("late_month", "晚于目标月份"),
        ("outside_root", "超出授权目录"),
    ],
)
def test_rollup_rejects_order_duplicate_hash_and_cross_year(
    tmp_path,
    mutation,
    message,
):
    _, source_snapshot = finalized_snapshot(tmp_path)
    snapshot_root = tmp_path / "snapshots"
    snapshot_root.mkdir()
    january = snapshot_variant(
        source_snapshot,
        snapshot_root / "01.json",
        month=1,
        batch_no="JAN",
    )
    february = snapshot_variant(
        source_snapshot,
        snapshot_root / "02.json",
        month=2,
        batch_no="FEB",
    )
    march = snapshot_variant(
        source_snapshot,
        snapshot_root / "03.json",
        month=3,
        batch_no="MAR",
    )
    snapshots = [january, february]
    target_month = 3
    if mutation == "out_of_order":
        snapshots = [february, january]
    elif mutation == "duplicate":
        snapshots = [january, copy.deepcopy(january)]
    elif mutation == "hash":
        snapshots[0] = {**january, "sha256": "0" * 64}
    elif mutation == "cross_year":
        january = snapshot_variant(
            source_snapshot,
            snapshot_root / "01-cross-year.json",
            month=1,
            batch_no="JAN",
            year=2025,
        )
        snapshots = [january]
    elif mutation == "size":
        snapshots[0] = {**january, "size": january["size"] + 1}
    elif mutation == "other_module":
        snapshots[0] = {**january, "module_id": "moisture"}
    elif mutation == "late_month":
        target_month = 2
    elif mutation == "outside_root":
        outside = tmp_path / "outside.json"
        outside.write_bytes(Path(january["path"]).read_bytes())
        snapshots[0] = {
            **january,
            "path": str(outside.resolve()),
            "size": outside.stat().st_size,
            "sha256": digest(outside),
        }
    output = tmp_path / "rollup"
    request = rollup_request(
        snapshot_root,
        output,
        snapshots=[*snapshots, march] if mutation != "out_of_order" else snapshots,
        target_month=target_month,
    )
    request_path = tmp_path / "request.json"
    response_path = output / "response.json"
    write_json(request_path, request)

    assert boiling_module.run_monthly_rollup(request_path, response_path) == 2
    response = json.loads(response_path.read_text(encoding="utf-8"))
    assert response["status"] == "failed"
    assert response["error"]["code"] == "contract_error"
    assert message in response["error"]["message"]
    assert response["network_call_count"] == 0


def test_module_owned_schemas_validate_examples(tmp_path):
    _, snapshot_path = finalized_snapshot(tmp_path)
    schema_dir = MODULE_ROOT / "schemas"
    snapshot_schema = json.loads(
        (schema_dir / "monthly-snapshot.schema.json").read_text(encoding="utf-8")
    )
    jsonschema.Draft202012Validator(snapshot_schema).validate(
        json.loads(snapshot_path.read_text(encoding="utf-8"))
    )
    learning_schema = json.loads(
        (schema_dir / "confirmation-learning.schema.json").read_text(encoding="utf-8")
    )
    description = boiling_module.learning_description_for_confirmation(
        {"confirmation_id": "C-EXPANSION-REVIEW-0123456789ab"},
        {"product_usage": "墙板"},
    )
    assert description is not None
    jsonschema.Draft202012Validator(learning_schema).validate(description)
    serialized = json.dumps(description, ensure_ascii=False)
    assert "脱敏竹材" not in serialized
    assert not re.search(r"(?i)[a-z]:[\\/]", serialized)


def test_learning_descriptor_is_embedded_in_public_confirmation(tmp_path):
    sample = tmp_path / "wall-panel.xlsx"
    make_sample(sample)
    workbook = openpyxl.load_workbook(sample)
    try:
        sheet = workbook["水煮实验"]
        sheet["F2"] = "墙板"
        for row in range(8, 14):
            sheet.cell(row, 7, 10.7)
        sheet.cell(14, 8, 0.07)
        workbook.save(sample)
    finally:
        workbook.close()
    draft = run_job(
        make_job(sample, tmp_path / "learning-output"),
        tmp_path / "learning-job.json",
    )
    prompt = next(
        item
        for item in draft["confirmations"]
        if item["confirmation_id"].startswith("C-EXPANSION-REVIEW-")
    )
    assert prompt["learning"]["learning_key"] == "boiling.expansion.rule-gap.v1"
    assert prompt["learning"]["context"]["product_usage"] == "墙板"
    assert prompt["learning"]["base_suggestion_id"] == "keep_review"
    assert all(set(item) == {"candidate_id", "label"} for item in prompt["learning"]["candidates"])
    module_schema = json.loads(
        (MODULE_ROOT / "schemas" / "confirmation-learning.schema.json").read_text(
            encoding="utf-8"
        )
    )
    central_schema = json.loads(
        (REPO_ROOT / "contracts" / "schemas" / "module-result.schema.json").read_text(
            encoding="utf-8"
        )
    )
    jsonschema.Draft202012Validator(module_schema).validate(prompt["learning"])
    jsonschema.Draft202012Validator(central_schema).validate(draft)
    assert all(item["type"] != "confirmation_learning" for item in draft["outputs"])


def test_source_and_optional_portable_rollup_entrypoints(tmp_path):
    _, source_snapshot = finalized_snapshot(tmp_path)
    snapshot_root = tmp_path / "snapshots"
    snapshot_root.mkdir()
    candidate = snapshot_variant(
        source_snapshot,
        snapshot_root / "01.json",
        month=1,
        batch_no="SOURCE-CLI",
    )
    output = tmp_path / "source-output"
    request = rollup_request(
        snapshot_root,
        output,
        snapshots=[candidate],
        target_month=1,
    )
    request_path = tmp_path / "request.json"
    response_path = output / "response.json"
    write_json(request_path, request)
    environment = os.environ.copy()
    environment["BAMBOO_PYTHON"] = sys.executable
    source_result = subprocess.run(
        [
            str(MODULE_ROOT / "run_module.bat"),
            "rollup",
            "--request",
            str(request_path),
            "--response",
            str(response_path),
        ],
        cwd=MODULE_ROOT.parents[1],
        env=environment,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    assert source_result.returncode == 0, source_result.stderr
    assert json.loads(response_path.read_text(encoding="utf-8"))["status"] == "success"

    portable = os.environ.get("BOILING_PORTABLE_RUNTIME")
    if not portable:
        pytest.skip("设置 BOILING_PORTABLE_RUNTIME 后复验便携入口")
    portable_output = tmp_path / "portable-output"
    portable_request = copy.deepcopy(request)
    portable_request["output_dir"] = str(portable_output.resolve())
    portable_request_path = tmp_path / "portable-request.json"
    portable_response_path = portable_output / "response.json"
    write_json(portable_request_path, portable_request)
    portable_result = subprocess.run(
        [
            portable,
            "rollup",
            "--request",
            str(portable_request_path),
            "--response",
            str(portable_response_path),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    assert portable_result.returncode == 0, portable_result.stderr
    assert (
        json.loads(portable_response_path.read_text(encoding="utf-8"))["status"]
        == "success"
    )
