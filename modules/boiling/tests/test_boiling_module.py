import datetime as dt
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

from core.confirmations.service import ConfirmationService
from core.contracts.confirmation import ConfirmationAnswer
from core.contracts.enums import ModuleId
from core.contracts.manifest import ModuleManifest
from core.contracts.result import ModuleResult
from core.modules.registry import ModuleRegistry, RegisteredModule
from core.platform.paths import AppPaths
from core.storage.database import Database
from core.storage.repositories.confirmations import ConfirmationRepository
from core.storage.repositories.tasks import TaskRepository
from core.tasks.service import TaskService

MODULE_ROOT = Path(__file__).resolve().parents[1]
MODULE_SPEC = importlib.util.spec_from_file_location(
    "boiling_module", MODULE_ROOT / "src" / "boiling_module.py"
)
assert MODULE_SPEC and MODULE_SPEC.loader
boiling_module = importlib.util.module_from_spec(MODULE_SPEC)
MODULE_SPEC.loader.exec_module(boiling_module)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def make_sample(path, include_samples=True):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "水煮实验"
    sheet.cell(2, 1, "TEST-001")
    sheet.cell(3, 3, "脱敏竹材")
    sheet.cell(3, 7, "100*100*10")
    sheet.cell(5, 3, dt.date(2026, 7, 16))
    if include_samples:
        sheet.cell(7, 1, "试件")
        sheet.cell(7, 2, "长度(mm)")
        sheet.cell(7, 3, "宽度(mm)")
        sheet.cell(7, 4, "煮前厚度(mm)")
        sheet.cell(7, 5, "重量(g)")
        sheet.cell(7, 6, "密度(g/cm³)")
        sheet.cell(7, 7, "煮后厚度(mm)")
        sheet.cell(7, 8, "膨胀率")
        sheet.cell(8, 1, "样1")
        sheet.cell(8, 2, 100)
        sheet.cell(8, 3, 100)
        sheet.cell(9, 5, 120)
        for row in range(8, 14):
            sheet.cell(row, 4, 10.0)
            sheet.cell(row, 7, 10.5)
        sheet.cell(14, 6, 1.2)
        sheet.cell(14, 8, 0.05)
        sheet.cell(17, 1, "检测结论：密度检测合格，2h膨胀率检测合格，浸渍剥离检测合格")
        sheet.cell(17, 8, "测试员")
    workbook.save(path)


def make_job(sample, output, phase="draft", **extra):
    job = {
        "schema_version": "1.0",
        "job_id": "TEST-JOB-001",
        "module_id": "boiling",
        "phase": phase,
        "run_mode": "single",
        "input_files": [str(sample)],
        "output_dir": str(output),
        "module_config_file": None,
        "ai_enabled": False,
        "allowed_knowledge_scopes": [],
        "parent_job_id": None,
        "confirmation_file": None,
        "service_response_file": None,
    }
    job.update(extra)
    return job


def write_ai_service_response(path, *, task_id, request_id, status="ok", text=None):
    response = {
        "request_id": request_id,
        "service": "ai",
        "status": status,
        "text": text,
    }
    path.write_text(
        json.dumps(
            {
                "schema_version": "1.0",
                "task_id": task_id,
                "module_id": "boiling",
                "responses": [response],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return path


def make_ai_response_payload(request, *, overall_summary="规则结果需要复核", evidence_refs=None):
    items = []
    for candidate in request["payload"]["candidates"]:
        items.append(
            {
                "candidate_id": candidate["candidate_id"],
                "summary": "该候选由确定性规则筛出。",
                "risk_level": "medium",
                "review_recommendation": "核对原始读数并由业务人员确认。",
                "evidence_refs": evidence_refs or candidate["evidence_refs"],
            }
        )
    return {
        "schema_version": "1.0",
        "overall_summary": overall_summary,
        "items": items,
        "limitations": ["AI 不改变公式、阈值和业务结论。"],
    }


def make_labeled_sample(
    path,
    *,
    sample_count=1,
    density=1.2,
    expansion=0.05,
    product_usage="家具",
    conclusion="密度检测合格，2h膨胀率检测合格，浸渍剥离检测合格",
    include_retest=False,
    source_density=None,
    source_expansion=None,
    densities=None,
    expansions=None,
    labels=None,
    batch_no="DEMO-001",
    specification="100×100×10",
):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "水煮实验"
    sheet.append(["批号", batch_no, "产品名称", "脱敏竹集成材", "规格型号", specification])
    sheet.append(["压制日期", dt.date(2026, 7, 10), "班长", "脱敏班长", "产品用途", product_usage])
    sheet.append(
        [
            "检测日期",
            dt.date(2026, 7, 16),
            "浸胶日期",
            dt.date(2026, 7, 11),
            "订单号",
            "ORDER-DEMO-001",
        ]
    )
    sheet.cell(7, 1, "试件")
    sheet.cell(7, 2, "长度(mm)")
    sheet.cell(7, 3, "宽度(mm)")
    sheet.cell(7, 4, "煮前厚度(mm)")
    sheet.cell(7, 5, "重量(g)")
    sheet.cell(7, 6, "密度(g/cm³)")
    sheet.cell(7, 7, "煮后厚度(mm)")
    sheet.cell(7, 8, "膨胀率")
    start = 8
    total = sample_count + (1 if include_retest else 0)
    density_values = list(densities or [density] * total)
    expansion_values = list(expansions or [expansion] * total)
    for index in range(total):
        label = (
            labels[index]
            if labels
            else "复测1"
            if include_retest and index == total - 1
            else f"样{index + 1}"
        )
        sample_start = start + index * 7
        sheet.cell(sample_start, 1, label)
        sheet.cell(sample_start, 2, 100)
        sheet.cell(sample_start, 3, 100)
        avg_before = 10.0
        sample_density = density_values[index]
        sample_expansion = expansion_values[index]
        avg_after = avg_before * (1 + sample_expansion)
        weight = sample_density * (100 * 100 * avg_before / 1000)
        sheet.cell(sample_start + 1, 5, weight)
        for row in range(sample_start, sample_start + 6):
            sheet.cell(row, 4, avg_before)
            sheet.cell(row, 7, avg_after)
        sheet.cell(
            sample_start + 6, 6, sample_density if source_density is None else source_density
        )
        sheet.cell(
            sample_start + 6, 8, sample_expansion if source_expansion is None else source_expansion
        )
    conclusion_row = start + total * 7 + 2
    sheet.cell(conclusion_row, 1, f"检测结论：{conclusion}")
    sheet.cell(conclusion_row, 7, "检测员")
    sheet.cell(conclusion_row, 8, "脱敏测试员")
    workbook.save(path)


def load_result(output):
    return json.loads((Path(output) / "result.json").read_text(encoding="utf-8"))


def excel_values(path):
    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        return [
            [cell.value for cell in row]
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
        ]
    finally:
        workbook.close()


def isolated_windows_environment(tmp_path, **extra):
    system_root = Path(os.environ.get("SYSTEMROOT", r"C:\Windows"))
    environment = {
        "SystemRoot": str(system_root),
        "ComSpec": str(system_root / "System32" / "cmd.exe"),
        "PATH": str(system_root / "System32"),
        "PATHEXT": ".COM;.EXE;.BAT;.CMD",
        "TEMP": str(tmp_path),
        "TMP": str(tmp_path),
    }
    environment.update(extra)
    return environment


def run_batch(launcher, arguments, environment):
    command_processor = environment["ComSpec"]
    command = subprocess.list2cmdline([str(launcher), *arguments])
    return subprocess.run(
        [command_processor, "/d", "/s", "/c", command],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=environment,
    )


@pytest.fixture(scope="module")
def fake_portable_runtime(tmp_path_factory):
    folder = tmp_path_factory.mktemp("fake-boiling-runtime")
    source = folder / "FakeRuntime.cs"
    executable = folder / "boiling_module.exe"
    source.write_text(
        """
using System;
public static class FakeRuntime {
    public static int Main(string[] args) {
        int exitCode = int.Parse(Environment.GetEnvironmentVariable("FAKE_RUNTIME_EXIT") ?? "0");
        if (exitCode == 0) {
            Console.WriteLine(
                "{\\\"type\\\":\\\"progress\\\",\\\"percent\\\":100," +
                "\\\"message\\\":\\\"boiling 0.1.0\\\"}"
            );
        } else {
            Console.Error.WriteLine("forced portable runtime failure");
        }
        return exitCode;
    }
}
""".strip(),
        encoding="utf-8",
    )
    escaped_source = str(source).replace("'", "''")
    escaped_executable = str(executable).replace("'", "''")
    completed = subprocess.run(
        [
            "powershell",
            "-NoProfile",
            "-Command",
            (
                f"Add-Type -Path '{escaped_source}' -OutputAssembly "
                f"'{escaped_executable}' -OutputType ConsoleApplication"
            ),
        ],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert completed.returncode == 0, completed.stderr
    assert executable.is_file()
    return executable


def test_extract_sample(tmp_path):
    sample = tmp_path / "sample.xlsx"
    make_sample(sample)
    record = boiling_module.extract_file(sample)
    assert record["product"] == "脱敏竹材"
    assert len(record["samples"]) == 1
    assert record["average_density_g_cm3"] == 1.2
    assert record["average_expansion_ratio"] == 0.05


def test_draft_job_generates_result_without_modifying_source(tmp_path):
    sample = tmp_path / "sample.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample, product_usage="")
    before = digest(sample)
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    assert before == digest(sample)
    assert (output / "boiling_summary.xlsx").is_file()
    assert (output / "boiling_analysis.xlsx").is_file()
    workbook = openpyxl.load_workbook(output / "boiling_summary.xlsx")
    summary_sheet = workbook["水煮检测汇总"]
    summary_headers = [cell.value for cell in summary_sheet[1]]
    assert summary_sheet.max_column == 36
    assert not {
        "缺失/待补字段",
        "命中规则ID",
        "密度来源单元格",
        "膨胀率来源单元格",
        "源表有效区域",
    }.intersection(summary_headers)
    assert summary_headers[12:22] == [
        "端1-1密度(g/cm³)",
        "端1-1膨胀率",
        "端1-2密度(g/cm³)",
        "端1-2膨胀率",
        "中1密度(g/cm³)",
        "中1膨胀率",
        "端2-1密度(g/cm³)",
        "端2-1膨胀率",
        "端2-2密度(g/cm³)",
        "端2-2膨胀率",
    ]
    assert [summary_sheet.cell(2, column).value for column in range(13, 23)] == [0] * 10
    assert summary_headers[-2:] == ["复算冲突数", "结论/规则冲突数"]
    chart_sheet = workbook["分析图表"]
    assert len(chart_sheet._charts) == 3
    result_table_row = chart_sheet.max_row - 2
    assert chart_sheet.cell(result_table_row, 1).value == "原表综合结论"
    assert chart_sheet.cell(result_table_row + 1, 1).value == "合格"
    assert chart_sheet.cell(result_table_row + 2, 1).value == "不合格"
    workbook.close()
    result = json.loads((output / "result.json").read_text(encoding="utf-8"))
    assert result["status"] == "needs_confirmation"
    assert result["summary"]["processed_count"] == 1
    assert result["confirmations"][0]["field"] == "product_usage"


def test_expansion_chart_aggregates_specimen_values_by_specification(tmp_path):
    first = tmp_path / "spec-a-first.xlsx"
    second = tmp_path / "spec-a-second.xlsx"
    third = tmp_path / "spec-b.xlsx"
    make_labeled_sample(
        first,
        batch_no="A-001",
        specification="100×100×10",
        sample_count=1,
        expansion=0.04,
    )
    make_labeled_sample(
        second,
        batch_no="A-002",
        specification="100×100×10",
        sample_count=3,
        expansion=0.08,
    )
    make_labeled_sample(
        third,
        batch_no="B-001",
        specification="120×80×12",
        sample_count=2,
        expansion=0.10,
    )
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    job = make_job(first, output)
    job["run_mode"] = "batch"
    job["input_files"] = [str(first), str(second), str(third)]
    job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")

    assert boiling_module.run_job(job_path) == 0
    workbook_path = output / "boiling_summary.xlsx"
    workbook = openpyxl.load_workbook(workbook_path)
    chart_sheet = workbook["分析图表"]
    assert chart_sheet["D1"].value == "规格型号"
    assert chart_sheet["E1"].value == "平均膨胀率"
    assert chart_sheet["D2"].value == "100×100×10"
    assert chart_sheet["E2"].value == pytest.approx(0.07)
    assert chart_sheet["D3"].value == "120×80×12"
    assert chart_sheet["E3"].value == pytest.approx(0.10)
    workbook.close()
    with zipfile.ZipFile(workbook_path) as archive:
        expansion_chart_xml = archive.read("xl/charts/chart2.xml").decode("utf-8")
    assert "各型号规格平均膨胀率对比" in expansion_chart_xml
    assert "各批次平均膨胀率对比" not in expansion_chart_xml


def test_finalize_applies_confirmation(tmp_path):
    sample = tmp_path / "sample.xlsx"
    draft_output = tmp_path / "draft"
    final_output = tmp_path / "final"
    draft_job_path = tmp_path / "draft-job.json"
    final_job_path = tmp_path / "final-job.json"
    confirmation_path = tmp_path / "confirmation.json"
    make_labeled_sample(sample, product_usage="")
    draft_job_path.write_text(
        json.dumps(make_job(sample, draft_output), ensure_ascii=False), encoding="utf-8"
    )
    assert boiling_module.run_job(draft_job_path) == 0
    draft_result = json.loads((draft_output / "result.json").read_text(encoding="utf-8"))
    confirmation_id = draft_result["confirmations"][0]["confirmation_id"]
    confirmation_path.write_text(
        json.dumps(
            {
                "schema_version": "1.0",
                "job_id": "TEST-JOB-001",
                "answers": [
                    {
                        "confirmation_id": confirmation_id,
                        "value": "家具",
                        "apply_to_batch": False,
                        "remember_mapping": False,
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    final_job = make_job(
        sample,
        final_output,
        phase="finalize",
        parent_job_id="TEST-JOB-001",
        confirmation_file=str(confirmation_path),
    )
    final_job["job_id"] = "TEST-JOB-002"
    final_job_path.write_text(json.dumps(final_job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(final_job_path) == 0
    result = json.loads((final_output / "result.json").read_text(encoding="utf-8"))
    assert result["status"] == "success"
    assert result["confirmations"] == []


def test_bad_workbook_is_structured_failure(tmp_path):
    sample = tmp_path / "broken.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    sample.write_bytes(b"not-an-excel-file")
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 3
    result = json.loads((output / "result.json").read_text(encoding="utf-8"))
    assert result["status"] == "failed"
    assert result["summary"]["failed_count"] == 1


def test_invalid_job_returns_two(tmp_path):
    job_path = tmp_path / "job.json"
    job_path.write_text(json.dumps({"module_id": "moisture"}), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 2


def test_empty_sheet_does_not_enter_normal_summary(tmp_path):
    sample = tmp_path / "empty.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_sample(sample, include_samples=False)
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    result = json.loads((output / "result.json").read_text(encoding="utf-8"))
    assert result["status"] == "success"
    assert result["summary"]["skipped_count"] == 0
    assert result["summary"]["processed_count"] == 1
    assert (output / "boiling_summary.xlsx").is_file()


def test_extracts_all_handoff_fields_and_recalculates(tmp_path):
    sample = tmp_path / "labeled.xlsx"
    make_labeled_sample(sample)
    record = boiling_module.extract_file(sample)
    assert record["batch_no"] == "DEMO-001"
    assert record["product"] == "脱敏竹集成材"
    assert record["specification"] == "100×100×10"
    assert record["pressing_date"] == "2026-07-10"
    assert record["team_leader"] == "脱敏班长"
    assert record["test_date"] == "2026-07-16"
    assert record["impregnation_date"] == "2026-07-11"
    assert record["order_info"] == "ORDER-DEMO-001"
    assert record["product_usage"] == "家具"
    assert record["samples"][0]["volume_cm3"] == 100
    assert record["samples"][0]["calculated_density_g_cm3"] == 1.2
    assert record["samples"][0]["calculated_expansion_ratio"] == 0.05


def test_variable_sample_count_and_retest_are_preserved(tmp_path):
    sample = tmp_path / "variable.xlsx"
    make_labeled_sample(sample, sample_count=3, include_retest=True)
    record = boiling_module.extract_file(sample)
    assert len(record["samples"]) == 4
    assert [item["sample_type"] for item in record["samples"]].count("retest") == 1


def test_product_usage_in_source_does_not_request_confirmation(tmp_path):
    sample = tmp_path / "usage.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample, product_usage="家具")
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    result = load_result(output)
    assert result["status"] == "success"
    assert result["confirmations"] == []
    assert result["service_requests"] == []
    assert {item["type"] for item in result["outputs"]} == {
        "summary_excel",
        "analysis_report",
        "analysis_report_markdown",
        "business_records",
    }


def test_ai_enabled_draft_builds_consent_required_privacy_minimal_request(tmp_path):
    sample = tmp_path / "private-source-name.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample, product_usage="家具", density=1.45)
    job = make_job(sample, output, ai_enabled=True)
    job["job_id"] = "AI-DRAFT-001"
    job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")

    assert boiling_module.run_job(job_path) == 0
    result = load_result(output)
    assert ModuleResult.model_validate(result).service_requests[0].service == "ai"
    assert len(result["service_requests"]) == 1
    request = result["service_requests"][0]
    assert request["service"] == "ai"
    assert request["requires_user_consent"] is True
    assert request["response_format"] == "json_object"
    assert request["request_id"].startswith("boiling-ai-")
    serialized = json.dumps(request["payload"], ensure_ascii=False)
    assert sample.name not in serialized
    assert "DEMO-001" not in serialized
    assert "ORDER-DEMO-001" not in serialized
    assert "脱敏班长" not in serialized
    assert not re.search(r"(?i)[a-z]:[\\/]", serialized)
    assert request["payload"]["candidate_sent"] == 1
    assert request["payload"]["candidates"][0]["candidate_id"].startswith("candidate-")
    assert request["payload"]["candidates"][0]["evidence_refs"]
    workbook = openpyxl.load_workbook(output / "boiling_analysis.xlsx")
    assert workbook.sheetnames[-1] == "AI辅助建议"
    assert workbook["AI辅助建议"]["B4"].value == "pending"
    workbook.close()


def test_ai_response_is_report_only_and_sanitized(tmp_path):
    sample = tmp_path / "source.xlsx"
    draft_output = tmp_path / "draft"
    final_output = tmp_path / "final"
    offline_output = tmp_path / "offline"
    make_labeled_sample(sample, product_usage="家具", density=1.45)

    draft_job = make_job(sample, draft_output, ai_enabled=True)
    draft_job["job_id"] = "AI-DRAFT-002"
    draft_job_path = tmp_path / "draft-job.json"
    draft_job_path.write_text(json.dumps(draft_job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(draft_job_path) == 0
    request = load_result(draft_output)["service_requests"][0]

    response_path = write_ai_service_response(
        tmp_path / "service-response.json",
        task_id="AI-DRAFT-002",
        request_id=request["request_id"],
        text=json.dumps(
            make_ai_response_payload(
                request,
                overall_summary="总体稳定。优先复核规则候选。不得读取 C:\\private\\customer.xlsx。",
            ),
            ensure_ascii=False,
        ),
    )
    confirmation_path = tmp_path / "confirmation.json"
    confirmation_path.write_text(
        json.dumps({"job_id": "AI-DRAFT-002", "answers": []}, ensure_ascii=False),
        encoding="utf-8",
    )
    final_job = make_job(
        sample,
        final_output,
        phase="finalize",
        parent_job_id="AI-DRAFT-002",
        confirmation_file=str(confirmation_path),
        service_response_file=str(response_path),
        ai_enabled=True,
    )
    final_job["job_id"] = "AI-FINAL-002"
    final_job_path = tmp_path / "final-job.json"
    final_job_path.write_text(json.dumps(final_job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(final_job_path) == 0

    result = load_result(final_output)
    assert result["status"] == "success"
    assert result["service_requests"] == []
    assert "AI辅助状态：ok" in result["analysis_summary"]
    markdown = (final_output / "boiling_analysis.md").read_text(encoding="utf-8")
    assert "总体稳定" in markdown
    assert "C:\\private" not in markdown
    assert "[已脱敏路径]" in markdown
    workbook = openpyxl.load_workbook(final_output / "boiling_analysis.xlsx")
    assert workbook["AI辅助建议"]["B4"].value == "ok"
    assert "总体稳定" in workbook["AI辅助建议"]["B6"].value
    workbook.close()

    offline_job = make_job(
        sample,
        offline_output,
        phase="finalize",
        parent_job_id="AI-DRAFT-002",
        confirmation_file=str(confirmation_path),
        ai_enabled=False,
    )
    offline_job["job_id"] = "AI-OFFLINE-002"
    offline_job_path = tmp_path / "offline-job.json"
    offline_job_path.write_text(json.dumps(offline_job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(offline_job_path) == 0
    assert (final_output / "boiling_records.json").read_text(encoding="utf-8") == (
        offline_output / "boiling_records.json"
    ).read_text(encoding="utf-8")


def test_denied_ai_response_degrades_to_offline_report(tmp_path):
    sample = tmp_path / "source.xlsx"
    draft_output = tmp_path / "draft"
    final_output = tmp_path / "final"
    make_labeled_sample(sample, product_usage="家具", density=1.45)
    draft_job = make_job(sample, draft_output, ai_enabled=True)
    draft_job["job_id"] = "AI-DRAFT-003"
    draft_job_path = tmp_path / "draft-job.json"
    draft_job_path.write_text(json.dumps(draft_job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(draft_job_path) == 0
    request = load_result(draft_output)["service_requests"][0]

    response_path = write_ai_service_response(
        tmp_path / "service-response.json",
        task_id="AI-DRAFT-003",
        request_id=request["request_id"],
        status="denied",
    )
    confirmation_path = tmp_path / "confirmation.json"
    confirmation_path.write_text(
        json.dumps({"job_id": "AI-DRAFT-003", "answers": []}, ensure_ascii=False),
        encoding="utf-8",
    )
    final_job = make_job(
        sample,
        final_output,
        phase="finalize",
        parent_job_id="AI-DRAFT-003",
        confirmation_file=str(confirmation_path),
        service_response_file=str(response_path),
        ai_enabled=False,
    )
    final_job["job_id"] = "AI-FINAL-003"
    final_job_path = tmp_path / "final-job.json"
    final_job_path.write_text(json.dumps(final_job, ensure_ascii=False), encoding="utf-8")

    assert boiling_module.run_job(final_job_path) == 0
    result = load_result(final_output)
    assert result["status"] == "success"
    assert result["service_requests"] == []
    assert "AI辅助状态：denied" in result["analysis_summary"]
    workbook = openpyxl.load_workbook(final_output / "boiling_analysis.xlsx")
    assert workbook["AI辅助建议"]["B4"].value == "denied"
    assert "离线规则报告" in workbook["AI辅助建议"]["B6"].value
    workbook.close()


def test_density_rule_uses_every_sample_and_difference(tmp_path):
    sample = tmp_path / "density.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample, density=1.45)
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    result = load_result(output)
    anomaly = next(item for item in result["anomalies"] if item["field"] == "density")
    assert anomaly["severity"] == "warning"
    assert "超出 1.1～1.3" in anomaly["message"]
    assert result["status"] == "needs_confirmation"


def test_expansion_observation_uses_product_usage_threshold(tmp_path):
    sample = tmp_path / "expansion.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample, expansion=0.10, product_usage="家具")
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    result = load_result(output)
    assert any(item["field"] == "average_expansion_ratio" for item in result["anomalies"])


def test_peel_failure_is_preserved_as_source_anomaly(tmp_path):
    sample = tmp_path / "peel.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample, conclusion="密度检测合格，浸渍剥离检测不合格")
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    result = load_result(output)
    assert any(
        item["field"] == "immersion_peel_result" and item["value"] == "不合格"
        for item in result["anomalies"]
    )


def test_source_and_recalculated_conflict_are_both_retained(tmp_path):
    sample = tmp_path / "conflict.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample, density=1.2, source_density=1.28, source_expansion=0.08)
    record = boiling_module.extract_file(sample)
    assert len(record["conflicts"]) == 2
    assert record["samples"][0]["source_density_g_cm3"] == 1.28
    assert record["samples"][0]["calculated_density_g_cm3"] == 1.2
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    result = load_result(output)
    assert any(item["rule_id"] == "BOILING-CALC-CONFLICT-001" for item in result["anomalies"])
    assert any(
        item["field"].startswith("密度｜conflict.xlsx｜样1")
        for item in result["confirmations"]
    )


def test_batch_mode_isolates_bad_file_and_skips_duplicate(tmp_path):
    sample = tmp_path / "sample.xlsx"
    duplicate = tmp_path / "duplicate.xlsx"
    broken = tmp_path / "broken.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample)
    duplicate.write_bytes(sample.read_bytes())
    broken.write_bytes(b"not an xlsx")
    job = make_job(sample, output)
    job["run_mode"] = "batch"
    job["input_files"] = [str(sample), str(duplicate), str(broken)]
    job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    result = load_result(output)
    assert result["status"] == "partial_success"
    assert result["summary"]["processed_count"] == 1
    assert result["summary"]["skipped_count"] == 1
    assert result["summary"]["failed_count"] == 1


def test_unsupported_extension_fails_without_output_artifacts(tmp_path):
    sample = tmp_path / "sample.xls"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    sample.write_bytes(b"legacy")
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 3
    result = load_result(output)
    assert result["status"] == "failed"
    assert result["outputs"] == []


def test_invalid_json_returns_two_without_creating_output(tmp_path):
    job_path = tmp_path / "job.json"
    job_path.write_text("{not-json", encoding="utf-8")
    assert boiling_module.run_job(job_path) == 2


def test_ai_toggle_only_emits_framework_request_when_enabled(tmp_path):
    sample = tmp_path / "sample.xlsx"
    make_labeled_sample(sample, density=1.45)
    for enabled in (False, True):
        output = tmp_path / f"output-{enabled}"
        job_path = tmp_path / f"job-{enabled}.json"
        job = make_job(sample, output, ai_enabled=enabled)
        job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        assert boiling_module.run_job(job_path) == 0
        requests = load_result(output)["service_requests"]
        assert bool(requests) is enabled
        if enabled:
            assert requests[0]["service"] == "ai"
            assert requests[0]["requires_user_consent"] is True


def test_ai_enabled_without_candidate_stays_offline(tmp_path):
    sample = tmp_path / "incomplete.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_sample(sample, include_samples=False)
    job = make_job(sample, output, ai_enabled=True)
    job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    result = load_result(output)
    assert result["service_requests"] == []
    workbook = openpyxl.load_workbook(output / "boiling_analysis.xlsx")
    assert "AI辅助建议" not in workbook.sheetnames
    workbook.close()


def test_ai_timeout_error_invalid_json_and_evidence_mismatch_degrade(tmp_path):
    sample = tmp_path / "candidate.xlsx"
    draft_output = tmp_path / "draft"
    draft_path = tmp_path / "draft-job.json"
    make_labeled_sample(sample, density=1.45)
    draft_job = make_job(sample, draft_output, ai_enabled=True)
    draft_job["job_id"] = "AI-DEGRADE-001"
    draft_path.write_text(json.dumps(draft_job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(draft_path) == 0
    request = load_result(draft_output)["service_requests"][0]
    confirmation_path = tmp_path / "confirmation.json"
    confirmation_path.write_text(
        json.dumps({"job_id": "AI-DEGRADE-001", "answers": []}, ensure_ascii=False),
        encoding="utf-8",
    )
    cases = [
        ("timeout", None, "timeout"),
        ("error", None, "error"),
        ("ok", "{not-json", "invalid"),
        (
            "ok",
            json.dumps(
                make_ai_response_payload(request, evidence_refs=["rule:UNKNOWN"]),
                ensure_ascii=False,
            ),
            "invalid",
        ),
    ]
    for index, (service_status, text, expected_status) in enumerate(cases, 1):
        response_path = write_ai_service_response(
            tmp_path / f"service-response-{index}.json",
            task_id="AI-DEGRADE-001",
            request_id=request["request_id"],
            status=service_status,
            text=text,
        )
        output = tmp_path / f"final-{index}"
        final_job = make_job(
            sample,
            output,
            phase="finalize",
            parent_job_id="AI-DEGRADE-001",
            confirmation_file=str(confirmation_path),
            service_response_file=str(response_path),
            ai_enabled=True,
        )
        final_job["job_id"] = f"AI-DEGRADE-FINAL-{index}"
        final_path = tmp_path / f"final-job-{index}.json"
        final_path.write_text(json.dumps(final_job, ensure_ascii=False), encoding="utf-8")
        assert boiling_module.run_job(final_path) == 0
        result = load_result(output)
        assert result["service_requests"] == []
        assert f"AI辅助状态：{expected_status}" in result["analysis_summary"]
        assert "离线规则报告" in (output / "boiling_analysis.md").read_text(
            encoding="utf-8"
        )


def test_outputs_are_inside_declared_output_directory(tmp_path):
    sample = tmp_path / "sample.xlsx"
    output = tmp_path / "declared-output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample)
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    result = load_result(output)
    declared = output.resolve()
    assert all(Path(item["path"]).resolve().is_relative_to(declared) for item in result["outputs"])


def test_real_template_four_and_five_sample_labels(tmp_path):
    five = tmp_path / "five.xlsx"
    four = tmp_path / "four.xlsx"
    make_labeled_sample(
        five,
        sample_count=5,
        labels=["端1-1", "端1-2", "中1", "端2-1", "端2-2"],
    )
    make_labeled_sample(four, sample_count=4, labels=[1, 2, 3, 4])
    five_book = openpyxl.load_workbook(five)
    five_book.active["J47"] = "脱敏扩展列"
    five_book.save(five)
    assert [sample["label"] for sample in boiling_module.extract_file(five)["samples"]] == [
        "端1-1",
        "端1-2",
        "中1",
        "端2-1",
        "端2-2",
    ]
    assert len(boiling_module.extract_file(four)["samples"]) == 4


def test_summary_lists_five_fixed_position_density_and_expansion_values(tmp_path):
    sample = tmp_path / "five.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    densities = [1.11, 1.12, 1.13, 1.14, 1.15]
    expansions = [0.01, 0.02, 0.03, 0.04, 0.05]
    make_labeled_sample(
        sample,
        sample_count=5,
        labels=["端1-1", "端1-2", "中1", "端2-1", "端2-2"],
        densities=densities,
        expansions=expansions,
    )
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")

    assert boiling_module.run_job(job_path) == 0
    workbook = openpyxl.load_workbook(output / "boiling_summary.xlsx")
    summary = workbook["水煮检测汇总"]
    expected_values = [
        value for pair in zip(densities, expansions, strict=True) for value in pair
    ]
    assert [summary.cell(2, column).value for column in range(13, 23)] == pytest.approx(
        expected_values
    )
    assert summary["M2"].number_format == "0.000"
    assert summary["N2"].number_format == "0.00%"
    workbook.close()
    analysis_workbook = openpyxl.load_workbook(output / "boiling_analysis.xlsx")
    assert analysis_workbook.sheetnames == [
        "分析报告",
        "需要关注",
        "未纳入统计",
        "阅读说明",
    ]
    analysis_sheet = analysis_workbook["分析报告"]
    assert analysis_sheet["A1"].value == "水煮实验分析报告"
    assert analysis_sheet["A5"].value == "整体状态"
    assert analysis_sheet["A6"].value == "本批次未发现异常"
    assert [cell.value for cell in analysis_sheet[9]] == [
        "来源文件",
        "批号",
        "产品",
        "产品用途",
        "平均密度(g/cm³)",
        "密度判定",
        "平均膨胀率",
        "自动等级",
        "需人工复核",
        "浸渍剥离",
        "业务状态",
        "阅读提示",
    ]
    attention_sheet = analysis_workbook["需要关注"]
    assert attention_sheet["A2"].value == "无"
    analysis_workbook.close()


def test_density_boundaries_and_difference_are_inclusive(tmp_path):
    rules = boiling_module.load_rules()
    cases = [
        ("inclusive.xlsx", [1.1, 1.2], True),
        ("outside.xlsx", [1.099, 1.2], False),
        ("difference.xlsx", [1.1, 1.201], False),
        ("upper.xlsx", [1.2, 1.3], True),
    ]
    for name, densities, expected in cases:
        path = tmp_path / name
        make_labeled_sample(path, sample_count=2, densities=densities)
        record = boiling_module.extract_file(path)
        boiling_module.evaluate_record(record, rules, digest(path))
        assert record["density_rule_pass"] is expected, name


def test_expansion_usage_boundaries_and_wall_review_gap(tmp_path):
    rules = boiling_module.load_rules()
    cases = [
        ("floor-green.xlsx", "地板", 0.065, True, "绿色优等候选", False),
        ("furniture-blue.xlsx", "家具", 0.08, True, "蓝色合格", False),
        ("floor-red.xlsx", "地板", 0.08001, False, "红色不合格", False),
        ("wall-blue.xlsx", "墙板", 0.12, True, "蓝色合格", False),
        ("railing-red.xlsx", "栏杆", 0.12001, False, "红色不合格", False),
        ("wall-gap.xlsx", "墙板", 0.07, None, "", True),
    ]
    for name, usage, expansion, expected_pass, expected_grade, expects_confirmation in cases:
        path = tmp_path / name
        make_labeled_sample(path, expansion=expansion, product_usage=usage)
        record = boiling_module.extract_file(path)
        _, confirmations = boiling_module.evaluate_record(record, rules, digest(path))
        assert record["expansion_numeric_pass"] is expected_pass, name
        assert record["expansion_grade"] == expected_grade, name
        expansion_reviews = [item for item in confirmations if item["field"] == "expansion_review"]
        assert bool(expansion_reviews) is expects_confirmation, name


def test_product_usage_confirmation_is_created_once_and_applied_to_batch(tmp_path):
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    make_labeled_sample(first, product_usage="")
    make_labeled_sample(second, product_usage="")
    draft_output = tmp_path / "draft"
    draft_job_path = tmp_path / "draft-job.json"
    draft_job = make_job(first, draft_output)
    draft_job["run_mode"] = "batch"
    draft_job["input_files"] = [str(first), str(second)]
    draft_job_path.write_text(json.dumps(draft_job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(draft_job_path) == 0
    draft_result = load_result(draft_output)
    prompts = [item for item in draft_result["confirmations"] if item["field"] == "product_usage"]
    assert len(prompts) == 1

    answer_path = tmp_path / "answer.json"
    answer_path.write_text(
        json.dumps(
            {
                "answers": [
                    {
                        "confirmation_id": prompts[0]["confirmation_id"],
                        "value": "家具",
                        "apply_to_batch": True,
                    }
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    final_output = tmp_path / "final"
    final_job_path = tmp_path / "final-job.json"
    final_job = make_job(
        first,
        final_output,
        phase="finalize",
        parent_job_id="TEST-JOB-001",
        confirmation_file=str(answer_path),
    )
    final_job["job_id"] = "TEST-JOB-002"
    final_job["run_mode"] = "batch"
    final_job["input_files"] = [str(first), str(second)]
    final_job_path.write_text(json.dumps(final_job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(final_job_path) == 0
    records = json.loads((final_output / "boiling_records.json").read_text(encoding="utf-8"))
    assert {record["product_usage"] for record in records["records"]} == {"家具"}
    result = load_result(final_output)
    assert result["status"] == "success"
    assert result["summary"]["skipped_count"] == 1
    assert result["confirmations"] == []


def test_normal_and_incomplete_files_share_one_batch(tmp_path):
    normal = tmp_path / "normal.xlsx"
    incomplete = tmp_path / "incomplete.xlsx"
    make_labeled_sample(normal)
    make_sample(incomplete, include_samples=False)
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    job = make_job(normal, output)
    job["run_mode"] = "batch"
    job["input_files"] = [str(normal), str(incomplete)]
    job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    result = load_result(output)
    assert len(result["file_results"]) == 2
    assert result["summary"]["processed_count"] == 2
    records = json.loads((output / "boiling_records.json").read_text(encoding="utf-8"))
    assert len(records["records"]) == 1
    assert len(records["excluded_records"]) == 1


def test_result_json_validates_module_result_contract(tmp_path):
    sample = tmp_path / "sample.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample)
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    ModuleResult.model_validate(load_result(output))


def test_41_file_sanitized_history_regression_and_source_hashes(tmp_path):
    source_dir = tmp_path / "history"
    source_dir.mkdir()
    paths = []
    for index in range(37):
        path = source_dir / f"FIVE-{index:02d}.xlsx"
        make_labeled_sample(
            path,
            sample_count=5,
            labels=["端1-1", "端1-2", "中1", "端2-1", "端2-2"],
            batch_no=f"FIVE-{index:02d}",
        )
        paths.append(path)
    for index in range(2):
        path = source_dir / f"FOUR-{index:02d}.xlsx"
        make_labeled_sample(
            path,
            sample_count=4,
            labels=[1, 2, 3, 4],
            batch_no=f"FOUR-{index:02d}",
        )
        paths.append(path)
    for index in range(2):
        incomplete = source_dir / f"INCOMPLETE-{index + 1}.xlsx"
        make_sample(incomplete, include_samples=False)
        incomplete_book = openpyxl.load_workbook(incomplete)
        incomplete_book.active["A2"] = f"INCOMPLETE-{index + 1}"
        incomplete_book.save(incomplete)
        incomplete_book.close()
        paths.append(incomplete)
    before = {path.name: digest(path) for path in paths}
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    job = make_job(paths[0], output)
    job["run_mode"] = "batch"
    job["input_files"] = [str(path) for path in paths]
    job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    result = load_result(output)
    assert result["status"] == "success"
    assert result["summary"]["input_count"] == 41
    assert len(result["file_results"]) == 41
    assert {item["status"] for item in result["file_results"]} == {"success"}
    assert result["summary"]["input_count"] == sum(
        result["summary"][key] for key in ("processed_count", "skipped_count", "failed_count")
    )
    records = json.loads((output / "boiling_records.json").read_text(encoding="utf-8"))
    assert sum(len(record["samples"]) for record in records["records"]) == 193
    assert len(records["records"]) == 39
    assert len(records["excluded_records"]) == 2
    report = (output / "boiling_analysis.md").read_text(encoding="utf-8")
    assert len(report.splitlines()) < 80
    assert "## 一、运行概况" in report
    assert "## 四、异常与待复核明细" in report
    assert "## 五、数据风险（按类别汇总）" in report
    assert "FIVE-00.xlsx" not in report
    assert "INCOMPLETE-1.xlsx" in report
    assert "Traceback (most recent call last)" not in report
    assert before == {path.name: digest(path) for path in paths}


def test_all_sanitized_samples_match_machine_expected_results(tmp_path):
    input_dir = MODULE_ROOT / "tests" / "fixtures" / "samples" / "input"
    expected_dir = MODULE_ROOT / "samples" / "expected"
    samples = sorted(input_dir.glob("*.xlsx"))
    assert len(samples) == 8
    for index, sample in enumerate(samples, 1):
        output = tmp_path / sample.stem
        job_path = tmp_path / f"job-{index}.json"
        job = make_job(sample, output)
        job["job_id"] = f"EXPECTED-{index:03d}"
        job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        assert boiling_module.run_job(job_path) == 0
        result = load_result(output)
        record = boiling_module.extract_file(sample)
        expected = json.loads(
            (expected_dir / f"{sample.stem}.expected.json").read_text(encoding="utf-8")
        )
        actual = {
            "expected_status": result["status"],
            "expected_processed_count": result["summary"]["processed_count"],
            "expected_failed_count": result["summary"]["failed_count"],
            "expected_anomaly_fields": sorted(item["field"] for item in result["anomalies"]),
            "expected_confirmation_fields": sorted(
                item["field"] for item in result["confirmations"]
            ),
            "expected_sample_count": len(record["samples"]),
            "expected_retest_count": sum(
                item["sample_type"] == "retest" for item in record["samples"]
            ),
            "expected_immersion_peel_result": record["immersion_peel_result"],
            "expected_average_density_g_cm3": record["average_density_g_cm3"],
            "expected_average_expansion_ratio": record["average_expansion_ratio"],
            "expected_conflict_count": len(record["conflicts"]),
        }
        for key, value in actual.items():
            assert expected[key] == value, f"{sample.name}: {key}"


def test_stdout_contains_only_progress_protocol_events(tmp_path):
    sample = tmp_path / "protocol.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample)
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")

    environment = os.environ.copy()
    environment.pop("PYTHONUTF8", None)
    environment.pop("PYTHONIOENCODING", None)
    completed = subprocess.run(
        [
            sys.executable,
            str(MODULE_ROOT / "src" / "boiling_module.py"),
            "run",
            "--job",
            str(job_path),
        ],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=environment,
    )

    assert completed.returncode == 0
    assert completed.stderr == ""
    events = [json.loads(line) for line in completed.stdout.splitlines()]
    assert events
    assert all(event["type"] == "progress" for event in events)
    assert events[-1]["percent"] == 100


def test_finalize_closes_usage_and_dependent_conclusion_conflicts_in_one_round(tmp_path):
    sample = tmp_path / "combined-conflicts.xlsx"
    make_labeled_sample(
        sample,
        density=1.45,
        expansion=0.10,
        product_usage="",
        conclusion="密度检测合格，2h膨胀率检测合格，浸渍剥离检测合格",
    )
    draft_output = tmp_path / "draft"
    draft_job_path = tmp_path / "draft-job.json"
    draft_job_path.write_text(
        json.dumps(make_job(sample, draft_output), ensure_ascii=False), encoding="utf-8"
    )
    assert boiling_module.run_job(draft_job_path) == 0
    draft_result = load_result(draft_output)
    assert any(item["field"] == "product_usage" for item in draft_result["confirmations"])
    assert any(
        item["field"].startswith("密度｜combined-conflicts.xlsx｜")
        for item in draft_result["confirmations"]
    )

    answers = []
    for prompt in draft_result["confirmations"]:
        value = "家具" if prompt["field"] == "product_usage" else prompt["options"][0]
        answers.append(
            {
                "confirmation_id": prompt["confirmation_id"],
                "value": value,
                "apply_to_batch": prompt["field"] == "product_usage",
                "remember_mapping": False,
            }
        )
    confirmation_path = tmp_path / "confirmation.json"
    confirmation_path.write_text(
        json.dumps({"job_id": "TEST-JOB-001", "answers": answers}, ensure_ascii=False),
        encoding="utf-8",
    )
    final_output = tmp_path / "final"
    final_job = make_job(
        sample,
        final_output,
        phase="finalize",
        parent_job_id="TEST-JOB-001",
        confirmation_file=str(confirmation_path),
    )
    final_job["job_id"] = "TEST-JOB-002"
    final_job_path = tmp_path / "final-job.json"
    final_job_path.write_text(json.dumps(final_job, ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(final_job_path) == 0
    result = load_result(final_output)

    assert result["status"] == "success"
    assert result["confirmations"] == []
    assert {item["field"] for item in result["anomalies"]} >= {"密度", "膨胀率"}


def test_repeated_input_and_answers_have_stable_ids_and_business_outputs(tmp_path):
    sample = tmp_path / "stable.xlsx"
    make_labeled_sample(sample, product_usage="")
    draft_results = []
    for index in (1, 2):
        output = tmp_path / f"draft-{index}"
        job = make_job(sample, output)
        job["job_id"] = f"DRAFT-{index}"
        job_path = tmp_path / f"draft-{index}.json"
        job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        assert boiling_module.run_job(job_path) == 0
        draft_results.append(load_result(output))
    assert [item["confirmation_id"] for item in draft_results[0]["confirmations"]] == [
        item["confirmation_id"] for item in draft_results[1]["confirmations"]
    ]

    confirmation_id = draft_results[0]["confirmations"][0]["confirmation_id"]
    confirmation_path = tmp_path / "stable-confirmation.json"
    confirmation_path.write_text(
        json.dumps(
            {
                "job_id": "DRAFT-1",
                "answers": [
                    {
                        "confirmation_id": confirmation_id,
                        "value": "家具",
                        "apply_to_batch": True,
                        "remember_mapping": False,
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    final_results = []
    final_outputs = []
    for index in (1, 2):
        output = tmp_path / f"final-{index}"
        job = make_job(
            sample,
            output,
            phase="finalize",
            parent_job_id="DRAFT-1",
            confirmation_file=str(confirmation_path),
        )
        job["job_id"] = f"FINAL-{index}"
        job_path = tmp_path / f"final-{index}.json"
        job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        assert boiling_module.run_job(job_path) == 0
        final_results.append(load_result(output))
        final_outputs.append(output)

    assert final_results[0]["status"] == final_results[1]["status"]
    assert [item["anomaly_id"] for item in final_results[0]["anomalies"]] == [
        item["anomaly_id"] for item in final_results[1]["anomalies"]
    ]
    assert (final_outputs[0] / "boiling_records.json").read_text(encoding="utf-8") == (
        final_outputs[1] / "boiling_records.json"
    ).read_text(encoding="utf-8")
    assert (final_outputs[0] / "boiling_analysis.md").read_text(encoding="utf-8") == (
        final_outputs[1] / "boiling_analysis.md"
    ).read_text(encoding="utf-8")
    assert excel_values(final_outputs[0] / "boiling_summary.xlsx") == excel_values(
        final_outputs[1] / "boiling_summary.xlsx"
    )
    assert excel_values(final_outputs[0] / "boiling_analysis.xlsx") == excel_values(
        final_outputs[1] / "boiling_analysis.xlsx"
    )


def test_formal_business_outputs_do_not_leak_absolute_windows_paths(tmp_path):
    sample = tmp_path / "path-scan.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(sample)
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    absolute_path = re.compile(r"(?i)[a-z]:[\\/]")

    for path in (output / "boiling_analysis.md", output / "boiling_records.json"):
        assert not absolute_path.search(path.read_text(encoding="utf-8")), path.name
    workbook = openpyxl.load_workbook(output / "boiling_summary.xlsx", data_only=False)
    try:
        leaked = [
            cell.value
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and absolute_path.search(cell.value)
        ]
    finally:
        workbook.close()
    assert leaked == []

def test_outputs_do_not_leak_absolute_paths_even_with_anomalies(tmp_path):
    sample = tmp_path / "anomaly-scan.xlsx"
    output = tmp_path / "output"
    job_path = tmp_path / "job.json"
    make_labeled_sample(
        sample,
        density=1.45,
        expansion=0.10,
        product_usage="",
        conclusion="密度检测合格，2h膨胀率检测合格，浸渍剥离检测合格",
    )
    job_path.write_text(json.dumps(make_job(sample, output), ensure_ascii=False), encoding="utf-8")
    assert boiling_module.run_job(job_path) == 0
    absolute_path = re.compile(r"(?i)[a-z]:[\\/]")
    result = json.loads((output / "result.json").read_text(encoding="utf-8"))
    assert result["status"] == "needs_confirmation"
    for path in (output / "boiling_analysis.md", output / "boiling_records.json"):
        text = path.read_text(encoding="utf-8")
        matches = absolute_path.findall(text)
        assert not matches, f"{path.name} leaks {len(matches)} absolute paths"
    workbook = openpyxl.load_workbook(output / "boiling_summary.xlsx", data_only=False)
    try:
        leaked = [
            cell.value
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and absolute_path.search(cell.value)
        ]
    finally:
        workbook.close()
    assert leaked == []
    # result.json must not leak the input file absolute path
    input_path = str(sample.absolute())
    result_text = (output / "result.json").read_text(encoding="utf-8")
    assert input_path not in result_text, (
        f"result.json leaks input path: {input_path}"
    )
def test_real_task_service_closes_after_one_confirmation_round(qtbot, tmp_path):
    sample = tmp_path / "task-service.xlsx"
    make_labeled_sample(sample, product_usage="")
    paths = AppPaths.for_root(tmp_path / "app-data")
    database = Database.open(paths)
    database.migrate()
    task_repository = TaskRepository(database)
    confirmation_repository = ConfirmationRepository(database)
    registry = ModuleRegistry(MODULE_ROOT.parent, tmp_path / "unused-mocks", use_mocks=False)
    manifest = ModuleManifest.model_validate_json(
        (MODULE_ROOT / "module.json").read_text(encoding="utf-8")
    )
    registry._modules[ModuleId.BOILING] = RegisteredModule(
        manifest=manifest,
        path=MODULE_ROOT,
        is_mock=False,
        health="healthy",
    )
    service = TaskService(
        paths=paths,
        registry=registry,
        task_repository=task_repository,
        confirmation_repository=confirmation_repository,
    )
    try:
        with qtbot.waitSignal(service.needs_confirmation, timeout=30_000) as waiting:
            job_id = service.create_and_start(ModuleId.BOILING, [sample])
        draft_result = waiting.args[1]
        prompt = next(item for item in draft_result.confirmations if item.field == "product_usage")
        confirmation_service = ConfirmationService(
            paths,
            confirmation_repository,
            task_repository,
        )
        bundle_path = confirmation_service.submit(
            job_id,
            [
                ConfirmationAnswer(
                    confirmation_id=prompt.confirmation_id,
                    value="家具",
                    apply_to_batch=True,
                )
            ],
        )
        with qtbot.waitSignal(service.finished, timeout=30_000) as finished:
            service.finalize(job_id, bundle_path)
        final_result = finished.args[0]
        assert final_result.status.value == "success"
        assert final_result.confirmations == []
        assert {item.type for item in final_result.outputs} == {
            "summary_excel",
            "analysis_report",
            "analysis_report_markdown",
            "business_records",
        }
    finally:
        service.shutdown()
        database.close()


def test_batch_without_python_or_portable_runtime_returns_four(tmp_path):
    module_dir = tmp_path / "boiling"
    module_dir.mkdir()
    launcher = module_dir / "run_module.bat"
    shutil.copy2(MODULE_ROOT / "run_module.bat", launcher)

    completed = run_batch(launcher, ["version"], isolated_windows_environment(tmp_path))

    assert completed.returncode == 4
    assert completed.stdout == ""
    error = json.loads(completed.stderr.strip())
    assert error == {"type": "error", "message": "Python 3.12 was not found"}


def test_batch_prefers_successful_portable_runtime(tmp_path, fake_portable_runtime):
    module_dir = tmp_path / "boiling"
    runtime_dir = module_dir / "runtime"
    runtime_dir.mkdir(parents=True)
    launcher = module_dir / "run_module.bat"
    shutil.copy2(MODULE_ROOT / "run_module.bat", launcher)
    shutil.copy2(fake_portable_runtime, runtime_dir / "boiling_module.exe")

    completed = run_batch(
        launcher,
        ["version"],
        isolated_windows_environment(tmp_path, FAKE_RUNTIME_EXIT="0"),
    )

    assert completed.returncode == 0
    assert "boiling 0.1.0" in completed.stdout
    assert completed.stderr == ""


def test_batch_preserves_portable_runtime_failure_exit_code(tmp_path, fake_portable_runtime):
    module_dir = tmp_path / "boiling"
    runtime_dir = module_dir / "runtime"
    runtime_dir.mkdir(parents=True)
    launcher = module_dir / "run_module.bat"
    shutil.copy2(MODULE_ROOT / "run_module.bat", launcher)
    shutil.copy2(fake_portable_runtime, runtime_dir / "boiling_module.exe")

    completed = run_batch(
        launcher,
        ["self-check"],
        isolated_windows_environment(tmp_path, FAKE_RUNTIME_EXIT="7"),
    )

    assert completed.returncode == 7
    assert completed.stdout == ""
    assert "forced portable runtime failure" in completed.stderr


def test_frozen_runtime_resolves_module_root_and_rules(monkeypatch, tmp_path):
    executable = tmp_path / "modules" / "boiling" / "runtime" / "boiling_module.exe"
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "executable", str(executable))

    module_root = boiling_module.resolve_module_root()

    assert module_root == executable.parent.parent
    assert module_root / "rules" / "default_rules.yaml" == (
        executable.parent.parent / "rules" / "default_rules.yaml"
    )


def test_portable_builder_is_repeatable_and_path_bounded():
    script = (MODULE_ROOT / "build_portable_runtime.ps1").read_text(encoding="utf-8")

    assert "Assert-RepositoryChild" in script
    assert "Portable runtime path escaped repository" in script
    assert "build\\boiling-runtime" in script
    assert "build\\boiling-runtime-verification" in script
    assert "dist\\BambooWorkbench\\modules\\boiling" in script
    assert "Remove-BuildTarget $RuntimeRoot" in script
    assert "--clean --noconfirm --onefile" in script
    assert "Boiling standalone portable runtime built and verified." in script


def test_generated_portable_artifacts_are_not_tracked():
    tracked = subprocess.run(
        ["git", "ls-files", "modules/boiling"],
        cwd=MODULE_ROOT.parents[1],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    ).stdout.splitlines()

    forbidden = [
        path
        for path in tracked
        if Path(path).suffix.lower() == ".exe"
        or any(part in {"runtime", "build", "dist"} for part in Path(path).parts)
    ]
    assert forbidden == []
