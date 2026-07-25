import argparse
import datetime as dt
import hashlib
import json
import math
import re
import sys
import tempfile
from copy import copy
from pathlib import Path


def configure_protocol_streams():
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if callable(reconfigure):
            reconfigure(encoding="utf-8", errors="replace")


configure_protocol_streams()

try:
    import openpyxl
    import yaml
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError as exc:  # pragma: no cover - exercised in a clean environment
    print(json.dumps({"type": "error", "message": str(exc)}), file=sys.stderr)
    raise SystemExit(4) from exc


MODULE_ID = "boiling"
MODULE_VERSION = "1.1.1"
MAX_AI_RECORDS = 50
MAX_AI_PAYLOAD_BYTES = 65_536
MAX_AI_RESPONSE_CHARS = 8_000


def resolve_module_root():
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent.parent
    return Path(__file__).resolve().parents[1]


MODULE_ROOT = resolve_module_root()
DEFAULT_RULES_FILE = MODULE_ROOT / "rules" / "default_rules.yaml"
ALLOWED_JOB_FIELDS = {
    "schema_version",
    "job_id",
    "module_id",
    "phase",
    "run_mode",
    "input_files",
    "output_dir",
    "module_config_file",
    "ai_enabled",
    "allowed_knowledge_scopes",
    "parent_job_id",
    "confirmation_file",
    "service_response_file",
}
FIELD_LABELS = {
    "batch_no": ("批号", "批次号", "批次"),
    "product": ("产品名称", "产品名", "品名"),
    "specification": ("规格型号", "型号规格", "产品规格", "规格"),
    "pressing_date": ("压制日期", "压制时间"),
    "team_leader": ("压机班长", "班组长", "班长"),
    "test_date": ("检测日期", "检验日期", "试验日期", "实验日期"),
    "impregnation_date": ("浸胶日期", "浸渍日期", "浸胶时间"),
    "order_info": ("订单信息", "订单编号", "订单号", "订单"),
    "product_usage": ("产品用途", "用途"),
    "record_status": ("记录状态", "完成状态", "表单状态"),
}
FIXED_FALLBACKS = {
    "batch_no": (2, 1),
    "product": (3, 3),
    "specification": (3, 7),
    "test_date": (5, 3),
}
NON_BLOCKING_METADATA_FIELDS = (
    "product",
    "specification",
    "pressing_date",
    "team_leader",
    "test_date",
    "impregnation_date",
    "order_info",
)
SAMPLE_PATTERN = re.compile(
    r"^(?:端\d+(?:-\d+)?|中\d+(?:-\d+)?|\d+|"
    r"(?:样(?:品|件)?|试件|复测)\s*[#№]?[A-Za-z0-9一二三四五六七八九十_-]+)$",
    re.I,
)
MISSING_PLACEHOLDERS = {"", "/", "／", "-", "—"}


def emit(message, percent=0):
    print(
        json.dumps(
            {"type": "progress", "percent": percent, "message": message},
            ensure_ascii=False,
        ),
        flush=True,
    )


def sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def stable_input_scope(input_files):
    """Return a path-independent identifier shared by draft and finalize snapshots."""
    fingerprints = []
    for raw_path in input_files:
        path = Path(raw_path)
        try:
            token = sha256_file(path) if path.is_file() else f"missing:{path.name}"
        except OSError:
            token = f"unreadable:{path.name}"
        fingerprints.append(token)
    material = "\n".join(sorted(fingerprints)).encode("utf-8")
    return hashlib.sha256(material).hexdigest()[:16]


def sanitize_business_value(value, key=None):
    """Remove machine-local paths from formal business deliverables."""
    if isinstance(value, dict):
        return {
            item_key: sanitize_business_value(item, item_key) for item_key, item in value.items()
        }
    if isinstance(value, list):
        return [sanitize_business_value(item, key) for item in value]
    if key == "source_file" and isinstance(value, str):
        return Path(value).name
    return value


def as_number(value):
    if value is None or value == "" or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def as_text(value):
    return "" if value is None else str(value).strip()


def is_missing_value(value):
    return as_text(value) in MISSING_PLACEHOLDERS


def as_date_text(value):
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%Y-%m-%d")
    text = as_text(value)
    if not text:
        return ""
    for pattern in (r"(20\d{2})[./年-](\d{1,2})[./月-](\d{1,2})",):
        match = re.search(pattern, text)
        if match:
            year, month, day = (int(part) for part in match.groups())
            try:
                return dt.date(year, month, day).isoformat()
            except ValueError:
                return text
    return text.replace("年", "-").replace("月", "-").replace("日", "")


def mean(values):
    clean = [value for value in values if value is not None]
    return sum(clean) / len(clean) if clean else None


def normalize_label(value):
    return re.sub(r"[\s:：()（）【】\[\]]", "", as_text(value)).lower()


def normalize_ratio(value):
    number = as_number(value)
    if number is None:
        return None
    return number / 100 if abs(number) > 1 else number


def number_format_precision(number_format, percentage=False):
    format_text = as_text(number_format).split(";")[0]
    match = re.search(r"0(?:\.(0+))?%?", format_text)
    if not match:
        return None
    precision = len(match.group(1) or "")
    return precision if not percentage else precision


def formatted_number(value, number_format, percentage=False):
    number = as_number(value)
    if number is None:
        return ""
    precision = number_format_precision(number_format, percentage)
    if precision is None:
        precision = 3 if not percentage else 2
    if percentage:
        return f"{number:.{precision}%}"
    return f"{number:.{precision}f}"


def compact_number(value, precision=4):
    number = as_number(value)
    return f"{number:.{precision}f}" if number is not None else "缺失"


def values_conflict(source_value, calculated_value, number_format, percentage=False):
    source_number = as_number(source_value)
    calculated_number = as_number(calculated_value)
    if source_number is None or calculated_number is None:
        return False
    precision = number_format_precision(number_format, percentage)
    if precision is None:
        tolerance = 0.0005 if not percentage else 0.00005
    else:
        scale = 100 if percentage else 1
        tolerance = 0.5 * (10 ** (-precision)) / scale
    return abs(source_number - calculated_number) > tolerance + 1e-12


def display_cell_value(value):
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%Y-%m-%d")
    return as_text(value)


def plausible_fixed_fallback(field, value):
    if value in (None, ""):
        return False
    if field.endswith("_date"):
        if isinstance(value, (dt.date, dt.datetime)):
            return True
        return bool(re.search(r"20\d{2}[./年-]\d{1,2}[./月-]\d{1,2}", as_text(value)))
    normalized = normalize_label(value)
    known_labels = {normalize_label(label) for labels in FIELD_LABELS.values() for label in labels}
    header_words = ("试件", "长度", "宽度", "重量", "密度", "膨胀率", "记录状态", "检测员")
    return normalized not in known_labels and not any(word in normalized for word in header_words)


def find_labeled_value(sheet, labels, max_rows=30, max_columns=12):
    normalized_labels = sorted(
        ((normalize_label(label), label) for label in labels),
        key=lambda item: len(item[0]),
        reverse=True,
    )
    for row in range(1, min(sheet.max_row, max_rows) + 1):
        for column in range(1, min(sheet.max_column, max_columns) + 1):
            raw = sheet.cell(row, column).value
            text = as_text(raw)
            normalized = normalize_label(text)
            if not normalized:
                continue
            for normalized_label, original_label in normalized_labels:
                if normalized == normalized_label:
                    for candidate_column in range(
                        column + 1,
                        min(sheet.max_column, column + 3) + 1,
                    ):
                        candidate = sheet.cell(row, candidate_column).value
                        if candidate not in (None, ""):
                            return candidate, sheet.cell(row, candidate_column).coordinate
                for separator in ("：", ":"):
                    prefix = original_label + separator
                    if text.replace(" ", "").startswith(prefix):
                        remainder = text.replace(" ", "")[len(prefix) :].strip()
                        if remainder:
                            return remainder, sheet.cell(row, column).coordinate
    return None, None


def find_measurement_header(sheet):
    for row in range(1, min(sheet.max_row, 80) + 1):
        headers = {
            column: normalize_label(sheet.cell(row=row, column=column).value)
            for column in range(1, min(sheet.max_column, 20) + 1)
        }
        values = list(headers.values())
        if not any(value in {"序号", "试件", "样件", "样品"} for value in values):
            continue
        if not any("长度" in value for value in values) or not any(
            "密度" in value for value in values
        ):
            continue
        label_column = next(
            column for column, value in headers.items() if value in {"序号", "试件", "样件", "样品"}
        )
        density_column = next(column for column, value in headers.items() if "密度" in value)
        expansion_column = next(
            (column for column, value in headers.items() if "膨胀率" in value),
            None,
        )
        thickness_columns = [column for column, value in headers.items() if "厚度" in value]
        before_column = next(
            (column for column in thickness_columns if column < density_column),
            None,
        )
        after_column = next(
            (column for column in thickness_columns if column > density_column),
            None,
        )
        return {
            "row": row,
            "label": label_column,
            "length": next(
                (column for column, value in headers.items() if "长度" in value),
                None,
            ),
            "width": next(
                (column for column, value in headers.items() if "宽度" in value),
                None,
            ),
            "before": before_column,
            "weight": next(
                (column for column, value in headers.items() if "重量" in value or "质量" in value),
                None,
            ),
            "density": density_column,
            "after": after_column,
            "expansion": expansion_column,
        }
    return None


def is_sample_label(value):
    text = as_text(value)
    return bool(text and SAMPLE_PATTERN.fullmatch(text))


def find_sample_markers(sheet, header=None):
    header = header or find_measurement_header(sheet)
    if not header:
        return []
    markers = []
    column = header["label"]
    for row in range(header["row"] + 1, sheet.max_row + 1):
        label = as_text(sheet.cell(row=row, column=column).value)
        if not is_sample_label(label):
            continue
        structural_values = [
            as_number(sheet.cell(row=row, column=header[key]).value)
            for key in ("length", "width", "density")
            if header.get(key)
        ]
        if any(value is not None for value in structural_values):
            markers.append({"row": row, "column": column, "label": label})
    return markers


def find_sample_rows(sheet):
    return [marker["row"] for marker in find_sample_markers(sheet)]


def worksheet_score(sheet):
    score = len(find_sample_markers(sheet)) * 100
    for labels in FIELD_LABELS.values():
        value, _ = find_labeled_value(sheet, labels)
        if value not in (None, ""):
            score += 2
    for row in range(1, min(sheet.max_row, 80) + 1):
        joined = " ".join(
            as_text(sheet.cell(row, column).value)
            for column in range(1, min(sheet.max_column, 12) + 1)
        )
        if "检测结论" in joined or "检验结论" in joined:
            score += 10
            break
    return score


def select_data_sheet(workbook):
    return max(workbook.worksheets, key=worksheet_score)


def first_numeric(value_getter, rows, column):
    for row in rows:
        number = as_number(value_getter(row, column))
        if number is not None:
            return number
    return None


def find_conclusion_row(sheet, start_row):
    for row in range(start_row, sheet.max_row + 1):
        joined = " ".join(
            as_text(sheet.cell(row, column).value)
            for column in range(1, min(sheet.max_column, 12) + 1)
        )
        if "检测结论" in joined or "检验结论" in joined:
            return row
    return None


def source_ref(path, sheet, cell_or_range=None):
    return {
        "source_file": str(Path(path).resolve()),
        "sheet": sheet,
        "cell_or_range": cell_or_range,
    }


def extract_file(path):
    formula_book = openpyxl.load_workbook(path, data_only=False, read_only=False)
    value_book = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        sheet = select_data_sheet(formula_book)
        value_sheet = value_book[sheet.title]
        def value(row, column):
            calculated = value_sheet.cell(row=row, column=column).value
            return (
                calculated if calculated is not None else sheet.cell(row=row, column=column).value
            )

        record = {
            "source_file": str(Path(path).resolve()),
            "source_sheet": sheet.title,
            "source_used_range": sheet.calculate_dimension(),
            "batch_no": "",
            "product": "",
            "specification": "",
            "pressing_date": "",
            "team_leader": "",
            "test_date": "",
            "impregnation_date": "",
            "order_info": "",
            "product_usage": "",
            "record_status": "",
            "field_sources": {},
            "samples": [],
            "source_conclusion": "",
            "tester": "",
            "missing_fields": [],
            "blocking_reasons": [],
            "conflicts": [],
            "conclusion_rule_conflicts": [],
            "rule_observations": {},
            "business_status": "信息不足",
        }
        for field, labels in FIELD_LABELS.items():
            raw, coordinate = find_labeled_value(sheet, labels)
            if raw in (None, "") and field in FIXED_FALLBACKS:
                row, column = FIXED_FALLBACKS[field]
                candidate = value(row, column)
                if plausible_fixed_fallback(field, candidate):
                    raw = candidate
                    coordinate = sheet.cell(row, column).coordinate
            record[field] = (
                as_date_text(raw) if field.endswith("_date") else display_cell_value(raw)
            )
            if coordinate:
                record["field_sources"][field] = coordinate
        record["batch_no"] = re.sub(r"^\s*(?:NO\.?|编号)\s*", "", record["batch_no"], flags=re.I)

        header = find_measurement_header(sheet)
        markers = find_sample_markers(sheet, header)
        conclusion_row = find_conclusion_row(sheet, 1)
        for sample_index, marker in enumerate(markers):
            start = marker["row"]
            next_start = (
                markers[sample_index + 1]["row"] if sample_index + 1 < len(markers) else None
            )
            block_end = next_start - 1 if next_start is not None else sheet.max_row
            if conclusion_row is not None and conclusion_row > start:
                block_end = min(block_end, conclusion_row - 1)
            block_end = min(block_end, start + 20)
            rows = list(range(start, block_end + 1))
            label_column = header["label"]
            measurement_rows = [
                row
                for row in rows
                if normalize_label(value(row, label_column)) not in {"平均值", "平均"}
                and any(
                    as_number(value(row, header[key])) is not None
                    for key in ("before", "after")
                    if header.get(key)
                )
            ]
            before_values = [
                as_number(value(row, header["before"]))
                for row in measurement_rows
                if header.get("before") and as_number(value(row, header["before"])) is not None
            ]
            after_values = [
                as_number(value(row, header["after"]))
                for row in measurement_rows
                if header.get("after") and as_number(value(row, header["after"])) is not None
            ]
            avg_before = mean(before_values)
            avg_after = mean(after_values)
            length = first_numeric(value, rows, header["length"]) if header.get("length") else None
            width = first_numeric(value, rows, header["width"]) if header.get("width") else None
            weight = first_numeric(value, rows, header["weight"]) if header.get("weight") else None
            density_rows = [
                row for row in rows if as_number(value(row, header["density"])) is not None
            ]
            expansion_rows = [
                row
                for row in rows
                if header.get("expansion")
                and normalize_ratio(value(row, header["expansion"])) is not None
            ]
            density_row = density_rows[0] if density_rows else None
            expansion_row = expansion_rows[0] if expansion_rows else None
            source_density = (
                as_number(value(density_row, header["density"])) if density_row else None
            )
            source_expansion = (
                normalize_ratio(value(expansion_row, header["expansion"]))
                if expansion_row and header.get("expansion")
                else None
            )
            density_cell = sheet.cell(density_row, header["density"]) if density_row else None
            expansion_cell = (
                sheet.cell(expansion_row, header["expansion"])
                if expansion_row and header.get("expansion")
                else None
            )
            volume_cm3 = (
                length * width * avg_before / 1000 if length and width and avg_before else None
            )
            calculated_density = weight / volume_cm3 if weight is not None and volume_cm3 else None
            calculated_expansion = (
                (avg_after - avg_before) / avg_before
                if avg_before not in (None, 0) and avg_after is not None
                else None
            )
            density = calculated_density if calculated_density is not None else source_density
            expansion = (
                calculated_expansion if calculated_expansion is not None else source_expansion
            )
            label = marker["label"]
            source_end_column = header.get("expansion") or header["density"]
            sample = {
                "label": label,
                "sample_type": "retest"
                if "复测" in label or "复测" in Path(path).stem
                else "normal",
                "length_mm": length,
                "width_mm": width,
                "weight_g": weight,
                "thickness_before_values_mm": before_values,
                "thickness_after_values_mm": after_values,
                "average_thickness_before_mm": avg_before,
                "average_thickness_after_mm": avg_after,
                "volume_cm3": volume_cm3,
                "source_density_g_cm3": source_density,
                "source_density_display": formatted_number(
                    source_density,
                    density_cell.number_format if density_cell else "",
                ),
                "source_density_formula": density_cell.value
                if density_cell and density_cell.data_type == "f"
                else None,
                "calculated_density_g_cm3": calculated_density,
                "density_g_cm3": density,
                "density_source_cell": density_cell.coordinate if density_cell else None,
                "source_expansion_ratio": source_expansion,
                "source_expansion_display": formatted_number(
                    source_expansion,
                    expansion_cell.number_format if expansion_cell else "",
                    percentage=True,
                ),
                "source_expansion_formula": (
                    expansion_cell.value
                    if expansion_cell and expansion_cell.data_type == "f"
                    else None
                ),
                "calculated_expansion_ratio": calculated_expansion,
                "expansion_ratio": expansion,
                "expansion_source_cell": expansion_cell.coordinate if expansion_cell else None,
                "source_range": (
                    f"{sheet.cell(start, label_column).coordinate}:"
                    f"{sheet.cell(block_end, source_end_column).coordinate}"
                ),
                "density_rule_result": "信息不足",
            }
            if values_conflict(
                source_density,
                calculated_density,
                density_cell.number_format if density_cell else "",
            ):
                record["conflicts"].append(
                    {
                        "field": "density_g_cm3",
                        "sample": label,
                        "source_value": source_density,
                        "source_display": sample["source_density_display"],
                        "calculated_value": calculated_density,
                        "difference": abs(source_density - calculated_density),
                        "cell_or_range": density_cell.coordinate,
                    }
                )
            if values_conflict(
                source_expansion,
                calculated_expansion,
                expansion_cell.number_format if expansion_cell else "",
                percentage=True,
            ):
                record["conflicts"].append(
                    {
                        "field": "expansion_ratio",
                        "sample": label,
                        "source_value": source_expansion,
                        "source_display": sample["source_expansion_display"],
                        "calculated_value": calculated_expansion,
                        "difference": abs(source_expansion - calculated_expansion),
                        "cell_or_range": expansion_cell.coordinate,
                    }
                )
            record["samples"].append(sample)

        for row in range(1, sheet.max_row + 1):
            for column in range(1, min(sheet.max_column, 20) + 1):
                text = as_text(sheet.cell(row, column).value)
                if "检测结论" not in text and "检验结论" not in text:
                    continue
                remainder = re.sub(r"^.*?(?:检测|检验)结论\s*[：:]?", "", text).strip()
                if remainder:
                    record["source_conclusion"] = remainder
                else:
                    for candidate_column in range(
                        column + 1, min(sheet.max_column, column + 4) + 1
                    ):
                        candidate = as_text(value(row, candidate_column))
                        if candidate and "检测员" not in candidate and "检验员" not in candidate:
                            record["source_conclusion"] = candidate
                            break
                record["field_sources"]["source_conclusion"] = sheet.cell(row, column).coordinate
                break
            if "source_conclusion" in record["field_sources"]:
                break
        tester_value, tester_coordinate = find_labeled_value(
            sheet,
            ("检测员", "检验员", "测试员"),
            max_rows=sheet.max_row,
        )
        record["tester"] = as_text(tester_value)
        if tester_coordinate:
            record["field_sources"]["tester"] = tester_coordinate

        densities = [
            sample["density_g_cm3"]
            for sample in record["samples"]
            if sample["density_g_cm3"] is not None
        ]
        expansions = [
            sample["expansion_ratio"]
            for sample in record["samples"]
            if sample["expansion_ratio"] is not None
        ]
        record["average_density_g_cm3"] = mean(densities)
        record["density_min_g_cm3"] = min(densities) if densities else None
        record["density_max_g_cm3"] = max(densities) if densities else None
        record["density_range_g_cm3"] = (
            record["density_max_g_cm3"] - record["density_min_g_cm3"] if densities else None
        )
        record["average_expansion_ratio"] = mean(expansions)
        conclusion = record["source_conclusion"]
        if re.search(r"浸渍剥离.{0,8}不合格", conclusion):
            record["immersion_peel_result"] = "不合格"
        elif re.search(r"浸渍剥离.{0,8}合格", conclusion):
            record["immersion_peel_result"] = "合格"
        else:
            record["immersion_peel_result"] = ""
        for field in NON_BLOCKING_METADATA_FIELDS:
            if is_missing_value(record[field]):
                record["missing_fields"].append(field)
        if is_missing_value(record["product_usage"]):
            record["missing_fields"].append("product_usage")
        if not record["source_conclusion"]:
            record["missing_fields"].append("source_conclusion")
        computable_samples = [
            sample
            for sample in record["samples"]
            if sample["density_g_cm3"] is not None and sample["expansion_ratio"] is not None
        ]
        record["computable_sample_count"] = len(computable_samples)
        if not record["samples"]:
            record["blocking_reasons"].append("未识别到测量试件")
        elif not computable_samples:
            record["blocking_reasons"].append("无可同时复算密度和膨胀率的试件")
        status_text = normalize_label(record["record_status"])
        if "未完成" in status_text or "空表" in status_text:
            record["blocking_reasons"].append("原表标记为未完成")
        if not record["source_conclusion"] and is_missing_value(record["test_date"]):
            record["blocking_reasons"].append("检测日期和检测结论均缺失，判定为未完成表")
        record["is_incomplete"] = bool(record["blocking_reasons"])
        return record
    finally:
        formula_book.close()
        value_book.close()


def load_rules(config_file=None):
    rules_path = DEFAULT_RULES_FILE
    if config_file:
        config_path = Path(config_file)
        data = yaml.safe_load(config_path.read_text(encoding="utf-8-sig")) or {}
        candidate = data.get("rules_file") if isinstance(data, dict) else None
        if candidate:
            candidate_path = Path(candidate)
            rules_path = (
                candidate_path
                if candidate_path.is_absolute()
                else config_path.parent / candidate_path
            )
        elif isinstance(data, dict) and ("density" in data or "expansion" in data):
            return data
    data = yaml.safe_load(rules_path.read_text(encoding="utf-8-sig")) or {}
    if not isinstance(data, dict):
        raise ValueError("规则文件根节点必须是对象")
    return data


def make_anomaly(
    record,
    fingerprint,
    suffix,
    severity,
    field,
    value,
    rule_id,
    message,
    suggestion,
    cell_or_range=None,
):
    return {
        "anomaly_id": f"A-{suffix}-{fingerprint[:12]}",
        "severity": severity,
        "field": field,
        "value": value,
        "rule_id": rule_id,
        "message": message,
        "suggestion": suggestion,
        "source": source_ref(record["source_file"], record["source_sheet"], cell_or_range),
    }


def conclusion_pass_claim(conclusion, keywords):
    for keyword in keywords:
        match = re.search(rf"{keyword}.{{0,16}}?(不合格|合格)", conclusion)
        if match:
            return match.group(1) == "合格"
    return None


def add_conclusion_rule_conflict(
    record,
    anomalies,
    confirmations,
    fingerprint,
    field,
    rule_id,
    source_claim,
    observed_outside,
    observed_value,
    conclusion_cell=None,
):
    if source_claim is None or observed_outside is None or source_claim == (not observed_outside):
        return
    source_text = "合格" if source_claim else "不合格"
    rule_text = "不合格" if observed_outside else "合格"
    source_name = Path(record["source_file"]).name
    display_field = f"{field}｜{source_name}｜原表{source_text}→规则{rule_text}（建议按规则）"
    if isinstance(observed_value, dict):
        observed_detail = (
            f"最小值 {compact_number(observed_value.get('minimum'))}、"
            f"最大值 {compact_number(observed_value.get('maximum'))}、"
            f"密度差 {compact_number(observed_value.get('difference'))}"
        )
    elif observed_value is not None:
        observed_detail = f"规则复核值 {observed_value:.2%}"
    else:
        observed_detail = "规则复核值缺失"
    conflict = {
        "field": field,
        "source_conclusion": source_text,
        "rule_observation": rule_text,
        "observed_value": observed_value,
        "rule_id": rule_id,
    }
    record["conclusion_rule_conflicts"].append(conflict)
    conflict_index = len(record["conclusion_rule_conflicts"])
    anomalies.append(
        make_anomaly(
            record,
            fingerprint,
            f"CONCLUSION-RULE-{conflict_index}",
            "warning",
            field,
            {
                "source_conclusion": source_text,
                "rule_observation": rule_text,
                "value": observed_value,
            },
            "BOILING-CONCLUSION-CONFLICT-001",
            f"原表{field}结论为{source_text}，但按本版规则复核为{rule_text}；两者均已保留。",
            "复核原表签字结论、测量数据和规则版本；人工确认前不得覆盖原表结论。",
        )
    )
    confirmations.append(
        {
            "confirmation_id": f"C-CONCLUSION-RULE-{field}-{fingerprint[:12]}",
            "field": display_field,
            "reason": (
                f"{source_name}：原表结论“{source_text}”与规则观察“{rule_text}”不一致；"
                f"{observed_detail}。建议按现行规则选择“{rule_text}”，原表结论仍会保留。"
            ),
            "suggested_value": rule_text,
            "options": [source_text, rule_text],
            "confidence": 0.0,
            "blocking": False,
            "batch_applicable": False,
            "rememberable": False,
            "source": source_ref(record["source_file"], record["source_sheet"], conclusion_cell),
        }
    )


def evaluate_record(record, rules, fingerprint, crack_inspection_confirmed=False):
    anomalies = []
    confirmations = []
    rule_hits = []
    density_rule = rules.get("density") or {}
    density_rule_id = density_rule.get("rule_id")
    minimum = as_number(density_rule.get("minimum_g_cm3"))
    maximum = as_number(density_rule.get("maximum_g_cm3"))
    maximum_difference = as_number(density_rule.get("maximum_difference_g_cm3"))
    density_samples = [
        sample for sample in record["samples"] if sample["density_g_cm3"] is not None
    ]
    outside_samples = []
    for sample in density_samples:
        value = sample["density_g_cm3"]
        outside = (minimum is not None and value < minimum - 1e-12) or (
            maximum is not None and value > maximum + 1e-12
        )
        sample["density_rule_result"] = "不合格" if outside else "合格"
        if outside:
            outside_samples.append(sample["label"])
    density_difference = record.get("density_range_g_cm3")
    difference_outside = (
        density_difference is not None
        and maximum_difference is not None
        and density_difference > maximum_difference + 1e-12
    )
    density_pass = bool(density_samples) and not outside_samples and not difference_outside
    record["density_rule_pass"] = density_pass if density_samples else None
    record["density_outside_samples"] = outside_samples
    record["rule_observations"]["density"] = (
        "合格" if density_pass else ("不合格" if density_samples else "信息不足")
    )
    rule_hits.append(
        {
            "rule_id": density_rule_id,
            "field": "density",
            "result": record["rule_observations"]["density"],
            "minimum_g_cm3": record.get("density_min_g_cm3"),
            "maximum_g_cm3": record.get("density_max_g_cm3"),
            "difference_g_cm3": density_difference,
            "outside_samples": outside_samples,
            "source_cells": [
                sample["density_source_cell"]
                for sample in density_samples
                if sample.get("density_source_cell")
            ],
        }
    )
    if density_samples and not density_pass:
        reasons = []
        if outside_samples:
            reasons.append("超出 1.1～1.3 g/cm³ 的试件：" + "、".join(outside_samples))
        if difference_outside:
            reasons.append(f"密度差 {density_difference:.6f} g/cm³ 大于 0.1 g/cm³")
        anomalies.append(
            make_anomaly(
                record,
                fingerprint,
                "DENSITY",
                "warning",
                "density",
                {
                    "minimum": record.get("density_min_g_cm3"),
                    "maximum": record.get("density_max_g_cm3"),
                    "difference": density_difference,
                    "outside_samples": outside_samples,
                },
                density_rule_id,
                "；".join(reasons),
                "核对各试件密度、来源单元格和密度差；边界 1.1、1.3、0.1 均按合格处理。",
            )
        )

    expansion_rule = rules.get("expansion") or {}
    expansion_rule_id = expansion_rule.get("rule_id")
    expansion_value = record.get("average_expansion_ratio")
    usage = as_text(record.get("product_usage"))
    excellent_max = normalize_ratio(expansion_rule.get("excellent_max_ratio"))
    numeric_pass = None
    expansion_grade = ""
    human_review = False
    crack_review_required = False
    rule_detail = ""
    if expansion_value is None:
        rule_detail = "信息不足"
    elif not usage:
        rule_detail = "待确认产品用途"
    elif usage in {"地板", "家具"}:
        pass_max = normalize_ratio(
            (expansion_rule.get("floor_furniture") or {}).get("pass_max_ratio")
        )
        if excellent_max is not None and expansion_value <= excellent_max + 1e-12:
            numeric_pass = True
            expansion_grade = "绿色优等候选"
            human_review = True
            crack_review_required = True
            rule_detail = "膨胀率≤6.5%，需结合开裂观察确认绿色优等"
        elif pass_max is not None and expansion_value <= pass_max + 1e-12:
            numeric_pass = True
            expansion_grade = "蓝色合格"
            human_review = True
            crack_review_required = True
            rule_detail = "6.5%＜膨胀率≤8%，膨胀率维度蓝色合格，仍需核对开裂观察"
        else:
            numeric_pass = False
            expansion_grade = "红色不合格"
            rule_detail = "膨胀率＞8%，红色不合格"
    elif usage in {"墙板", "栏杆"}:
        blue_min = normalize_ratio((expansion_rule.get("wall_railing") or {}).get("blue_min_ratio"))
        pass_max = normalize_ratio((expansion_rule.get("wall_railing") or {}).get("pass_max_ratio"))
        if excellent_max is not None and expansion_value <= excellent_max + 1e-12:
            numeric_pass = True
            expansion_grade = "绿色优等候选"
            human_review = True
            crack_review_required = True
            rule_detail = "膨胀率≤6.5%，需结合开裂观察确认绿色优等"
        elif blue_min is not None and expansion_value < blue_min - 1e-12:
            human_review = True
            rule_detail = "墙板/栏杆膨胀率＞6.5%且＜8%，v0.1 规则缺口转人工复核"
        elif pass_max is not None and expansion_value <= pass_max + 1e-12:
            numeric_pass = True
            expansion_grade = "蓝色合格"
            human_review = True
            crack_review_required = True
            rule_detail = "8%≤膨胀率≤12%，膨胀率维度蓝色合格，仍需核对开裂观察"
        else:
            numeric_pass = False
            expansion_grade = "红色不合格"
            rule_detail = "膨胀率＞12%，红色不合格"
    else:
        human_review = True
        rule_detail = f"产品用途“{usage}”没有自动判定规则"

    if crack_review_required and crack_inspection_confirmed:
        human_review = False
        rule_detail = rule_detail.replace(
            "需结合开裂观察确认绿色优等",
            "开裂已人工检查且符合绿色优等要求",
        ).replace(
            "仍需核对开裂观察",
            "开裂已人工检查且符合对应等级",
        )

    record["expansion_numeric_pass"] = numeric_pass
    record["expansion_grade"] = expansion_grade
    record["expansion_human_review_required"] = human_review
    record["expansion_crack_review_applicable"] = crack_review_required
    record["crack_inspection_confirmed"] = bool(crack_inspection_confirmed)
    record["crack_inspection_status"] = (
        "已检查且符合对应等级"
        if record["crack_inspection_confirmed"]
        else ("待确认" if expansion_value is not None else "不适用")
    )
    record["rule_observations"]["expansion"] = rule_detail
    rule_hits.append(
        {
            "rule_id": expansion_rule_id,
            "field": "expansion",
            "result": rule_detail,
            "product_usage": usage,
            "average_expansion_ratio": expansion_value,
            "human_review_required": human_review,
            "crack_review_applicable": crack_review_required,
            "crack_inspection_confirmed": record["crack_inspection_confirmed"],
            "source_cells": [
                sample["expansion_source_cell"]
                for sample in record["samples"]
                if sample.get("expansion_source_cell")
            ],
        }
    )
    if numeric_pass is False:
        anomalies.append(
            make_anomaly(
                record,
                fingerprint,
                "EXPANSION",
                "warning",
                "average_expansion_ratio",
                expansion_value,
                expansion_rule_id,
                rule_detail,
                "复核产品用途、平均膨胀率和来源单元格，并保留原表结论供人工复核。",
            )
        )
    if human_review and usage and numeric_pass is None:
        reason = rule_detail
        confirmations.append(
            {
                "confirmation_id": f"C-EXPANSION-REVIEW-{fingerprint[:12]}",
                "field": "expansion_review",
                "reason": reason,
                "suggested_value": None,
                "options": ["开裂观察符合对应等级", "开裂观察不符合对应等级", "继续待复核"],
                "confidence": 0.0,
                "blocking": False,
                "batch_applicable": False,
                "rememberable": False,
                "source": source_ref(record["source_file"], record["source_sheet"]),
            }
        )
    record["rule_hits"] = rule_hits
    conclusion = record.get("source_conclusion", "")
    add_conclusion_rule_conflict(
        record,
        anomalies,
        confirmations,
        fingerprint,
        "密度",
        density_rule.get("rule_id"),
        conclusion_pass_claim(conclusion, ("密度",)),
        (not density_pass) if density_samples else None,
        {
            "minimum": record.get("density_min_g_cm3"),
            "maximum": record.get("density_max_g_cm3"),
            "difference": density_difference,
        },
        conclusion_cell=None,
    )
    add_conclusion_rule_conflict(
        record,
        anomalies,
        confirmations,
        fingerprint,
        "膨胀率",
        expansion_rule.get("rule_id"),
        conclusion_pass_claim(conclusion, ("膨胀率", "膨胀")),
        (not numeric_pass) if numeric_pass is not None else None,
        expansion_value,
        conclusion_cell=None,
    )
    if record.get("immersion_peel_result") == "不合格":
        anomalies.append(
            make_anomaly(
                record,
                fingerprint,
                "PEEL",
                "warning",
                "immersion_peel_result",
                "不合格",
                "BOILING-PEEL-SOURCE-001",
                "原表记录浸渍剥离检测不合格。",
                "保留原表结论并转人工复核；本模块不替代检测人员作正式判定。",
            )
        )
    for index, conflict in enumerate(record.get("conflicts", []), 1):
        source_name = Path(record["source_file"]).name
        source_display = conflict.get("source_display") or compact_number(
            conflict["source_value"]
        )
        if conflict["field"] == "expansion_ratio":
            field_name = "膨胀率"
            calculated_display = f"{conflict['calculated_value']:.2%}"
        else:
            field_name = "密度"
            calculated_display = f"{conflict['calculated_value']:.3f}"
        display_field = (
            f"{field_name}｜{source_name}｜{conflict['sample']}"
            f"（{conflict.get('cell_or_range') or '来源位置未知'}）｜"
            f"原表{source_display}→复算{calculated_display}（建议复算值）"
        )
        anomalies.append(
            make_anomaly(
                record,
                fingerprint,
                f"CONFLICT-{index}",
                "warning",
                conflict["field"],
                {"source": conflict["source_value"], "calculated": conflict["calculated_value"]},
                "BOILING-CALC-CONFLICT-001",
                f"{conflict['sample']} 的原表值与模块复算值不一致。",
                "核对公式、单位和原始测量记录；在确认前同时保留两个值。",
                conflict.get("cell_or_range"),
            )
        )
        confirmations.append(
            {
                "confirmation_id": f"C-CONFLICT-{index}-{fingerprint[:12]}",
                "field": display_field,
                "reason": (
                    f"{source_name} 的 {conflict['sample']}："
                    f"原表值 {source_display} 与复算值 {calculated_display} 冲突；"
                    "建议采用复算值，"
                    "两者仍会在报告中保留。"
                ),
                "suggested_value": conflict["calculated_value"],
                "options": [conflict["source_value"], conflict["calculated_value"]],
                "confidence": 0.5,
                "blocking": False,
                "batch_applicable": False,
                "rememberable": False,
                "source": source_ref(
                    record["source_file"], record["source_sheet"], conflict.get("cell_or_range")
                ),
            }
        )
    return anomalies, confirmations


def style_header(cells):
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in cells:
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_result_cell(cell, value):
    """Apply the delivery-standard color while keeping the text conclusion visible."""
    text = str(value or "")
    if "红色" in text or text in {"异常", "不合格"}:
        color = "F4CCCC"
    elif "蓝色" in text:
        color = "CFE2F3"
    elif "绿色" in text or text in {"正常", "合格"}:
        color = "D9EAD3"
    elif text in {"待确认", "是"} or "待确认" in text or "人工复核" in text:
        color = "FFF2CC"
    elif text == "信息不足":
        color = "D9D9D9"
    else:
        return
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(name="微软雅黑", color="000000")


def write_summary(records, output_path, rules, excluded_records=None):
    excluded_records = excluded_records or []
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "水煮检测汇总"
    headers = [
        "来源文件",
        "批号",
        "产品",
        "规格",
        "压制日期",
        "班长",
        "检测日期",
        "浸胶日期",
        "订单信息",
        "产品用途",
        "试件数",
        "复测试件数",
        "端1-1密度(g/cm³)",
        "端1-1膨胀率",
        "端1-2密度(g/cm³)",
        "端1-2膨胀率",
        "中1密度(g/cm³)",
        "中1膨胀率",
        "端2-1密度(g/cm³)",
        "端2-1膨胀率",
        "端2-2密度(g/cm³)",
        "端2-2膨胀率",
        "平均密度(g/cm³)",
        "最小密度(g/cm³)",
        "最大密度(g/cm³)",
        "密度差(g/cm³)",
        "密度规则",
        "平均膨胀率",
        "膨胀率规则",
        "自动等级",
        "需人工复核",
        "浸渍剥离",
        "原表结论",
        "业务状态",
        "复算冲突数",
        "结论/规则冲突数",
    ]
    summary.append(headers)
    style_header(summary[1])
    for record in records:
        samples_by_label = {}
        for sample in record["samples"]:
            samples_by_label.setdefault(str(sample.get("label") or "").strip(), sample)
        position_values = []
        for label in ("端1-1", "端1-2", "中1", "端2-1", "端2-2"):
            sample = samples_by_label.get(label, {})
            position_values.extend(
                [
                    sample.get("density_g_cm3") or 0,
                    sample.get("expansion_ratio") or 0,
                ]
            )
        summary.append(
            [
                Path(record["source_file"]).name,
                record["batch_no"],
                record["product"],
                record["specification"],
                record["pressing_date"],
                record["team_leader"],
                record["test_date"],
                record["impregnation_date"],
                record["order_info"],
                record["product_usage"],
                len(record["samples"]),
                sum(1 for sample in record["samples"] if sample["sample_type"] == "retest"),
                *position_values,
                record["average_density_g_cm3"],
                record["density_min_g_cm3"],
                record["density_max_g_cm3"],
                record["density_range_g_cm3"],
                record["rule_observations"].get("density", ""),
                record["average_expansion_ratio"],
                record["rule_observations"].get("expansion", ""),
                record.get("expansion_grade", ""),
                "是" if record.get("expansion_human_review_required") else "否",
                record["immersion_peel_result"],
                record["source_conclusion"],
                record["business_status"],
                len(record["conflicts"]),
                len(record["conclusion_rule_conflicts"]),
            ]
        )
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    summary.sheet_view.showGridLines = False
    for column in ("M", "O", "Q", "S", "U", "W", "X", "Y", "Z"):
        for cell in summary[column][1:]:
            cell.number_format = "0.000"
    for column in ("N", "P", "R", "T", "V", "AB"):
        for cell in summary[column][1:]:
            cell.number_format = "0.00%"
    for row in range(2, summary.max_row + 1):
        style_result_cell(summary[f"AA{row}"], summary[f"AA{row}"].value)
        style_result_cell(summary[f"AD{row}"], summary[f"AD{row}"].value)
        style_result_cell(summary[f"AE{row}"], summary[f"AE{row}"].value)
        style_result_cell(summary[f"AH{row}"], summary[f"AH{row}"].value)
    widths = {
        "A": 24,
        "B": 16,
        "C": 18,
        "D": 16,
        "E": 12,
        "F": 14,
        "G": 12,
        "H": 12,
        "I": 18,
        "J": 12,
        "K": 10,
        "L": 10,
        "M": 17,
        "N": 14,
        "O": 17,
        "P": 14,
        "Q": 17,
        "R": 14,
        "S": 17,
        "T": 14,
        "U": 17,
        "V": 14,
        "W": 15,
        "X": 15,
        "Y": 15,
        "Z": 15,
        "AA": 12,
        "AB": 14,
        "AC": 50,
        "AD": 16,
        "AE": 12,
        "AF": 12,
        "AG": 48,
        "AH": 12,
        "AI": 12,
        "AJ": 16,
    }
    for column, width in widths.items():
        summary.column_dimensions[column].width = width

    excluded_sheet = workbook.create_sheet("未纳入正常统计")
    excluded_sheet.append(
        [
            "来源文件",
            "批号",
            "产品",
            "规格",
            "业务状态",
            "缺失字段",
            "阻断原因",
            "处理说明",
        ]
    )
    style_header(excluded_sheet[1])
    for record in excluded_records:
        excluded_sheet.append(
            [
                Path(record["source_file"]).name,
                record["batch_no"],
                record["product"],
                record["specification"],
                "信息不足",
                ", ".join(record["missing_fields"]),
                "；".join(record.get("blocking_reasons", [])),
                "空表、未完成表或无可计算试件；未进入正常汇总和图表统计。",
            ]
        )
    excluded_sheet.auto_filter.ref = excluded_sheet.dimensions
    excluded_sheet.sheet_view.showGridLines = False
    for row in range(2, excluded_sheet.max_row + 1):
        style_result_cell(excluded_sheet[f"E{row}"], excluded_sheet[f"E{row}"].value)
    for column, width in {
        "A": 26,
        "B": 16,
        "C": 18,
        "D": 18,
        "E": 14,
        "F": 34,
        "G": 48,
        "H": 54,
    }.items():
        excluded_sheet.column_dimensions[column].width = width

    details = workbook.create_sheet("试件明细")
    detail_headers = [
        "来源文件",
        "批号",
        "试件",
        "类型",
        "长(mm)",
        "宽(mm)",
        "重量(g)",
        "煮前平均厚度(mm)",
        "煮后平均厚度(mm)",
        "体积(cm³)",
        "原表密度",
        "原表显示密度",
        "复算密度",
        "采用密度",
        "密度规则",
        "密度来源单元格",
        "原表膨胀率",
        "原表显示膨胀率",
        "复算膨胀率",
        "采用膨胀率",
        "膨胀率来源单元格",
        "来源区域",
    ]
    details.append(detail_headers)
    style_header(details[1])
    for record in records:
        for sample in record["samples"]:
            details.append(
                [
                    Path(record["source_file"]).name,
                    record["batch_no"],
                    sample["label"],
                    sample["sample_type"],
                    sample["length_mm"],
                    sample["width_mm"],
                    sample["weight_g"],
                    sample["average_thickness_before_mm"],
                    sample["average_thickness_after_mm"],
                    sample["volume_cm3"],
                    sample["source_density_g_cm3"],
                    sample["source_density_display"],
                    sample["calculated_density_g_cm3"],
                    sample["density_g_cm3"],
                    sample["density_rule_result"],
                    sample["density_source_cell"],
                    sample["source_expansion_ratio"],
                    sample["source_expansion_display"],
                    sample["calculated_expansion_ratio"],
                    sample["expansion_ratio"],
                    sample["expansion_source_cell"],
                    sample["source_range"],
                ]
            )
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    details.sheet_view.showGridLines = False
    for column in ("Q", "S", "T"):
        for cell in details[column][1:]:
            cell.number_format = "0.00%"
    for column in (5, 6, 7, 8, 9, 10, 11, 13, 14):
        for cell in details.iter_cols(min_col=column, max_col=column, min_row=2):
            for item in cell:
                item.number_format = "0.000"
    for row in range(2, details.max_row + 1):
        style_result_cell(details[f"O{row}"], details[f"O{row}"].value)
    for column, width in {
        "A": 24,
        "B": 15,
        "C": 10,
        "D": 10,
        "H": 18,
        "I": 18,
        "J": 14,
        "K": 14,
        "L": 16,
        "M": 14,
        "N": 14,
        "O": 12,
        "P": 18,
        "Q": 14,
        "R": 18,
        "S": 14,
        "T": 14,
        "U": 20,
        "V": 16,
    }.items():
        details.column_dimensions[column].width = width

    rule_sheet = workbook.create_sheet("规则说明")
    rule_sheet.append(["规则版本", rules.get("rule_version", "")])
    rule_sheet.append(["规则状态", rules.get("status", "")])
    rule_sheet.append(
        [
            "密度正式规则",
            "每个有效试件 1.1～1.3 g/cm³（含边界），且最大值-最小值≤0.1 g/cm³（含边界）。",
        ]
    )
    rule_sheet.append(["绿色优等", "平均膨胀率≤6.5%，并须结合表面、侧面和端头开裂观察人工确认。"])
    rule_sheet.append(["地板/家具", "6.5%＜平均膨胀率≤8% 为蓝色合格；＞8% 为红色不合格。"])
    rule_sheet.append(["墙板/栏杆", "8%≤平均膨胀率≤12% 为蓝色合格；＞12% 为红色不合格。"])
    rule_sheet.append(["v0.1 人工边界", "墙板/栏杆平均膨胀率＞6.5%且＜8% 不自动标色，转人工复核。"])
    rule_sheet.append(["用途缺失", "不得猜测产品用途；生成一次可批量应用的确认项。"])
    rule_sheet.column_dimensions["A"].width = 20
    rule_sheet.column_dimensions["B"].width = 80
    rule_sheet.sheet_view.showGridLines = False
    style_header(rule_sheet[1])

    if records:
        chart_sheet = workbook.create_sheet("分析图表")
        chart_sheet.append(["批号/文件", "平均密度(g/cm³)"])
        for record in records:
            chart_sheet.append(
                [
                    record["batch_no"] or Path(record["source_file"]).stem,
                    record["average_density_g_cm3"],
                ]
            )
        specification_expansions = {}
        for record in records:
            specification = as_text(record.get("specification")) or "未填写"
            values = [
                sample.get("expansion_ratio")
                for sample in record.get("samples", [])
                if sample.get("expansion_ratio") is not None
            ]
            if values:
                specification_expansions.setdefault(specification, []).extend(values)
        if not specification_expansions:
            specification_expansions["无可计算数据"] = []
        chart_sheet.cell(1, 4, "规格型号")
        chart_sheet.cell(1, 5, "平均膨胀率")
        for row_index, (specification, values) in enumerate(
            specification_expansions.items(), start=2
        ):
            chart_sheet.cell(row_index, 4, specification)
            chart_sheet.cell(row_index, 5, mean(values))
        style_header([chart_sheet.cell(1, 1), chart_sheet.cell(1, 2)])
        style_header([chart_sheet.cell(1, 4), chart_sheet.cell(1, 5)])
        chart_sheet.column_dimensions["A"].width = 24
        chart_sheet.column_dimensions["B"].width = 18
        chart_sheet.column_dimensions["D"].width = 24
        chart_sheet.column_dimensions["E"].width = 18
        for cell in chart_sheet["E"][1 : len(specification_expansions) + 1]:
            cell.number_format = "0.00%"
        batch_categories = Reference(
            chart_sheet, min_col=1, min_row=2, max_row=len(records) + 1
        )
        density_chart = BarChart()
        density_chart.type = "col"
        density_chart.title = "各批次平均密度对比"
        density_chart.y_axis.title = "g/cm³"
        density_chart.y_axis.numFmt = "0.0"
        density_chart.legend = None
        density_chart.add_data(
            Reference(chart_sheet, min_col=2, min_row=1, max_row=len(records) + 1),
            titles_from_data=True,
        )
        density_chart.set_categories(batch_categories)
        density_chart.height = 8
        density_chart.width = 15
        chart_sheet.add_chart(density_chart, "G2")
        expansion_chart = BarChart()
        expansion_chart.type = "col"
        expansion_chart.title = "各型号规格平均膨胀率对比"
        expansion_chart.y_axis.title = "比例"
        expansion_chart.y_axis.numFmt = "0.0%"
        expansion_chart.legend = None
        expansion_chart.add_data(
            Reference(
                chart_sheet,
                min_col=5,
                min_row=1,
                max_row=len(specification_expansions) + 1,
            ),
            titles_from_data=True,
        )
        expansion_chart.set_categories(
            Reference(
                chart_sheet,
                min_col=4,
                min_row=2,
                max_row=len(specification_expansions) + 1,
            )
        )
        expansion_chart.height = 8
        expansion_chart.width = 15
        chart_sheet.add_chart(expansion_chart, "G18")

        result_counts = {"合格": 0, "不合格": 0}
        for record in records:
            source_conclusion = str(record.get("source_conclusion") or "")
            if "不合格" in source_conclusion:
                result_counts["不合格"] += 1
            elif "合格" in source_conclusion:
                result_counts["合格"] += 1
        result_table_row = len(records) + 4
        chart_sheet.cell(result_table_row, 1, "原表综合结论")
        chart_sheet.cell(result_table_row, 2, "文件数")
        chart_sheet.cell(result_table_row + 1, 1, "合格")
        chart_sheet.cell(result_table_row + 1, 2, result_counts["合格"])
        chart_sheet.cell(result_table_row + 2, 1, "不合格")
        chart_sheet.cell(result_table_row + 2, 2, result_counts["不合格"])
        style_header(chart_sheet[result_table_row])

        result_chart = PieChart()
        result_chart.title = "原表综合结论：合格与不合格"
        result_chart.add_data(
            Reference(
                chart_sheet,
                min_col=2,
                min_row=result_table_row,
                max_row=result_table_row + 2,
            ),
            titles_from_data=True,
        )
        result_chart.set_categories(
            Reference(
                chart_sheet,
                min_col=1,
                min_row=result_table_row + 1,
                max_row=result_table_row + 2,
            )
        )
        result_chart.dataLabels = DataLabelList()
        result_chart.dataLabels.showPercent = True
        result_chart.dataLabels.showLeaderLines = True
        result_chart.height = 8
        result_chart.width = 15
        chart_sheet.add_chart(result_chart, "G34")
        chart_sheet.sheet_view.showGridLines = False

    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin_gray)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row > 1:
                    cell.border = border
        sheet.freeze_panes = sheet.freeze_panes or "A2"
    workbook.save(output_path)


def write_analysis_workbook(
    records,
    anomalies,
    confirmations,
    output_path,
    rules,
    excluded_records=None,
    ai_review=None,
):
    """Write a concise, human-oriented report without replacing audit outputs."""
    excluded_records = excluded_records or []
    workbook = openpyxl.Workbook()
    overview = workbook.active
    overview.title = "分析报告"
    overview.sheet_view.showGridLines = False

    title_fill = PatternFill("solid", fgColor="17365D")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    note_fill = PatternFill("solid", fgColor="FFF2CC")
    thin_gray = Side(style="thin", color="D9E2F3")
    all_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    overview.merge_cells("A1:L2")
    overview["A1"] = "水煮实验分析报告"
    overview["A1"].font = Font(name="微软雅黑", size=20, bold=True, color="FFFFFF")
    overview["A1"].fill = title_fill
    overview["A1"].alignment = Alignment(horizontal="center", vertical="center")
    overview.merge_cells("A3:L3")
    overview["A3"] = (
        "本报告用于快速阅读规则复核结果；人工复核项须由有权限人员确认，不替代最终业务签字。"
    )
    overview["A3"].fill = note_fill
    overview["A3"].font = Font(name="微软雅黑", italic=True, color="7F6000")
    overview["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    status_counts = {"正常": 0, "异常": 0, "待确认": 0}
    for record in records:
        status = record.get("business_status", "")
        if status in status_counts:
            status_counts[status] += 1
    if confirmations:
        overall_status = "待人工确认"
    elif anomalies:
        overall_status = "发现异常/观察项"
    elif excluded_records:
        overall_status = "部分文件信息不足"
    else:
        overall_status = "本批次未发现异常"

    metric_labels = [
        "整体状态",
        "纳入统计",
        "正常",
        "异常",
        "待确认",
        "信息不足",
    ]
    metric_values = [
        overall_status,
        len(records),
        status_counts["正常"],
        status_counts["异常"],
        status_counts["待确认"],
        len(excluded_records),
    ]
    for index, (label, value) in enumerate(zip(metric_labels, metric_values, strict=True)):
        start_column = index * 2 + 1
        label_cell = overview.cell(5, start_column, label)
        value_cell = overview.cell(6, start_column, value)
        overview.merge_cells(
            start_row=5,
            start_column=start_column,
            end_row=5,
            end_column=start_column + 1,
        )
        overview.merge_cells(
            start_row=6,
            start_column=start_column,
            end_row=6,
            end_column=start_column + 1,
        )
        label_cell.fill = section_fill
        label_cell.font = Font(name="微软雅黑", bold=True, color="1F4E78")
        value_cell.font = Font(name="微软雅黑", size=14, bold=True)
        label_cell.alignment = value_cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        label_cell.border = value_cell.border = all_border
        style_result_cell(value_cell, value)

    overview.merge_cells("A8:L8")
    overview["A8"] = "批次结果一览"
    overview["A8"].fill = section_fill
    overview["A8"].font = Font(name="微软雅黑", size=12, bold=True, color="1F4E78")
    overview["A8"].alignment = Alignment(vertical="center")
    overview_headers = [
        "来源文件",
        "批号",
        "产品",
        "产品用途",
        "平均密度(g/cm³)",
        "密度判定",
        "平均膨胀率",
        "自动等级",
        "需人工复核",
        "浸渍剥离",
        "业务状态",
        "阅读提示",
    ]
    overview.append(overview_headers)
    style_header(overview[9])
    for record in records:
        hints = []
        if record.get("conflicts"):
            hints.append(f"复算冲突 {len(record['conflicts'])} 项")
        if record.get("conclusion_rule_conflicts"):
            hints.append(f"结论冲突 {len(record['conclusion_rule_conflicts'])} 项")
        if record.get("expansion_human_review_required"):
            hints.append("膨胀率需人工复核")
        overview.append(
            [
                Path(record["source_file"]).name,
                record.get("batch_no", ""),
                record.get("product", ""),
                record.get("product_usage", ""),
                record.get("average_density_g_cm3"),
                record.get("rule_observations", {}).get("density", ""),
                record.get("average_expansion_ratio"),
                record.get("expansion_grade", ""),
                "是" if record.get("expansion_human_review_required") else "否",
                record.get("immersion_peel_result", ""),
                record.get("business_status", ""),
                "；".join(hints) or "无特别提示",
            ]
        )
    if not records:
        overview.append(["无可纳入正常统计的文件"] + [""] * 11)
    overview.freeze_panes = "A10"
    overview.auto_filter.ref = f"A9:L{overview.max_row}"
    for row in range(10, overview.max_row + 1):
        overview[f"E{row}"].number_format = "0.000"
        overview[f"G{row}"].number_format = "0.00%"
        for column in ("F", "H", "I", "J", "K"):
            style_result_cell(overview[f"{column}{row}"], overview[f"{column}{row}"].value)
    overview.row_dimensions[1].height = 28
    overview.row_dimensions[2].height = 28
    overview.row_dimensions[3].height = 32
    for column, width in {
        "A": 28,
        "B": 16,
        "C": 18,
        "D": 14,
        "E": 18,
        "F": 34,
        "G": 16,
        "H": 18,
        "I": 14,
        "J": 18,
        "K": 14,
        "L": 32,
    }.items():
        overview.column_dimensions[column].width = width

    attention = workbook.create_sheet("需要关注")
    attention.sheet_view.showGridLines = False
    attention_headers = [
        "类别",
        "严重程度",
        "来源文件",
        "工作表",
        "单元格/区域",
        "业务字段",
        "原因与处理提示",
        "稳定标识",
    ]
    attention.append(attention_headers)
    style_header(attention[1])
    for item in anomalies:
        source = item.get("source", {})
        attention.append(
            [
                "异常/观察",
                item.get("severity", ""),
                Path(source.get("source_file", "")).name,
                source.get("sheet", ""),
                source.get("cell_or_range", ""),
                item.get("field", ""),
                item.get("message", ""),
                item.get("anomaly_id", ""),
            ]
        )
    for item in confirmations:
        source = item.get("source", {})
        attention.append(
            [
                "人工确认",
                "待确认",
                Path(source.get("source_file", "")).name,
                source.get("sheet", ""),
                source.get("cell_or_range", ""),
                item.get("field", ""),
                item.get("reason", ""),
                item.get("confirmation_id", ""),
            ]
        )
    if not anomalies and not confirmations:
        attention.append(["无", "", "", "", "", "", "本批次没有异常、观察或确认项。", ""])
    attention.freeze_panes = "A2"
    attention.auto_filter.ref = attention.dimensions
    for row in range(2, attention.max_row + 1):
        style_result_cell(attention[f"B{row}"], attention[f"B{row}"].value)
    for column, width in {
        "A": 14,
        "B": 14,
        "C": 28,
        "D": 18,
        "E": 18,
        "F": 18,
        "G": 72,
        "H": 34,
    }.items():
        attention.column_dimensions[column].width = width

    excluded = workbook.create_sheet("未纳入统计")
    excluded.sheet_view.showGridLines = False
    excluded.append(["来源文件", "批号", "产品", "缺失字段", "阻断原因", "处理建议"])
    style_header(excluded[1])
    for record in excluded_records:
        excluded.append(
            [
                Path(record["source_file"]).name,
                record.get("batch_no", ""),
                record.get("product", ""),
                "、".join(record.get("missing_fields", [])),
                "；".join(record.get("blocking_reasons", [])),
                "补全源表后重新运行；该文件当前未参与批次统计。",
            ]
        )
    if not excluded_records:
        excluded.append(["无", "", "", "", "本批次没有因信息不足而排除的文件。", ""])
    excluded.freeze_panes = "A2"
    excluded.auto_filter.ref = excluded.dimensions
    for column, width in {"A": 28, "B": 16, "C": 20, "D": 34, "E": 60, "F": 48}.items():
        excluded.column_dimensions[column].width = width

    notes = workbook.create_sheet("阅读说明")
    notes.sheet_view.showGridLines = False
    notes.append(["项目", "说明"])
    notes.append(["规则版本", rules.get("rule_version", "")])
    notes.append(["正常", "可计算且未发现异常或待确认项。"])
    notes.append(["异常", "发现测量、剥离或规则观察异常，需要业务人员复核。"])
    notes.append(["待确认", "产品用途、数据冲突或规则边界需要人工确认。"])
    notes.append(["信息不足", "空表、未完成表或无可计算试件，未纳入正常统计。"])
    notes.append(["自动等级", "按当前模块规则自动计算，仅用于辅助判断。"])
    notes.append(["责任边界", "框架兼容通过不代表业务结论正确；最终结论由有权限人员确认。"])
    style_header(notes[1])
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 90

    if ai_review:
        ai_sheet = workbook.create_sheet("AI辅助建议")
        ai_sheet.sheet_view.showGridLines = False
        ai_sheet.merge_cells("A1:H2")
        ai_sheet["A1"] = "AI 辅助建议（不作为业务结论）"
        ai_sheet["A1"].fill = title_fill
        ai_sheet["A1"].font = Font(name="微软雅黑", size=16, bold=True, color="FFFFFF")
        ai_sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ai_sheet.merge_cells("A3:H3")
        ai_sheet["A3"] = (
            "本页只整理已脱敏规则候选的辅助复核建议；原因均为待验证假设，"
            "最终判断以原始记录、确定性规则和有权限人员复核为准。"
        )
        ai_sheet["A3"].fill = note_fill
        ai_sheet["A3"].font = Font(name="微软雅黑", italic=True, color="7F6000")
        ai_sheet["A3"].alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        ai_sheet["A4"] = "AI 状态"
        ai_sheet["B4"] = ai_review.get("status", "")
        ai_sheet["C4"] = "建议项数量"
        ai_sheet["D4"] = len(ai_review.get("items", []))
        ai_sheet["E4"] = "输出性质"
        ai_sheet.merge_cells("F4:H4")
        ai_sheet["F4"] = "辅助复核，不作为业务结论"
        ai_sheet["A6"] = "总体说明"
        ai_sheet.merge_cells("B6:H8")
        ai_sheet["B6"] = ai_review.get("overall_summary") or ai_review.get("text", "")
        style_header([ai_sheet["A4"], ai_sheet["C4"], ai_sheet["E4"], ai_sheet["A6"]])
        ai_sheet["B4"].fill = note_fill
        ai_sheet["D4"].fill = note_fill
        ai_sheet["F4"].fill = note_fill
        ai_sheet["B6"].alignment = Alignment(vertical="top", wrap_text=True)
        ai_sheet.row_dimensions[3].height = 32
        for row in range(6, 9):
            ai_sheet.row_dimensions[row].height = 26

        structured_items = ai_review.get("items", [])
        table_headers = [
            "序号",
            "优先级",
            "候选标识",
            "风险说明",
            "待验证原因",
            "核对清单",
            "复测建议",
            "证据引用",
        ]
        for column, value in enumerate(table_headers, 1):
            ai_sheet.cell(10, column, value)
        style_header(ai_sheet[10])
        if structured_items:
            priority_fills = {
                "高": PatternFill("solid", fgColor="F4CCCC"),
                "中": PatternFill("solid", fgColor="FFF2CC"),
                "低": PatternFill("solid", fgColor="D9EAD3"),
            }
            for index, item in enumerate(structured_items, 1):
                ai_sheet.append(
                    [
                        index,
                        item["priority"],
                        item["candidate_id"],
                        item["risk_summary"],
                        "\n".join(f"• {value}" for value in item["possible_causes"]),
                        "\n".join(f"□ {value}" for value in item["check_items"]),
                        (
                            "\n".join(
                                f"□ {value}" for value in item["retest_recommendations"]
                            )
                            or "无新增复测建议"
                        ),
                        "\n".join(item["evidence_refs"]),
                    ]
                )
                current_row = ai_sheet.max_row
                ai_sheet[f"B{current_row}"].fill = priority_fills[item["priority"]]
                ai_sheet[f"B{current_row}"].font = Font(
                    name="微软雅黑", bold=True, color="7F6000"
                )
                max_lines = max(
                    1,
                    len(item["possible_causes"]),
                    len(item["check_items"]),
                    len(item["retest_recommendations"]),
                    len(item["evidence_refs"]),
                )
                ai_sheet.row_dimensions[current_row].height = min(150, 36 + max_lines * 18)
            ai_sheet.auto_filter.ref = f"A10:H{ai_sheet.max_row}"
        else:
            ai_sheet.merge_cells("A11:H13")
            ai_sheet["A11"] = (
                "当前没有可展示的结构化 AI 建议。"
                "如状态不是 ok，请按总体说明继续使用完整离线规则报告。"
            )
            ai_sheet["A11"].fill = section_fill
            ai_sheet["A11"].alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )

        limitations = ai_review.get("limitations", [])
        if limitations:
            limitation_row = ai_sheet.max_row + 2
            ai_sheet.merge_cells(
                start_row=limitation_row,
                start_column=1,
                end_row=limitation_row,
                end_column=8,
            )
            ai_sheet.cell(limitation_row, 1, "使用限制")
            ai_sheet.cell(limitation_row, 1).fill = section_fill
            ai_sheet.cell(limitation_row, 1).font = Font(name="微软雅黑", bold=True)
            content_row = limitation_row + 1
            ai_sheet.merge_cells(
                start_row=content_row,
                start_column=1,
                end_row=content_row + 1,
                end_column=8,
            )
            ai_sheet.cell(
                content_row,
                1,
                "\n".join(f"• {value}" for value in limitations),
            )
            ai_sheet.cell(content_row, 1).alignment = Alignment(
                vertical="top", wrap_text=True
            )
            ai_sheet.row_dimensions[content_row].height = max(36, len(limitations) * 20)

        for column, width in {
            "A": 9,
            "B": 11,
            "C": 25,
            "D": 32,
            "E": 34,
            "F": 38,
            "G": 34,
            "H": 28,
        }.items():
            ai_sheet.column_dimensions[column].width = width
        ai_sheet.freeze_panes = "A11"

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                font = copy(cell.font)
                font.name = "微软雅黑"
                cell.font = font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if sheet.title != "分析报告" or cell.row >= 9:
                    cell.border = all_border
        sheet.freeze_panes = sheet.freeze_panes or "A2"
    workbook.save(output_path)


def write_analysis_report(
    records,
    anomalies,
    confirmations,
    output_path,
    rules,
    excluded_records=None,
    ai_review=None,
    run_summary=None,
    file_results=None,
    errors=None,
):
    excluded_records = excluded_records or []
    run_summary = run_summary or {}
    file_results = file_results or []
    errors = errors or []
    all_records = [*records, *excluded_records]

    def markdown_text(value):
        return as_text(value).replace("|", "｜").replace("\n", " ").strip()

    def risk_category(item):
        field = as_text(item.get("field"))
        rule_id = as_text(item.get("rule_id"))
        combined = f"{field} {rule_id} {item.get('message', '')}"
        if "密度" in combined or "DENSITY" in combined:
            return "密度"
        if "膨胀" in combined or "EXPANSION" in combined:
            return "膨胀率"
        if "剥离" in combined or "PEEL" in combined:
            return "浸渍剥离"
        if "product_usage" in combined or "用途" in combined:
            return "产品用途"
        if "结论" in combined or "CONCLUSION" in combined:
            return "原表结论与规则"
        if "INCOMPLETE" in combined or "completion" in combined or "完整" in combined:
            return "信息完整性"
        return "数据一致性"

    business_counts = {"正常": 0, "异常": 0, "信息不足": 0, "待确认": 0}
    for record in all_records:
        status = record.get("business_status", "信息不足")
        business_counts[status] = business_counts.get(status, 0) + 1

    density_values = [
        sample.get("density_g_cm3")
        for record in records
        for sample in record.get("samples", [])
        if sample.get("density_g_cm3") is not None
    ]
    expansion_values = [
        sample.get("expansion_ratio")
        for record in records
        for sample in record.get("samples", [])
        if sample.get("expansion_ratio") is not None
    ]
    density_ranges = [
        record.get("density_range_g_cm3")
        for record in records
        if record.get("density_range_g_cm3") is not None
    ]
    peel_attention = sum(
        1
        for record in records
        if as_text(record.get("immersion_peel_result"))
        and "合格" not in as_text(record.get("immersion_peel_result"))
    )
    density_value_text = (
        f"{min(density_values):.3f}～{max(density_values):.3f}"
        if density_values
        else "无可计算数据"
    )
    density_range_text = (
        f"{max(density_ranges):.3f}" if density_ranges else "无可计算数据"
    )
    expansion_value_text = (
        f"{min(expansion_values):.2%}～{max(expansion_values):.2%}"
        if expansion_values
        else "无可计算数据"
    )

    input_count = run_summary.get("input_count", len(file_results) or len(all_records))
    processed_count = run_summary.get("processed_count", len(all_records))
    failed_count = run_summary.get("failed_count", 0)
    skipped_count = run_summary.get("skipped_count", 0)
    lines = [
        "# 水煮实验分析报告",
        "",
        "## 一、运行概况",
        "",
        "| 输入文件 | 已处理 | 处理失败 | 重复跳过 | 未纳入正常统计 | 待确认 |",
        "|---:|---:|---:|---:|---:|---:|",
        f"| {input_count} | {processed_count} | {failed_count} | {skipped_count} | "
        f"{len(excluded_records)} | {len(confirmations)} |",
        "",
        f"> 规则版本：`{rules.get('rule_version', '')}`。业务异常表示检测结果需要关注，"
        "不等同于文件处理失败。",
        "",
        "## 二、总体结论",
        "",
        "| 正常 | 异常 | 信息不足 | 待确认 | 规则/计算冲突 |",
        "|---:|---:|---:|---:|---:|",
        f"| {business_counts.get('正常', 0)} | {business_counts.get('异常', 0)} | "
        f"{business_counts.get('信息不足', 0)} | {business_counts.get('待确认', 0)} | "
        f"{sum(1 for item in anomalies if 'CONFLICT' in as_text(item.get('rule_id')))} |",
        "",
        "## 三、关键指标汇总",
        "",
        "| 指标 | 本批结果 | 单位/说明 |",
        "|---|---:|---|",
        f"| 试件密度范围 | {density_value_text} | g/cm³ |",
        f"| 文件内最大密度差 | {density_range_text} | g/cm³ |",
        f"| 试件膨胀率范围 | {expansion_value_text} | 百分比 |",
        f"| 浸渍剥离需关注文件 | {peel_attention} | 文件数 |",
    ]

    issue_files = {}
    severity_order = {"critical": 3, "warning": 2, "info": 1}
    for item in anomalies:
        source = item.get("source", {})
        filename = Path(source.get("source_file", "")).name or "未知文件"
        bucket = issue_files.setdefault(
            filename,
            {"severity": "info", "categories": set(), "reasons": [], "sources": set()},
        )
        if severity_order.get(item.get("severity"), 0) > severity_order.get(
            bucket["severity"], 0
        ):
            bucket["severity"] = item.get("severity", "info")
        bucket["categories"].add(risk_category(item))
        message = markdown_text(item.get("message"))
        if message and message not in bucket["reasons"]:
            bucket["reasons"].append(message)
        location = " / ".join(
            value
            for value in (
                markdown_text(source.get("sheet")),
                markdown_text(source.get("cell_or_range")),
            )
            if value
        )
        if location:
            bucket["sources"].add(location)
    for record in excluded_records:
        filename = Path(record.get("source_file", "")).name or "未知文件"
        bucket = issue_files.setdefault(
            filename,
            {"severity": "critical", "categories": set(), "reasons": [], "sources": set()},
        )
        bucket["categories"].add("信息完整性")
        reason = markdown_text("；".join(record.get("blocking_reasons", [])))
        if reason and reason not in bucket["reasons"]:
            bucket["reasons"].append(reason)
    for item in confirmations:
        source = item.get("source", {})
        filename = Path(source.get("source_file", "")).name or "未知文件"
        bucket = issue_files.setdefault(
            filename,
            {"severity": "warning", "categories": set(), "reasons": [], "sources": set()},
        )
        bucket["categories"].add("待人工确认")
        reason = markdown_text(item.get("reason"))
        if reason and reason not in bucket["reasons"]:
            bucket["reasons"].append(reason)
    for item in file_results:
        if item.get("status") != "failed":
            continue
        filename = Path(item.get("source_file", "")).name or "未知文件"
        bucket = issue_files.setdefault(
            filename,
            {"severity": "critical", "categories": set(), "reasons": [], "sources": set()},
        )
        bucket["categories"].add("文件处理失败")
        reason = markdown_text(item.get("error"))
        if reason and reason not in bucket["reasons"]:
            bucket["reasons"].append(reason)

    lines.extend(
        [
            "",
            "## 四、异常与待复核明细",
            "",
            "| 文件 | 最高级别 | 问题类别 | 主要原因 | 来源 |",
            "|---|---|---|---|---|",
        ]
    )
    ordered_issues = sorted(
        issue_files.items(),
        key=lambda entry: (-severity_order.get(entry[1]["severity"], 0), entry[0]),
    )
    if ordered_issues:
        for filename, item in ordered_issues[:25]:
            reasons = "；".join(item["reasons"][:2])
            if len(item["reasons"]) > 2:
                reasons += f"；另有{len(item['reasons']) - 2}项"
            sources = "；".join(sorted(item["sources"])[:2]) or "详见结构化记录"
            lines.append(
                f"| {markdown_text(filename)} | {item['severity']} | "
                f"{markdown_text('、'.join(sorted(item['categories'])))} | "
                f"{markdown_text(reasons)} | {markdown_text(sources)} |"
            )
        if len(ordered_issues) > 25:
            lines.append(
                f"\n> 其余 {len(ordered_issues) - 25} 个需关注文件请在 `boiling_analysis.xlsx` "
                "的“需要关注”工作表或 `boiling_records.json` 中查看。"
            )
    else:
        lines.append("| 无 | - | 本批次未发现异常或待复核事项 | - | - |")

    risk_groups = {}
    for item in anomalies:
        category = risk_category(item)
        group = risk_groups.setdefault(category, {"count": 0, "files": set(), "examples": []})
        group["count"] += 1
        filename = Path(item.get("source", {}).get("source_file", "")).name
        if filename:
            group["files"].add(filename)
        message = markdown_text(item.get("message"))
        if message and message not in group["examples"]:
            group["examples"].append(message)
    lines.extend(
        [
            "",
            "## 五、数据风险（按类别汇总）",
            "",
            "| 类别 | 项目数 | 涉及文件 | 代表性说明 |",
            "|---|---:|---:|---|",
        ]
    )
    if risk_groups:
        for category, group in sorted(risk_groups.items()):
            lines.append(
                f"| {category} | {group['count']} | {len(group['files'])} | "
                f"{markdown_text('；'.join(group['examples'][:2]))} |"
            )
    else:
        lines.append("| 无 | 0 | 0 | 本批次未发现数据风险 |")

    if ai_review:
        lines.extend(
            [
                "",
                "## 六、AI 辅助分析（不作为业务结论）",
                "",
                f"- 状态：`{ai_review.get('status', '')}`",
                "- AI 不能改变公式、阈值、自动等级、业务状态或人工确认结论。",
                "",
                ai_review.get("text", ""),
            ]
        )
    next_steps = []
    if business_counts.get("信息不足", 0):
        next_steps.append("补齐空表、未完成表或缺失的核心测量数据后重新运行。")
    if confirmations:
        next_steps.append("完成产品用途或冲突项的人工确认，再执行 finalize。")
    if "密度" in risk_groups or "膨胀率" in risk_groups:
        next_steps.append("优先核对异常试件原始读数、单位和来源单元格，必要时安排复测。")
    if "浸渍剥离" in risk_groups:
        next_steps.append("复核浸渍剥离原表结论和试件状态，由业务人员签署最终结论。")
    if errors or failed_count:
        next_steps.append("修复损坏或结构无法识别的文件后单独重跑，避免遗漏。")
    if not next_steps:
        next_steps.append("本批次未发现需立即处理事项；按既定抽检与归档流程执行。")
    next_section = "七" if ai_review else "六"
    appendix_section = "八" if ai_review else "七"
    lines.extend(["", f"## {next_section}、建议的下一步", ""])
    lines.extend(f"{index}. {step}" for index, step in enumerate(next_steps, 1))
    lines.extend(
        [
            "",
            f"## {appendix_section}、技术附录",
            "",
            (
                "- 完整异常、确认项、稳定标识、规则 ID 和来源单元格："
                "`boiling_analysis.xlsx` 的“需要关注”工作表。"
            ),
            "- 全部来源、试件明细、规则命中和冲突审计字段：`boiling_records.json`。",
            "- 本报告不直接倾倒原始 JSON、技术 warning 或 traceback。",
        ]
    )
    Path(output_path).write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_business_records(records, excluded_records, output_path, rules):
    payload = {
        "schema_version": "1.0",
        "module_id": MODULE_ID,
        "module_version": MODULE_VERSION,
        "rule_version": rules.get("rule_version"),
        "status_definitions": {
            "正常": "可计算且未发现异常或待确认项。",
            "异常": "发现测量、剥离或规则观察异常，需要业务复核。",
            "信息不足": "空表、未完成表或无可计算试件，未纳入正常统计。",
            "待确认": "产品用途或数据/结论冲突需要人工确认。",
        },
        "records": sanitize_business_value(records),
        "excluded_records": sanitize_business_value(excluded_records),
    }
    Path(output_path).write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def load_json(path, label):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"无法读取{label}: {exc}") from exc


def validate_job(job):
    errors = []
    if not isinstance(job, dict):
        return ["任务JSON根节点必须是对象"]
    unknown = set(job) - ALLOWED_JOB_FIELDS
    if unknown:
        errors.append("任务包含未知字段: " + ", ".join(sorted(unknown)))
    required = {
        "schema_version",
        "job_id",
        "module_id",
        "phase",
        "run_mode",
        "input_files",
        "output_dir",
    }
    missing = sorted(required - set(job))
    if missing:
        errors.append("任务缺少字段: " + ", ".join(missing))
    if job.get("schema_version") != "1.0":
        errors.append("schema_version 必须为 1.0")
    if job.get("module_id") != MODULE_ID:
        errors.append("module_id 必须为 boiling")
    if not re.fullmatch(r"[A-Za-z0-9-]{6,64}", str(job.get("job_id", ""))):
        errors.append("job_id 格式不合法")
    if job.get("phase") not in {"draft", "finalize"}:
        errors.append("phase 必须为 draft 或 finalize")
    if job.get("run_mode") not in {"single", "batch"}:
        errors.append("本模块只支持 single 或 batch")
    files = job.get("input_files")
    if (
        not isinstance(files, list)
        or not files
        or not all(isinstance(item, str) and item for item in files)
    ):
        errors.append("input_files 必须是非空字符串数组")
    if job.get("run_mode") == "single" and isinstance(files, list) and len(files) != 1:
        errors.append("single 模式必须且只能提供一个文件")
    if not isinstance(job.get("output_dir"), str) or not job.get("output_dir"):
        errors.append("output_dir 必须是非空路径")
    if job.get("phase") == "finalize":
        if not job.get("parent_job_id"):
            errors.append("finalize 阶段缺少 parent_job_id")
        if not job.get("confirmation_file"):
            errors.append("finalize 阶段缺少 confirmation_file")
    if not isinstance(job.get("ai_enabled"), bool):
        errors.append("ai_enabled 必须为布尔值")
    for key in ("module_config_file", "confirmation_file", "service_response_file"):
        path_value = job.get(key)
        if path_value is not None and not isinstance(path_value, str):
            errors.append(f"{key} 必须是字符串或 null")
    return errors


def confirmation_answers(job):
    if job["phase"] != "finalize":
        return {"by_id": {}, "batch_usage": ""}
    data = load_json(job["confirmation_file"], "人工确认文件")
    answers = data.get("answers")
    if not isinstance(answers, list):
        raise ValueError("人工确认文件缺少 answers 数组")
    by_id = {
        item.get("confirmation_id"): item.get("value")
        for item in answers
        if isinstance(item, dict) and item.get("confirmation_id")
    }
    batch_usage = ""
    for item in answers:
        if not isinstance(item, dict):
            continue
        confirmation_id = as_text(item.get("confirmation_id"))
        if confirmation_id.startswith("C-PRODUCT-USAGE-BATCH-") and item.get("value"):
            batch_usage = as_text(item.get("value"))
            break
        if (
            confirmation_id.startswith("C-PRODUCT-USAGE-")
            and item.get("apply_to_batch")
            and item.get("value")
        ):
            batch_usage = as_text(item.get("value"))
            break
    return {"by_id": by_id, "batch_usage": batch_usage}


def result_template(job, started, rule_version):
    return {
        "schema_version": "1.0",
        "job_id": job["job_id"],
        "module_id": MODULE_ID,
        "module_version": MODULE_VERSION,
        "status": "failed",
        "summary": {
            "input_count": len(job["input_files"]),
            "processed_count": 0,
            "skipped_count": 0,
            "failed_count": 0,
            "anomaly_count": 0,
            "confirmation_count": 0,
        },
        "anomalies": [],
        "confirmations": [],
        "outputs": [],
        "file_results": [],
        "errors": [],
        "service_requests": [],
        "rule_version": rule_version,
        "analysis_summary": None,
        "started_at": started.isoformat(timespec="seconds"),
        "finished_at": started.isoformat(timespec="seconds"),
    }


def product_usage_confirmation(record, scope_id):
    return {
        "confirmation_id": f"C-PRODUCT-USAGE-BATCH-{scope_id}",
        "field": "product_usage",
        "reason": (
            "本批部分或全部输入文件未提供产品用途；"
            "请确认一次并批量应用（地板/家具/墙板/栏杆/其他），"
            "模块不得自行猜测膨胀率标准。"
        ),
        "suggested_value": None,
        "options": ["地板", "家具", "墙板", "栏杆", "其他"],
        "confidence": 0.0,
        "blocking": True,
        "batch_applicable": True,
        "rememberable": True,
        "source": source_ref(record["source_file"], record["source_sheet"]),
    }


def crack_inspection_confirmation(record, scope_id):
    return {
        "confirmation_id": f"C-CRACK-INSPECTION-BATCH-{scope_id}",
        "field": "人工开裂检查",
        "reason": (
            "本批次存在需要结合开裂观察确认膨胀率等级的记录。"
            "“是”表示已经人工检查，并确认表面、侧面和端头开裂情况符合对应等级；"
            "选择后只消除开裂观察提示，不改变密度、膨胀率数值、自动等级或其他复核项。"
        ),
        "suggested_value": "否（尚未完成或不符合对应等级）",
        "options": [
            "是（已检查且符合对应等级）",
            "否（尚未完成或不符合对应等级）",
        ],
        "confidence": 1.0,
        "blocking": True,
        "batch_applicable": True,
        "rememberable": False,
        "source": source_ref(record["source_file"], record["source_sheet"]),
    }


def crack_inspection_is_confirmed(value):
    return as_text(value).startswith("是")


def unresolved_finalize_anomaly(record, fingerprint, confirmation):
    confirmation_id = confirmation["confirmation_id"]
    stable_suffix = hashlib.sha256(confirmation_id.encode("utf-8")).hexdigest()[:12]
    return {
        "anomaly_id": f"A-FINALIZE-{stable_suffix}",
        "severity": "warning",
        "field": confirmation.get("field", "manual_review"),
        "value": confirmation.get("suggested_value"),
        "rule_id": "BOILING-FINALIZE-NONBLOCKING-001",
        "message": confirmation.get("reason", "最终阶段出现需要人工复核的观察项。"),
        "suggestion": "本次结果已保留为非阻断观察项；如需调整业务结论，请在业务复核流程中处理。",
        "source": confirmation.get("source")
        or source_ref(record["source_file"], record["source_sheet"]),
    }


def ai_root_job_id(job):
    return as_text(job.get("parent_job_id") or job["job_id"])


def ai_metric_name(item):
    combined = f"{item.get('field', '')} {item.get('rule_id', '')} {item.get('message', '')}"
    if "密度" in combined or "DENSITY" in combined:
        return "密度"
    if "膨胀" in combined or "EXPANSION" in combined:
        return "膨胀率"
    if "剥离" in combined or "PEEL" in combined:
        return "浸渍剥离"
    if "用途" in combined or "product_usage" in combined:
        return "产品用途"
    if "结论" in combined or "CONCLUSION" in combined:
        return "原表结论与规则冲突"
    return "数据一致性"


def ai_safe_value(value):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    return sanitize_ai_text(value)[:120]


def ai_deviation_direction(item):
    message = as_text(item.get("message"))
    if any(word in message for word in ("超过", "偏高", "高于", "上限")):
        return "above"
    if any(word in message for word in ("低于", "偏低", "不足", "下限")):
        return "below"
    if "冲突" in message:
        return "conflict"
    return "review"


def ai_record_summary(record, anomalies):
    source_path = Path(record["source_file"])
    fingerprint = sha256_file(source_path)
    related_anomalies = []
    for item in anomalies:
        item_source = Path(item.get("source", {}).get("source_file", ""))
        if item_source.name != source_path.name:
            continue
        evidence_refs = []
        if item.get("anomaly_id"):
            evidence_refs.append(f"anomaly:{item['anomaly_id']}")
        if item.get("rule_id"):
            evidence_refs.append(f"rule:{item['rule_id']}")
        related_anomalies.append(
            {
                "severity": as_text(item.get("severity")),
                "metric_name": ai_metric_name(item),
                "observed_value": ai_safe_value(item.get("value")),
                "rule_id": as_text(item.get("rule_id")),
                "deviation_direction": ai_deviation_direction(item),
                "evidence_refs": evidence_refs,
            }
        )
    has_candidate = bool(related_anomalies) or bool(record.get("conflicts")) or bool(
        record.get("conclusion_rule_conflicts")
    )
    if not has_candidate:
        return None
    evidence_refs = sorted(
        {
            reference
            for item in related_anomalies
            for reference in item.get("evidence_refs", [])
        }
    )
    if not evidence_refs:
        evidence_refs = [f"business_status:{as_text(record.get('business_status')) or '待复核'}"]
    return {
        "candidate_id": f"candidate-{fingerprint[:16]}",
        "product_usage": as_text(record.get("product_usage")) or "未确认",
        "sample_count": len(record.get("samples", [])),
        "average_density_g_cm3": record.get("average_density_g_cm3"),
        "density_min_g_cm3": record.get("density_min_g_cm3"),
        "density_max_g_cm3": record.get("density_max_g_cm3"),
        "density_range_g_cm3": record.get("density_range_g_cm3"),
        "density_rule": as_text(record.get("rule_observations", {}).get("density")),
        "average_expansion_ratio": record.get("average_expansion_ratio"),
        "expansion_rule": as_text(record.get("rule_observations", {}).get("expansion")),
        "automatic_grade": as_text(record.get("expansion_grade")),
        "manual_review_required": bool(record.get("expansion_human_review_required")),
        "immersion_peel_result": as_text(record.get("immersion_peel_result")),
        "business_status": as_text(record.get("business_status")),
        "calculation_conflict_count": len(record.get("conflicts", [])),
        "conclusion_conflict_count": len(record.get("conclusion_rule_conflicts", [])),
        "observations": related_anomalies[:20],
        "evidence_refs": evidence_refs,
    }


def ai_payload_size(payload):
    return len(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode(
            "utf-8"
        )
    )


def build_ai_service_request(job, records, anomalies, rules):
    if not records:
        return None
    root_job_id = ai_root_job_id(job)
    scope_id = stable_input_scope(job["input_files"])
    identity = hashlib.sha256(
        f"{root_job_id}:boiling:ai:{scope_id}:{rules.get('rule_version', '')}".encode()
    ).hexdigest()[:24]
    candidates = [
        candidate
        for record in records
        if (candidate := ai_record_summary(record, anomalies)) is not None
    ]
    if not candidates:
        return None
    summaries = []
    for candidate in candidates[:MAX_AI_RECORDS]:
        proposed = [*summaries, candidate]
        payload = {
            "schema_version": "1.0",
            "module_id": MODULE_ID,
            "rule_version": rules.get("rule_version", ""),
            "instruction": (
                "仅基于脱敏候选生成中文辅助说明。逐项给出风险说明、待验证原因、"
                "可执行核对清单、复测建议和优先级；possible_causes 中的内容必须明确"
                "作为待验证假设，不得写成事实。不得修改或否定模块规则结论，不得"
                "推断缺失数据，不得给出最终质量签字结论；evidence_refs 只能引用"
                "对应候选已提供的证据标识。"
            ),
            "candidates": proposed,
            "candidate_total": len(candidates),
            "candidate_sent": len(proposed),
            "candidate_omitted": len(candidates) - len(proposed),
            "response_schema": ai_response_schema(),
        }
        if ai_payload_size(payload) > MAX_AI_PAYLOAD_BYTES:
            break
        summaries = proposed
    payload = {
        "schema_version": "1.0",
        "module_id": MODULE_ID,
        "rule_version": rules.get("rule_version", ""),
        "instruction": (
            "仅基于脱敏候选生成中文辅助说明。逐项给出风险说明、待验证原因、"
            "可执行核对清单、复测建议和优先级；possible_causes 中的内容必须明确"
            "作为待验证假设，不得写成事实。不得修改或否定模块规则结论，不得"
            "推断缺失数据，不得给出最终质量签字结论；evidence_refs 只能引用"
            "对应候选已提供的证据标识。"
        ),
        "candidates": summaries,
        "candidate_total": len(candidates),
        "candidate_sent": len(summaries),
        "candidate_omitted": len(candidates) - len(summaries),
        "response_schema": ai_response_schema(),
    }
    return {
        "request_id": f"boiling-ai-{identity}",
        "service": "ai",
        "scope": None,
        "purpose": "生成水煮实验脱敏辅助分析建议",
        "payload": payload,
        "requires_user_consent": True,
        "response_format": "json_object",
    }


def ai_response_schema():
    return {
        "type": "object",
        "additionalProperties": False,
        "required": ["schema_version", "overall_summary", "items", "limitations"],
        "properties": {
            "schema_version": {"const": "1.0"},
            "overall_summary": {"type": "string", "maxLength": 1200},
            "items": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "required": [
                        "candidate_id",
                        "risk_summary",
                        "possible_causes",
                        "check_items",
                        "retest_recommendations",
                        "priority",
                        "evidence_refs",
                    ],
                    "properties": {
                        "candidate_id": {"type": "string"},
                        "risk_summary": {"type": "string", "maxLength": 800},
                        "possible_causes": {
                            "type": "array",
                            "items": {"type": "string", "maxLength": 300},
                            "minItems": 1,
                            "maxItems": 5,
                        },
                        "check_items": {
                            "type": "array",
                            "items": {"type": "string", "maxLength": 300},
                            "minItems": 1,
                            "maxItems": 8,
                        },
                        "retest_recommendations": {
                            "type": "array",
                            "items": {"type": "string", "maxLength": 300},
                            "maxItems": 5,
                        },
                        "priority": {"enum": ["high", "medium", "low"]},
                        "evidence_refs": {
                            "type": "array",
                            "items": {"type": "string"},
                            "minItems": 1,
                        },
                    },
                },
            },
            "limitations": {
                "type": "array",
                "items": {"type": "string", "maxLength": 500},
                "maxItems": 8,
            },
        },
    }


def sanitize_ai_text(value):
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", as_text(value)).strip()
    text = re.sub(r"(?i)[a-z]:[\\/][^\s，。；;]+", "[已脱敏路径]", text)
    return text[:MAX_AI_RESPONSE_CHARS]


def valid_ai_string_list(value, *, min_items, max_items, max_length):
    return (
        isinstance(value, list)
        and min_items <= len(value) <= max_items
        and all(
            isinstance(item, str) and bool(item.strip()) and len(item) <= max_length
            for item in value
        )
    )


def validate_and_render_ai_response(response_text, request):
    try:
        payload = json.loads(as_text(response_text))
    except json.JSONDecodeError:
        return None
    if not isinstance(payload, dict) or set(payload) != {
        "schema_version",
        "overall_summary",
        "items",
        "limitations",
    }:
        return None
    if payload.get("schema_version") != "1.0":
        return None
    if (
        not isinstance(payload.get("overall_summary"), str)
        or len(payload["overall_summary"]) > 1200
    ):
        return None
    if not isinstance(payload.get("items"), list) or not isinstance(
        payload.get("limitations"), list
    ):
        return None
    if not valid_ai_string_list(
        payload["limitations"], min_items=0, max_items=8, max_length=500
    ):
        return None
    candidates = request.get("payload", {}).get("candidates", [])
    allowed = {
        item.get("candidate_id"): set(item.get("evidence_refs", []))
        for item in candidates
        if isinstance(item, dict) and item.get("candidate_id")
    }
    seen = set()
    rendered_items = []
    structured_items = []
    required_keys = {
        "candidate_id",
        "risk_summary",
        "possible_causes",
        "check_items",
        "retest_recommendations",
        "priority",
        "evidence_refs",
    }
    priority_labels = {"high": "高", "medium": "中", "low": "低"}
    for item in payload["items"]:
        if not isinstance(item, dict) or set(item) != required_keys:
            return None
        candidate_id = item.get("candidate_id")
        evidence_refs = item.get("evidence_refs")
        if (
            not isinstance(candidate_id, str)
            or candidate_id not in allowed
            or candidate_id in seen
            or not isinstance(item.get("risk_summary"), str)
            or not item["risk_summary"].strip()
            or len(item["risk_summary"]) > 800
            or not valid_ai_string_list(
                item.get("possible_causes"),
                min_items=1,
                max_items=5,
                max_length=300,
            )
            or not valid_ai_string_list(
                item.get("check_items"),
                min_items=1,
                max_items=8,
                max_length=300,
            )
            or not valid_ai_string_list(
                item.get("retest_recommendations"),
                min_items=0,
                max_items=5,
                max_length=300,
            )
            or item.get("priority") not in priority_labels
            or not isinstance(evidence_refs, list)
            or not evidence_refs
            or not all(isinstance(reference, str) for reference in evidence_refs)
            or not set(evidence_refs).issubset(allowed[candidate_id])
        ):
            return None
        seen.add(candidate_id)
        possible_causes = [
            f"  - {sanitize_ai_text(value)}" for value in item["possible_causes"]
        ]
        check_items = [f"  - {sanitize_ai_text(value)}" for value in item["check_items"]]
        retest_recommendations = [
            f"  - {sanitize_ai_text(value)}" for value in item["retest_recommendations"]
        ] or ["  - 无新增复测建议"]
        rendered_items.append(
            "\n".join(
                [
                    f"### {candidate_id}",
                    f"- 优先级：{priority_labels[item['priority']]}",
                    f"- 风险说明：{sanitize_ai_text(item['risk_summary'])}",
                    "- 待验证原因（以下均为假设，需人工核实）：",
                    *possible_causes,
                    "- 核对清单：",
                    *check_items,
                    "- 复测建议：",
                    *retest_recommendations,
                    f"- 证据引用：{', '.join(evidence_refs)}",
                ]
            )
        )
        structured_items.append(
            {
                "candidate_id": candidate_id,
                "priority": priority_labels[item["priority"]],
                "risk_summary": sanitize_ai_text(item["risk_summary"]),
                "possible_causes": [
                    sanitize_ai_text(value) for value in item["possible_causes"]
                ],
                "check_items": [
                    sanitize_ai_text(value) for value in item["check_items"]
                ],
                "retest_recommendations": [
                    sanitize_ai_text(value) for value in item["retest_recommendations"]
                ],
                "evidence_refs": list(evidence_refs),
            }
        )
    if seen != set(allowed):
        return None
    lines = [sanitize_ai_text(payload["overall_summary"]), "", *rendered_items]
    limitations = [sanitize_ai_text(item) for item in payload["limitations"] if item.strip()]
    if limitations:
        lines.extend(["", "### 使用限制", *[f"- {item}" for item in limitations]])
    return {
        "text": sanitize_ai_text("\n".join(lines)),
        "overall_summary": sanitize_ai_text(payload["overall_summary"]),
        "items": structured_items,
        "limitations": limitations,
    }


def read_ai_service_response(job, request):
    if request is None:
        return None
    if job["phase"] == "draft":
        return {
            "status": "pending",
            "text": "等待用户查看脱敏预览、确认联网授权并完成 AI 服务响应。",
        }
    path_value = job.get("service_response_file")
    if not path_value:
        return {"status": "unavailable", "text": "未收到 AI 服务响应，已使用离线规则报告。"}
    try:
        payload = load_json(path_value, "AI服务响应文件")
    except ValueError:
        return {"status": "invalid", "text": "AI 服务响应无法读取，已使用离线规则报告。"}
    if (
        not isinstance(payload, dict)
        or payload.get("schema_version") != "1.0"
        or payload.get("task_id") != ai_root_job_id(job)
        or payload.get("module_id") != MODULE_ID
        or not isinstance(payload.get("responses"), list)
    ):
        return {"status": "invalid", "text": "AI 服务响应合同不匹配，已使用离线规则报告。"}
    matches = [
        item
        for item in payload["responses"]
        if isinstance(item, dict)
        and item.get("request_id") == request["request_id"]
        and item.get("service") == "ai"
    ]
    if len(matches) != 1:
        return {"status": "invalid", "text": "AI 服务响应标识不匹配，已使用离线规则报告。"}
    response = matches[0]
    status = response.get("status")
    if status == "ok":
        validated = validate_and_render_ai_response(response.get("text"), request)
        if validated:
            return {"status": "ok", **validated}
        return {
            "status": "invalid",
            "text": "AI 返回内容为空、不是严格 JSON 或证据引用不匹配，已使用离线规则报告。",
        }
    status_messages = {
        "denied": "用户未授权发送脱敏摘要，已使用离线规则报告。",
        "timeout": "AI 服务响应超时，已使用离线规则报告。",
        "error": "AI 服务调用失败，已使用离线规则报告。",
    }
    return {
        "status": status if status in status_messages else "invalid",
        "text": status_messages.get(status, "AI 服务响应状态不合法，已使用离线规则报告。"),
    }


def run_job(job_path):
    try:
        job = load_json(job_path, "任务JSON")
    except ValueError as exc:
        print(str(exc), file=sys.stderr)
        return 2
    validation_errors = validate_job(job)
    if validation_errors:
        print("; ".join(validation_errors), file=sys.stderr)
        return 2
    try:
        answers = confirmation_answers(job)
        rules = load_rules(job.get("module_config_file"))
    except (OSError, ValueError, yaml.YAMLError) as exc:
        print(str(exc), file=sys.stderr)
        return 2

    started = dt.datetime.now().astimezone()
    result = result_template(job, started, rules.get("rule_version"))
    output_dir = Path(job["output_dir"]).resolve()
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        print(f"无法创建输出目录: {exc}", file=sys.stderr)
        return 3
    result_path = output_dir / "result.json"
    records = []
    excluded_records = []
    seen = set()
    usage_prompt_added = False
    crack_prompt_added = False
    is_finalize = job["phase"] == "finalize"
    scope_id = stable_input_scope(job["input_files"])
    answer_map = answers["by_id"]
    batch_usage = answers["batch_usage"]
    crack_confirmation_id = f"C-CRACK-INSPECTION-BATCH-{scope_id}"
    crack_answer = as_text(answer_map.get(crack_confirmation_id))
    crack_answered = bool(crack_answer)
    crack_inspection_confirmed = crack_inspection_is_confirmed(crack_answer)
    emit("任务开始", 0)
    for index, raw_path in enumerate(job["input_files"], 1):
        source = Path(raw_path).resolve()
        fingerprint = "unavailable"
        status = "failed"
        extracted_records = 0
        missing_fields = []
        error = None
        anomaly_count = 0
        try:
            if source.suffix.lower() != ".xlsx" or source.name.startswith("~$"):
                raise ValueError("不支持的文件或 Excel 临时文件")
            if not source.is_file():
                raise ValueError("输入文件不存在")
            fingerprint = sha256_file(source)
            if fingerprint in seen:
                status = "partial_success"
                result["summary"]["skipped_count"] += 1
                error = "重复文件，已跳过"
            else:
                seen.add(fingerprint)
                record = extract_file(source)
                missing_fields = record["missing_fields"]
                if record["is_incomplete"]:
                    status = "success"
                    record["business_status"] = "信息不足"
                    anomalies = [
                        make_anomaly(
                            record,
                            fingerprint,
                            "INCOMPLETE",
                            "critical",
                            "workbook_completion",
                            record["blocking_reasons"],
                            "BOILING-INPUT-COMPLETENESS-001",
                            "；".join(record["blocking_reasons"]),
                            "复核空表、未完成表或模板结构变化；补齐后重新运行。",
                        )
                    ]
                    result["anomalies"].extend(anomalies)
                    anomaly_count += len(anomalies)
                    excluded_records.append(record)
                    extracted_records = 1
                    result["summary"]["processed_count"] += 1
                else:
                    finalize_anomalies = []
                    if is_missing_value(record["product_usage"]):
                        record["product_usage"] = ""
                    batch_confirmation_id = f"C-PRODUCT-USAGE-BATCH-{scope_id}"
                    # Individual product_usage confirmations not generated - batch-level is used.
                    usage_answer = (
                        batch_usage
                        or answer_map.get(batch_confirmation_id)
                    )
                    usage_confirmation_pending = False
                    if usage_answer:
                        record["product_usage"] = as_text(usage_answer)
                        record["missing_fields"] = [
                            field for field in record["missing_fields"] if field != "product_usage"
                        ]
                    elif is_missing_value(record["product_usage"]):
                        if not is_finalize and not usage_prompt_added:
                            result["confirmations"].append(
                                product_usage_confirmation(record, scope_id)
                            )
                            usage_prompt_added = True
                        if is_finalize:
                            missing_usage = product_usage_confirmation(record, scope_id)
                            finalize_anomalies.append(
                                unresolved_finalize_anomaly(record, fingerprint, missing_usage)
                            )
                        else:
                            usage_confirmation_pending = True
                    anomalies, confirmations = evaluate_record(
                        record,
                        rules,
                        fingerprint,
                        crack_inspection_confirmed=crack_inspection_confirmed,
                    )
                    crack_confirmation_relevant = bool(
                        record.get("expansion_crack_review_applicable")
                        or (
                            not as_text(record.get("product_usage"))
                            and record.get("average_expansion_ratio") is not None
                        )
                    )
                    crack_confirmation_pending = bool(
                        crack_confirmation_relevant and not crack_answered
                    )
                    if crack_confirmation_relevant:
                        if crack_answered and not crack_inspection_confirmed:
                            record["crack_inspection_status"] = (
                                "已确认尚未完成或不符合对应等级"
                            )
                        elif not is_finalize and not crack_prompt_added:
                            result["confirmations"].append(
                                crack_inspection_confirmation(record, scope_id)
                            )
                            crack_prompt_added = True
                        elif is_finalize and not crack_answered:
                            finalize_anomalies.append(
                                unresolved_finalize_anomaly(
                                    record,
                                    fingerprint,
                                    crack_inspection_confirmation(record, scope_id),
                                )
                            )
                    pending_confirmations = [
                        item for item in confirmations if item["confirmation_id"] not in answer_map
                    ]
                    if is_finalize:
                        finalize_anomalies.extend(
                            unresolved_finalize_anomaly(record, fingerprint, item)
                            for item in pending_confirmations
                        )
                        pending_confirmations = []
                    else:
                        result["confirmations"].extend(pending_confirmations)
                    anomalies.extend(finalize_anomalies)
                    result["anomalies"].extend(anomalies)
                    anomaly_count += len(anomalies)
                    if (
                        usage_confirmation_pending
                        or crack_confirmation_pending
                        or pending_confirmations
                    ):
                        record["business_status"] = "待确认"
                    elif anomalies:
                        record["business_status"] = "异常"
                    else:
                        record["business_status"] = "正常"
                    records.append(record)
                    extracted_records = 1
                    result["summary"]["processed_count"] += 1
                    status = "success"
        except Exception as exc:  # isolate a bad workbook from the remaining batch
            error = str(exc)
            result["summary"]["failed_count"] += 1
            result["errors"].append(f"{source}: {exc}")
        result["file_results"].append(
            {
                "source_file": str(source),
                "fingerprint": fingerprint,
                "status": status,
                "extracted_records": extracted_records,
                "anomaly_count": anomaly_count,
                "missing_fields": missing_fields,
                "error": error,
            }
        )
        emit(
            f"已处理 {index}/{len(job['input_files'])}: {source.name}",
            round(index / len(job["input_files"]) * 80),
        )

    ai_request = None
    ai_review = None
    ai_request = build_ai_service_request(
        job,
        records,
        result["anomalies"],
        rules,
    )
    if ai_request is not None:
        if is_finalize:
            ai_review = read_ai_service_response(job, ai_request)
        else:
            result["service_requests"].append(ai_request)
            ai_review = read_ai_service_response(job, ai_request)

    result["summary"]["anomaly_count"] = len(result["anomalies"])
    result["summary"]["confirmation_count"] = len(result["confirmations"])
    if records or excluded_records:
        summary_path = output_dir / "boiling_summary.xlsx"
        analysis_workbook_path = output_dir / "boiling_analysis.xlsx"
        report_path = output_dir / "boiling_analysis.md"
        records_path = output_dir / "boiling_records.json"
        try:
            write_summary(records, summary_path, rules, excluded_records)
            write_analysis_workbook(
                records,
                result["anomalies"],
                result["confirmations"],
                analysis_workbook_path,
                rules,
                excluded_records,
                ai_review,
            )
            write_analysis_report(
                records,
                result["anomalies"],
                result["confirmations"],
                report_path,
                rules,
                excluded_records,
                ai_review,
                result["summary"],
                result["file_results"],
                result["errors"],
            )
            write_business_records(records, excluded_records, records_path, rules)
            result["outputs"].extend(
                [
                    {
                        "type": "summary_excel",
                        "path": str(summary_path),
                        "label": "水煮实验汇总",
                        "mime_type": (
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        ),
                    },
                    {
                        "type": "analysis_report",
                        "path": str(analysis_workbook_path),
                        "label": "水煮实验分析报告",
                        "mime_type": (
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        ),
                    },
                    {
                        "type": "analysis_report_markdown",
                        "path": str(report_path),
                        "label": "水煮实验分析记录（Markdown）",
                        "mime_type": "text/markdown",
                    },
                    {
                        "type": "business_records",
                        "path": str(records_path),
                        "label": "水煮实验结构化业务记录",
                        "mime_type": "application/json",
                    },
                ]
            )
        except (OSError, ValueError) as exc:
            result["errors"].append(f"生成汇总文件失败: {exc}")

    if result["confirmations"]:
        result["status"] = "needs_confirmation"
    elif result["summary"]["failed_count"] and result["summary"]["processed_count"]:
        result["status"] = "partial_success"
    elif result["summary"]["failed_count"] or result["errors"]:
        result["status"] = "failed"
    else:
        result["status"] = "success"
    result["analysis_summary"] = (
        f"处理{result['summary']['processed_count']}个文件，失败{result['summary']['failed_count']}个，"
        f"未纳入正常统计{len(excluded_records)}个，异常/观察{result['summary']['anomaly_count']}项，"
        f"待确认{result['summary']['confirmation_count']}项。"
    )
    if ai_review:
        result["analysis_summary"] += f" AI辅助状态：{ai_review['status']}。"
    result["finished_at"] = dt.datetime.now().astimezone().isoformat(timespec="seconds")
    try:
        result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError as exc:
        print(f"无法写入 result.json: {exc}", file=sys.stderr)
        return 3
    emit(f"任务完成: {result['status']}", 100)
    return 3 if result["status"] == "failed" else 0


def self_check():
    with tempfile.TemporaryDirectory() as folder:
        sample_path = Path(folder) / "self-check.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "水煮实验"
        sheet.append(["水煮实验检测记录（自检）"])
        sheet.append(["批号", "TEST-001", "产品名称", "脱敏测试产品", "规格", "100*100*10"])
        sheet.append(["压制日期", dt.date(2026, 7, 15), "班长", "脱敏班长", "产品用途", "家具"])
        sheet.append(
            [
                "检测日期",
                dt.date(2026, 7, 16),
                "浸胶日期",
                dt.date(2026, 7, 14),
                "订单号",
                "ORDER-001",
            ]
        )
        sheet.append(
            [
                "试件",
                "长度(mm)",
                "宽度(mm)",
                "煮前厚度(mm)",
                "重量(g)",
                "密度(g/cm³)",
                "煮后厚度(mm)",
                "膨胀率",
            ]
        )
        sheet.cell(6, 1, "样1")
        sheet.cell(6, 2, 100)
        sheet.cell(6, 3, 100)
        sheet.cell(7, 5, 120)
        for row in range(6, 12):
            sheet.cell(row, 4, 10.0)
            sheet.cell(row, 7, 10.5)
        sheet.cell(12, 6, 1.2)
        sheet.cell(12, 8, 0.05)
        sheet.cell(14, 1, "检测结论：密度检测合格，2h膨胀率检测合格，浸渍剥离检测合格")
        workbook.save(sample_path)
        record = extract_file(sample_path)
        if record["product"] != "脱敏测试产品" or len(record["samples"]) != 1:
            print("self-check failed", file=sys.stderr)
            return 4
    emit("self-check passed", 100)
    return 0


def main():
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("version")
    subparsers.add_parser("self-check")
    run_parser = subparsers.add_parser("run")
    run_parser.add_argument("--job", required=True)
    args = parser.parse_args()
    if args.command == "version":
        emit(f"{MODULE_ID} {MODULE_VERSION}", 100)
        return 0
    if args.command == "self-check":
        return self_check()
    return run_job(args.job)


if __name__ == "__main__":
    raise SystemExit(main())
