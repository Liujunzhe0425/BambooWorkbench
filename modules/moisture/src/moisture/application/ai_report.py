"""Independent, consent-gated AI analysis report orchestration."""

from __future__ import annotations

import hashlib
import json
import os
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from types import SimpleNamespace

from jinja2 import Environment, FileSystemLoader, StrictUndefined
from moisture.adapters.experience.data_root import resolve_runtime_experience_root
from moisture.adapters.experience.local_approved_repository import (
    ExperienceRepositoryError,
    LocalApprovedExperienceRepository,
)
from moisture.adapters.history.repository import HistoryRepository
from moisture.adapters.rules.data_root import load_runtime_rule_snapshot
from moisture.adapters.rules.package_loader import RulePackageLoader
from moisture.application.module_config import ModuleConfig
from moisture.application.run_result import InternalMoistureRunResult, InternalOutput
from moisture.contract.job_models import JobRequest
from moisture.contract.mapper import map_run_result
from moisture.contract.private_schemas import (
    PrivateSchemaError,
    load_private_schema,
    validate_private_payload,
)
from moisture.contract.service_response_reader import read_service_responses
from moisture.infrastructure.atomic_write import AtomicResultWriter, write_json_atomic
from moisture.infrastructure.hashing import sha256_file
from moisture.reporting.workbook_validation import REQUIRED_SHEET_NAMES
from moisture.runtime_paths import module_root
from openpyxl import load_workbook

MAX_PAYLOAD_BYTES = 65_536
_LEGACY_MODEL_SHEET = "模型分析记录"
_ORANGE = "FFF4B183"
_GREEN = "FFA9D18E"
_FORBIDDEN_CONCLUSIONS = ("正式不合格", "责任认定", "处罚结论", "放行结论")
_INSTRUCTION = (
    "Return only a JSON object matching response_schema. Base every statement on the "
    "provided deterministic summaries and evidence references. Do not change moisture "
    "values, rule results or colors. Do not issue formal nonconformity, responsibility, "
    "penalty or release conclusions. The disclaimer must state that the report is only "
    "auxiliary analysis and does not constitute a formal quality conclusion."
)


class AIReportError(ValueError):
    """The independent AI report input or response is invalid."""


@dataclass(frozen=True, slots=True)
class AIReportDraft:
    request: dict[str, object]
    context_hash: str
    current_hash: str
    current_display_name: str
    current_path: Path
    history_ids: tuple[str, ...]
    experiences: tuple[dict[str, str], ...]
    payload: dict[str, object]


def _canonical_bytes(payload: object) -> bytes:
    return json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode(
        "utf-8"
    )


def _decimal(value: object) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        parsed = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None
    return parsed if parsed.is_finite() else None


def _date_text(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and value:
        return value[:10]
    return None


def _pseudonym(value: object, content_hash: str) -> str:
    raw = "" if value is None else str(value)
    digest = hashlib.sha256(f"{content_hash}:{raw}".encode()).hexdigest()[:10]
    return f"supplier-{digest}"


def _regular_xlsx(path: Path) -> Path:
    if path.suffix.lower() != ".xlsx" or path.is_symlink() or not path.is_file():
        raise AIReportError("current summary must be a regular XLSX file")
    try:
        resolved = path.resolve(strict=True)
    except OSError as error:
        raise AIReportError("current summary cannot be resolved") from error
    if resolved.is_symlink():
        raise AIReportError("current summary must not be a link")
    return resolved


def _summary_context(path: Path, *, history_id: str | None = None) -> dict[str, object]:
    source = _regular_xlsx(path)
    before = sha256_file(source)
    try:
        workbook = load_workbook(source, read_only=False, data_only=True)
    except Exception as error:
        raise AIReportError("summary workbook cannot be opened") from error
    try:
        sheet_names = tuple(workbook.sheetnames)
        supported_layouts = (
            REQUIRED_SHEET_NAMES,
            (*REQUIRED_SHEET_NAMES, _LEGACY_MODEL_SHEET),
        )
        if sheet_names not in supported_layouts:
            raise AIReportError("summary workbook is not a supported four/five-sheet format")
        sheet = workbook[REQUIRED_SHEET_NAMES[0]]
        if sheet.max_column != 19:
            raise AIReportError("summary workbook has an unsupported main-table layout")
        batches: list[dict[str, object]] = []
        dates: list[str] = []
        means: list[Decimal] = []
        review_items: list[dict[str, object]] = []
        for row_index in range(2, sheet.max_row + 1):
            if sheet.cell(row_index, 1).value is None:
                continue
            test_date = _date_text(sheet.cell(row_index, 5).value)
            if test_date:
                dates.append(test_date)
            mean = _decimal(sheet.cell(row_index, 19).value)
            if mean is not None:
                means.append(mean)
            batch_ref = (
                "batch-"
                + hashlib.sha256(f"{before}:{row_index}:{test_date or ''}".encode()).hexdigest()[
                    :16
                ]
            )
            color_rgb = sheet.cell(row_index, 19).fill.fgColor.rgb
            color = "orange" if color_rgb == _ORANGE else "green" if color_rgb == _GREEN else "none"
            batch = {
                "batch_ref": batch_ref,
                "supplier_ref": _pseudonym(sheet.cell(row_index, 2).value, before),
                "material": str(sheet.cell(row_index, 3).value or "unknown")[:100],
                "test_date": test_date,
                "minimum": str(_decimal(sheet.cell(row_index, 17).value) or ""),
                "maximum": str(_decimal(sheet.cell(row_index, 16).value) or ""),
                "range": str(_decimal(sheet.cell(row_index, 18).value) or ""),
                "mean": str(mean or ""),
                "deterministic_color": color,
            }
            batches.append(batch)
            if color != "none":
                review_items.append(
                    {
                        **batch,
                        "evidence_refs": [f"batch:{batch_ref}"],
                    }
                )
        average = sum(means, Decimal("0")) / len(means) if means else None
        metrics = {
            "mean_of_batch_means": str(average) if average is not None else None,
            "high_suspect_count": sum(item["deterministic_color"] == "orange" for item in batches),
            "low_suspect_count": sum(item["deterministic_color"] == "green" for item in batches),
        }
        context: dict[str, object] = {
            "content_hash": before,
            "batch_count": len(batches),
            "date_range": {
                "start": min(dates) if dates else None,
                "end": max(dates) if dates else None,
            },
            "metrics": metrics,
        }
        if history_id is None:
            context["review_items"] = review_items
        else:
            context["history_id"] = history_id
        return context
    finally:
        workbook.close()
        if sha256_file(source) != before:
            raise AIReportError("summary workbook changed while being read")


def _history_contexts(
    history_ids: tuple[str, ...], history_root: Path | None
) -> list[dict[str, object]]:
    if not history_ids:
        return []
    if history_root is None:
        raise AIReportError("history data directory is unavailable")
    repository = HistoryRepository(history_root)
    return [
        _summary_context(repository.resolve_active_summary(history_id), history_id=history_id)
        for history_id in sorted(history_ids)
    ]


def _experience_contexts(
    selections: tuple[object, ...], experience_root: Path
) -> tuple[list[dict[str, object]], tuple[dict[str, str], ...]]:
    if not selections:
        return [], ()
    try:
        repository = LocalApprovedExperienceRepository(experience_root)
        repository.list_approved()
    except (ExperienceRepositoryError, OSError, ValueError) as error:
        raise AIReportError("approved experience repository is unavailable") from error
    contexts: list[dict[str, object]] = []
    metadata: list[dict[str, str]] = []
    for selection in selections:
        try:
            reference = repository.get(selection.entry_id, selection.version)
        except (ExperienceRepositoryError, OSError, ValueError) as error:
            raise AIReportError("selected approved experience is unavailable") from error
        if reference is None or reference.content_hash != selection.sha256:
            raise AIReportError("selected approved experience does not match its hash")
        excerpt = reference.quoted_reference[:2000]
        contexts.append(
            {
                "entry_id": reference.experience_id,
                "version": reference.version,
                "sha256": reference.content_hash,
                "excerpt": excerpt,
            }
        )
        metadata.append(
            {
                "entry_id": reference.experience_id,
                "title": reference.title,
                "version": reference.version,
                "scope": ", ".join(reference.applicable_materials),
                "effective_from": reference.effective_at.date().isoformat(),
                "effective_to": "",
                "sha256": reference.content_hash,
            }
        )
    contexts.sort(key=lambda item: (str(item["entry_id"]), str(item["version"])))
    metadata.sort(key=lambda item: (item["entry_id"], item["version"]))
    return contexts, tuple(metadata)


def _rule_context(
    root: Path,
    *,
    rule_version: str | None = None,
    rules_sha256: str | None = None,
    use_managed_root: bool = True,
) -> dict[str, str]:
    snapshot = (
        load_runtime_rule_snapshot(
            rule_version=rule_version,
            rules_sha256=rules_sha256,
        )
        if use_managed_root
        else RulePackageLoader().load(root / "rules" / "active")
    )
    descriptions = []
    for definition in snapshot.definitions:
        threshold = definition.threshold or definition.upper or definition.lower
        target = definition.field or "combined"
        operation = definition.operator or definition.mode or ""
        descriptions.append(f"{definition.rule_id}:{target}:{operation}:{threshold or ''}")
    return {
        "version": snapshot.rule_version,
        "sha256": snapshot.rules_sha256,
        "summary": "; ".join(descriptions)[:4000] or "No active numeric rule definitions.",
    }


def _counts(total: int, sent: int) -> dict[str, int]:
    return {"total": total, "sent": sent, "omitted": total - sent}


def _payload_size(payload: object) -> int:
    return len(_canonical_bytes(payload))


def build_ai_report_draft(
    job: JobRequest,
    config: ModuleConfig,
    *,
    history_root: Path | None = None,
    experience_root: Path | None = None,
    asset_root: Path | None = None,
) -> AIReportDraft:
    """Build one deterministic, privacy-minimal AI service request."""

    root = asset_root or module_root()
    current_path = job.input_files[config.current_summary_input_index or 0]
    current = _summary_context(current_path)
    history = _history_contexts(config.history_ids, history_root)
    selected_experience_root = experience_root
    if selected_experience_root is None:
        selected_experience_root = (
            resolve_runtime_experience_root()
            if config.experience_selections
            else root / "experience"
        )
    experiences, experience_metadata = _experience_contexts(
        config.experience_selections, selected_experience_root
    )
    selection = config.rule_selection
    rule = _rule_context(
        root,
        rule_version=selection.version if selection else None,
        rules_sha256=selection.sha256 if selection else None,
        use_managed_root=asset_root is None,
    )
    base_context = {
        "question": config.question,
        "current_summary": current,
        "history_summaries": history,
        "experiences": experiences,
        "rule_snapshot": rule,
    }
    context_hash = hashlib.sha256(_canonical_bytes(base_context)).hexdigest()
    response_schema = load_private_schema("ai-analysis-response.schema.json", root)
    payload: dict[str, object] = {
        "schema_version": "1.0",
        "context_hash": context_hash,
        **base_context,
        "selection_counts": {
            "history": _counts(len(history), len(history)),
            "experiences": _counts(len(experiences), len(experiences)),
            "review_items": _counts(len(current["review_items"]), len(current["review_items"])),
        },
        "instruction": _INSTRUCTION,
        "response_schema": response_schema,
    }
    totals = {
        "history": len(history),
        "experiences": len(experiences),
        "review_items": len(current["review_items"]),
    }
    while _payload_size(payload) > MAX_PAYLOAD_BYTES:
        if payload["history_summaries"]:
            payload["history_summaries"].pop()
            key = "history"
        elif payload["experiences"]:
            payload["experiences"].pop()
            key = "experiences"
        elif payload["current_summary"]["review_items"]:
            payload["current_summary"]["review_items"].pop()
            key = "review_items"
        else:
            raise AIReportError("AI analysis payload exceeds the UTF-8 byte budget")
        sent = len(
            payload["current_summary"]["review_items"]
            if key == "review_items"
            else payload[f"{key if key == 'experiences' else 'history_summaries'}"]
        )
        payload["selection_counts"][key] = _counts(totals[key], sent)

    request_identity = hashlib.sha256(
        _canonical_bytes({"root_job_id": job.root_job_id, "context_hash": context_hash})
    ).hexdigest()[:24]
    request: dict[str, object] = {
        "request_id": f"msr-ai-report-{request_identity}",
        "service": "ai",
        "scope": None,
        "purpose": "moisture_summary_analysis",
        "payload": payload,
        "requires_user_consent": True,
        "response_format": "json_object",
    }
    validate_private_payload("ai-analysis-request.schema.json", request, root)
    return AIReportDraft(
        request=request,
        context_hash=context_hash,
        current_hash=str(current["content_hash"]),
        current_display_name=current_path.name,
        current_path=current_path,
        history_ids=config.history_ids,
        experiences=experience_metadata,
        payload=payload,
    )


def _history_root_from_environment() -> Path | None:
    value = os.environ.get("BAMBOO_MOISTURE_DATA_DIR")
    return None if not value else Path(value) / "history"


def _file_result(job: JobRequest, draft: AIReportDraft | None, error: str | None = None):
    source = draft.current_path if draft else job.input_files[0]
    return (
        {
            "source_file": str(source),
            "fingerprint": draft.current_hash if draft else "",
            "status": "failed" if error else "success",
            "extracted_records": 0 if error else 1,
            "missing_fields": [],
            "anomaly_count": 0,
            "error": error,
        },
    )


def _publish_run(
    job: JobRequest,
    started_at: datetime,
    *,
    draft: AIReportDraft | None,
    outputs: tuple[InternalOutput, ...] = (),
    requests: tuple[dict[str, object], ...] = (),
    errors: tuple[str, ...] = (),
    failed: bool = False,
) -> int:
    internal = InternalMoistureRunResult(
        job_id=job.job_id,
        root_job_id=job.root_job_id,
        phase=job.phase.value,
        input_count=len(job.input_files),
        processed_count=0 if failed else 1,
        skipped_count=0,
        failed_count=1 if failed else 0,
        outputs=outputs,
        file_results=_file_result(job, draft, errors[0] if failed and errors else None),
        service_requests=requests,
        analysis_summary=(
            "独立 AI 汇总报告请求已准备，等待用户授权。"
            if job.phase.value == "draft"
            else "独立 AI 汇总报告处理完成。"
        ),
        errors=errors,
        business_failed=failed,
        started_at=started_at,
        finished_at=datetime.now(UTC),
    )
    mapped = map_run_result(internal)
    AtomicResultWriter(job.output_dir).write(mapped.result)
    return mapped.exit_code


def run_ai_report_draft(job: JobRequest, config: ModuleConfig) -> int:
    started_at = datetime.now(UTC)
    try:
        draft = build_ai_report_draft(job, config, history_root=_history_root_from_environment())
    except (AIReportError, PrivateSchemaError, ValueError, OSError):
        return _publish_run(
            job,
            started_at,
            draft=None,
            errors=("AI_REPORT_CONTEXT_INVALID",),
            failed=True,
        )
    return _publish_run(job, started_at, draft=draft, requests=(draft.request,))


def _validate_response(text: str | None, draft: AIReportDraft) -> dict[str, object]:
    if text is None or len(text.encode("utf-8")) > MAX_PAYLOAD_BYTES:
        raise AIReportError("AI response is missing or oversized")
    try:
        payload = json.loads(text, object_pairs_hook=_reject_duplicate_keys)
    except json.JSONDecodeError as error:
        raise AIReportError("AI response is not JSON") from error
    validate_private_payload("ai-analysis-response.schema.json", payload)
    if payload["context_hash"] != draft.context_hash:
        raise AIReportError("AI response context hash does not match")
    serialized = json.dumps(payload, ensure_ascii=False)
    if any(phrase in serialized for phrase in _FORBIDDEN_CONCLUSIONS):
        raise AIReportError("AI response contains a forbidden formal conclusion")
    if "辅助" not in payload["disclaimer"] or "正式质量结论" not in payload["disclaimer"]:
        raise AIReportError("AI response disclaimer is insufficient")
    batch_refs = {item["batch_ref"] for item in draft.payload["current_summary"]["review_items"]}
    allowed_evidence = {f"batch:{item}" for item in batch_refs}
    allowed_evidence.update(f"history:{item}" for item in draft.history_ids)
    allowed_evidence.update(
        f"experience:{item['entry_id']}@{item['version']}" for item in draft.experiences
    )
    for item in payload["review_batches"]:
        if item["batch_ref"] not in batch_refs:
            raise AIReportError("AI response refers to an unknown batch")
        if not set(item["evidence_refs"]).issubset(allowed_evidence):
            raise AIReportError("AI response contains unknown evidence")
    if not set(payload["evidence_refs"]).issubset(allowed_evidence):
        raise AIReportError("AI response contains unknown evidence")
    return payload


def _reject_duplicate_keys(pairs: list[tuple[str, object]]) -> dict[str, object]:
    payload: dict[str, object] = {}
    for key, value in pairs:
        if key in payload:
            raise AIReportError(f"AI response contains duplicate key: {key}")
        payload[key] = value
    return payload


def _namespace(**values: object) -> SimpleNamespace:
    return SimpleNamespace(**values)


def _report_view(draft: AIReportDraft, response: dict[str, object]) -> SimpleNamespace:
    current = draft.payload["current_summary"]
    metrics = current["metrics"]
    review_lookup = {item["batch_ref"]: item for item in current["review_items"]}
    review_batches = []
    for decision in response["review_batches"]:
        source = review_lookup[decision["batch_ref"]]
        review_batches.append(
            _namespace(
                record_id=decision["batch_ref"],
                supplier=source["supplier_ref"],
                material=source["material"],
                mean=source["mean"],
                display_color_label=source["deterministic_color"],
                screen_status="deterministic suspect",
                ai_recommendation=decision["reason"],
            )
        )
    evidence = tuple(
        _namespace(
            ref_id=ref,
            source_type=ref.split(":", 1)[0],
            description="Selected, validated analysis evidence.",
            relevant_period=None,
            remarks=None,
        )
        for ref in response["evidence_refs"]
    )
    catalog = _namespace(
        schema_version="1.0",
        entry_count=len(draft.experiences),
        entries=tuple(_namespace(**entry) for entry in draft.experiences),
    )
    date_range = current["date_range"]
    history = draft.payload["history_summaries"]
    history_dates = [item["date_range"] for item in history]
    return _namespace(
        generated_at=datetime.now(UTC).isoformat(),
        period_start=date_range["start"] or "unknown",
        period_end=date_range["end"] or "unknown",
        current_batch_count=current["batch_count"],
        rule_triggered_count=(metrics["high_suspect_count"] + metrics["low_suspect_count"]),
        high_suspect_count=metrics["high_suspect_count"],
        low_suspect_count=metrics["low_suspect_count"],
        experience_citation_count=len(draft.experiences),
        knowledge_citation_count=0,
        summary_text=response["overall_summary"],
        history_available=bool(history),
        history_period_start=(history_dates[0]["start"] if history else None),
        history_period_end=(history_dates[-1]["end"] if history else None),
        history_batch_count=sum(item["batch_count"] for item in history),
        history_rule_triggered_count=0,
        historical_changes=(
            (
                _namespace(
                    metric="AI historical comparison",
                    current_value="current period",
                    history_value="selected history",
                    direction="analysis",
                    delta="n/a",
                    interpretation=response["historical_comparison"],
                ),
            )
            if history
            else ()
        ),
        batches_needing_review=tuple(review_batches),
        possible_causes=tuple(
            _namespace(
                heading=f"Possible cause {index}",
                related_records=(),
                likelihood="requires human verification",
                basis_points=(cause,),
                evidence_refs=tuple(response["evidence_refs"]),
            )
            for index, cause in enumerate(response["possible_causes"], start=1)
        ),
        recommended_steps=tuple(
            _namespace(
                index=index,
                title=f"Review step {index}",
                related_records=(),
                description=step,
                expected_outcome=None,
            )
            for index, step in enumerate(response["recommendations"], start=1)
        ),
        evidence_references=evidence,
        approved_catalog_snapshot=catalog if draft.experiences else None,
        limitations=(response["disclaimer"],),
    )


def _write_report(output_dir: Path, view: SimpleNamespace) -> Path:
    environment = Environment(
        loader=FileSystemLoader(module_root() / "templates"),
        autoescape=False,
        undefined=StrictUndefined,
        keep_trailing_newline=True,
    )
    environment.globals.clear()
    rendered = environment.get_template("ai_analysis_report.md.j2").render(report=view)
    destination = output_dir / "ai_analysis_report.md"
    temporary = output_dir / ".ai_analysis_report.tmp"
    output_dir.mkdir(parents=True, exist_ok=True)
    try:
        with temporary.open("w", encoding="utf-8", newline="\n") as stream:
            stream.write(rendered if rendered.endswith("\n") else rendered + "\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)
    return destination


def _analysis_result(
    draft: AIReportDraft, status: str, degradation_reason: str | None
) -> dict[str, object]:
    return {
        "schema_version": "1.0",
        "status": status,
        "request_id": draft.request["request_id"],
        "context_hash": draft.context_hash,
        "current_summary": {
            "display_name": draft.current_display_name,
            "sha256": draft.current_hash,
        },
        "history_ids": list(draft.history_ids),
        "experience_selections": [
            {
                "entry_id": item["entry_id"],
                "version": item["version"],
                "sha256": item["sha256"],
            }
            for item in draft.experiences
        ],
        "report_file": "ai_analysis_report.md" if status == "success" else None,
        "degradation_reason": degradation_reason,
    }


def run_ai_report_finalize(job: JobRequest, config: ModuleConfig) -> int:
    started_at = datetime.now(UTC)
    try:
        draft = build_ai_report_draft(job, config, history_root=_history_root_from_environment())
    except (AIReportError, PrivateSchemaError, ValueError, OSError):
        return _publish_run(
            job,
            started_at,
            draft=None,
            errors=("AI_REPORT_CONTEXT_INVALID",),
            failed=True,
        )
    expected = {str(draft.request["request_id"]): "ai"}
    response_result = read_service_responses(
        job.service_response_file,
        root_job_id=job.root_job_id,
        expected_requests=expected,
    )
    status = "invalid_response"
    degradation = "invalid_or_missing_response"
    response_payload: dict[str, object] | None = None
    if response_result.envelope is not None and not response_result.errors:
        response = response_result.envelope.responses[0]
        if response.status == "ok":
            try:
                response_payload = _validate_response(response.text, draft)
            except (AIReportError, PrivateSchemaError):
                degradation = "invalid_response"
            else:
                status = "success"
                degradation = None
        else:
            status = response.status
            degradation = f"provider_{response.status}"

    result_payload = _analysis_result(draft, status, degradation)
    validate_private_payload("ai-analysis-result.schema.json", result_payload)
    outputs: list[InternalOutput] = []
    completed: list[Path] = []
    try:
        if response_payload is not None:
            report_path = _write_report(job.output_dir, _report_view(draft, response_payload))
            completed.append(report_path)
            outputs.append(
                InternalOutput("markdown", report_path, "AI 辅助分析报告", "text/markdown")
            )
        result_path = job.output_dir / "ai_analysis_result.json"
        write_json_atomic(result_path, result_payload)
        completed.append(result_path)
        outputs.append(InternalOutput("json", result_path, "AI 分析结构化结果", "application/json"))
    except Exception:
        for artifact in completed:
            artifact.unlink(missing_ok=True)
        return _publish_run(
            job,
            started_at,
            draft=draft,
            errors=("AI_REPORT_ARTIFACT_WRITE_FAILED",),
            failed=True,
        )
    errors = () if status == "success" else (f"AI_REPORT_DEGRADED: {degradation}",)
    return _publish_run(
        job,
        started_at,
        draft=draft,
        outputs=tuple(outputs),
        errors=errors,
    )
