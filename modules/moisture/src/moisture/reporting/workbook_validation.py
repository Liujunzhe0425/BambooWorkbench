"""Minimum post-close validation for a generated moisture workbook."""

from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook

REQUIRED_SHEET_NAMES = (
    "竹丝含水率统计",
    "异常分析明细",
    "趋势分析",
    "来源追溯",
    "模型分析记录",
)


class WorkbookValidationError(ValueError):
    """The closed workbook cannot be accepted as a formal artifact."""


def validate_workbook(path: Path) -> None:
    """Reopen *path* and verify the fixed report skeleton."""

    try:
        with closing(load_workbook(path, data_only=False, read_only=False)) as workbook:
            if tuple(workbook.sheetnames) != REQUIRED_SHEET_NAMES:
                raise WorkbookValidationError("workbook does not contain the required sheets")
            for sheet_name in REQUIRED_SHEET_NAMES:
                worksheet = workbook[sheet_name]
                if worksheet.max_row < 1 or worksheet.max_column < 1:
                    raise WorkbookValidationError(f"worksheet has no header: {sheet_name}")
                if not any(cell.value for cell in worksheet[1]):
                    raise WorkbookValidationError(f"worksheet has an empty header: {sheet_name}")
    except WorkbookValidationError:
        raise
    except Exception as error:
        raise WorkbookValidationError("generated workbook cannot be reopened") from error
