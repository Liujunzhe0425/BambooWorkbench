"""Read-only Excel source parser migrated from the stable Task 1—9 implementation."""

from __future__ import annotations

import re
from collections import OrderedDict
from contextlib import ExitStack
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from moisture.adapters.excel.template_detector import (
    TemplateCandidate,
    find_template_candidates,
    is_numeric,
    scan_sample_rows,
)
from moisture.domain.formula_verification import verify_moisture
from moisture.domain.models import BatchRecord, DepartmentGroup, SampleMeasurement
from moisture.infrastructure.hashing import sha256_file
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet


class WorkbookStructureError(ValueError):
    """The workbook does not contain one usable moisture table."""

    code = "WORKBOOK_STRUCTURE_INVALID"


class AmbiguousWorkbookStructure(WorkbookStructureError):
    """Multiple worksheets are equally plausible moisture sources."""

    code = "WORKBOOK_STRUCTURE_AMBIGUOUS"


class SampleLimitExceeded(WorkbookStructureError):
    """The source contains more samples than the domain permits."""

    code = "SAMPLE_LIMIT_EXCEEDED"


class SourceFileModified(ValueError):
    """The source workbook changed while it was being parsed."""

    code = "SOURCE_FILE_MODIFIED"


_LABEL_ALIASES = {
    "supplier": {"供应商", "供货商", "供应单位", "送样单位"},
    "material_type": {"物料类型", "物料名称", "品名", "材料"},
    "arrival_date": {"来料日期", "到货日期", "进厂日期"},
    "test_date": {"检测日期", "检验日期", "测试日期"},
    "total_weight": {"总重量", "总重", "来料重量", "重量"},
}


def _as_decimal(value: object) -> Decimal:
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError) as error:
        raise WorkbookStructureError(f"non-numeric mass value: {value!r}") from error


def _raw_text(value: object | None) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _clean_text(value: object | None) -> str | None:
    raw = _raw_text(value)
    return raw.strip() if raw is not None else None


def _label_and_inline_value(value: object) -> tuple[str, str | None]:
    text = str(value)
    separator = re.search(r"[:：]", text)
    if separator is None:
        return text.strip(), None
    inline_value = text[separator.end() :]
    return text[: separator.start()].strip(), inline_value if inline_value.strip() else None


def _extract_metadata(worksheet: Worksheet, header_row: int) -> dict[str, object | None]:
    result: dict[str, object | None] = {field: None for field in _LABEL_ALIASES}
    for row in range(1, header_row):
        for column in range(1, worksheet.max_column + 1):
            cell_value = worksheet.cell(row, column).value
            if cell_value is None:
                continue
            label, inline_value = _label_and_inline_value(cell_value)
            for field, aliases in _LABEL_ALIASES.items():
                if result[field] is not None or label not in aliases:
                    continue
                if inline_value is not None:
                    result[field] = inline_value
                    continue
                for value_column in range(column + 1, min(column + 4, worksheet.max_column + 1)):
                    candidate = worksheet.cell(row, value_column).value
                    if candidate is not None:
                        result[field] = candidate
                        break
    return result


_FULL_DATE_PATTERNS = (
    re.compile(r"^(?P<year>\d{4})[-/.](?P<month>\d{1,2})[-/.](?P<day>\d{1,2})$"),
    re.compile(r"^(?P<year>\d{4})年(?P<month>\d{1,2})月(?P<day>\d{1,2})日?$"),
)
_PARTIAL_DATE_PATTERNS = (
    re.compile(r"^(?P<month>\d{1,2})[-/.](?P<day>\d{1,2})$"),
    re.compile(r"^(?P<month>\d{1,2})月(?P<day>\d{1,2})日?$"),
)


def _explicit_year(value: object | None) -> int | None:
    if isinstance(value, (date, datetime)):
        return value.year
    if value is None:
        return None
    text = str(value).strip()
    for pattern in _FULL_DATE_PATTERNS:
        match = pattern.match(text)
        if match:
            return int(match.group("year"))
    return None


def _parse_date(value: object | None, fallback_year: int | None) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for pattern in _FULL_DATE_PATTERNS:
        match = pattern.match(text)
        if match:
            return date(
                int(match.group("year")), int(match.group("month")), int(match.group("day"))
            )
    for pattern in _PARTIAL_DATE_PATTERNS:
        match = pattern.match(text)
        if match:
            if fallback_year is None:
                raise WorkbookStructureError("default_year is required for month/day dates")
            return date(fallback_year, int(match.group("month")), int(match.group("day")))
    raise WorkbookStructureError(f"unsupported date value: {text!r}")


def _parse_weight_kg(value: object | None) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, bool):
        raise WorkbookStructureError(f"unsupported weight value: {value!r}")
    text = str(value).strip()
    match = re.search(r"[-+]?\d[\d,，]*(?:\.\d+)?", text)
    if match is None:
        raise WorkbookStructureError(f"unsupported weight value: {text!r}")
    amount = Decimal(match.group(0).replace(",", "").replace("，", ""))
    suffix = text[match.end() :].lower()
    if re.search(r"(?:^|\s)(?:t|吨)(?:\s|$|[（(])", suffix):
        return amount * Decimal("1000")
    if re.search(r"(?:^|\s)(?:g|克)(?:\s|$|[（(])", suffix) and not re.search(
        r"kg|公斤|千克", suffix
    ):
        return amount / Decimal("1000")
    return amount


def _select_candidate(
    structure_worksheets: list[Worksheet], value_worksheets: list[Worksheet]
) -> TemplateCandidate:
    candidates = find_template_candidates(structure_worksheets, value_worksheets)
    if not candidates:
        raise WorkbookStructureError("no worksheet contains the required moisture headers")
    best_count = max(candidate.valid_sample_count for candidate in candidates)
    best = [candidate for candidate in candidates if candidate.valid_sample_count == best_count]
    if len(best) != 1:
        raise AmbiguousWorkbookStructure("multiple worksheet structures are equally plausible")
    if best[0].valid_sample_count == 0:
        raise WorkbookStructureError("moisture table contains no valid samples")
    return best[0]


class ExcelSourceParser:
    def parse(self, path: Path, default_year: int | None = None) -> BatchRecord:
        source_hash = sha256_file(path)
        try:
            record = self._parse_record(path, default_year, source_hash)
        except Exception as parsing_error:
            self._raise_if_source_modified(path, source_hash, cause=parsing_error)
            if isinstance(parsing_error, WorkbookStructureError):
                raise
            raise WorkbookStructureError("unable to read source workbook") from parsing_error
        self._raise_if_source_modified(path, source_hash)
        return record

    def _parse_record(
        self, path: Path, default_year: int | None, source_hash: str
    ) -> BatchRecord:
        with ExitStack() as cleanup:
            formula_book = load_workbook(path, data_only=False, read_only=False)
            cleanup.callback(formula_book.close)
            value_book = load_workbook(path, data_only=True, read_only=False)
            cleanup.callback(value_book.close)
            candidate = _select_candidate(
                list(formula_book.worksheets), list(value_book.worksheets)
            )
            formula_sheet = formula_book[candidate.sheet_name]
            value_sheet = value_book[candidate.sheet_name]
            metadata = _extract_metadata(formula_sheet, candidate.header_row)
            samples, departments = self._parse_samples(formula_sheet, value_sheet, candidate)

        explicit_years = {
            year
            for value in (metadata["arrival_date"], metadata["test_date"])
            if (year := _explicit_year(value)) is not None
        }
        inferred_year = next(iter(explicit_years)) if len(explicit_years) == 1 else None
        date_year = inferred_year if inferred_year is not None else default_year

        return BatchRecord(
            record_id=path.stem,
            samples=samples,
            source_file=path,
            source_file_hash=source_hash,
            source_sheet=candidate.sheet_name,
            supplier_raw=_raw_text(metadata["supplier"]),
            supplier=_clean_text(metadata["supplier"]),
            material_type=_clean_text(metadata["material_type"]),
            arrival_date_raw=_raw_text(metadata["arrival_date"]),
            arrival_date=_parse_date(metadata["arrival_date"], date_year),
            test_date_raw=_raw_text(metadata["test_date"]),
            test_date=_parse_date(metadata["test_date"], date_year),
            total_weight_raw=_raw_text(metadata["total_weight"]),
            total_weight_kg=_parse_weight_kg(metadata["total_weight"]),
            departments=departments,
        )

    @staticmethod
    def _raise_if_source_modified(
        path: Path, source_hash: str, cause: Exception | None = None
    ) -> None:
        try:
            final_hash = sha256_file(path)
        except OSError as hash_error:
            error = SourceFileModified("source workbook became unavailable while being parsed")
            raise error from (cause if cause is not None else hash_error)
        if final_hash != source_hash:
            error = SourceFileModified("source workbook changed while it was being parsed")
            if cause is not None:
                raise error from cause
            raise error

    def _parse_samples(
        self,
        formula_sheet: Worksheet,
        value_sheet: Worksheet,
        candidate: TemplateCandidate,
    ) -> tuple[tuple[SampleMeasurement, ...], tuple[DepartmentGroup, ...]]:
        positions = candidate.positions
        samples: list[SampleMeasurement] = []
        department_indices: OrderedDict[str, list[int]] = OrderedDict()
        current_department: str | None = None
        for scanned_row in scan_sample_rows(
            formula_sheet, value_sheet, candidate.header_row, positions
        ):
            row = scanned_row.row
            department_cell = formula_sheet.cell(row, positions["department"])
            if not isinstance(department_cell, MergedCell) and department_cell.value is not None:
                current_department = str(department_cell.value).strip()

            if not scanned_row.is_valid:
                continue
            if not current_department:
                raise WorkbookStructureError(f"sample row {row} has no department")

            initial_mass, dry_mass, foil_mass = (
                _as_decimal(value) for value in scanned_row.mass_values
            )
            moisture_value = value_sheet.cell(row, positions["moisture"]).value
            source_moisture = _as_decimal(moisture_value) if is_numeric(moisture_value) else None
            verification = verify_moisture(
                initial_mass, dry_mass, foil_mass, source_moisture
            )
            sample_index = len(samples) + 1
            samples.append(
                SampleMeasurement(
                    sample_index=sample_index,
                    department_raw=current_department,
                    department=current_department,
                    source_row=row,
                    initial_mass=initial_mass,
                    dry_mass=dry_mass,
                    foil_mass=foil_mass,
                    source_moisture_value=source_moisture,
                    verified_moisture_value=verification.verified_value,
                    value_used_for_statistics=verification.value_used,
                )
            )
            department_indices.setdefault(current_department, []).append(sample_index)

        if len(samples) > 10:
            raise SampleLimitExceeded(f"workbook contains {len(samples)} valid samples")
        if not samples:
            raise WorkbookStructureError("moisture table contains no valid samples")

        departments = tuple(
            DepartmentGroup(name=name, sample_indices=tuple(indices))
            for name, indices in department_indices.items()
        )
        return tuple(samples), departments
