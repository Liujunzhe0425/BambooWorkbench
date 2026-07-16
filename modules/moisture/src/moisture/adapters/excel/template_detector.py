"""Dynamic worksheet and header detection migrated from the stable parser."""

from collections.abc import Iterator
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from openpyxl.worksheet.worksheet import Worksheet

REQUIRED_HEADERS = {
    "department": {"取样部门", "部门"},
    "initial_mass": {"初始质量"},
    "dry_mass": {"烘干质量"},
    "foil_mass": {"锡纸", "锡纸质量"},
    "moisture": {"相对含水率", "含水率"},
}
SECTION_BOUNDARIES = {"备注", "复核", "审核", "统计", "平均值", "结论", "以下空白"}


@dataclass(frozen=True, slots=True)
class TemplateCandidate:
    sheet_name: str
    header_row: int
    positions: dict[str, int]
    valid_sample_count: int


@dataclass(frozen=True, slots=True)
class ScannedSampleRow:
    row: int
    mass_values: tuple[object, object, object]
    is_valid: bool


def score_header_row(values: list[object]) -> dict[str, int] | None:
    normalized = [str(value).strip() if value is not None else "" for value in values]
    positions: dict[str, int] = {}
    for field, aliases in REQUIRED_HEADERS.items():
        matches = [index for index, value in enumerate(normalized) if value in aliases]
        if len(matches) != 1:
            return None
        positions[field] = matches[0] + 1
    return positions


def is_numeric(value: object) -> bool:
    if value is None or isinstance(value, bool):
        return False
    try:
        Decimal(str(value))
    except (InvalidOperation, ValueError):
        return False
    return True


def is_section_boundary(worksheet: Worksheet, row: int) -> bool:
    for cell in worksheet[row]:
        if cell.value is None:
            continue
        text = str(cell.value).strip()
        if any(
            text == boundary or text.startswith(f"{boundary}：")
            for boundary in SECTION_BOUNDARIES
        ):
            return True
    return False


def scan_sample_rows(
    structure_worksheet: Worksheet,
    value_worksheet: Worksheet,
    header_row: int,
    positions: dict[str, int],
) -> Iterator[ScannedSampleRow]:
    empty_run = 0
    mass_columns = (
        positions["initial_mass"],
        positions["dry_mass"],
        positions["foil_mass"],
    )
    for row in range(header_row + 1, structure_worksheet.max_row + 1):
        if is_section_boundary(structure_worksheet, row):
            break
        mass_values = tuple(value_worksheet.cell(row, column).value for column in mass_columns)
        valid = all(is_numeric(value) for value in mass_values)
        yield ScannedSampleRow(
            row=row,
            mass_values=(mass_values[0], mass_values[1], mass_values[2]),
            is_valid=valid,
        )
        if valid:
            empty_run = 0
        else:
            empty_run += 1
            if empty_run >= 10:
                break


def count_valid_samples(
    structure_worksheet: Worksheet,
    value_worksheet: Worksheet,
    header_row: int,
    positions: dict[str, int],
) -> int:
    return sum(
        scanned_row.is_valid
        for scanned_row in scan_sample_rows(
            structure_worksheet, value_worksheet, header_row, positions
        )
    )


def find_template_candidates(
    structure_worksheets: list[Worksheet], value_worksheets: list[Worksheet]
) -> list[TemplateCandidate]:
    candidates: list[TemplateCandidate] = []
    value_sheets_by_name = {worksheet.title: worksheet for worksheet in value_worksheets}
    for structure_worksheet in structure_worksheets:
        value_worksheet = value_sheets_by_name[structure_worksheet.title]
        for row in range(1, min(30, structure_worksheet.max_row) + 1):
            positions = score_header_row(
                [
                    structure_worksheet.cell(row, column).value
                    for column in range(1, structure_worksheet.max_column + 1)
                ]
            )
            if positions is None:
                continue
            candidates.append(
                TemplateCandidate(
                    sheet_name=structure_worksheet.title,
                    header_row=row,
                    positions=positions,
                    valid_sample_count=count_valid_samples(
                        structure_worksheet, value_worksheet, row, positions
                    ),
                )
            )
    return candidates
