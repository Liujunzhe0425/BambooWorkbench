"""Standard summary workbook reader — v1.0.0 5-sheet and v1.0.1 4-sheet."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from moisture.domain.enums import (
    ConfirmationState,
    DisplayColor,
    ReviewStatus,
    ScreenStatus,
)
from moisture.domain.models import (
    AIReviewDecision,
    BatchRecord,
    BatchStatistics,
    SampleMeasurement,
    ScreeningDecision,
)
from moisture.ports.history_summary import (
    HistoricalSummarySnapshot,
    SummaryFormat,
)
from moisture.ports.writers import ExperienceCitation, WorkbookBatch
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# Sheet names
# ---------------------------------------------------------------------------
_SHEET_SUMMARY = "竹丝含水率统计"
_SHEET_ANOMALY = "异常分析明细"
_SHEET_TREND = "趋势分析"
_SHEET_TRACE = "来源追溯"
_SHEET_MODEL = "模型分析记录"  # v1.0.0 only

_V1_0_1_SHEETS = (_SHEET_SUMMARY, _SHEET_ANOMALY, _SHEET_TREND, _SHEET_TRACE)
_V1_0_0_SHEETS = (*_V1_0_1_SHEETS, _SHEET_MODEL)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _col(headers: list[str | None], name: str) -> int:
    return headers.index(name)


def _pct(value: object | None) -> Decimal | None:
    """Convert an openpyxl numeric cell (fraction or None) to Decimal or None."""
    if value is None:
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def _date(value: object | None) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _int(value: object | None) -> int | None:
    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _str(value: object | None) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _read_row(ws: Worksheet, row: int, max_col: int) -> list[object | None]:
    return [ws.cell(row, col + 1).value for col in range(max_col)]


def _sample_rows(value: object | None) -> dict[int, int]:
    rows: dict[int, int] = {}
    for sample_index, item in enumerate(_str(value).split(","), start=1):
        try:
            rows[sample_index] = int(item.strip())
        except ValueError:
            continue
    return rows


def _department_map(value: object | None) -> dict[int, str]:
    departments: dict[int, str] = {}
    for item in _str(value).split(";"):
        label, separator, department = item.strip().partition(":")
        sample_text = label.removeprefix("样本")
        if separator and sample_text.isdigit():
            departments[int(sample_text)] = department.strip()
    return departments


def _experience_references(value: object | None) -> tuple[ExperienceCitation, ...]:
    references: list[ExperienceCitation] = []
    for item in _str(value).split(";"):
        entry_id, separator, version = item.strip().rpartition(" v")
        if separator and entry_id and version:
            references.append(
                ExperienceCitation(
                    entry_id=entry_id,
                    version=version,
                    title="",
                    source="",
                )
            )
    return tuple(references)


def _rule_confirmation_status(value: object | None) -> str:
    normalized = _str(value)
    return "confirmed" if normalized in {"confirmed", "已确认"} else "pending_confirmation"


# ---------------------------------------------------------------------------
# Row-level parsers
# ---------------------------------------------------------------------------

_MAIN_HEADERS = [
    "序号",
    "供应商",
    "材料类型",
    "到货日期",
    "检测日期",
    *(f"样本{i}" for i in range(1, 11)),
    "最大值",
    "最小值",
    "区间值",
    "均值",
]

_TREND_HEADERS = [
    "序号",
    "批次编号",
    "检测日期",
    "供应商",
    "均值",
    "区间值",
    "初筛下限",
    "初筛上限",
    "标准差",
    "中位数",
    "第一四分位数",
    "第三四分位数",
    "四分位距",
    "部门均值最大差",
]

_TRACE_HEADERS = [
    "批次编号",
    "来源文件",
    "文件哈希",
    "来源工作表",
    "样本行",
    "总重",
    "样本部门映射",
    "规则版本",
    "规则SHA-256",
    "规则适用范围",
    "规则确认状态",
    "经验引用",
]

_ANOMALY_HEADERS = [
    "批次编号",
    "供应商",
    "检测日期",
    "目标字段",
    "观测值",
    "筛查结果",
    "原因",
    "复核建议",
]


def _parse_main_row(
    headers: list[str | None],
    row_data: list[object | None],
) -> dict[str, object | None]:
    sample_values: dict[int, Decimal | None] = {}
    for i in range(1, 11):
        sample_values[i] = _pct(row_data[_col(headers, f"样本{i}")])
    return {
        "supplier": _str(row_data[_col(headers, "供应商")]),
        "material_type": _str(row_data[_col(headers, "材料类型")]) or "竹丝",
        "arrival_date": _date(row_data[_col(headers, "到货日期")]),
        "test_date": _date(row_data[_col(headers, "检测日期")]),
        "sample_values": sample_values,
        "maximum": _pct(row_data[_col(headers, "最大值")]),
        "minimum": _pct(row_data[_col(headers, "最小值")]),
        "range_value": _pct(row_data[_col(headers, "区间值")]),
        "mean": _pct(row_data[_col(headers, "均值")]),
    }


def _parse_trend_row(
    headers: list[str | None],
    row_data: list[object | None],
) -> dict[str, object | None]:
    return {
        "record_id": _str(row_data[_col(headers, "批次编号")]),
        "test_date": _date(row_data[_col(headers, "检测日期")]),
        "supplier": _str(row_data[_col(headers, "供应商")]),
        "mean": _pct(row_data[_col(headers, "均值")]),
        "range_value": _pct(row_data[_col(headers, "区间值")]),
        "lower_threshold": _pct(row_data[_col(headers, "初筛下限")]),
        "upper_threshold": _pct(row_data[_col(headers, "初筛上限")]),
        "standard_deviation": _pct(row_data[_col(headers, "标准差")]),
        "median": _pct(row_data[_col(headers, "中位数")]),
        "first_quartile": _pct(row_data[_col(headers, "第一四分位数")]),
        "third_quartile": _pct(row_data[_col(headers, "第三四分位数")]),
        "interquartile_range": _pct(row_data[_col(headers, "四分位距")]),
        "max_dept_diff": _pct(row_data[_col(headers, "部门均值最大差")]),
    }


def _parse_trace_row(
    headers: list[str | None],
    row_data: list[object | None],
) -> dict[str, object | None]:
    return {
        "record_id": _str(row_data[_col(headers, "批次编号")]),
        "source_file": _str(row_data[_col(headers, "来源文件")]),
        "source_hash": _str(row_data[_col(headers, "文件哈希")]),
        "source_sheet": _str(row_data[_col(headers, "来源工作表")]),
        "sample_rows": _str(row_data[_col(headers, "样本行")]),
        "total_weight": row_data[_col(headers, "总重")],
        "department_map": _str(row_data[_col(headers, "样本部门映射")]),
        "rule_version": _str(row_data[_col(headers, "规则版本")]),
        "rule_hash": _str(row_data[_col(headers, "规则SHA-256")]),
        "rule_scope": _str(row_data[_col(headers, "规则适用范围")]),
        "rule_status": _str(row_data[_col(headers, "规则确认状态")]),
        "experience": _str(row_data[_col(headers, "经验引用")]),
    }


def _parse_anomaly_row(
    headers: list[str | None],
    row_data: list[object | None],
) -> dict[str, object | None]:
    return {
        "record_id": _str(row_data[_col(headers, "批次编号")]),
        "target_field": _str(row_data[_col(headers, "目标字段")]),
        "screen_result": _str(row_data[_col(headers, "筛查结果")]),
        "reason": _str(row_data[_col(headers, "原因")]),
        "suggestion": _str(row_data[_col(headers, "复核建议")]),
    }


# ---------------------------------------------------------------------------
# Row header validation
# ---------------------------------------------------------------------------


def _validate_headers(sheet_name: str, actual: list[str], expected: list[str]) -> None:
    missing = [h for h in expected if h not in actual]
    if missing:
        raise ValueError(f"'{sheet_name}' is missing required columns: {', '.join(missing)}")


# ---------------------------------------------------------------------------
# Public reader
# ---------------------------------------------------------------------------


class StandardSummaryReader:
    """Read a module-generated standard summary workbook and produce a
    validated ``HistoricalSummarySnapshot``.

    Supports v1.0.1 4-sheet and v1.0.0 5-sheet (legacy) formats.  The
    extra v1.0.0 ``模型分析记录`` sheet is validated but its AI content
    is ignored — only the deterministic four-sheet data is reconstructed.
    """

    def read(self, source: Path) -> HistoricalSummarySnapshot:
        if not source.is_file():
            raise FileNotFoundError(f"Summary workbook not found: {source}")
        if source.suffix.lower() != ".xlsx":
            raise ValueError("Summary workbook must be a .xlsx file")

        try:
            workbook = load_workbook(source, data_only=True, read_only=True)
        except Exception as error:
            raise ValueError("Summary workbook must be a valid .xlsx file") from error
        sheet_names = tuple(workbook.sheetnames)

        if sheet_names == _V1_0_1_SHEETS:
            format_version = SummaryFormat.V1_0_1_FOUR_SHEET
        elif sheet_names == _V1_0_0_SHEETS:
            format_version = SummaryFormat.V1_0_0_FIVE_SHEET
        else:
            raise ValueError(
                "Workbook does not match standard v1.0.1 (4-sheet) or "
                f"v1.0.0 (5-sheet) layout. Found: {list(sheet_names)}"
            )

        # Read all sheets
        main_ws = workbook[_SHEET_SUMMARY]
        anomaly_ws = workbook[_SHEET_ANOMALY]
        trend_ws = workbook[_SHEET_TREND]
        trace_ws = workbook[_SHEET_TRACE]

        # Trend sheet: record IDs and statistics
        trend_headers = [_str(c.value) for c in trend_ws[1]]
        _validate_headers(_SHEET_TREND, trend_headers, _TREND_HEADERS)
        trend_rows = [
            _parse_trend_row(trend_headers, _read_row(trend_ws, r, trend_ws.max_column))
            for r in range(2, trend_ws.max_row + 1)
        ]

        # Trace sheet: source/rule metadata
        trace_headers = [_str(c.value) for c in trace_ws[1]]
        _validate_headers(_SHEET_TRACE, trace_headers, _TRACE_HEADERS)
        trace_by_id: dict[str, dict[str, object | None]] = {}
        for r in range(2, trace_ws.max_row + 1):
            tr = _parse_trace_row(trace_headers, _read_row(trace_ws, r, trace_ws.max_column))
            trace_by_id[_str(tr["record_id"])] = tr

        # Main table: aligned with trend by position
        main_headers = [_str(c.value) for c in main_ws[1]]
        _validate_headers(_SHEET_SUMMARY, main_headers, _MAIN_HEADERS)
        main_rows = [
            _parse_main_row(main_headers, _read_row(main_ws, r, main_ws.max_column))
            for r in range(2, main_ws.max_row + 1)
        ]

        # Anomaly sheet: screening decisions
        anomaly_headers = [_str(c.value) for c in anomaly_ws[1]]
        _validate_headers(_SHEET_ANOMALY, anomaly_headers, _ANOMALY_HEADERS)
        anomaly_by_id: dict[str, dict[str, object | None]] = {}
        for r in range(2, anomaly_ws.max_row + 1):
            ar = _parse_anomaly_row(
                anomaly_headers, _read_row(anomaly_ws, r, anomaly_ws.max_column)
            )
            anomaly_by_id[_str(ar["record_id"])] = ar

        # Model sheet (v1.0.0 only): validate structure but ignore AI content
        if format_version == SummaryFormat.V1_0_0_FIVE_SHEET:
            model_ws = workbook[_SHEET_MODEL]
            if model_ws.max_row < 1:
                raise ValueError("模型分析记录 sheet has no header")

        workbook.close()

        # -----------------------------------------------------------------
        # Reconstruct batches
        # -----------------------------------------------------------------
        batches: list[WorkbookBatch] = []
        for idx, trend in enumerate(trend_rows):
            record_id = _str(trend["record_id"])
            trace = trace_by_id.get(record_id, {})
            main = main_rows[idx] if idx < len(main_rows) else {}
            anomaly = anomaly_by_id.get(record_id, {})
            sample_rows = _sample_rows(trace.get("sample_rows"))
            departments = _department_map(trace.get("department_map"))

            # Samples from main table
            sv = main.get("sample_values") or {}
            samples: list[SampleMeasurement] = []
            for si in range(1, 11):
                val = sv.get(si)
                if val is not None:
                    samples.append(
                        SampleMeasurement(
                            sample_index=si,
                            department_raw=departments.get(si, ""),
                            department=departments.get(si, ""),
                            source_row=sample_rows.get(si, 1),
                            initial_mass=Decimal("0"),
                            dry_mass=Decimal("0"),
                            foil_mass=Decimal("0"),
                            source_moisture_value=val,
                            verified_moisture_value=val,
                            value_used_for_statistics=val,
                        )
                    )

            if not samples:
                raise ValueError(f"Batch '{record_id}' has no readable samples")

            # Statistics
            stats = BatchStatistics(
                sample_count=len(samples),
                mean=main.get("mean") or Decimal("0"),
                total_average=main.get("mean") or Decimal("0"),
                standard_deviation=trend.get("standard_deviation"),
                minimum=main.get("minimum") or Decimal("0"),
                maximum=main.get("maximum") or Decimal("0"),
                range_value=main.get("range_value") or Decimal("0"),
                median=trend.get("median"),
                first_quartile=trend.get("first_quartile"),
                third_quartile=trend.get("third_quartile"),
                interquartile_range=trend.get("interquartile_range"),
                department_averages={},
                maximum_department_average_difference=trend.get("max_dept_diff"),
            )

            # Screening decision from anomaly sheet
            screen_text = _str(anomaly.get("screen_result"))
            display_color = DisplayColor.NONE
            screen_status = ScreenStatus.NORMAL
            if "偏高" in screen_text and ("疑似" in screen_text or "high" in screen_text.lower()):
                screen_status = ScreenStatus.HIGH
                display_color = DisplayColor.ORANGE
            elif "偏低" in screen_text and ("疑似" in screen_text or "low" in screen_text.lower()):
                screen_status = ScreenStatus.LOW
                display_color = DisplayColor.GREEN

            screening = ScreeningDecision(
                status=screen_status,
                display_color=display_color,
                target_field=_str(anomaly.get("target_field")) or None,
                reason=_str(anomaly.get("reason")),
            )

            source_file = _str(trace.get("source_file"))
            record = BatchRecord(
                record_id=record_id,
                samples=tuple(samples),
                source_file=Path(source_file) if source_file else Path("."),
                source_file_hash=_str(trace.get("source_hash")),
                source_sheet=_str(trace.get("source_sheet")),
                supplier=_str(main.get("supplier") or trend.get("supplier")),
                material_type=_str(main.get("material_type")) or "竹丝",
                arrival_date=main.get("arrival_date"),
                test_date=main.get("test_date"),
                total_weight_kg=_pct(trace.get("total_weight")),
            )

            batch = WorkbookBatch(
                record=record,
                statistics=stats,
                screening_decision=screening,
                ai_review=AIReviewDecision(status=ReviewStatus.NOT_COMPLETED),
                final_display_color=display_color,
                confirmation_state=ConfirmationState.NOT_REQUIRED,
                rule_version=_str(trace.get("rule_version")),
                rule_hash=_str(trace.get("rule_hash")),
                rule_applicable_scope=_str(trace.get("rule_scope")),
                rule_confirmation_status=_rule_confirmation_status(trace.get("rule_status")),
                lower_threshold=trend.get("lower_threshold"),
                upper_threshold=trend.get("upper_threshold"),
                review_suggestion=_str(anomaly.get("suggestion")),
                experience_references=_experience_references(trace.get("experience")),
            )
            batches.append(batch)

        # Source identity
        source_bytes = source.read_bytes()
        source_sha256 = __import__("hashlib").sha256(source_bytes).hexdigest()

        return HistoricalSummarySnapshot(
            format_version=format_version,
            source_sha256=source_sha256,
            source_display_name=source.name,
            batches=tuple(batches),
        )
