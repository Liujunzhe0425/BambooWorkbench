from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

MODULE_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(MODULE_ROOT / "src"))
sys.path.insert(0, str(MODULE_ROOT / "tests"))

from moisture.adapters.rules.local_editable_repository import (  # noqa: E402
    LocalEditableRuleRepository,
)
from moisture.application.draft_orchestrator import run_draft  # noqa: E402
from moisture.application.finalize_orchestrator import run_finalize  # noqa: E402
from moisture.application.management import execute_management  # noqa: E402
from moisture.contract.job_models import (  # noqa: E402
    JobPhase,
    JobRequest,
    ModuleId,
    RunMode,
)

pytest_plugins = ("fixtures.workbooks.workbook_factory",)

_ORANGE = "FFF4B183"


def _manifest(version: str) -> dict[str, object]:
    return {
        "schema_version": "1.0",
        "package_id": f"runtime-{version}",
        "rule_version": version,
        "rules_file": "rules.yaml",
        "rules_sha256": "0" * 64,
        "source": "synthetic runtime snapshot",
        "applicable_scope": "原竹丝；全年",
        "customer_confirmation_status": "confirmed",
        "released_at": "2026-07-24T00:00:00+08:00",
        "module_version": "1.2.0",
    }


def _rules(threshold: str) -> str:
    return f"""definitions:
  - rule_id: RUNTIME_HIGH
    rule_type: numeric_threshold
    field: total_average
    operator: ">"
    threshold: "{threshold}"
    color: orange
    applicable_materials: [原竹丝]
    applicable_seasons: [all]
    customer_confirmation_status: confirmed
"""


def _config(
    path: Path,
    *,
    version: str | None = None,
    sha256: str | None = None,
) -> Path:
    payload: dict[str, object] = {
        "schema_version": "1.0",
        "operation": "summary",
        "default_year": 2026,
    }
    if version is not None and sha256 is not None:
        payload["rule_selection"] = {"version": version, "sha256": sha256}
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def _job(
    source: Path,
    output: Path,
    config: Path,
    *,
    phase: JobPhase,
) -> JobRequest:
    return JobRequest(
        "1.0",
        "rule-snapshot-draft" if phase is JobPhase.DRAFT else "rule-snapshot-final",
        ModuleId.MOISTURE,
        phase,
        RunMode.SINGLE,
        (source,),
        output,
        parent_job_id="rule-snapshot-draft" if phase is JobPhase.FINALIZE else None,
        module_config_file=config,
    )


def _is_orange(path: Path) -> bool:
    workbook = load_workbook(path)
    try:
        sheet = workbook["竹丝含水率统计"]
        return sheet.cell(2, 19).fill.fgColor.rgb == _ORANGE
    finally:
        workbook.close()


def _source(tmp_path: Path, workbook_factory) -> Path:
    return workbook_factory(
        tmp_path / "source.xlsx",
        moisture_values=(0.60,),
        mass_values=((100, 40, 0),),
        material_type="原竹丝",
        arrival_date="2026-07-01",
        test_date="2026-07-02",
    )


def test_pinned_draft_finalize_ignore_later_active_publication(
    tmp_path: Path, workbook_factory, monkeypatch: pytest.MonkeyPatch
) -> None:
    data_dir = tmp_path / "module-data"
    monkeypatch.setenv("BAMBOO_MOISTURE_DATA_DIR", str(data_dir))
    repository = LocalEditableRuleRepository(data_dir / "rules")
    active_response = execute_management(
        {
            "schema_version": "1.1",
            "request_id": "snapshot-query",
            "module_id": "moisture",
            "operation": "get_active_rule_identity",
            "data_dir": str(data_dir.resolve()),
            "parameters": {},
        }
    )
    pinned = active_response["result"]["active_rule"]
    source = _source(tmp_path, workbook_factory)
    config = _config(
        tmp_path / "pinned.json",
        version=pinned["version"],
        sha256=pinned["sha256"],
    )

    draft_dir = tmp_path / "draft"
    assert run_draft(_job(source, draft_dir, config, phase=JobPhase.DRAFT)) == 0
    assert _is_orange(draft_dir / "moisture_summary.xlsx")

    next_draft = repository.create_draft(_manifest("runtime-v2"), _rules("0.65"))
    repository.publish(next_draft.rule_version, next_draft.package_hash)

    final_dir = tmp_path / "final"
    assert run_finalize(_job(source, final_dir, config, phase=JobPhase.FINALIZE)) == 0
    assert _is_orange(final_dir / "moisture_summary.xlsx")
    final_result = json.loads(
        (final_dir / "result.json").read_text(encoding="utf-8")
    )
    assert final_result["job_id"] == "rule-snapshot-final"
    assert final_result["confirmations"] == []
    assert final_result["service_requests"] == []

    new_dir = tmp_path / "new-task"
    active_config = _config(tmp_path / "active.json")
    assert run_draft(_job(source, new_dir, active_config, phase=JobPhase.DRAFT)) == 0
    assert not _is_orange(new_dir / "moisture_summary.xlsx")


@pytest.mark.parametrize("failure", ["missing", "hash", "retired", "pointer"])
def test_invalid_or_retired_snapshot_never_falls_back(
    tmp_path: Path,
    workbook_factory,
    monkeypatch: pytest.MonkeyPatch,
    failure: str,
) -> None:
    data_dir = tmp_path / "module-data"
    monkeypatch.setenv("BAMBOO_MOISTURE_DATA_DIR", str(data_dir))
    repository = LocalEditableRuleRepository(data_dir / "rules")
    current = repository.list_current()[0]
    source = _source(tmp_path, workbook_factory)

    if failure == "missing":
        config = _config(
            tmp_path / "config.json",
            version="missing-v1",
            sha256="a" * 64,
        )
    elif failure == "hash":
        config = _config(
            tmp_path / "config.json",
            version=current.rule_version,
            sha256="a" * 64,
        )
    elif failure == "retired":
        repository.retire_current(current.rules_sha256)
        config = _config(
            tmp_path / "config.json",
            version=current.rule_version,
            sha256=current.rules_sha256,
        )
    else:
        (data_dir / "rules" / "active.json").write_text("{}", encoding="utf-8")
        config = _config(tmp_path / "config.json")

    output = tmp_path / "failed"
    run_draft(_job(source, output, config, phase=JobPhase.DRAFT))
    result = json.loads((output / "result.json").read_text(encoding="utf-8"))

    assert result["status"] == "failed"
    assert result["outputs"] == []
    assert any(
        "PIPELINE_INITIALIZATION_FAILED" in error for error in result["errors"]
    )
