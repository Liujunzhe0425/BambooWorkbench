"""Integration tests for HistoryRepository."""

from __future__ import annotations

import os
import sys
from contextlib import suppress
from pathlib import Path

import pytest

MODULE_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(MODULE_ROOT / "src"))

from moisture.adapters.history.repository import (  # noqa: E402
    HistoryImportError,
    HistoryRepository,
    HistorySecurityError,
)

# ---------------------------------------------------------------------------
#  XLSX fixture helpers
# ---------------------------------------------------------------------------


def _make_four_sheet_xlsx(path: Path, *, seed: str = "") -> None:
    """Create a valid four-sheet analysis summary workbook."""
    from openpyxl import Workbook

    wb = Workbook()
    # Remove default sheet and create the four required sheets
    wb.remove(wb.active)

    _add_stats_sheet(wb, seed)
    _add_stub_sheet(wb, "异常分析明细")
    _add_stub_sheet(wb, "趋势分析")
    _add_stub_sheet(wb, "来源追溯")

    wb.save(path)
    wb.close()


def _add_stats_sheet(wb, seed: str) -> None:
    """Create the 竹丝含水率统计 sheet with core columns."""
    ws = wb.create_sheet("竹丝含水率统计", 0)
    # Header row with core column keywords
    ws["A1"] = "检测日期"
    ws["B1"] = "供应商名称"
    ws["C1"] = "含水率均值"
    ws["D1"] = "样本数量"
    # Some data rows
    ws["A2"] = f"2026-07-{10 + len(seed)}"
    ws["B2"] = f"测试供应商{seed}"
    ws["C2"] = 0.15
    ws["D2"] = 5
    ws["A3"] = "2026-07-12"
    ws["B3"] = "测试供应商B"
    ws["C3"] = 0.18
    ws["D3"] = 3


def _add_stub_sheet(wb, name: str) -> None:
    """Add a placeholder sheet with the given name."""
    ws = wb.create_sheet(name)
    ws["A1"] = "placeholder"


def _make_five_sheet_xlsx(path: Path) -> None:
    """Create a supported legacy five-sheet v1.0.0 workbook."""
    _make_four_sheet_xlsx(path)
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.create_sheet("模型分析记录")
    ws["A1"] = "old model analysis"
    wb.save(path)
    wb.close()


def _make_non_xlsx(path: Path) -> None:
    """Write a file that is not a valid XLSX."""
    path.write_bytes(b"this is not an xlsx file")


# ---------------------------------------------------------------------------
#  Tests
# ---------------------------------------------------------------------------


class TestImportValidWorkbook:
    """Successful imports of four-sheet workbooks."""

    def test_import_returns_history_id(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "valid.xlsx"
        _make_four_sheet_xlsx(src, seed="A")

        history_id = repo.import_summary(src, "测试导入A")

        assert isinstance(history_id, str)
        assert len(history_id) == 32  # uuid4 hex
        assert history_id.isalnum()

    def test_import_creates_index_and_copy(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "valid.xlsx"
        _make_four_sheet_xlsx(src, seed="B")

        history_id = repo.import_summary(src, "测试导入B")

        # Index file exists
        index_path = storage / "history-index.json"
        assert index_path.is_file()

        # Copy exists
        copy_path = storage / f"{history_id}.xlsx"
        assert copy_path.is_file()

    def test_imported_entry_appears_in_list(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "valid.xlsx"
        _make_four_sheet_xlsx(src, seed="C")

        history_id = repo.import_summary(src, "列表测试")

        entries = repo.list_entries()
        assert len(entries) == 1
        assert entries[0]["history_id"] == history_id
        assert entries[0]["display_name"] == "列表测试"
        assert entries[0]["status"] == "active"
        assert entries[0]["format_version"] == "1.0.1-four-sheet"
        assert entries[0]["storage_name"] == f"{history_id}.xlsx"

    def test_entry_is_retrievable_by_id(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "valid.xlsx"
        _make_four_sheet_xlsx(src, seed="D")

        history_id = repo.import_summary(src, "获取测试")

        entry = repo.get_entry(history_id)
        assert entry is not None
        assert entry["history_id"] == history_id
        assert entry["display_name"] == "获取测试"

    def test_get_nonexistent_returns_none(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)

        assert repo.get_entry("nonexistent-id") is None


class TestDuplicateImport:
    """Same file imported twice returns the original id."""

    def test_same_file_returns_existing_id(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "dup.xlsx"
        _make_four_sheet_xlsx(src, seed="dup")

        first_id = repo.import_summary(src, "首次导入")
        second_id = repo.import_summary(src, "重复导入")

        assert first_id == second_id

    def test_same_file_does_not_duplicate_copy(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "dup2.xlsx"
        _make_four_sheet_xlsx(src, seed="dup2")

        repo.import_summary(src, "首次")
        repo.import_summary(src, "重复")

        entries = repo.list_entries()
        assert len(entries) == 1

    def test_different_content_creates_new_entry(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src1 = tmp_path / "a.xlsx"
        src2 = tmp_path / "b.xlsx"
        _make_four_sheet_xlsx(src1, seed="X")
        _make_four_sheet_xlsx(src2, seed="Y")

        id1 = repo.import_summary(src1, "文件A")
        id2 = repo.import_summary(src2, "文件B")

        assert id1 != id2
        assert len(repo.list_entries()) == 2


class TestLegacyFiveSheet:
    """Legacy deterministic data remains available after upgrading."""

    def test_five_sheet_is_imported_and_old_model_sheet_is_ignored(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "five-sheet.xlsx"
        _make_five_sheet_xlsx(src)

        history_id = repo.import_summary(src, "五表旧格式")

        entry = repo.get_entry(history_id)
        assert entry is not None
        assert entry["format_version"] == "1.0.0-five-sheet"
        assert repo.resolve_active_summary(history_id).is_file()


class TestRejectInvalidFormats:
    """Unknown workbook formats are rejected with clear errors."""

    def test_non_xlsx_rejected(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "not-xlsx.bin"
        _make_non_xlsx(src)

        with pytest.raises(HistoryImportError):
            repo.import_summary(src, "非XLSX文件")

    def test_missing_sheet_rejected(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "three-sheet.xlsx"
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        _add_stub_sheet(wb, "竹丝含水率统计")
        _add_stub_sheet(wb, "异常分析明细")
        _add_stub_sheet(wb, "趋势分析")
        # missing 来源追溯
        wb.save(src)
        wb.close()

        with pytest.raises(HistoryImportError, match="来源追溯"):
            repo.import_summary(src, "缺表")

    def test_extra_sheet_rejected(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "extra-sheet.xlsx"
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        _add_stub_sheet(wb, "竹丝含水率统计")
        _add_stub_sheet(wb, "异常分析明细")
        _add_stub_sheet(wb, "趋势分析")
        _add_stub_sheet(wb, "来源追溯")
        _add_stub_sheet(wb, "额外表")
        wb.save(src)
        wb.close()

        with pytest.raises(HistoryImportError, match="unexpected"):
            repo.import_summary(src, "多余表")


class TestRetire:
    """Retiring entries."""

    def test_retire_changes_status(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "retire.xlsx"
        _make_four_sheet_xlsx(src, seed="retire")

        history_id = repo.import_summary(src, "待退役")
        assert repo.get_entry(history_id)["status"] == "active"

        repo.retire_entry(history_id)
        assert repo.get_entry(history_id)["status"] == "retired"

    def test_retire_idempotent(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "retire2.xlsx"
        _make_four_sheet_xlsx(src, seed="retire2")

        history_id = repo.import_summary(src, "二次退役")
        repo.retire_entry(history_id)
        repo.retire_entry(history_id)  # should not raise

        assert repo.get_entry(history_id)["status"] == "retired"

    def test_retire_nonexistent_raises(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)

        with pytest.raises(KeyError, match="not found"):
            repo.retire_entry("nonexistent")

    def test_list_with_status_filter(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)

        src1 = tmp_path / "active.xlsx"
        src2 = tmp_path / "retired_file.xlsx"
        _make_four_sheet_xlsx(src1, seed="act")
        _make_four_sheet_xlsx(src2, seed="ret")

        active_id = repo.import_summary(src1, "活跃条目")
        retired_id = repo.import_summary(src2, "退役条目")
        repo.retire_entry(retired_id)

        actives = repo.list_entries(status_filter="active")
        retireds = repo.list_entries(status_filter="retired")

        assert len(actives) == 1
        assert actives[0]["history_id"] == active_id
        assert len(retireds) == 1
        assert retireds[0]["history_id"] == retired_id


class TestIndexAtomicWrite:
    """Index file integrity across write operations."""

    def test_index_not_corrupted_by_failed_import(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)

        # Seed a valid entry first
        src1 = tmp_path / "good.xlsx"
        _make_four_sheet_xlsx(src1, seed="good")
        good_id = repo.import_summary(src1, "有效条目")

        # Remember the index content
        index_path = storage / "history-index.json"
        original_text = index_path.read_text(encoding="utf-8")

        # Attempt a failing import (bad file)
        bad_src = tmp_path / "bad.bin"
        _make_non_xlsx(bad_src)
        with suppress(HistoryImportError):
            repo.import_summary(bad_src, "应该失败")

        # Index should be unchanged
        current_text = index_path.read_text(encoding="utf-8")
        assert current_text == original_text

        # The valid entry is still there
        entries = repo.list_entries()
        assert len(entries) == 1
        assert entries[0]["history_id"] == good_id

    def test_new_repo_starts_with_empty_index(self, tmp_path: Path) -> None:
        storage = tmp_path / "fresh"
        repo = HistoryRepository(storage)

        entries = repo.list_entries()
        assert entries == []

        index_path = storage / "history-index.json"
        assert not index_path.exists()  # created lazily

    def test_index_survives_multiple_operations(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)

        src = tmp_path / "multi.xlsx"
        _make_four_sheet_xlsx(src, seed="multi")
        history_id = repo.import_summary(src, "多次操作")

        # Retire
        repo.retire_entry(history_id)
        assert repo.get_entry(history_id)["status"] == "retired"

        # Import again (new content)
        src2 = tmp_path / "multi2.xlsx"
        _make_four_sheet_xlsx(src2, seed="multi2")
        repo.import_summary(src2, "第二次")

        entries = repo.list_entries()
        assert len(entries) == 2


class TestPathSecurity:
    """Security boundary checks."""

    def test_symlink_source_rejected(self, tmp_path: Path) -> None:

        storage = tmp_path / "history"
        repo = HistoryRepository(storage)

        if os.name != "nt":
            real = tmp_path / "real.xlsx"
            _make_four_sheet_xlsx(real, seed="sym")
            link = tmp_path / "link.xlsx"
            link.symlink_to(real)

            with pytest.raises(HistorySecurityError):
                repo.import_summary(link, "符号链接")
        else:
            # On Windows, creating symlinks typically requires admin,
            # so we skip but verify that non-xlsx still fails correctly.
            src = tmp_path / "not-xlsx.bin"
            _make_non_xlsx(src)
            with pytest.raises(HistoryImportError):
                repo.import_summary(src, "非XLSX")


class TestBatchCountInference:
    """Batch count is inferred from the stats sheet."""

    def test_batch_count_recorded(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "batch.xlsx"
        _make_four_sheet_xlsx(src, seed="batch")

        history_id = repo.import_summary(src, "批次计数")
        entry = repo.get_entry(history_id)
        # We created 2 data rows in _make_four_sheet_xlsx
        assert entry["batch_count"] == 2


class TestDisplayNameValidation:
    """Display name must be 1-200 characters."""

    def test_empty_display_name_rejected(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "name.xlsx"
        _make_four_sheet_xlsx(src, seed="name")

        with pytest.raises(HistoryImportError, match="display_name"):
            repo.import_summary(src, "")

    def test_overlong_display_name_rejected(self, tmp_path: Path) -> None:
        storage = tmp_path / "history"
        repo = HistoryRepository(storage)
        src = tmp_path / "name2.xlsx"
        _make_four_sheet_xlsx(src, seed="name2")

        with pytest.raises(HistoryImportError, match="display_name"):
            repo.import_summary(src, "x" * 201)
