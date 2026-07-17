"""Behavior tests migrated from the stable Task 1—9 parser suite."""

from __future__ import annotations

import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

MODULE_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(MODULE_ROOT / "src"))
sys.path.insert(0, str(MODULE_ROOT / "tests"))

from moisture.adapters.excel.source_parser import (  # noqa: E402
    AmbiguousWorkbookStructure,
    ExcelSourceParser,
    SampleLimitExceeded,
    SourceFileModified,
    WorkbookStructureError,
)
from moisture.infrastructure.hashing import sha256_file  # noqa: E402

pytest_plugins = ("fixtures.workbooks.workbook_factory",)


@pytest.mark.parametrize("sample_count", [1, 2, 6, 10])
def test_parses_one_to_ten_samples_dynamically(
    tmp_path: Path, workbook_factory, sample_count: int
) -> None:
    path = workbook_factory(tmp_path / f"{sample_count}.xlsx", departments={"品质": sample_count})

    record = ExcelSourceParser().parse(path, default_year=2026)

    assert len(record.samples) == sample_count
    assert [sample.sample_index for sample in record.samples] == list(range(1, sample_count + 1))


def test_reads_dynamic_departments_and_source_rows(tmp_path: Path, workbook_factory) -> None:
    path = workbook_factory(
        tmp_path / "six.xlsx",
        departments={"二厂": 3, "品质": 2, "新部门": 1},
        sheet_name="检测记录",
        start_row=5,
    )

    record = ExcelSourceParser().parse(path, default_year=2026)

    assert [group.name for group in record.departments] == ["二厂", "品质", "新部门"]
    assert [group.sample_indices for group in record.departments] == [(1, 2, 3), (4, 5), (6,)]
    assert record.source_sheet == "检测记录"
    assert [sample.source_row for sample in record.samples] == [6, 7, 8, 9, 10, 11]


def test_eleven_samples_fails_without_truncation(tmp_path: Path, workbook_factory) -> None:
    path = workbook_factory(tmp_path / "eleven.xlsx", departments={"品质": 11})

    with pytest.raises(SampleLimitExceeded, match="11 valid samples") as error:
        ExcelSourceParser().parse(path, default_year=2026)

    assert error.value.code == "SAMPLE_LIMIT_EXCEEDED"


def test_detects_shifted_header_on_non_default_sheet(tmp_path: Path, workbook_factory) -> None:
    path = workbook_factory(
        tmp_path / "shifted.xlsx",
        departments={"二厂": 2},
        sheet_name="水分原始记录",
        start_row=27,
    )

    record = ExcelSourceParser().parse(path, default_year=2026)

    assert record.source_sheet == "水分原始记录"
    assert [sample.source_row for sample in record.samples] == [28, 29]


def test_preserves_source_and_verified_formula_values(tmp_path: Path, workbook_factory) -> None:
    path = workbook_factory(
        tmp_path / "formula.xlsx",
        departments={"品质": 2},
        moisture_formula_cached_values=[Decimal("0.1"), Decimal("0.5")],
    )

    record = ExcelSourceParser().parse(path, default_year=2026)

    first = record.samples[0]
    assert first.source_moisture_value == Decimal("0.1")
    assert first.verified_moisture_value != first.source_moisture_value
    assert first.value_used_for_statistics == first.source_moisture_value


def test_formula_without_cache_uses_independent_verification(
    tmp_path: Path, workbook_factory
) -> None:
    path = workbook_factory(tmp_path / "no-cache.xlsx")

    sample = ExcelSourceParser().parse(path, default_year=2026).samples[0]

    expected = (sample.initial_mass - sample.dry_mass) / (sample.initial_mass - sample.foil_mass)
    assert sample.source_moisture_value is None
    assert sample.verified_moisture_value == expected
    assert sample.value_used_for_statistics == expected


def test_reads_labelled_metadata_and_preserves_raw_values(tmp_path: Path, workbook_factory) -> None:
    path = workbook_factory(tmp_path / "metadata.xlsx")

    record = ExcelSourceParser().parse(path, default_year=2026)

    assert record.supplier_raw == "  青山供应商  "
    assert record.supplier == "青山供应商"
    assert record.material_type == "竹丝"
    assert record.arrival_date_raw == "7月12日"
    assert record.arrival_date == date(2026, 7, 12)
    assert record.test_date == date(2026, 7, 13)
    assert record.total_weight_raw == "约 1,234.5 kg（含包装）"
    assert record.total_weight_kg == Decimal("1234.5")


def test_legacy_monthly_layout_uses_parent_year_and_title_metadata(
    tmp_path: Path, workbook_factory
) -> None:
    path = workbook_factory(
        tmp_path / "2026年4月份" / "legacy.xlsx",
        start_row=3,
        departments={"二厂": 4, "品质": 4},
        legacy_title="脱敏供应方原竹丝水分检测",
        total_weight="13270kg",
    )

    record = ExcelSourceParser().parse(path)

    assert record.supplier == "脱敏供应方"
    assert record.material_type == "原竹丝"
    assert record.arrival_date == date(2026, 4, 7)
    assert record.test_date == date(2026, 4, 8)
    assert record.total_weight_kg == Decimal("13270")
    assert len(record.samples) == 8


def test_legacy_monthly_layout_without_year_context_stays_invalid(
    tmp_path: Path, workbook_factory
) -> None:
    path = workbook_factory(
        tmp_path / "legacy.xlsx",
        start_row=3,
        legacy_title="脱敏供应方竹丝水分检测",
    )

    with pytest.raises(WorkbookStructureError, match="default_year"):
        ExcelSourceParser().parse(path)


def test_conflicting_parent_year_contexts_are_rejected(
    tmp_path: Path, workbook_factory
) -> None:
    path = workbook_factory(
        tmp_path / "2025年" / "2026年4月份" / "legacy.xlsx",
        start_row=3,
        legacy_title="脱敏供应方竹丝水分检测",
    )

    with pytest.raises(WorkbookStructureError, match="conflicting year contexts"):
        ExcelSourceParser().parse(path)


def test_explicit_default_year_has_priority_over_parent_context(
    tmp_path: Path, workbook_factory
) -> None:
    path = workbook_factory(
        tmp_path / "2025年4月份" / "legacy.xlsx",
        start_row=3,
        legacy_title="脱敏供应方竹丝水分检测",
    )

    record = ExcelSourceParser().parse(path, default_year=2026)

    assert record.arrival_date == date(2026, 4, 7)
    assert record.test_date == date(2026, 4, 8)


def test_explicit_metadata_has_priority_over_legacy_title(
    tmp_path: Path, workbook_factory
) -> None:
    from openpyxl import load_workbook

    path = workbook_factory(
        tmp_path / "2026年4月份" / "legacy.xlsx",
        start_row=3,
        legacy_title="标题供应方原竹丝水分检测",
    )
    workbook = load_workbook(path)
    worksheet = workbook.active
    worksheet["B2"] = "供应商"
    worksheet["C2"] = "显式供应方"
    worksheet["H2"] = "物料类型"
    worksheet["I2"] = "显式材料"
    workbook.save(path)
    workbook.close()

    record = ExcelSourceParser().parse(path)

    assert record.supplier == "显式供应方"
    assert record.material_type == "显式材料"


def test_formatting_residue_does_not_create_samples(tmp_path: Path, workbook_factory) -> None:
    path = workbook_factory(tmp_path / "formatted.xlsx", departments={"品质": 2})

    record = ExcelSourceParser().parse(path, default_year=2026)

    assert len(record.samples) == 2


def test_missing_required_header_is_structured(tmp_path: Path, workbook_factory) -> None:
    from openpyxl import load_workbook

    path = workbook_factory(tmp_path / "missing-header.xlsx")
    workbook = load_workbook(path)
    workbook.active["D4"] = "未知字段"
    workbook.save(path)
    workbook.close()

    with pytest.raises(WorkbookStructureError) as error:
        ExcelSourceParser().parse(path, default_year=2026)

    assert error.value.code == "WORKBOOK_STRUCTURE_INVALID"


def test_ambiguous_sheet_returns_issue_not_guess(tmp_path: Path, workbook_factory) -> None:
    from openpyxl import load_workbook

    path = workbook_factory(tmp_path / "ambiguous.xlsx")
    workbook = load_workbook(path)
    duplicate = workbook.copy_worksheet(workbook.active)
    duplicate.title = "副本"
    workbook.save(path)
    workbook.close()

    with pytest.raises(AmbiguousWorkbookStructure) as error:
        ExcelSourceParser().parse(path, default_year=2026)

    assert error.value.code == "WORKBOOK_STRUCTURE_AMBIGUOUS"


def test_corrupt_workbook_is_structured(tmp_path: Path) -> None:
    path = tmp_path / "corrupt.xlsx"
    path.write_bytes(b"not an xlsx archive")

    with pytest.raises(WorkbookStructureError) as error:
        ExcelSourceParser().parse(path, default_year=2026)

    assert error.value.code == "WORKBOOK_STRUCTURE_INVALID"


def test_parser_does_not_modify_source_workbook(tmp_path: Path, workbook_factory) -> None:
    path = workbook_factory(tmp_path / "readonly.xlsx", departments={"品质": 3})
    before = sha256_file(path)

    record = ExcelSourceParser().parse(path, default_year=2026)

    assert sha256_file(path) == before
    assert record.source_file_hash == before
    assert record.source_file == path


def test_source_modification_overrides_original_parser_error(
    tmp_path: Path, workbook_factory, monkeypatch
) -> None:
    path = workbook_factory(tmp_path / "modified.xlsx", departments={"品质": 11})
    original = ExcelSourceParser._parse_samples

    def parse_samples_then_modify(self, formula_sheet, value_sheet, candidate):
        try:
            return original(self, formula_sheet, value_sheet, candidate)
        except SampleLimitExceeded:
            with path.open("ab") as source:
                source.write(b"changed-during-parse")
            raise

    monkeypatch.setattr(ExcelSourceParser, "_parse_samples", parse_samples_then_modify)

    with pytest.raises(SourceFileModified) as error:
        ExcelSourceParser().parse(path, default_year=2026)

    assert error.value.code == "SOURCE_FILE_MODIFIED"
    assert isinstance(error.value.__cause__, SampleLimitExceeded)
