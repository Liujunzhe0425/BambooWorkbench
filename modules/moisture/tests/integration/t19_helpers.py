from __future__ import annotations

import hashlib
import json
from pathlib import Path

from moisture.application.deterministic_pipeline import DeterministicPipeline
from moisture.contract.job_models import JobPhase, JobRequest, ModuleId, RunMode
from openpyxl import load_workbook


def rule_root(root: Path, direction: str, *, marker: str = "v1") -> Path:
    high = direction == "high"
    rules = (
        "definitions:\n"
        f"  - rule_id: {'RULE-HIGH' if high else 'RULE-LOW'}\n"
        "    rule_type: numeric_threshold\n"
        "    field: total_average\n"
        f'    operator: "{">" if high else "<"}"\n'
        f'    threshold: "{"0.55" if high else "0.47"}"\n'
        f"    color: {'orange' if high else 'green'}\n"
        f"    message: governed-{marker}\n"
        "    customer_confirmation_status: confirmed\n"
    )
    package = root / "rules" / "active"
    package.mkdir(parents=True)
    rules_path = package / "rules.yaml"
    rules_path.write_text(rules, encoding="utf-8", newline="\n")
    digest = hashlib.sha256(rules_path.read_bytes()).hexdigest()
    manifest = {
        "schema_version": "1.0",
        "package_id": f"confirmed-{direction}-{marker}",
        "rule_version": f"fixture-{marker}",
        "rules_file": "rules.yaml",
        "rules_sha256": digest,
        "source": "T19 fixture",
        "applicable_scope": "test",
        "customer_confirmation_status": "confirmed",
        "released_at": "2026-07-16T00:00:00+08:00",
        "module_version": "0.1.1",
    }
    (package / "rule-package.json").write_text(
        json.dumps(manifest, ensure_ascii=False), encoding="utf-8"
    )
    return root


def governed_workbook(workbook_factory, path: Path, direction: str) -> Path:
    source = workbook_factory(path, departments={"品质": 1}, moisture_values=(0.50,))
    workbook = load_workbook(source)
    sheet = workbook.active
    sheet["B5"] = 100
    sheet["C5"] = 40 if direction == "high" else 60
    sheet["D5"] = 0
    sheet["E5"] = 0.60 if direction == "high" else 0.40
    workbook.save(source)
    workbook.close()
    return source


def job(
    source: Path,
    output: Path,
    *,
    phase: JobPhase,
    service_file: Path | None = None,
    allowed_scopes: tuple[str, ...] = (),
) -> JobRequest:
    return JobRequest(
        "1.0",
        "draft-root" if phase is JobPhase.DRAFT else "finalize-current",
        ModuleId.MOISTURE,
        phase,
        RunMode.SINGLE,
        (source,),
        output,
        parent_job_id="draft-root" if phase is JobPhase.FINALIZE else None,
        service_response_file=service_file,
        allowed_knowledge_scopes=allowed_scopes,
    )


def pipeline(root: Path) -> DeterministicPipeline:
    return DeterministicPipeline(module_root=root)


def valid_ai_response(request: dict[str, object]) -> dict[str, object]:
    payload = request["payload"]
    candidate = payload["candidates"][0]
    text = json.dumps(
        {
            "schema_version": "1.0",
            "rule_hash": payload["rule_hash"],
            "decisions": [
                {
                    "candidate_id": candidate["candidate_id"],
                    "direction": candidate["direction"],
                    "field": candidate["field"],
                    "observed_value": candidate["observed_value"],
                    "decision": "escalate",
                    "confidence": 0.86,
                    "reason": "structured evidence",
                    "evidence_refs": candidate["evidence_refs"],
                }
            ],
        },
        ensure_ascii=False,
    )
    return {
        "request_id": request["request_id"],
        "service": "ai",
        "status": "ok",
        "text": text,
    }


def write_envelope(path: Path, responses: list[dict[str, object]]) -> Path:
    path.write_text(
        json.dumps(
            {
                "schema_version": "1.0",
                "task_id": "draft-root",
                "module_id": "moisture",
                "responses": responses,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return path


def workbook_contains_color(path: Path, color: str) -> bool:
    expected = {
        "orange": "FFF4B183",
        "green": "FFA9D18E",
        "red": "FFFF0000",
        "blue": "FF5B9BD5",
        "none": "00000000",
    }[color]
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return any(
            cell.fill.fgColor.rgb == expected
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
    finally:
        workbook.close()
