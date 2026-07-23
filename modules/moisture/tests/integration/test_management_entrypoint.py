from __future__ import annotations

import json
import os
import subprocess
import sys
from itertools import count
from pathlib import Path

from openpyxl import Workbook

MODULE_ROOT = Path(__file__).resolve().parents[2]
ENTRYPOINT = MODULE_ROOT / "run_management.bat"
_WORKSPACE_IDS = count()


def _summary(path: Path, *, legacy: bool = False) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    statistics = workbook.create_sheet("竹丝含水率统计")
    statistics.append(["检测日期", "供应商", "均值"])
    statistics.append(["2026-07-01", "脱敏供应商", 0.50])
    for name in ("异常分析明细", "趋势分析", "来源追溯"):
        workbook.create_sheet(name)["A1"] = "脱敏测试"
    if legacy:
        workbook.create_sheet("模型分析记录")["A1"] = "旧版辅助记录"
    workbook.save(path)
    workbook.close()
    return path


def _call(
    tmp_path: Path,
    data_dir: Path,
    operation: str,
    *,
    input_file: Path | None = None,
    history_id: str | None = None,
) -> tuple[subprocess.CompletedProcess[str], dict[str, object]]:
    workspace = tmp_path / f"workspace-{operation}-{next(_WORKSPACE_IDS)}"
    workspace.mkdir()
    request_path = (workspace / "request.json").resolve()
    response_path = (workspace / "response.json").resolve()
    request = {
        "schema_version": "1.0",
        "request_id": f"request-{operation.replace('_', '-')}",
        "module_id": "moisture",
        "operation": operation,
        "data_dir": str(data_dir.resolve()),
        "input_file": str(input_file.resolve()) if input_file else None,
        "history_id": history_id,
    }
    request_path.write_text(json.dumps(request, ensure_ascii=False), encoding="utf-8")
    environment = os.environ.copy()
    environment["BAMBOO_MODULE_PYTHON"] = sys.executable
    environment["BAMBOO_MOISTURE_DATA_DIR"] = str(data_dir.resolve())
    completed = subprocess.run(
        [
            str(ENTRYPOINT),
            "manage",
            "--request",
            str(request_path),
            "--response",
            str(response_path),
        ],
        cwd=tmp_path,
        env=environment,
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=30,
    )
    response = (
        json.loads(response_path.read_text(encoding="utf-8")) if response_path.exists() else {}
    )
    return completed, response


def _assert_progress_only(completed: subprocess.CompletedProcess[str]) -> None:
    completed.stdout.encode("ascii")
    events = [json.loads(line) for line in completed.stdout.splitlines()]
    assert [event["percent"] for event in events] == [0, 100]
    assert all(event["type"] == "progress" for event in events)
    assert "Traceback" not in completed.stderr


def test_management_entrypoint_imports_lists_and_retires_four_and_five_sheet(
    tmp_path: Path,
) -> None:
    data_dir = tmp_path / "module-data"
    data_dir.mkdir()
    current = _summary(tmp_path / "current.xlsx")
    legacy = _summary(tmp_path / "legacy.xlsx", legacy=True)

    for source in (current, legacy):
        completed, response = _call(
            tmp_path,
            data_dir,
            "import_history_summary",
            input_file=source,
        )
        assert completed.returncode == 0
        assert response["status"] == "ok"
        assert response["imported_history_id"]
        _assert_progress_only(completed)

    completed, response = _call(tmp_path, data_dir, "list_history_summaries")
    assert completed.returncode == 0
    assert len(response["history_items"]) == 2
    assert all("\\" not in item["display_name"] for item in response["history_items"])
    _assert_progress_only(completed)

    history_id = response["history_items"][0]["history_id"]
    completed, response = _call(
        tmp_path,
        data_dir,
        "retire_history_summary",
        history_id=history_id,
    )
    assert completed.returncode == 0
    assert response["status"] == "ok"
    _assert_progress_only(completed)

    _, response = _call(tmp_path, data_dir, "list_history_summaries")
    retired = next(item for item in response["history_items"] if item["history_id"] == history_id)
    assert retired["active"] is False


def test_management_entrypoint_lists_only_approved_experiences(tmp_path: Path) -> None:
    data_dir = tmp_path / "module-data"
    data_dir.mkdir()

    completed, response = _call(tmp_path, data_dir, "list_approved_experiences")

    assert completed.returncode == 0
    assert response["status"] == "ok"
    assert response["experience_items"]
    assert all(item["sha256"] for item in response["experience_items"])
    _assert_progress_only(completed)


def test_management_expected_failure_is_structured_and_path_free(tmp_path: Path) -> None:
    data_dir = tmp_path / "module-data"
    data_dir.mkdir()
    invalid = tmp_path / "invalid.xlsx"
    invalid.write_bytes(b"not an xlsx")

    completed, response = _call(
        tmp_path,
        data_dir,
        "import_history_summary",
        input_file=invalid,
    )

    assert completed.returncode == 0
    assert response["status"] == "error"
    assert response["error_code"] == "MANAGEMENT_OPERATION_FAILED"
    assert str(tmp_path) not in json.dumps(response, ensure_ascii=False)
    _assert_progress_only(completed)
