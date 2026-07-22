from __future__ import annotations

# ruff: noqa: E402, I001

import re
import sys
from contextlib import closing
from dataclasses import replace
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pytest

MODULE_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(MODULE_ROOT / "src"))

from moisture.domain.enums import (
    ConfirmationState,
    DisplayColor,
    ReviewStatus,
    ScreenStatus,
)
from moisture.domain.models import (
    AIReviewDecision,
    BatchRecord,
    BatchStatistics,
    SampleMeasurement,
    ScreeningDecision,
)
from moisture.ports.writers import ExperienceCitation, WorkbookBatch, WorkbookReport
from moisture.reporting import xlsx_writer
from moisture.reporting.xlsx_writer import COLOR_HEX, write_summary_xlsx

SHEET_NAMES = [
    "竹丝含水率统计",
    "异常分析明细",
    "趋势分析",
    "来源追溯",
    "模型分析记录",
]


def _record(record_id: str, source: Path, values: tuple[str, ...]) -> BatchRecord:
    samples = tuple(
        SampleMeasurement(
            sample_index=index,
            department_raw="品质" if index == 1 else "二厂",
            department="品质" if index == 1 else "二厂",
            source_row=index + 4,
            initial_mass=Decimal("100"),
            dry_mass=Decimal("52"),
            foil_mass=Decimal("4"),
            source_moisture_value=Decimal(value),
            verified_moisture_value=Decimal(value),
            value_used_for_statistics=Decimal(value),
        )
        for index, value in enumerate(values, start=1)
    )
    return BatchRecord(
        record_id=record_id,
        samples=samples,
        source_file=source,
        source_file_hash=("a" if record_id == "R1" else "b") * 64,
        source_sheet="检测数据",
        supplier="供应商甲",
        material_type="竹丝",
        arrival_date=date(2026, 7, 12),
        test_date=date(2026, 7, 13),
        total_weight_kg=Decimal("1234.5"),
    )


def _statistics(values: tuple[str, ...]) -> BatchStatistics:
    decimals = [Decimal(value) for value in values]
    mean = sum(decimals, start=Decimal("0")) / Decimal(len(decimals))
    return BatchStatistics(
        sample_count=len(decimals),
        mean=mean,
        total_average=mean,
        standard_deviation=Decimal("0.014142"),
        minimum=min(decimals),
        maximum=max(decimals),
        range_value=max(decimals) - min(decimals),
        median=mean,
        first_quartile=min(decimals),
        third_quartile=max(decimals),
        interquartile_range=max(decimals) - min(decimals),
        department_averages={"品质": decimals[0], "二厂": decimals[-1]},
        maximum_department_average_difference=max(decimals) - min(decimals),
    )


@pytest.fixture
def report(tmp_path: Path) -> WorkbookReport:
    batches: list[WorkbookBatch] = []
    settings = (
        ("R1", ("0.60", "0.62"), ScreenStatus.HIGH, DisplayColor.ORANGE, "mean"),
        ("R2", ("0.45", "0.46"), ScreenStatus.LOW, DisplayColor.GREEN, "minimum"),
        ("R3", ("0.57", "0.58"), ScreenStatus.HIGH, DisplayColor.RED, "maximum"),
        ("R4", ("0.43", "0.44"), ScreenStatus.LOW, DisplayColor.BLUE, "sample_1"),
    )
    for record_id, values, status, color, target in settings:
        source = tmp_path / f"{record_id}.xlsx"
        source.write_bytes(f"unchanged-{record_id}".encode())
        record = _record(record_id, source, values)
        batches.append(
            WorkbookBatch(
                record=record,
                statistics=_statistics(values),
                screening_decision=ScreeningDecision(
                    status=status,
                    display_color=color
                    if color in {DisplayColor.ORANGE, DisplayColor.GREEN}
                    else DisplayColor.ORANGE,
                    target_field=target,
                    reason="确定性初筛疑似",
                ),
                ai_review=AIReviewDecision(
                    status=(
                        ReviewStatus.REVIEW_HIGH
                        if color == DisplayColor.RED
                        else ReviewStatus.REVIEW_LOW
                        if color == DisplayColor.BLUE
                        else ReviewStatus.NOT_COMPLETED
                    ),
                    display_color=color
                    if color in {DisplayColor.RED, DisplayColor.BLUE}
                    else DisplayColor.NONE,
                    confirmation_state=ConfirmationState.PENDING_HUMAN_REVIEW,
                    target_field=target if color in {DisplayColor.RED, DisplayColor.BLUE} else None,
                    confidence=Decimal("0.86")
                    if color in {DisplayColor.RED, DisplayColor.BLUE}
                    else None,
                    rationale="高度疑似，非正式错误结论"
                    if color in {DisplayColor.RED, DisplayColor.BLUE}
                    else None,
                ),
                final_display_color=color,
                confirmation_state=ConfirmationState.PENDING_HUMAN_REVIEW,
                rule_version="rules-v1",
                rule_hash="c" * 64,
                rule_applicable_scope="竹丝/原竹丝；全年",
                rule_confirmation_status="confirmed",
                lower_threshold=Decimal("0.47"),
                upper_threshold=Decimal("0.55"),
                review_suggestion="建议人工复核，不构成正式判定",
                experience_references=(
                    ExperienceCitation(
                        "EXP-001", "1.2.0", "夏季竹丝复核经验", "experience/exp-001.md"
                    ),
                ),
                ai_degradation_reason="AI 未授权"
                if color in {DisplayColor.ORANGE, DisplayColor.GREEN}
                else None,
            )
        )
    return WorkbookReport(
        started_at=datetime(2026, 7, 13, 8, 0, tzinfo=UTC),
        batches=tuple(batches),
    )


MAIN_TABLE_HEADERS_20 = [
    "序号", "供应商", "材料类型", "到货日期", "检测日期",
    "样本1", "样本2", "样本3", "样本4", "样本5",
    "样本6", "样本7", "样本8", "样本9", "样本10",
    "最大值", "最小值", "区间值", "均值", "筛查结果",
]

MAIN_TABLE_EXCLUDED_COLUMNS = [
    "批次编号", "总重(kg)", "样本部门映射",
    "标准差", "中位数", "第一四分位数", "第三四分位数", "四分位距",
    "部门均值最大差", "初筛下限", "初筛上限",
    "规则版本", "规则SHA-256", "规则适用范围", "规则确认状态",
    "初筛状态", "AI复判状态", "置信度", "最终显示色", "人工确认状态",
    "复核建议", "来源文件", "来源工作表",
]


def _column(headers: list[str | None], name: str) -> int:
    return headers.index(name) + 1


def test_main_table_has_exactly_20_columns_and_correct_headers(
    tmp_path: Path, report: WorkbookReport
) -> None:
    output = write_summary_xlsx(report, tmp_path / "output")
    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        summary = workbook[SHEET_NAMES[0]]
        headers = [cell.value for cell in summary[1]]
        assert headers == MAIN_TABLE_HEADERS_20


def test_main_table_freeze_and_filter_are_constrained(
    tmp_path: Path, report: WorkbookReport
) -> None:
    output = write_summary_xlsx(report, tmp_path / "output")
    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        summary = workbook[SHEET_NAMES[0]]
        assert summary.freeze_panes == "F2"
        assert summary.auto_filter.ref == "A1:T1"


def test_main_table_excludes_technical_and_rule_columns(
    tmp_path: Path, report: WorkbookReport
) -> None:
    output = write_summary_xlsx(report, tmp_path / "output")
    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        summary = workbook[SHEET_NAMES[0]]
        headers = [cell.value for cell in summary[1]]
        for excluded in MAIN_TABLE_EXCLUDED_COLUMNS:
            assert excluded not in headers, (
                f"'{excluded}' must not appear in main table headers"
            )


def test_main_table_blank_sample_cells_are_empty(
    tmp_path: Path, report: WorkbookReport
) -> None:
    output = write_summary_xlsx(report, tmp_path / "output")
    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        summary = workbook[SHEET_NAMES[0]]
        for row in range(2, summary.max_row + 1):
            for col in range(6, 16):  # columns F-O (样本1—样本10)
                cell = summary.cell(row, col)
                if cell.value is None:
                    continue
                assert cell.value != 0, (
                    f"Empty sample cell ({row},{col}) must be None, not 0"
                )
                assert cell.value != "", (
                    f"Empty sample cell ({row},{col}) must be None, not empty string"
                )


def test_main_table_screen_result_text_mapping(
    tmp_path: Path, report: WorkbookReport
) -> None:
    output = write_summary_xlsx(report, tmp_path / "output")
    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        summary = workbook[SHEET_NAMES[0]]
        headers = [cell.value for cell in summary[1]]
        screen_col = _column(headers, "筛查结果")
        screen_values = {
            summary.cell(row, screen_col).value
            for row in range(2, summary.max_row + 1)
        }
        assert screen_values.issubset({"正常", "偏高疑似", "偏低疑似", "需人工确认"})
        assert "偏高疑似" in screen_values  # at least one HIGH case present


def test_screen_result_uses_data_quality_not_generic_pending_confirmation(
    tmp_path: Path, report: WorkbookReport
) -> None:
    first, second = report.batches[:2]
    normal_pending = replace(
        first,
        screening_decision=ScreeningDecision(
            status=ScreenStatus.NORMAL,
            display_color=DisplayColor.NONE,
            target_field=None,
            reason="确定性结果正常",
        ),
        final_display_color=DisplayColor.NONE,
        confirmation_state=ConfirmationState.PENDING_HUMAN_REVIEW,
    )
    data_quality = replace(
        second,
        screening_decision=ScreeningDecision(
            status=ScreenStatus.DATA_QUALITY,
            display_color=DisplayColor.NONE,
            target_field=None,
            reason="数据质量需要人工确认",
        ),
        final_display_color=DisplayColor.NONE,
    )
    focused_report = WorkbookReport(
        started_at=report.started_at,
        batches=(normal_pending, data_quality),
    )

    output = write_summary_xlsx(focused_report, tmp_path / "output")
    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        summary = workbook[SHEET_NAMES[0]]
        headers = [cell.value for cell in summary[1]]
        screen_col = _column(headers, "筛查结果")
        assert summary.cell(2, screen_col).value == "正常"
        assert summary.cell(3, screen_col).value == "需人工确认"


def test_main_table_material_type_is_normalized_to_yuanzhusi(
    tmp_path: Path, report: WorkbookReport
) -> None:
    output = write_summary_xlsx(report, tmp_path / "output")
    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        summary = workbook[SHEET_NAMES[0]]
        headers = [cell.value for cell in summary[1]]
        mat_col = _column(headers, "材料类型")
        for row in range(2, summary.max_row + 1):
            cell_value = summary.cell(row, mat_col).value
            assert cell_value == "原竹丝", (
                f"Material type must be '原竹丝', got '{cell_value}'"
            )


def test_workbook_contains_five_required_sheets_and_key_values(
    tmp_path: Path, report: WorkbookReport
) -> None:
    output = write_summary_xlsx(report, tmp_path / "output")

    assert output == tmp_path / "output" / "moisture_summary.xlsx"
    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        assert workbook.sheetnames == SHEET_NAMES
        summary = workbook[SHEET_NAMES[0]]
        headers = [cell.value for cell in summary[1]]

        # Core business values still present
        assert summary.cell(2, _column(headers, "均值")).value == pytest.approx(0.61)
        assert summary.cell(2, _column(headers, "均值")).number_format == "0.00%"

        # Relocated: 来源追溯 fields
        trace = workbook["来源追溯"]
        trace_headers = [cell.value for cell in trace[1]]
        assert trace.cell(2, _column(trace_headers, "来源文件")).value == "R1.xlsx"
        assert trace.cell(2, _column(trace_headers, "来源工作表")).value == "检测数据"
        assert trace.cell(2, _column(trace_headers, "规则SHA-256")).value == "c" * 64
        assert trace.cell(2, _column(trace_headers, "规则适用范围")).value == "竹丝/原竹丝；全年"
        assert trace.cell(2, _column(trace_headers, "规则确认状态")).value == "已确认"
        assert trace.cell(2, _column(trace_headers, "总重")).value == pytest.approx(1234.5)
        assert trace.cell(2, _column(trace_headers, "总重")).number_format == "#,##0.0"
        assert str(tmp_path) not in " ".join(
            str(cell.value)
            for row_ in trace.iter_rows()
            for cell in row_
            if cell.value is not None
        )

        # Relocated: 趋势分析 fields
        trend = workbook["趋势分析"]
        trend_headers = [cell.value for cell in trend[1]]
        assert trend.cell(2, _column(trend_headers, "初筛下限")).value == pytest.approx(0.47)
        assert trend.cell(2, _column(trend_headers, "初筛上限")).value == pytest.approx(0.55)

        # Relocated: 模型分析记录 fields
        model = workbook["模型分析记录"]
        model_headers = [cell.value for cell in model[1]]
        assert model.cell(2, _column(model_headers, "初筛状态")).value == "偏高疑似"
        assert model.cell(2, _column(model_headers, "AI复判状态")).value == "辅助复判未完成"
        assert model.cell(2, _column(model_headers, "最终显示色")).value == "橙色"
        assert model.cell(2, _column(model_headers, "人工确认状态")).value == "待人工复核"

        # Customer-visible text hygiene
        customer_text = " ".join(
            str(cell.value)
            for worksheet in workbook.worksheets
            if worksheet.title != "来源追溯"
            for row_ in worksheet.iter_rows()
            for cell in row_
            if cell.value is not None
        )
        assert "not_required" not in customer_text
        assert "pending_human_review" not in customer_text
        assert not re.search(r"\bNone\b", customer_text)


def test_review_colors_match_contract_semantics(tmp_path: Path, report: WorkbookReport) -> None:
    output = write_summary_xlsx(report, tmp_path / "output")
    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        sheet = workbook[SHEET_NAMES[0]]
        headers = [cell.value for cell in sheet[1]]
        screen_col = _column(headers, "筛查结果")
        targets = ("均值", "最小值", "最大值", "样本1")
        colors = ("orange", "green", "red", "blue")
        for row, target, color in zip(range(2, 6), targets, colors, strict=True):
            # Target indicator cell keeps its color
            assert (
                sheet.cell(row, _column(headers, target)).fill.fgColor.rgb
                == f"FF{COLOR_HEX[color]}"
            )
            # 筛查结果 cell must also carry the final display color
            assert (
                sheet.cell(row, screen_col).fill.fgColor.rgb
                == f"FF{COLOR_HEX[color]}"
            )


def test_workbook_has_named_charts_and_plain_untrusted_text(
    tmp_path: Path, report: WorkbookReport
) -> None:
    first = report.batches[0]
    unsafe_record = replace(first.record, supplier='=HYPERLINK("https://invalid.example","x")')
    unsafe_report = WorkbookReport(
        started_at=report.started_at,
        batches=(replace(first, record=unsafe_record), *report.batches[1:]),
    )
    output = write_summary_xlsx(unsafe_report, tmp_path / "output")

    with ZipFile(output) as archive:
        chart_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("xl/charts/chart")
        )
    assert "来料平均含水率散点图" in chart_xml
    assert "区间值趋势图" in chart_xml

    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        sheet = workbook[SHEET_NAMES[0]]
        headers = [cell.value for cell in sheet[1]]
        supplier = sheet.cell(2, headers.index("供应商") + 1)
        assert supplier.value == '=HYPERLINK("https://invalid.example","x")'
        assert supplier.data_type == "s"
        error_pattern = re.compile(r"#(?:REF!|DIV/0!|VALUE!|NAME\?|N/A)")
        assert all(
            not (isinstance(cell.value, str) and error_pattern.search(cell.value))
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
        )


def test_trend_charts_do_not_overlap_extended_statistics_table(
    tmp_path: Path, report: WorkbookReport
) -> None:
    output = write_summary_xlsx(report, tmp_path / "output")
    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        trend = workbook[SHEET_NAMES[2]]
        assert trend.max_column == 14
        assert len(trend._charts) == 2
        assert all(chart.anchor._from.col >= 15 for chart in trend._charts)


def test_trend_rows_are_sorted_by_detection_date(
    tmp_path: Path, report: WorkbookReport
) -> None:
    dates = (date(2026, 4, 29), date(2026, 4, 8), date(2026, 4, 10), date(2026, 4, 9))
    unsorted = WorkbookReport(
        started_at=report.started_at,
        batches=tuple(
            replace(batch, record=replace(batch.record, test_date=test_date))
            for batch, test_date in zip(report.batches, dates, strict=True)
        ),
    )

    output = write_summary_xlsx(unsorted, tmp_path / "output")

    with closing(openpyxl.load_workbook(output, data_only=False)) as workbook:
        trend = workbook[SHEET_NAMES[2]]
        assert [trend.cell(row, 2).value for row in range(2, 6)] == ["R2", "R4", "R3", "R1"]
        assert [trend.cell(row, 3).value.date() for row in range(2, 6)] == sorted(dates)


def test_workbook_is_deterministic(tmp_path: Path, report: WorkbookReport) -> None:
    first = write_summary_xlsx(report, tmp_path / "first")
    second = write_summary_xlsx(report, tmp_path / "second")

    assert first.read_bytes() == second.read_bytes()


def test_xlsx_is_replaced_atomically(
    tmp_path: Path, report: WorkbookReport, monkeypatch: pytest.MonkeyPatch
) -> None:
    output_dir = tmp_path / "output"
    output_dir.mkdir()
    final = output_dir / "moisture_summary.xlsx"
    final.write_bytes(b"previous-valid-report")

    def fail_validation(path: Path) -> None:
        assert path.parent == output_dir
        assert path.name.startswith(".moisture_summary-")
        assert path.suffix == ".xlsx"
        raise ValueError("synthetic validation failure")

    monkeypatch.setattr(xlsx_writer, "validate_workbook", fail_validation)

    with pytest.raises(ValueError, match="synthetic validation failure"):
        write_summary_xlsx(report, output_dir)

    assert final.read_bytes() == b"previous-valid-report"
    assert not list(output_dir.glob(".moisture_summary-*.xlsx"))


def test_generation_failure_leaves_no_formal_or_temporary_output(
    tmp_path: Path, report: WorkbookReport, monkeypatch: pytest.MonkeyPatch
) -> None:
    def fail_summary(*args: object) -> None:
        raise RuntimeError("synthetic generation failure")

    monkeypatch.setattr(xlsx_writer, "_write_summary", fail_summary)
    output_dir = tmp_path / "output"

    with pytest.raises(RuntimeError, match="synthetic generation failure"):
        write_summary_xlsx(report, output_dir)

    assert not (output_dir / "moisture_summary.xlsx").exists()
    assert not list(output_dir.glob(".moisture_summary-*.xlsx"))


def test_writer_does_not_modify_input_files(tmp_path: Path, report: WorkbookReport) -> None:
    before = {
        batch.record.source_file: batch.record.source_file.read_bytes() for batch in report.batches
    }

    write_summary_xlsx(report, tmp_path / "output")

    assert {path: path.read_bytes() for path in before} == before
