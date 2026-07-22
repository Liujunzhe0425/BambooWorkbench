from __future__ import annotations

import json
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

MODULE_ROOT = Path(__file__).resolve().parents[2]
REPOSITORY_ROOT = MODULE_ROOT.parents[1]
LAUNCHER = MODULE_ROOT / "run_module.bat"
sys.path.insert(0, str(MODULE_ROOT / "src"))
sys.path.insert(0, str(MODULE_ROOT / "tests"))

from moisture.application.deterministic_pipeline import DeterministicPipeline  # noqa: E402
from moisture.application.draft_orchestrator import run_draft  # noqa: E402
from moisture.contract.job_models import JobRequest  # noqa: E402

pytest_plugins = ("fixtures.workbooks.workbook_factory",)


def _job(
    source: Path,
    output_dir: Path,
    *,
    ai_enabled: bool = False,
    module_config_file: Path | None = None,
) -> JobRequest:
    return JobRequest.from_payload(
        {
            "schema_version": "1.0",
            "job_id": "job-t14-offline",
            "module_id": "moisture",
            "phase": "draft",
            "run_mode": "single",
            "input_files": [str(source.resolve())],
            "output_dir": str(output_dir.resolve()),
            "module_config_file": (
                str(module_config_file.resolve()) if module_config_file else None
            ),
            "ai_enabled": ai_enabled,
        }
    )


def test_ai_disabled_still_produces_xlsx_markdown_and_result_last(
    tmp_path: Path, workbook_factory, monkeypatch
) -> None:
    source = workbook_factory(tmp_path / "source.xlsx", moisture_values=(0.50,))
    output_dir = tmp_path / "output"
    replacements: list[str] = []
    import os

    real_replace = os.replace

    def recording_replace(source_path, target_path):
        replacements.append(Path(target_path).name)
        return real_replace(source_path, target_path)

    monkeypatch.setattr(os, "replace", recording_replace)
    exit_code = run_draft(_job(source, output_dir, ai_enabled=False))

    assert exit_code == 0
    assert replacements[-1] == "result.json"
    payload = json.loads((output_dir / "result.json").read_text(encoding="utf-8"))
    assert payload["status"] == "needs_confirmation"
    assert len(payload["confirmations"]) == 1
    assert payload["service_requests"] == []
    assert [item["path"] for item in payload["outputs"]] == [
        str((output_dir / "moisture_summary.xlsx").resolve()),
        str((output_dir / "analysis_report.md").resolve()),
    ]
    workbook = load_workbook(output_dir / "moisture_summary.xlsx", read_only=True)
    assert workbook.sheetnames == [
        "竹丝含水率统计",
        "异常分析明细",
        "趋势分析",
        "来源追溯",
        "模型分析记录",
    ]
    workbook.close()


def test_all_failed_still_writes_schema_valid_result_and_returns_zero(tmp_path: Path) -> None:
    source = tmp_path / "corrupt.xlsx"
    source.write_bytes(b"not a workbook")
    output_dir = tmp_path / "output"

    exit_code = run_draft(_job(source, output_dir))

    assert exit_code == 0
    payload = json.loads((output_dir / "result.json").read_text(encoding="utf-8"))
    assert payload["status"] == "failed"
    assert payload["outputs"] == []
    assert payload["summary"]["failed_count"] == 1


def test_module_config_default_year_enables_month_day_workbook(
    tmp_path: Path, workbook_factory
) -> None:
    source = workbook_factory(
        tmp_path / "source.xlsx",
        arrival_date="7月12日",
        test_date="7月13日",
    )
    config = tmp_path / "module-config.json"
    config.write_text('{"default_year": 2026}', encoding="utf-8")
    output_dir = tmp_path / "output"

    exit_code = run_draft(_job(source, output_dir, module_config_file=config))

    assert exit_code == 0
    payload = json.loads((output_dir / "result.json").read_text(encoding="utf-8"))
    assert payload["status"] == "success"
    assert payload["summary"]["processed_count"] == 1
    assert {Path(item["path"]).name for item in payload["outputs"]} == {
        "moisture_summary.xlsx",
        "analysis_report.md",
    }


def test_missing_module_config_requests_one_batch_year_without_discarding_record(
    tmp_path: Path, workbook_factory
) -> None:
    source = workbook_factory(
        tmp_path / "source.xlsx",
        arrival_date="7月12日",
        test_date="7月13日",
    )
    output_dir = tmp_path / "output"

    assert run_draft(_job(source, output_dir)) == 0

    payload = json.loads((output_dir / "result.json").read_text(encoding="utf-8"))
    assert payload["status"] == "needs_confirmation"
    assert payload["summary"]["processed_count"] == 1
    assert payload["summary"]["failed_count"] == 0
    assert len(payload["confirmations"]) == 1
    prompt = payload["confirmations"][0]
    assert prompt["field"] == "default_year"
    assert prompt["options"] == []
    assert prompt["batch_applicable"] is True


@pytest.mark.parametrize(
    "content",
    (
        "not-json",
        "[]",
        '{"default_year": true}',
        '{"default_year": 1800}',
        '{"default_year": 2026, "unknown": 1}',
        '{"default_year": 2025, "default_year": 2026}',
    ),
)
def test_invalid_module_config_writes_explicit_legal_failed_result(
    tmp_path: Path, workbook_factory, content: str
) -> None:
    source = workbook_factory(tmp_path / "source.xlsx")
    config = tmp_path / "module-config.json"
    config.write_text(content, encoding="utf-8")
    output_dir = tmp_path / "output"

    assert run_draft(_job(source, output_dir, module_config_file=config)) == 0

    payload = json.loads((output_dir / "result.json").read_text(encoding="utf-8"))
    assert payload["status"] == "failed"
    assert payload["summary"]["failed_count"] == 1
    assert payload["errors"] == ["MODULE_CONFIG_INVALID"]
    assert payload["file_results"][0]["error"] == "MODULE_CONFIG_INVALID"


def test_formal_launcher_consumes_default_year_without_modifying_input(
    tmp_path: Path, workbook_factory
) -> None:
    source = workbook_factory(
        tmp_path / "source.xlsx",
        arrival_date="7月12日",
        test_date="7月13日",
    )
    before = source.read_bytes()
    config = tmp_path / "module-config.json"
    config.write_text('{"default_year": 2026}', encoding="utf-8")
    output_dir = tmp_path / "formal-output"
    job = {
        "schema_version": "1.0",
        "job_id": "job-t14-config-launcher",
        "module_id": "moisture",
        "phase": "draft",
        "run_mode": "single",
        "input_files": [str(source.resolve())],
        "output_dir": str(output_dir.resolve()),
        "module_config_file": str(config.resolve()),
        "ai_enabled": False,
        "allowed_knowledge_scopes": [],
        "parent_job_id": None,
        "confirmation_file": None,
        "service_response_file": None,
    }
    job_path = tmp_path / "job.json"
    job_path.write_text(json.dumps(job), encoding="utf-8")

    completed = subprocess.run(
        [str(LAUNCHER), "run", "--job", str(job_path.resolve())],
        cwd=REPOSITORY_ROOT,
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )

    assert completed.returncode == 0, completed.stderr
    payload = json.loads((output_dir / "result.json").read_text(encoding="utf-8"))
    assert payload["status"] == "success"
    assert source.read_bytes() == before


def test_high_low_normal_use_rule_snapshot_bounds_and_distinct_suggestions(
    tmp_path: Path, workbook_factory
) -> None:
    sources = (
        workbook_factory(tmp_path / "high.xlsx", moisture_values=(0.56,) * 6),
        workbook_factory(tmp_path / "low.xlsx", moisture_values=(0.46,) * 8),
        workbook_factory(tmp_path / "normal.xlsx", moisture_values=(0.50,) * 6),
    )

    outcome = DeterministicPipeline(MODULE_ROOT).run(sources, default_year=2026)

    assert outcome.processed_count == 3
    assert [batch.screening_decision.status.value for batch in outcome.batches] == [
        "high",
        "low",
        "normal",
    ]
    assert all(batch.lower_threshold == Decimal("0.47") for batch in outcome.batches)
    assert all(batch.upper_threshold == Decimal("0.55") for batch in outcome.batches)
    assert "高于55.00%" in outcome.batches[0].review_suggestion
    assert "低于47.00%" in outcome.batches[1].review_suggestion
    assert outcome.batches[2].review_suggestion == "未命中已确认筛查规则"
    assert all(batch.rule_version == "default-v2-approved" for batch in outcome.batches)
    assert all(batch.rule_hash == outcome.rule_package.rules_sha256 for batch in outcome.batches)
    assert all(batch.rule_applicable_scope == "竹丝/原竹丝；全年" for batch in outcome.batches)
    assert all(batch.rule_confirmation_status == "confirmed" for batch in outcome.batches)
