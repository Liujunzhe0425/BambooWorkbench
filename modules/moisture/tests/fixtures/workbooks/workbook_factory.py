"""Synthetic workbook fixtures migrated from the stable Task 1—9 tests."""

from __future__ import annotations

import re
from collections.abc import Callable, Mapping, Sequence
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill


def _rewrite_formula_caches(path: Path, cached_values: Mapping[str, object]) -> None:
    temporary_path = path.with_suffix(".cached-values.tmp")
    with ZipFile(path, "r") as source_archive:
        entries = [(item, source_archive.read(item.filename)) for item in source_archive.infolist()]

    with ZipFile(temporary_path, "w", compression=ZIP_DEFLATED) as target_archive:
        for item, content in entries:
            if item.filename == "xl/worksheets/sheet1.xml":
                for coordinate, cached_value in cached_values.items():
                    pattern = re.compile(
                        rb'(<c\s+[^>]*r="'
                        + re.escape(coordinate.encode("ascii"))
                        + rb'"[^>]*>\s*<f[^>]*>.*?</f>\s*)(?:<v\s*/>|<v>.*?</v>)',
                        re.DOTALL,
                    )
                    replacement = rb"\g<1><v>" + str(cached_value).encode("ascii") + rb"</v>"
                    content, replacement_count = pattern.subn(replacement, content, count=1)
                    if replacement_count != 1:
                        raise AssertionError(
                            f"formula cell {coordinate} was not found in worksheet XML"
                        )
            target_archive.writestr(item, content)
    path.unlink()
    temporary_path.rename(path)


@pytest.fixture
def workbook_factory() -> Callable[..., Path]:
    def create(
        path: Path,
        *,
        departments: Mapping[str, int] | None = None,
        sheet_name: str = "Sheet1",
        start_row: int = 4,
        moisture_values: Sequence[object] | None = None,
        moisture_formula_cached_values: Sequence[object] | None = None,
        mass_formulas_with_cached_values: bool = False,
        supplier: object = "  青山供应商  ",
        material_type: object = "竹丝",
        arrival_date: object = "7月12日",
        test_date: object = "2026-07-13",
        total_weight: object = "约 1,234.5 kg（含包装）",
    ) -> Path:
        groups = departments or {"品质": 1}
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name

        metadata = (
            ("供应商", supplier),
            ("物料类型", material_type),
            ("来料日期", arrival_date),
            ("检测日期", test_date),
            ("总重量", total_weight),
        )
        for column_offset, (label, value) in enumerate(metadata):
            column = 1 + column_offset * 2
            worksheet.cell(1, column, label)
            worksheet.cell(1, column + 1, value)

        headers = ("取样部门", "初始质量", "烘干质量", "锡纸质量", "相对含水率")
        for column, header in enumerate(headers, start=1):
            worksheet.cell(start_row, column, header)

        sample_index = 0
        formula_caches: dict[str, object] = {}
        row = start_row + 1
        for department, count in groups.items():
            first_row = row
            for _ in range(count):
                sample_index += 1
                initial = 100 + sample_index
                dry = 52 + sample_index
                foil = 4
                worksheet.cell(row, 1, department if row == first_row else None)
                if mass_formulas_with_cached_values:
                    for column, value in ((2, initial), (3, dry), (4, foil)):
                        cell = worksheet.cell(row, column, f"={value}")
                        formula_caches[cell.coordinate] = value
                else:
                    worksheet.cell(row, 2, initial)
                    worksheet.cell(row, 3, dry)
                    worksheet.cell(row, 4, foil)
                if moisture_formula_cached_values is not None:
                    cell = worksheet.cell(row, 5, f"=(B{row}-C{row})/(B{row}-D{row})")
                    formula_caches[cell.coordinate] = moisture_formula_cached_values[
                        sample_index - 1
                    ]
                elif moisture_values is None:
                    worksheet.cell(row, 5, f"=(B{row}-C{row})/(B{row}-D{row})")
                else:
                    worksheet.cell(row, 5, moisture_values[sample_index - 1])
                row += 1
            if count > 1:
                worksheet.merge_cells(
                    start_row=first_row,
                    start_column=1,
                    end_row=row - 1,
                    end_column=1,
                )

        blank_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        for blank_row in range(start_row, row + 12):
            for column in range(8, 16):
                worksheet.cell(blank_row, column).fill = blank_fill

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        workbook.close()
        if formula_caches:
            _rewrite_formula_caches(path, formula_caches)
        return path

    return create
