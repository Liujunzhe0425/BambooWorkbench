from __future__ import annotations

# ruff: noqa: E402, I001

import hashlib
import json
import sys
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook

MODULE_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(MODULE_ROOT / "src"))

from moisture.application.monthly_rollup import render_monthly_rollup
from moisture.contract.monthly_rollup import (
    MonthlyRollupError,
    MonthlyRollupRequest,
    MonthlySnapshotReference,
)
from moisture.contract.monthly_snapshot import MonthlySnapshot, canonical_snapshot_bytes

FIXTURE = MODULE_ROOT / "tests" / "fixtures" / "monthly" / "snapshot-example.json"
RULE_HASH = "a" * 64


def _snapshot_payload(*, month: int, record_id: str, day: int, mean: str) -> dict:
    payload = deepcopy(json.loads(FIXTURE.read_text(encoding="utf-8")))
    payload["period"] = {"year": 2026, "month": month}
    record = payload["records"][0]
    record["record_id"] = record_id
    record["business_fingerprint"] = hashlib.sha256(record_id.encode()).hexdigest()
    record["arrival_date"] = f"2026-{month:02d}-{day:02d}"
    record["test_date"] = f"2026-{month:02d}-{day + 1:02d}"
    record["supplier"] = f"供应方-{record_id}"
    record["traceability"]["source_display_name"] = f"{record_id}.xlsx"
    record["traceability"]["source_file_hash"] = hashlib.sha256(
        f"source-{record_id}".encode()
    ).hexdigest()
    for sample in record["samples"]:
        sample["value"] = mean
    for field in ("mean", "total_average", "minimum", "maximum", "median"):
        record["statistics"][field] = mean
    record["statistics"]["range_value"] = "0"
    return payload


def _write_snapshot(
    root: Path,
    *,
    month: int,
    record_id: str,
    day: int,
    mean: str,
) -> MonthlySnapshotReference:
    snapshot = MonthlySnapshot.from_payload(
        _snapshot_payload(month=month, record_id=record_id, day=day, mean=mean)
    )
    content = canonical_snapshot_bytes(snapshot)
    path = root / f"{month:02d}.json"
    path.write_bytes(content)
    return MonthlySnapshotReference(
        month=month,
        path=path.resolve(),
        sha256=hashlib.sha256(content).hexdigest(),
    )


def _request(
    tmp_path: Path,
    *,
    months: tuple[int, ...] = (1, 3),
) -> MonthlyRollupRequest:
    root = (tmp_path / "snapshots").resolve()
    output = (tmp_path / "output").resolve()
    root.mkdir(parents=True)
    output.mkdir(parents=True)
    references = tuple(
        _write_snapshot(
            root,
            month=month,
            record_id=f"batch-{month:02d}",
            day=5,
            mean=f"0.{45 + month:02d}",
        )
        for month in months
    )
    return MonthlyRollupRequest(
        year=2026,
        month=months[-1],
        snapshot_root=root,
        output_dir=output,
        business_config_version="default-v2-approved",
        business_config_sha256=RULE_HASH,
        snapshots=references,
    )


def test_gap_months_render_complete_four_sheet_outputs_and_business_records(
    tmp_path: Path,
) -> None:
    request = _request(tmp_path)
    result = render_monthly_rollup(request)

    assert result.used_months == (1, 3)
    assert result.record_count == 2
    assert [artifact.path for artifact in result.artifacts] == [
        "moisture_summary.xlsx",
        "analysis_report.md",
        "moisture_business_records.json",
    ]
    assert all(
        (request.output_dir / artifact.path).is_file()
        and len(artifact.sha256) == 64
        for artifact in result.artifacts
    )

    workbook = load_workbook(request.output_dir / "moisture_summary.xlsx")
    assert workbook.sheetnames == ["竹丝含水率统计", "异常分析明细", "趋势分析", "来源追溯"]
    summary = workbook["竹丝含水率统计"]
    assert [summary.cell(row=row, column=2).value for row in (2, 3)] == [
        "供应方-batch-01",
        "供应方-batch-03",
    ]
    workbook.close()

    business = json.loads(
        (request.output_dir / "moisture_business_records.json").read_text("utf-8")
    )
    assert business["used_months"] == [1, 3]
    assert [record["record_id"] for record in business["records"]] == [
        "batch-01",
        "batch-03",
    ]
    assert str(tmp_path) not in json.dumps(result.to_payload(), ensure_ascii=False)


def test_request_rejects_network_duplicate_disordered_and_late_months(
    tmp_path: Path,
) -> None:
    request = _request(tmp_path)
    common = {
        "year": request.year,
        "month": request.month,
        "snapshot_root": request.snapshot_root,
        "output_dir": request.output_dir,
        "business_config_version": request.business_config_version,
        "business_config_sha256": request.business_config_sha256,
    }

    with pytest.raises(MonthlyRollupError) as offline:
        MonthlyRollupRequest(
            **common,
            snapshots=request.snapshots,
            network_allowed=True,
        )
    assert offline.value.code == "MONTHLY_ROLLUP_OFFLINE_VIOLATION"

    with pytest.raises(MonthlyRollupError) as duplicate:
        MonthlyRollupRequest(
            **common,
            snapshots=(request.snapshots[0], request.snapshots[0]),
        )
    assert duplicate.value.code == "MONTHLY_ROLLUP_DUPLICATE_MONTH"

    with pytest.raises(MonthlyRollupError) as unordered:
        MonthlyRollupRequest(
            **common,
            snapshots=tuple(reversed(request.snapshots)),
        )
    assert unordered.value.code == "MONTHLY_ROLLUP_MONTH_ORDER_INVALID"

    with pytest.raises(MonthlyRollupError) as late:
        MonthlyRollupRequest(
            **(common | {"month": 1}),
            snapshots=request.snapshots,
        )
    assert late.value.code == "MONTHLY_ROLLUP_MONTH_AFTER_TARGET"

    with pytest.raises(MonthlyRollupError) as target_missing:
        MonthlyRollupRequest(
            **(common | {"month": 2}),
            snapshots=(request.snapshots[0],),
        )
    assert target_missing.value.code == "MONTHLY_ROLLUP_REQUEST_INVALID"


def test_hash_year_declared_month_and_path_escape_are_rejected(tmp_path: Path) -> None:
    request = _request(tmp_path)
    first = request.snapshots[0]

    bad_hash = MonthlyRollupRequest(
        year=request.year,
        month=request.month,
        snapshot_root=request.snapshot_root,
        output_dir=request.output_dir,
        business_config_version=request.business_config_version,
        business_config_sha256=request.business_config_sha256,
        snapshots=(
            MonthlySnapshotReference(first.month, first.path, "f" * 64),
            request.snapshots[1],
        ),
    )
    with pytest.raises(MonthlyRollupError) as hash_error:
        render_monthly_rollup(bad_hash)
    assert hash_error.value.code == "MONTHLY_ROLLUP_SNAPSHOT_HASH_MISMATCH"

    outside = (tmp_path / "outside.json").resolve()
    outside.write_bytes(first.path.read_bytes())
    escaped = MonthlyRollupRequest(
        year=request.year,
        month=1,
        snapshot_root=request.snapshot_root,
        output_dir=request.output_dir,
        business_config_version=request.business_config_version,
        business_config_sha256=request.business_config_sha256,
        snapshots=(
            MonthlySnapshotReference(
                1,
                outside,
                hashlib.sha256(outside.read_bytes()).hexdigest(),
            ),
        ),
    )
    with pytest.raises(MonthlyRollupError) as path_error:
        render_monthly_rollup(escaped)
    assert path_error.value.code == "MONTHLY_ROLLUP_SNAPSHOT_INVALID"

    wrong_year = json.loads(first.path.read_text("utf-8"))
    wrong_year["period"]["year"] = 2025
    wrong_year["records"][0]["arrival_date"] = "2025-01-05"
    wrong_year["records"][0]["test_date"] = "2025-01-06"
    rebuilt = MonthlySnapshot.from_payload(wrong_year)
    first.path.write_bytes(canonical_snapshot_bytes(rebuilt))
    cross_year = MonthlyRollupRequest(
        year=request.year,
        month=request.month,
        snapshot_root=request.snapshot_root,
        output_dir=request.output_dir,
        business_config_version=request.business_config_version,
        business_config_sha256=request.business_config_sha256,
        snapshots=(
            MonthlySnapshotReference(
                1,
                first.path,
                hashlib.sha256(first.path.read_bytes()).hexdigest(),
            ),
            request.snapshots[1],
        ),
    )
    with pytest.raises(MonthlyRollupError) as year_error:
        render_monthly_rollup(cross_year)
    assert year_error.value.code == "MONTHLY_ROLLUP_YEAR_MISMATCH"


def test_artifact_failure_cleans_all_formal_outputs(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    request = _request(tmp_path, months=(1,))
    from moisture.application import monthly_rollup

    def fail_markdown(*_args: object, **_kwargs: object) -> Path:
        raise OSError("simulated report failure")

    monkeypatch.setattr(monthly_rollup, "write_analysis_markdown", fail_markdown)
    with pytest.raises(MonthlyRollupError) as error:
        render_monthly_rollup(request)
    assert error.value.code == "MONTHLY_ROLLUP_ARTIFACT_WRITE_FAILED"
    assert not any((request.output_dir / name).exists() for name in (
        "moisture_summary.xlsx",
        "analysis_report.md",
        "moisture_business_records.json",
    ))


def test_replacing_earlier_month_recomputes_from_supplied_snapshots_only(
    tmp_path: Path,
) -> None:
    original = _request(tmp_path / "original")
    replacement = _request(tmp_path / "replacement")
    replacement_snapshot = MonthlySnapshot.from_payload(
        _snapshot_payload(
            month=1,
            record_id="batch-01",
            day=5,
            mean="0.61",
        )
    )
    replacement_bytes = canonical_snapshot_bytes(replacement_snapshot)
    replacement.snapshots[0].path.write_bytes(replacement_bytes)
    replacement = MonthlyRollupRequest(
        year=replacement.year,
        month=replacement.month,
        snapshot_root=replacement.snapshot_root,
        output_dir=replacement.output_dir,
        business_config_version=replacement.business_config_version,
        business_config_sha256=replacement.business_config_sha256,
        snapshots=(
            MonthlySnapshotReference(
                month=1,
                path=replacement.snapshots[0].path,
                sha256=hashlib.sha256(replacement_bytes).hexdigest(),
            ),
            replacement.snapshots[1],
        ),
    )

    render_monthly_rollup(original)
    render_monthly_rollup(replacement)
    original_records = json.loads(
        (original.output_dir / "moisture_business_records.json").read_text("utf-8")
    )["records"]
    replacement_records = json.loads(
        (replacement.output_dir / "moisture_business_records.json").read_text("utf-8")
    )["records"]

    assert original_records[0]["statistics"]["mean"] == "0.46"
    assert replacement_records[0]["statistics"]["mean"] == "0.61"
    assert original_records[1] == replacement_records[1]


def test_rollup_source_has_no_pipeline_ai_knowledge_or_management_dependency() -> None:
    source = (
        MODULE_ROOT / "src" / "moisture" / "application" / "monthly_rollup.py"
    ).read_text(encoding="utf-8")
    forbidden = (
        "deterministic_pipeline",
        "confirmation_builder",
        "ai_report",
        "knowledge",
        "management",
        "httpx",
        "provider",
    )
    assert not any(name in source for name in forbidden)
