# ruff: noqa: E402, I001

from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from PySide6.QtTest import QSignalSpy

REPO_ROOT = Path(__file__).resolve().parents[4]
MODULE_ROOT = REPO_ROOT / "modules" / "moisture"
sys.path.insert(0, str(MODULE_ROOT / "src"))
sys.path.insert(0, str(MODULE_ROOT / "tests"))

from core.confirmations.service import ConfirmationService
from core.contracts.confirmation import ConfirmationAnswer
from core.contracts.enums import ModuleId
from core.knowledge.learning import LearningEvidenceRepository
from core.modules.registry import ModuleRegistry
from core.platform.paths import AppPaths
from core.storage.database import Database
from core.storage.repositories.confirmations import ConfirmationRepository
from core.storage.repositories.modules import ModuleRepository
from core.storage.repositories.monthly import MonthlyRepository
from core.storage.repositories.output_roots import OutputRootRepository
from core.storage.repositories.tasks import TaskRepository
from core.tasks.monthly_service import MonthlyRollupService
from core.tasks.service import TaskService
from moisture.adapters.history.summary_reader import StandardSummaryReader
from moisture.application.draft_orchestrator import run_draft
from moisture.contract.job_models import (
    JobPhase as ModuleJobPhase,
    JobRequest as ModuleJobRequest,
    ModuleId as ModuleContractId,
    RunMode as ModuleRunMode,
)
from moisture.infrastructure.hashing import sha256_file

pytest_plugins = ("fixtures.workbooks.workbook_factory",)


def _service(tmp_path: Path):
    paths = AppPaths.for_root(tmp_path / "app-data")
    database = Database.open(paths)
    database.migrate()
    task_repository = TaskRepository(database)
    confirmation_repository = ConfirmationRepository(database)
    registry = ModuleRegistry(
        REPO_ROOT / "modules",
        REPO_ROOT / "mock_modules",
        use_mocks=False,
        repository=ModuleRepository(database),
    )
    registry.scan()
    output_roots = OutputRootRepository(database, paths)
    output_roots.initialize_default()
    monthly_service = MonthlyRollupService(
        paths,
        registry,
        task_repository,
        confirmation_repository,
        MonthlyRepository(database),
        LearningEvidenceRepository(database),
        output_roots,
    )
    service = TaskService(
        paths=paths,
        registry=registry,
        task_repository=task_repository,
        confirmation_repository=confirmation_repository,
        monthly_rollup_service=monthly_service,
    )
    return paths, database, task_repository, confirmation_repository, service


def _base_summary(source: Path, output_dir: Path) -> Path:
    job = ModuleJobRequest(
        schema_version="1.0",
        job_id="a-h1-base-summary",
        module_id=ModuleContractId.MOISTURE,
        phase=ModuleJobPhase.DRAFT,
        run_mode=ModuleRunMode.SINGLE,
        input_files=(source,),
        output_dir=output_dir,
    )
    assert run_draft(job) == 0
    return output_dir / "moisture_summary.xlsx"


def test_real_task_service_snapshots_and_registers_formal_outputs(
    qtbot, tmp_path: Path, workbook_factory
) -> None:
    paths, database, task_repository, _, service = _service(tmp_path)
    try:
        source = workbook_factory(tmp_path / "source.xlsx")
        before = source.read_bytes()
        finished = QSignalSpy(service.finished)
        failed = QSignalSpy(service.failed)

        job_id = service.create_and_start(ModuleId.MOISTURE, [source])
        qtbot.waitUntil(lambda: finished.count() + failed.count() == 1, timeout=30_000)

        assert failed.count() == 0
        task = task_repository.get(job_id)
        assert task["status"] == "success"
        assert Path(str(task["result_path"])).is_file()
        outputs = task_repository.list_outputs(job_id)
        assert {Path(str(item["path"])).name for item in outputs} == {
            "moisture_summary.xlsx",
            "analysis_report.md",
        }
        assert all(Path(str(item["path"])).is_file() for item in outputs)
        snapshots = task_repository.list_files(job_id)
        assert len(snapshots) == 1
        assert Path(str(snapshots[0]["snapshot_path"])).read_bytes() == before
        assert source.read_bytes() == before
        assert service.runners == {}
    finally:
        service.shutdown()
        database.close()


def test_real_task_service_append_summary_conflict_finalize_and_isolation(
    qtbot, tmp_path: Path, workbook_factory
) -> None:
    paths, database, task_repository, confirmation_repository, service = _service(tmp_path)
    try:
        old_raw = workbook_factory(
            tmp_path / "old" / "same-batch.xlsx",
            moisture_values=(0.50,),
            mass_values=((100, 50, 0),),
            arrival_date="2026-04-08",
            test_date="2026-04-09",
        )
        base = _base_summary(old_raw, tmp_path / "base-output")
        conflicting = workbook_factory(
            tmp_path / "new" / "same-batch.xlsx",
            moisture_values=(0.52,),
            mass_values=((100, 48, 0),),
            arrival_date="2026-04-08",
            test_date="2026-04-09",
        )
        later = workbook_factory(
            tmp_path / "new" / "later-batch.xlsx",
            moisture_values=(0.53,),
            mass_values=((100, 47, 0),),
            arrival_date="2026-04-10",
            test_date="2026-04-11",
        )
        duplicate = tmp_path / "new" / "duplicate-source.xlsx"
        shutil.copyfile(old_raw, duplicate)
        corrupt = tmp_path / "new" / "corrupt.xlsx"
        corrupt.write_bytes(b"not an xlsx workbook")
        inputs = [base, conflicting, later, duplicate, corrupt]
        before = {path: sha256_file(path) for path in inputs}
        config = {
            "schema_version": "1.0",
            "operation": "append_summary",
            "inputs": [
                {"role": "base_summary", "input_index": 0},
                *[
                    {"role": "new_raw", "input_index": index}
                    for index in range(1, len(inputs))
                ],
            ],
        }
        needs_confirmation = QSignalSpy(service.needs_confirmation)
        finished = QSignalSpy(service.finished)
        failed = QSignalSpy(service.failed)

        job_id = service.create_and_start(
            ModuleId.MOISTURE,
            inputs,
            module_config=config,
            period_year=2026,
            period_month=4,
        )
        qtbot.waitUntil(
            lambda: needs_confirmation.count() + failed.count() == 1,
            timeout=60_000,
        )

        assert failed.count() == 0
        task_dir = paths.tasks / job_id
        draft = json.loads(
            (task_dir / "output-draft" / "result.json").read_text(encoding="utf-8")
        )
        assert draft["status"] == "needs_confirmation"
        assert [item["field"] for item in draft["confirmations"]] == [
            "history_record_conflict"
        ]
        assert draft["confirmations"][0]["options"] == ["keep_existing", "use_new"]
        assert draft["summary"]["skipped_count"] == 1
        assert draft["summary"]["failed_count"] == 1

        bundle = ConfirmationService(
            paths,
            confirmation_repository,
            task_repository,
        ).submit(
            job_id,
            [
                ConfirmationAnswer(
                    confirmation_id=draft["confirmations"][0]["confirmation_id"],
                    value="use_new",
                    apply_to_batch=False,
                    remember_mapping=False,
                )
            ],
        )
        service.finalize(job_id, bundle)
        qtbot.waitUntil(
            lambda: finished.count() + failed.count() == 1,
            timeout=60_000,
        )

        assert failed.count() == 0
        final = json.loads(
            (task_dir / "output-final" / "result.json").read_text(encoding="utf-8")
        )
        assert final["job_id"] == job_id
        assert final["status"] == "partial_success"
        assert final["confirmations"] == []
        assert final["service_requests"] == []
        assert final["summary"]["skipped_count"] == 1
        assert final["summary"]["failed_count"] == 1
        summary_path = task_dir / "output-final" / "moisture_summary.xlsx"
        appended = StandardSummaryReader().read(summary_path)
        assert [batch.record.record_id for batch in appended.batches] == [
            "same-batch",
            "later-batch",
        ]
        assert [batch.record.test_date.isoformat() for batch in appended.batches] == [
            "2026-04-09",
            "2026-04-11",
        ]
        assert str(appended.batches[0].statistics.mean) == "0.52"
        workbook = load_workbook(summary_path, read_only=True)
        try:
            assert workbook.sheetnames == [
                "竹丝含水率统计",
                "异常分析明细",
                "趋势分析",
                "来源追溯",
            ]
        finally:
            workbook.close()
        assert {path: sha256_file(path) for path in inputs} == before
        snapshots = task_repository.list_files(job_id)
        assert len(snapshots) == len(inputs)
        assert all(Path(str(item["snapshot_path"])).is_file() for item in snapshots)
        formal_outputs = {
            Path(str(item["path"])).name
            for item in task_repository.list_outputs(job_id)
        }
        assert formal_outputs == {
            "moisture_summary.xlsx",
            "analysis_report.md",
            "moisture_business_records.json",
        }
        assert Path(final["monthly_snapshot"]["path"]).name not in formal_outputs
        assert task_repository.get(job_id)["status"] == "partial_success"
        assert service.runners == {}
    finally:
        service.shutdown()
        database.close()


def test_real_task_service_year_confirmation_does_not_request_ai(
    qtbot, tmp_path: Path, workbook_factory
) -> None:
    paths, database, task_repository, confirmation_repository, service = _service(tmp_path)
    try:
        source = workbook_factory(
            tmp_path / "source.xlsx",
            moisture_values=(0.50,),
            mass_values=((100, 50, 0),),
            arrival_date="4月8日",
            test_date="4月9日",
        )
        before = source.read_bytes()
        needs_confirmation = QSignalSpy(service.needs_confirmation)
        finished = QSignalSpy(service.finished)
        failed = QSignalSpy(service.failed)

        job_id = service.create_and_start(
            ModuleId.MOISTURE,
            [source],
            period_year=2026,
            period_month=4,
        )
        qtbot.waitUntil(
            lambda: needs_confirmation.count() + failed.count() == 1,
            timeout=30_000,
        )
        assert failed.count() == 0
        task_dir = paths.tasks / job_id
        draft = json.loads((task_dir / "output-draft" / "result.json").read_text(encoding="utf-8"))
        assert [item["field"] for item in draft["confirmations"]] == ["default_year"]
        assert not any(item["service"] == "ai" for item in draft["service_requests"])

        confirmation = ConfirmationService(paths, confirmation_repository, task_repository).submit(
            job_id,
            [
                ConfirmationAnswer(
                    confirmation_id=draft["confirmations"][0]["confirmation_id"],
                    value="2026",
                    apply_to_batch=True,
                    remember_mapping=False,
                )
            ],
        )
        service.finalize(job_id, confirmation)
        qtbot.waitUntil(lambda: finished.count() + failed.count() == 1, timeout=30_000)

        assert failed.count() == 0
        final = json.loads((task_dir / "output-final" / "result.json").read_text(encoding="utf-8"))
        assert final["job_id"] == job_id
        assert final["status"] == "success"
        assert final["confirmations"] == []
        assert final["service_requests"] == []
        assert not any("ai_" in error.lower() for error in final["errors"])
        assert {Path(item["path"]).name for item in final["outputs"]} == {
            "moisture_summary.xlsx",
            "analysis_report.md",
        }
        formal_outputs = {
            Path(str(item["path"])).name
            for item in task_repository.list_outputs(job_id)
        }
        assert formal_outputs == {
            "moisture_summary.xlsx",
            "analysis_report.md",
            "moisture_business_records.json",
        }
        assert Path(final["monthly_snapshot"]["path"]).name not in formal_outputs
        assert source.read_bytes() == before
        assert service.runners == {}
    finally:
        service.shutdown()
        database.close()
