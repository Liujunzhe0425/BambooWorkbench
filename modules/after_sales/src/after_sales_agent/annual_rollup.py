from __future__ import annotations

import ast
import re
from collections.abc import Iterable
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from after_sales_agent.manual_review import (
    effective_after_sales_fee,
    effective_responsible_factory,
)

ISSUE_TYPES = ("性能", "外观", "加工")
TITLE_FILL = PatternFill("solid", fgColor="0F6B55")
HEADER_FILL = PatternFill("solid", fgColor="DCEFE7")
SUBHEADER_FILL = PatternFill("solid", fgColor="EEF7F3")
FEE_FILL = PatternFill("solid", fgColor="FFF29A")
TOTAL_FILL = PatternFill("solid", fgColor="DCE8F1")
FAILED_FILL = PatternFill("solid", fgColor="FDE8E7")
WHITE_FONT = Font(name="微软雅黑", color="FFFFFF", bold=True)
NORMAL_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="96A6A0")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_annual_rollup_workbook(
    records: list[dict[str, Any]], output_path: Path, rules: dict[str, Any] | None = None
) -> None:
    """Create the independent, reference-style annual rollup workbook.

    One ledger row is emitted for every input record, including failed records.
    Summary formulas read only rows marked ``纳入汇总=是`` and use the complaint /
    repair date year, so same-numbered months from different years never mix.
    """

    del rules  # The records have already been classified by the analysis engine.
    normalized = [_normalize_record(record, index) for index, record in enumerate(records, 1)]
    normalized.sort(key=_record_sort_key)

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    summary = workbook.create_sheet("汇总表", 0)
    groups = _group_records_by_year(normalized)
    detail_sheets: dict[int | str, tuple[str, int, int]] = {}
    for year, year_records in groups:
        sheet_name = _detail_sheet_name(year)
        detail = workbook.create_sheet(sheet_name)
        start_row, end_row = _write_reference_detail(detail, year_records, year)
        detail_sheets[year] = (sheet_name, start_row, end_row)
    _write_reference_summary(summary, groups, detail_sheets)
    audit = workbook.create_sheet("输入文件对账")
    _write_audit(audit, normalized)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _group_records_by_year(
    records: list[dict[str, Any]],
) -> list[tuple[int | str, list[dict[str, Any]]]]:
    years: list[int | str] = sorted(
        {record["year"] for record in records if isinstance(record["year"], int)}
    )
    if any(not isinstance(record["year"], int) for record in records):
        years.append("年份未注明")
    if not years:
        years = ["年份未注明"]
    return [(year, [record for record in records if record["year"] == year]) for year in years]


def _detail_sheet_name(year: int | str) -> str:
    if isinstance(year, int):
        return f"{str(year)[-2:]}年售后统计表"
    return "年份未注明售后统计表"


def _write_reference_detail(
    ws: Any, records: list[dict[str, Any]], year: int | str
) -> tuple[int, int]:
    """Write the detailed register in the exact visual structure of the supplied template."""

    title = f"{year}年售后统计表" if isinstance(year, int) else "年份未注明售后统计表"
    headers = [
        "序号",
        "受理日期",
        "订单号",
        "销售员",
        "产品规格",
        "客户地点",
        "发货时间",
        "品牌",
        "投诉问题",
        "原因分析",
        "处理方案",
        "结案时间",
        "售后费用",
        "责任工厂",
        "投诉类型",
        "投诉类型",
        "是否结案",
        "结案时间",
    ]
    black = Side(style="thin", color="000000")
    border = Border(left=black, right=black, top=black, bottom=black)
    body_font = Font(name="宋体", size=10, color="000000")

    ws.merge_cells("A1:S1")
    ws["A1"] = title
    ws["A1"].font = Font(name="宋体", size=16, color="000000")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for column, header in enumerate(headers, 2):
        cell = ws.cell(2, column, header)
        cell.font = Font(name="宋体", size=11, color="000000")
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 31

    data_start = 3
    for sequence, record in enumerate(records, 1):
        row = data_start + sequence - 1
        complaint_date = record["complaint_date"]
        values = [
            sequence,
            complaint_date,
            record["order_no"],
            record["salesperson"],
            " ".join(part for part in (record["product"], record["spec"]) if part),
            record["location"],
            _delivery_date_value(record["delivery_date"]),
            record["brand"],
            record["feedback"] or record["application"],
            record["issue_reason"],
            record["solution"],
            None,
            record["after_sales_fee"],
            _responsibility_text(record),
            _complaint_type_label(record),
            record["primary_issue_type"] if record["primary_issue_type"] in ISSUE_TYPES else "性能",
            record["closed"],
            None,
        ]
        for column, value in enumerate(values, 2):
            cell = ws.cell(row, column, value)
            cell.font = body_font
            cell.border = border
            centered_columns = {2, 3, 5, 7, 8, 9, 13, 14, 16, 17, 18, 19}
            cell.alignment = Alignment(
                horizontal="center" if column in centered_columns else "left",
                vertical="center",
                wrap_text=column in {7, 10, 11, 12, 15, 16},
            )
        if isinstance(complaint_date, date):
            ws.cell(row, 3).number_format = "yyyy-m-d"
        delivery_date = values[6]
        if isinstance(delivery_date, date):
            ws.cell(row, 8).number_format = "yyyy-m-d"
        ws.cell(row, 14).number_format = "#,##0.00"
        ws.cell(row, 20, record["month"])
        ws.cell(row, 21, record["source_file"])
        ws.cell(row, 22, record["included"])
        longest = max((len(_text(value)) for value in values[8:15]), default=0)
        ws.row_dimensions[row].height = min(110, max(52, 36 + longest / 22 * 9))

    data_end = max(data_start, data_start + len(records) - 1)
    if not records:
        for column in range(2, 20):
            cell = ws.cell(data_start, column)
            cell.border = border
        ws.cell(data_start, 2, "本年度没有可写入记录")

    widths = {
        "A": 3,
        "B": 6,
        "C": 12,
        "D": 24,
        "E": 10,
        "F": 20,
        "G": 14,
        "H": 13,
        "I": 12,
        "J": 43,
        "K": 48,
        "L": 48,
        "M": 12,
        "N": 12,
        "O": 18,
        "P": 15,
        "Q": 10,
        "R": 10,
        "S": 12,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for column in ("T", "U", "V"):
        ws.column_dimensions[column].hidden = True
    ws["T2"] = "统计月份"
    ws["U2"] = "来源文件"
    ws["V2"] = "纳入汇总"

    category_validation = DataValidation(type="list", formula1='"性能,外观,加工"')
    closed_validation = DataValidation(type="list", formula1='"是,否"')
    ws.add_data_validation(category_validation)
    ws.add_data_validation(closed_validation)
    category_validation.add(f"Q{data_start}:Q{data_end}")
    closed_validation.add(f"R{data_start}:R{data_end}")

    ws.auto_filter.ref = f"B2:S{data_end}"
    ws.freeze_panes = "J3"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:2"
    return data_start, data_end


def _write_reference_summary(
    ws: Any,
    groups: list[tuple[int | str, list[dict[str, Any]]]],
    detail_sheets: dict[int | str, tuple[str, int, int]],
) -> None:
    """Write one reference-style brand/month table per complaint year."""

    black = Side(style="thin", color="000000")
    border = Border(left=black, right=black, top=black, bottom=black)
    body_font = Font(name="宋体", size=10, color="000000")
    row = 1
    for block_index, (year, records) in enumerate(groups):
        sheet_name, data_start, data_end = detail_sheets[year]
        title = f"{year}年售后统计表" if isinstance(year, int) else "年份未注明售后统计表"
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=16)
        title_cell = ws.cell(row, 1, title)
        title_cell.font = Font(name="宋体", size=16, color="000000")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 28

        header_top = row + 1
        header_month = row + 2
        for column in (1, 2, 3, 16):
            ws.merge_cells(
                start_row=header_top,
                start_column=column,
                end_row=header_month,
                end_column=column,
            )
        ws.merge_cells(
            start_row=header_top, start_column=4, end_row=header_top, end_column=15
        )
        ws.cell(header_top, 1, "序号")
        ws.cell(header_top, 2, "品牌")
        ws.cell(header_top, 3, "类型")
        ws.cell(header_top, 4, "月份")
        ws.cell(header_top, 16, "合计")
        for month in range(1, 13):
            ws.cell(header_month, 3 + month, month)
        for current_row in (header_top, header_month):
            for column in range(1, 17):
                cell = ws.cell(current_row, column)
                cell.font = body_font
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")

        included = [record for record in records if record["included"] == "是"]
        brands = sorted({record["brand"] for record in included}) or ["未知"]
        data_row = row + 3
        for sequence, brand in enumerate(brands, 1):
            brand_start = data_row
            ws.merge_cells(
                start_row=brand_start,
                start_column=1,
                end_row=brand_start + 3,
                end_column=1,
            )
            ws.merge_cells(
                start_row=brand_start,
                start_column=2,
                end_row=brand_start + 3,
                end_column=2,
            )
            ws.cell(brand_start, 1, sequence)
            ws.cell(brand_start, 2, brand)
            for offset, issue_type in enumerate(ISSUE_TYPES):
                current = brand_start + offset
                ws.cell(current, 3, issue_type)
                for month in range(1, 13):
                    ws.cell(
                        current,
                        3 + month,
                        _reference_case_formula(
                            sheet_name,
                            data_start,
                            data_end,
                            brand_start,
                            current,
                            header_month,
                            month,
                        ),
                    )
                ws.cell(current, 16, f"=SUM(D{current}:O{current})")

            fee_row = brand_start + 3
            ws.cell(fee_row, 3, "费用")
            for month in range(1, 13):
                ws.cell(
                    fee_row,
                    3 + month,
                    _reference_fee_formula(
                        sheet_name,
                        data_start,
                        data_end,
                        brand_start,
                        header_month,
                        month,
                    ),
                )
            ws.cell(fee_row, 16, f"=SUM(D{fee_row}:O{fee_row})")
            for column in range(3, 17):
                ws.cell(fee_row, column).fill = PatternFill("solid", fgColor="FFF200")
                if column >= 4:
                    ws.cell(fee_row, column).number_format = "#,##0.00"
            data_row += 4

        total_rows = (("售后小计", "count"), ("售后费用", "fee"), ("关闭数", "closed"))
        for label, kind in total_rows:
            ws.merge_cells(
                start_row=data_row, start_column=1, end_row=data_row, end_column=2
            )
            ws.cell(data_row, 1, label)
            for month in range(1, 13):
                ws.cell(
                    data_row,
                    3 + month,
                    _reference_total_formula(
                        sheet_name, data_start, data_end, header_month, month, kind
                    ),
                )
            ws.cell(data_row, 16, f"=SUM(D{data_row}:O{data_row})")
            if kind == "fee":
                for column in range(4, 17):
                    ws.cell(data_row, column).number_format = "#,##0.00"
            data_row += 1

        table_end = data_row - 1
        for current_row in range(header_top, table_end + 1):
            for column in range(1, 17):
                cell = ws.cell(current_row, column)
                cell.font = body_font
                cell.border = border
                if column not in {2, 3}:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
        note_row = data_row
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=16)
        ws.cell(note_row, 1, "备注：费用不明确时保持空白并进入人工复核；已确认的费用会计入统计。")
        ws.cell(note_row, 1).font = Font(name="宋体", size=9, color="000000")
        ws.cell(note_row, 1).border = border
        ws.cell(note_row, 1).alignment = Alignment(horizontal="left", vertical="center")
        row = note_row + 3

        if block_index == 0:
            reference = [
                ["性能", "外观", "加工"],
                ["开裂", "色差", "尺寸加工"],
                ["发霉变形", "掉漆划痕", "孔位槽位"],
                ["分层强度", "表面观感", "图纸工艺"],
            ]
            for row_offset, values in enumerate(reference, 2):
                for column_offset, value in enumerate(values, 19):
                    cell = ws.cell(row_offset, column_offset, value)
                    cell.font = body_font
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {"A": 8, "B": 14, "C": 12, "P": 14, "S": 14, "T": 14, "U": 14}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for column in range(4, 16):
        ws.column_dimensions[get_column_letter(column)].width = 10
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "D4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _reference_case_formula(
    sheet_name: str,
    start: int,
    end: int,
    brand_row: int,
    issue_row: int,
    month_row: int,
    month: int,
) -> str:
    month_column = get_column_letter(3 + month)
    return (
        f'=COUNTIFS(\'{sheet_name}\'!$I${start}:$I${end},$B{brand_row},'
        f"'{sheet_name}'!$Q${start}:$Q${end},$C{issue_row},"
        f"'{sheet_name}'!$T${start}:$T${end},{month_column}${month_row},"
        f"'{sheet_name}'!$V${start}:$V${end},\"是\")"
    )


def _reference_fee_formula(
    sheet_name: str,
    start: int,
    end: int,
    brand_row: int,
    month_row: int,
    month: int,
) -> str:
    month_column = get_column_letter(3 + month)
    return (
        f'=SUMIFS(\'{sheet_name}\'!$N${start}:$N${end},'
        f"'{sheet_name}'!$I${start}:$I${end},$B{brand_row},"
        f"'{sheet_name}'!$T${start}:$T${end},{month_column}${month_row},"
        f"'{sheet_name}'!$V${start}:$V${end},\"是\")"
    )


def _reference_total_formula(
    sheet_name: str,
    start: int,
    end: int,
    month_row: int,
    month: int,
    kind: str,
) -> str:
    month_column = get_column_letter(3 + month)
    month_criteria = (
        f"'{sheet_name}'!$T${start}:$T${end},{month_column}${month_row},"
        f"'{sheet_name}'!$V${start}:$V${end},\"是\""
    )
    if kind == "fee":
        return f"=SUMIFS('{sheet_name}'!$N${start}:$N${end},{month_criteria})"
    if kind == "closed":
        return (
            f"=COUNTIFS('{sheet_name}'!$R${start}:$R${end},\"是\",{month_criteria})"
        )
    return f"=COUNTIFS({month_criteria})"


def _delivery_date_value(value: Any) -> date | str:
    text = _text(value)
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        serial = int(float(text))
        if 30000 <= serial <= 80000:
            return date.fromordinal(date(1899, 12, 30).toordinal() + serial)
    match = re.fullmatch(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})(?:日)?", text)
    if match:
        try:
            return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
        except ValueError:
            pass
    return text


def _responsibility_text(record: dict[str, Any]) -> str:
    return effective_responsible_factory(record)


def _complaint_type_label(record: dict[str, Any]) -> str:
    text = " ".join((record["feedback"], record["issue_reason"]))
    keyword_groups = {
        "性能": ("开裂", "裂纹", "发霉", "霉变", "分层", "变形", "收缩", "膨胀", "破裂"),
        "外观": ("色差", "颜色", "掉漆", "划痕", "漆面", "表面", "观感"),
        "加工": ("尺寸", "孔位", "开槽", "切割", "涂胶", "图纸", "加工", "工艺"),
    }
    primary = (
        record["primary_issue_type"]
        if record["primary_issue_type"] in ISSUE_TYPES
        else "性能"
    )
    matches = [keyword for keyword in keyword_groups[primary] if keyword in text]
    return "、".join(matches[:2]) if matches else f"{primary}问题"


def _normalize_record(record: dict[str, Any], fallback_index: int) -> dict[str, Any]:
    year, month, day, typed_date, display_date = _date_parts(record)
    issue_types = _issue_types(record)
    primary = str(record.get("primary_issue_type") or "").strip()
    primary = (
        primary
        if primary in ISSUE_TYPES
        else next((item for item in ISSUE_TYPES if item in issue_types), "其他")
    )
    included = bool(record.get("included_in_summary", True))
    status = _resolution_status(record)
    analysis_status = _analysis_status(record, included)
    source = Path(str(record.get("source_file") or f"输入文件{fallback_index}")).name
    product = _text(record.get("product"))
    spec = _text(record.get("spec"))

    return {
        "input_index": _int(record.get("input_index")) or fallback_index,
        "complaint_date": typed_date or display_date,
        "complaint_date_display": display_date,
        "year": year if year is not None else "年份未注明",
        "month": month,
        "day": day,
        "order_no": _text(record.get("order_no")),
        "salesperson": _text(record.get("salesperson")),
        "product": product,
        "spec": spec,
        "customer": _text(record.get("customer")),
        "location": _text(record.get("location")),
        "delivery_date": _text(record.get("delivery_date")),
        "brand": _text(record.get("brand")) or "未知",
        "feedback": _text(record.get("feedback")),
        "application": _text(record.get("application")),
        "investigation": _text(record.get("investigation")),
        "solution": _text(record.get("selected_solution")) or _text(record.get("solution")),
        "responsibility": _text(record.get("responsibility")),
        "responsible_factory": effective_responsible_factory(record) or "未知",
        "responsibility_reviewed": bool(record.get("responsibility_reviewed")),
        "issue_reason": _issue_reason(record),
        "after_sales_fee": _after_sales_fee(record),
        "enterprise_cost": _amount(record, "enterprise_cost_amount", "fee_amount", "cost"),
        "internal_cost": _amount(record, "internal_accountability_amount"),
        "external_cost": _amount(record, "external_borne_amount"),
        "specific_issue": _specific_issue(record),
        "issue_types": issue_types,
        "primary_issue_type": primary,
        "performance": 1 if "性能" in issue_types else 0,
        "appearance": 1 if "外观" in issue_types else 0,
        "processing": 1 if "加工" in issue_types else 0,
        "resolution_status": status,
        "closed": "是" if status == "已解决" else "否",
        "source_file": source,
        "analysis_status": analysis_status,
        "analysis_error": _text(record.get("analysis_error")),
        "included": "是" if included else "否",
    }


def _write_ledger(ws: Any, records: list[dict[str, Any]]) -> tuple[int, int]:
    headers = [
        "序号",
        "受理/报修日期",
        "统计年份",
        "统计月份",
        "订单号",
        "销售员",
        "产品",
        "规格",
        "客户/项目",
        "客户地点",
        "发货时间",
        "品牌",
        "客户反馈内容",
        "方案申请明细",
        "售后负责人调查情况",
        "售后解决方案",
        "责任落实",
        "售后问题分析",
        "最终售后费用",
        "内部责任扣款",
        "外部承担金额",
        "具体问题",
        "原因分类",
        "性能",
        "外观",
        "加工",
        "处理状态",
        "是否结案",
        "来源文件",
        "分析状态",
        "纳入汇总",
    ]
    last_col = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, "售后案件完整台账（每个输入文件固定保留一行）")
    ws.cell(1, 1).fill = TITLE_FILL
    ws.cell(1, 1).font = Font(name="微软雅黑", size=18, bold=True, color="FFFFFF")
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(
        2,
        1,
        "按提出质量问题/报修日期排序；日期不明的文件置于末尾。费用未知保持空白，不按 0 元统计。",
    )
    ws.cell(2, 1).fill = SUBHEADER_FILL
    ws.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 25

    header_row = 4
    for column, title in enumerate(headers, 1):
        cell = ws.cell(header_row, column, title)
        cell.fill = HEADER_FILL
        cell.font = Font(name="微软雅黑", bold=True, color="173B31")
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 34

    for row_number, record in enumerate(records, header_row + 1):
        values = [
            record["input_index"],
            record["complaint_date"],
            record["year"],
            record["month"],
            record["order_no"],
            record["salesperson"],
            record["product"],
            record["spec"],
            record["customer"],
            record["location"],
            record["delivery_date"],
            record["brand"],
            record["feedback"],
            record["application"],
            record["investigation"],
            record["solution"],
            record["responsibility"],
            record["issue_reason"],
            record["after_sales_fee"],
            record["internal_cost"],
            record["external_cost"],
            record["specific_issue"],
            "、".join(record["issue_types"]) or "其他",
            record["performance"],
            record["appearance"],
            record["processing"],
            record["resolution_status"],
            record["closed"],
            record["source_file"],
            record["analysis_status"],
            record["included"],
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_number, column, value)
            cell.font = NORMAL_FONT
            cell.border = GRID
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if isinstance(record["complaint_date"], date):
            ws.cell(row_number, 2).number_format = "yyyy-mm-dd"
        for column in (19, 20, 21):
            ws.cell(row_number, column).number_format = '#,##0.00" 元"'
        for column in (24, 25, 26, 28, 31):
            ws.cell(row_number, column).alignment = Alignment(
                horizontal="center", vertical="center"
            )
        if record["included"] == "否":
            for column in range(1, last_col + 1):
                ws.cell(row_number, column).fill = FAILED_FILL
        longest = max((len(_text(value)) for value in values[12:23]), default=0)
        ws.row_dimensions[row_number].height = min(120, max(34, 18 + longest / 18 * 14))

    data_end = max(header_row + 1, header_row + len(records))
    if not records:
        ws.cell(header_row + 1, 1, "本批没有可写入的输入记录")
        ws.merge_cells(
            start_row=header_row + 1, start_column=1, end_row=header_row + 1, end_column=last_col
        )
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{data_end}"
    ws.freeze_panes = "M5"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    widths = {
        "A": 8,
        "B": 15,
        "C": 13,
        "D": 10,
        "E": 18,
        "F": 13,
        "G": 18,
        "H": 18,
        "I": 24,
        "J": 18,
        "K": 15,
        "L": 14,
        "M": 38,
        "N": 38,
        "O": 42,
        "P": 42,
        "Q": 34,
        "R": 42,
        "S": 16,
        "T": 16,
        "U": 16,
        "V": 34,
        "W": 16,
        "X": 9,
        "Y": 9,
        "Z": 9,
        "AA": 13,
        "AB": 11,
        "AC": 34,
        "AD": 15,
        "AE": 11,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    return header_row + 1, data_end


def _write_summary(ws: Any, records: list[dict[str, Any]], ledger_bounds: tuple[int, int]) -> None:
    data_start, data_end = ledger_bounds
    ws.merge_cells("A1:Q1")
    ws["A1"] = "售后年度统计汇总表"
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = Font(name="微软雅黑", size=19, bold=True, color="FFFFFF")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells("A2:Q2")
    ws["A2"] = (
        "按报修年份、品牌和月份统计；问题分类由智能分析结果生成，一案可同时属于性能、外观和加工。"
    )
    ws["A2"].fill = SUBHEADER_FILL
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

    headers = ["序号", "年份", "品牌", "类型", *[f"{month}月" for month in range(1, 13)], "合计"]
    for column, title in enumerate(headers, 1):
        cell = ws.cell(4, column, title)
        cell.fill = HEADER_FILL
        cell.font = Font(name="微软雅黑", bold=True, color="173B31")
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 28

    included = [record for record in records if record["included"] == "是"]
    years: list[int | str] = sorted(
        {record["year"] for record in included if isinstance(record["year"], int)}
    )
    if any(record["year"] == "年份未注明" for record in included):
        years.append("年份未注明")
    if not years:
        years = ["年份未注明"]

    row = 5
    sequence = 1
    issue_flag_columns = {"性能": "X", "外观": "Y", "加工": "Z"}
    for year in years:
        year_records = [record for record in included if record["year"] == year]
        brands = sorted({record["brand"] for record in year_records}) or ["未知"]
        for brand in brands:
            first_brand_row = row
            for issue_type in ISSUE_TYPES:
                ws.cell(row, 1, sequence if row == first_brand_row else "")
                ws.cell(row, 2, year)
                ws.cell(row, 3, brand)
                ws.cell(row, 4, issue_type)
                for month in range(1, 13):
                    ws.cell(
                        row,
                        4 + month,
                        _case_formula(
                            data_start, data_end, row, month, issue_flag_columns[issue_type]
                        ),
                    )
                ws.cell(
                    row,
                    17,
                    _case_formula(data_start, data_end, row, None, issue_flag_columns[issue_type]),
                )
                row += 1
            ws.cell(row, 1, "")
            ws.cell(row, 2, year)
            ws.cell(row, 3, brand)
            ws.cell(row, 4, "费用")
            for month in range(1, 13):
                ws.cell(row, 4 + month, _cost_formula(data_start, data_end, row, month))
            ws.cell(row, 17, _cost_formula(data_start, data_end, row, None))
            for column in range(1, 18):
                ws.cell(row, column).fill = FEE_FILL
            for column in range(5, 18):
                ws.cell(row, column).number_format = '#,##0.00" 元"'
            row += 1
            sequence += 1

        for label, status in (
            ("案件总数", None),
            ("售后费用", "费用"),
            ("已解决数量", "已解决"),
            ("处理中数量", "处理中"),
            ("未解决数量", "未解决"),
        ):
            ws.cell(row, 2, year)
            ws.cell(row, 3, "年度合计")
            ws.cell(row, 4, label)
            for month in range(1, 13):
                ws.cell(
                    row, 4 + month, _year_total_formula(data_start, data_end, row, month, status)
                )
            ws.cell(row, 17, _year_total_formula(data_start, data_end, row, None, status))
            for column in range(1, 18):
                ws.cell(row, column).fill = FEE_FILL if status == "费用" else TOTAL_FILL
            if status == "费用":
                for column in range(5, 18):
                    ws.cell(row, column).number_format = '#,##0.00" 元"'
            row += 1
        row += 1

    table_end = max(5, row - 1)
    for row_number in range(5, table_end + 1):
        for column in range(1, 18):
            cell = ws.cell(row_number, column)
            cell.font = NORMAL_FONT
            cell.border = GRID
            if column not in (3, 4):
                cell.alignment = Alignment(horizontal="center", vertical="center")

    _write_category_reference(ws, records, data_start, data_end)
    ws.freeze_panes = "E5"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A4:Q{table_end}"
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 15
    for column in range(5, 18):
        ws.column_dimensions[get_column_letter(column)].width = 12
    ws.column_dimensions["S"].width = 14
    ws.column_dimensions["T"].width = 13
    ws.column_dimensions["U"].width = 14
    ws.column_dimensions["V"].width = 48
    ws.column_dimensions["W"].width = 30
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _write_category_reference(
    ws: Any, records: list[dict[str, Any]], data_start: int, data_end: int
) -> None:
    ws.merge_cells("S1:W1")
    ws["S1"] = "售后原因分类概览"
    ws["S1"].fill = TITLE_FILL
    ws["S1"].font = WHITE_FONT
    ws["S1"].alignment = Alignment(horizontal="center", vertical="center")
    for column, title in enumerate(
        ("年份", "原因分类", "案件数", "代表性具体原因", "涉及品牌"), 19
    ):
        cell = ws.cell(4, column, title)
        cell.fill = HEADER_FILL
        cell.font = Font(name="微软雅黑", bold=True, color="173B31")
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    included = [record for record in records if record["included"] == "是"]
    years: list[int | str] = sorted(
        {record["year"] for record in included if isinstance(record["year"], int)}
    )
    if any(record["year"] == "年份未注明" for record in included):
        years.append("年份未注明")
    row = 5
    flag_columns = {"性能": "X", "外观": "Y", "加工": "Z"}
    for year in years:
        for issue_type in ISSUE_TYPES:
            matching = [r for r in included if r["year"] == year and issue_type in r["issue_types"]]
            ws.cell(row, 19, year)
            ws.cell(row, 20, issue_type)
            ws.cell(
                row,
                21,
                f"=COUNTIFS('售后台账'!$C${data_start}:$C${data_end},S{row},'售后台账'!${flag_columns[issue_type]}${data_start}:${flag_columns[issue_type]}${data_end},1,'售后台账'!$AE${data_start}:$AE${data_end},\"是\")",
            )
            ws.cell(row, 22, _representative_text(r["issue_reason"] for r in matching))
            ws.cell(row, 23, "、".join(sorted({r["brand"] for r in matching})) or "—")
            for column in range(19, 24):
                cell = ws.cell(row, column)
                cell.border = GRID
                cell.font = NORMAL_FONT
                cell.alignment = Alignment(
                    vertical="top", wrap_text=True, horizontal="center" if column < 22 else "left"
                )
            ws.row_dimensions[row].height = 42
            row += 1


def _write_audit(ws: Any, records: list[dict[str, Any]]) -> None:
    headers = [
        "输入序号",
        "来源文件",
        "分析状态",
        "失败原因",
        "纳入汇总",
        "报修日期",
        "品牌",
        "原因分类",
        "售后问题分析",
        "最终售后费用",
    ]
    ws.merge_cells("A1:J1")
    ws["A1"] = "输入文件完整性对账"
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = Font(name="微软雅黑", size=17, bold=True, color="FFFFFF")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:J2")
    ws["A2"] = f"本批输入记录 {len(records)} 份；下表应与上传文件逐一对应，失败文件不会被静默删除。"
    ws["A2"].fill = SUBHEADER_FILL
    for column, title in enumerate(headers, 1):
        cell = ws.cell(4, column, title)
        cell.fill = HEADER_FILL
        cell.font = Font(name="微软雅黑", bold=True, color="173B31")
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number, record in enumerate(sorted(records, key=lambda item: item["input_index"]), 5):
        values = [
            record["input_index"],
            record["source_file"],
            record["analysis_status"],
            record["analysis_error"],
            record["included"],
            record["complaint_date"],
            record["brand"],
            "、".join(record["issue_types"]) or "其他",
            record["issue_reason"],
            record["after_sales_fee"],
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_number, column, value)
            cell.font = NORMAL_FONT
            cell.border = GRID
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if isinstance(record["complaint_date"], date):
            ws.cell(row_number, 6).number_format = "yyyy-mm-dd"
        ws.cell(row_number, 10).number_format = '#,##0.00" 元"'
        if record["included"] == "否":
            for column in range(1, 11):
                ws.cell(row_number, column).fill = FAILED_FILL
        ws.row_dimensions[row_number].height = 38
    end_row = max(5, len(records) + 4)
    ws.auto_filter.ref = f"A4:J{end_row}"
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    widths = (10, 42, 16, 42, 12, 16, 18, 18, 48, 18)
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width


def _case_formula(start: int, end: int, row: int, month: int | None, flag_column: str) -> str:
    criteria = [
        f"'售后台账'!$C${start}:$C${end}",
        f"$B{row}",
        f"'售后台账'!$L${start}:$L${end}",
        f"$C{row}",
        f"'售后台账'!${flag_column}${start}:${flag_column}${end}",
        "1",
        f"'售后台账'!$AE${start}:$AE${end}",
        '"是"',
    ]
    if month is not None:
        criteria.extend([f"'售后台账'!$D${start}:$D${end}", f"{month}"])
    return "=COUNTIFS(" + ",".join(criteria) + ")"


def _cost_formula(start: int, end: int, row: int, month: int | None) -> str:
    arguments = [
        f"'售后台账'!$S${start}:$S${end}",
        f"'售后台账'!$C${start}:$C${end}",
        f"$B{row}",
        f"'售后台账'!$L${start}:$L${end}",
        f"$C{row}",
        f"'售后台账'!$AE${start}:$AE${end}",
        '"是"',
    ]
    if month is not None:
        arguments.extend([f"'售后台账'!$D${start}:$D${end}", f"{month}"])
    return "=SUMIFS(" + ",".join(arguments) + ")"


def _year_total_formula(
    start: int, end: int, row: int, month: int | None, status: str | None
) -> str:
    if status == "费用":
        arguments = [
            f"'售后台账'!$S${start}:$S${end}",
            f"'售后台账'!$C${start}:$C${end}",
            f"$B{row}",
            f"'售后台账'!$AE${start}:$AE${end}",
            '"是"',
        ]
        if month is not None:
            arguments.extend([f"'售后台账'!$D${start}:$D${end}", f"{month}"])
        return "=SUMIFS(" + ",".join(arguments) + ")"
    criteria = [
        f"'售后台账'!$C${start}:$C${end}",
        f"$B{row}",
        f"'售后台账'!$AE${start}:$AE${end}",
        '"是"',
    ]
    if status:
        criteria.extend([f"'售后台账'!$AA${start}:$AA${end}", f'"{status}"'])
    if month is not None:
        criteria.extend([f"'售后台账'!$D${start}:$D${end}", f"{month}"])
    return "=COUNTIFS(" + ",".join(criteria) + ")"


def _date_parts(
    record: dict[str, Any],
) -> tuple[int | None, int | None, int | None, date | None, str]:
    raw = record.get("complaint_date") or record.get("complaint_date_raw") or ""
    if isinstance(raw, datetime):
        raw = raw.date()
    if isinstance(raw, date):
        return raw.year, raw.month, raw.day, raw, raw.isoformat()
    text = _text(raw)
    if text:
        iso = re.search(r"(?<!\d)(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
        if iso:
            try:
                parsed = date(int(iso.group(1)), int(iso.group(2)), int(iso.group(3)))
                return parsed.year, parsed.month, parsed.day, parsed, text
            except ValueError:
                pass
        month_day = re.search(r"(?<!\d)(\d{1,2})\s*月\s*(\d{1,2})?", text)
        if month_day:
            month = _int(month_day.group(1))
            day = _int(month_day.group(2)) if month_day.group(2) else None
            if month and 1 <= month <= 12:
                year = _int(record.get("year"))
                if year and day:
                    try:
                        parsed = date(year, month, day)
                        return year, month, day, parsed, text
                    except ValueError:
                        pass
                return year, month, day, None, text
    year = _int(record.get("year"))
    month = _int(record.get("month"))
    if month is not None and not 1 <= month <= 12:
        month = None
    return year, month, None, None, text


def _record_sort_key(record: dict[str, Any]) -> tuple[Any, ...]:
    year = record["year"] if isinstance(record["year"], int) else 9999
    month = record["month"] if isinstance(record["month"], int) else 99
    day = record["day"] if isinstance(record["day"], int) else 32
    return year, month, day, record["input_index"]


def _issue_types(record: dict[str, Any]) -> list[str]:
    raw = record.get("issue_types")
    values: list[str] = []
    if isinstance(raw, (list, tuple, set)):
        values = [_text(item) for item in raw]
    elif isinstance(raw, str) and raw.strip():
        try:
            parsed = ast.literal_eval(raw)
            if isinstance(parsed, (list, tuple, set)):
                values = [_text(item) for item in parsed]
            else:
                values = re.split(r"[、,，/|;；\s]+", raw)
        except (ValueError, SyntaxError):
            values = re.split(r"[、,，/|;；\s]+", raw)
    primary = _text(record.get("primary_issue_type"))
    if primary:
        values.append(primary)
    result = [
        issue_type for issue_type in ISSUE_TYPES if any(issue_type in value for value in values)
    ]
    if result:
        return result
    text = " ".join(
        _text(record.get(key))
        for key in ("issue_reason_summary", "feedback", "application", "investigation", "solution")
    )
    keyword_groups = {
        "性能": (
            "开裂",
            "裂缝",
            "变形",
            "起翘",
            "发霉",
            "霉变",
            "分层",
            "脱层",
            "强度",
            "含水率",
            "松动",
            "异响",
            "塌陷",
            "腐朽",
            "虫蛀",
        ),
        "外观": (
            "色差",
            "颜色",
            "漆面",
            "掉漆",
            "划痕",
            "污渍",
            "发白",
            "黄变",
            "表面",
            "光泽",
            "露底",
        ),
        "加工": (
            "尺寸",
            "切割",
            "开槽",
            "孔位",
            "拼接",
            "图纸",
            "加工",
            "工艺",
            "倒角",
            "长度",
            "宽度",
            "厚度",
            "锯切",
            "涂胶",
            "胶合",
        ),
    }
    return [
        category
        for category in ISSUE_TYPES
        if any(word in text for word in keyword_groups[category])
    ]


def _issue_reason(record: dict[str, Any]) -> str:
    value = _text(record.get("issue_reason_summary"))
    if value:
        return value
    for key in ("investigation", "application", "feedback"):
        value = _text(record.get(key))
        if value:
            return value
    error = _text(record.get("analysis_error"))
    return f"文件分析失败：{error}" if error else "原表未写明售后原因"


def _specific_issue(record: dict[str, Any]) -> str:
    for key in ("feedback", "issue_reason_summary", "investigation"):
        value = _text(record.get(key))
        if value:
            return value
    return "原表未写明具体问题"


def _resolution_status(record: dict[str, Any]) -> str:
    status = _text(record.get("resolution_status"))
    if status in {"已解决", "处理中", "未解决"}:
        return status
    case_status = _text(record.get("case_status"))
    if case_status in {"已关闭", "已解决"} or record.get("closed") is True:
        return "已解决"
    if case_status == "处理中":
        return "处理中"
    return "未解决"


def _analysis_status(record: dict[str, Any], included: bool) -> str:
    if not included or _text(record.get("analysis_error")):
        return "分析失败"
    source = _text(record.get("issue_analysis_source"))
    status = _text(record.get("analysis_status"))
    if source == "api" or "智能" in status or "AI" in status:
        return "智能分析完成"
    return "规则分析完成"


def _amount(record: dict[str, Any], *keys: str) -> float | None:
    for key in keys:
        if key not in record or record.get(key) in (None, ""):
            continue
        value = record.get(key)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return round(float(value), 2)
        match = re.search(r"-?\d+(?:\.\d+)?", _text(value).replace(",", ""))
        if match:
            return round(float(match.group()), 2)
    return None


def _after_sales_fee(record: dict[str, Any]) -> float | None:
    return effective_after_sales_fee(record)


def _representative_text(values: Iterable[str]) -> str:
    unique: list[str] = []
    for value in values:
        cleaned = _text(value)
        if cleaned and cleaned not in unique:
            unique.append(cleaned)
        if len(unique) == 3:
            break
    return "；".join(unique) if unique else "—"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return "、".join(_text(item) for item in value if _text(item))
    if isinstance(value, dict):
        return "；".join(f"{_text(key)}：{_text(item)}" for key, item in value.items())
    return re.sub(r"\s+", " ", str(value)).strip()


def _int(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None
