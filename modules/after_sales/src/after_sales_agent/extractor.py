from __future__ import annotations

import hashlib
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.datetime import from_excel

from after_sales_agent.llm_analyzer import analyze_issue_with_api, api_is_configured
from after_sales_agent.normalizer import build_date_warnings, determine_fee_breakdown

MONEY_RE = re.compile(r"(?<!\d)(\d+(?:\.\d+)?)\s*元")
FORMULA_RESULT_RE = re.compile(r"=\s*(\d+(?:\.\d+)?)\s*元")
FULL_DATE_RE = re.compile(r"(20\d{2})[./年-]\s*(\d{1,2})[./月-]\s*(\d{1,2})")
MONTH_DAY_RE = re.compile(r"(?<!\d)(\d{1,2})\s*月\s*(\d{1,2})?\s*(?:日|号)?")
ORDER_DATE_RE = re.compile(r"(20\d{2})(\d{2})(\d{2})")


def process_workbook(
    path: Path,
    rules: dict[str, Any],
    api_config: dict[str, Any] | None = None,
) -> tuple[dict[str, Any] | None, list[dict[str, Any]], list[dict[str, Any]]]:
    anomalies: list[dict[str, Any]] = []
    confirmations: list[dict[str, Any]] = []
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = _select_after_sales_sheet(workbook)
    if sheet is None:
        anomalies.append(_anomaly(path, None, "workbook", "critical", "未找到售后方案申请表工作表"))
        return None, anomalies, confirmations

    fields = extract_fields(sheet, rules)
    raw_workbook_text = extract_workbook_text(workbook)
    if raw_workbook_text:
        fields["raw_text"] = raw_workbook_text
    if not fields:
        anomalies.append(
            _anomaly(path, sheet.title, "fields", "critical", "工作簿中没有可分析的文本")
        )
        return None, anomalies, confirmations
    business_fields = [key for key in fields if key not in {"raw_text", "__source_cells"}]
    if not business_fields and "raw_text" in fields:
        confirmations.append(
            _confirmation(
                path,
                "template",
                "未识别到标准售后字段，已把工作簿原文整体交给AI分析",
                "通用Excel原文分析",
                [],
                0.6,
                False,
            )
        )

    # 文件名是品牌等业务信息的合法显式来源，也要交给模型；客户名称绝不作为品牌。
    fields["source_file"] = path.name
    all_text = " ".join([path.name, *[str(v) for v in fields.values()]])
    brand, brand_confidence, brand_evidence = infer_brand(all_text, rules)
    issue_types, issue_evidence = infer_issue_types(all_text, rules)
    issue_primary = infer_primary_issue_type(fields, issue_types)
    issue_reason_summary = ""
    issue_confidence = 0.85 if issue_types else 0.0
    issue_analysis_source = "rules"
    analysis_note = "离线规则分析完成"
    api_analysis: dict[str, Any] | None = None
    strict_api = bool((api_config or {}).get("strict"))

    if api_is_configured(api_config):
        try:
            api_analysis = analyze_issue_with_api(fields, rules, api_config)
            api_brand = normalize_brand(api_analysis.get("brand_name"))
            if api_brand != "未知":
                brand = api_brand
                brand_confidence = api_analysis.get("confidence", 0.0)
                brand_evidence = api_analysis.get("brand_reason") or "api"
            issue_types = api_analysis["issue_types"]
            issue_primary = api_analysis["primary_issue_type"]
            issue_reason_summary = api_analysis["reason_summary"]
            issue_evidence = {"api": api_analysis["evidence"]}
            issue_confidence = api_analysis["confidence"]
            issue_analysis_source = "api"
            analysis_note = "已使用授权 AI 服务辅助分析"
            if (
                issue_confidence < 0.65
                or not issue_reason_summary
                or not api_analysis.get("evidence")
            ):
                confirmations.append(
                    _confirmation(
                        path,
                        "issue_types",
                        "API分类置信度较低，建议人工复核",
                        issue_types,
                        rules.get("issue_order", []),
                        issue_confidence,
                        False,
                    )
                )
        except Exception:
            if strict_api:
                raise
            analysis_note = "AI 辅助不可用，离线规则分析完成"
    elif strict_api:
        raise RuntimeError("API key is required in strict AI mode")
    else:
        analysis_note = "离线规则分析完成"

    if not issue_types:
        issue_types = ["其他"]
        issue_primary = "其他"
        confirmations.append(
            _confirmation(
                path,
                "issue_types",
                "未能自动判断问题类型，已归入其他",
                None,
                rules.get("issue_order", []),
                0.0,
                True,
            )
        )

    if not issue_reason_summary:
        issue_reason_summary = fallback_reason_summary(fields, issue_types)

    complaint_date_raw = fields.get("complaint_date", "")
    complaint_date, year, month, date_confidence = infer_date(fields, path)
    date_analysis_source = "rules"
    if month is None:
        confirmations.append(
            _confirmation(
                path,
                "complaint_date",
                "提出质量问题/报修日期为空或无法识别；未使用订单、交货或文件日期代替",
                complaint_date_raw or None,
                [],
                0.0,
                False,
            )
        )
    elif year is None:
        confirmations.append(
            _confirmation(
                path,
                "complaint_date",
                "提出质量问题/报修日期未注明年份；已保留原月日并归入“年份未注明”统计",
                complaint_date,
                [],
                date_confidence,
                False,
            )
        )

    # 品牌和客户是两个不同字段。品牌没有明确证据时统一记为“未知”，
    # 但这不是阻断项，原因、费用、解决状态和汇总仍然继续生成。
    brand = normalize_brand(brand)
    # 模型只提供候选费用和语义证据；最终金额必须由后端确定性规则复核。
    fee_breakdown = determine_fee_breakdown(fields, rules)
    fee = fee_breakdown.enterprise_cost
    internal_fee = fee_breakdown.internal_accountability
    external_fee = fee_breakdown.external_borne
    cost = fee.amount
    cost_included = bool(fee.included and fee.amount is not None)
    cost_evidence = fee.evidence
    cost_items = fee.items
    cost_status = fee.status
    cost_analysis_source = "deterministic"
    api_cost = api_analysis.get("cost_amount") if api_analysis else None
    api_cost_evidence = api_analysis.get("cost_reason") if api_analysis else ""
    if api_cost is not None and cost is not None and abs(float(api_cost) - float(cost)) > 0.01:
        confirmations.append(
            _confirmation(
                path,
                "enterprise_cost_amount",
                f"模型候选费用 {float(api_cost):.2f} 元与后端复核金额 {float(cost):.2f} 元不一致",
                cost,
                [api_cost],
                float(api_analysis.get("confidence") or 0.0),
                False,
            )
        )
    if fee.status in {"not_found", "needs_review"}:
        confirmations.append(
            _confirmation(
                path,
                "enterprise_cost_amount",
                fee.evidence or "未找到可确认的企业售后成本；保持空值，不按0元统计",
                None,
                [],
                0.0,
                False,
            )
        )
    for field, category_fee in [
        ("internal_accountability_amount", internal_fee),
        ("external_borne_amount", external_fee),
    ]:
        if category_fee.status == "needs_review" or (
            field == "external_borne_amount"
            and category_fee.status == "transferred"
            and category_fee.amount is None
        ):
            confirmations.append(
                _confirmation(
                    path,
                    field,
                    category_fee.evidence or "该类金额存在冲突或歧义，需人工复核",
                    None,
                    [],
                    0.0,
                    False,
                )
            )

    resolution_analysis_source = "rules"
    resolution_status, resolution_reason = infer_resolution_status(fields, rules)
    if api_analysis:
        if api_analysis.get("resolution_status"):
            resolution_status = str(api_analysis["resolution_status"])
        elif api_analysis.get("closed") is True:
            resolution_status = "已解决"
        resolution_reason = (
            api_analysis.get("resolution_reason")
            or api_analysis.get("closed_reason")
            or resolution_reason
        )
        resolution_analysis_source = "api"
    closed = resolution_status == "已解决"
    warnings = list(fee.warnings)
    warnings.extend(
        build_date_warnings(complaint_date, fields.get("delivery_date"), fields.get("order_no"))
    )
    if brand == "未知":
        warnings.append("品牌未知，需人工确认；客户名称未被用作品牌")
    if issue_confidence < 0.7:
        warnings.append("主原因置信度低于0.70")
    if year is None:
        warnings.append("提出质量问题/报修日期未注明年份")
    needs_review = bool(
        warnings
        or fee.status in {"not_found", "needs_review"}
        or internal_fee.status == "needs_review"
        or external_fee.status == "needs_review"
    )
    source_cells = (
        fields.get("__source_cells", {}) if isinstance(fields.get("__source_cells"), dict) else {}
    )
    case_status = "已关闭" if closed else ("处理中" if resolution_status == "处理中" else "待确认")

    record = {
        "source_file": path.name,
        "sheet": sheet.title,
        "file_sha256": file_sha256(path),
        "customer": fields.get("customer", ""),
        "location": fields.get("location", ""),
        "salesperson": fields.get("salesperson", ""),
        "product": fields.get("product", ""),
        "spec": fields.get("spec", ""),
        "order_no": fields.get("order_no", ""),
        "delivery_date": fields.get("delivery_date", ""),
        "complaint_date": complaint_date,
        "complaint_date_raw": complaint_date_raw,
        "year": year,
        "month": month,
        "date_analysis_source": date_analysis_source,
        "brand": brand,
        "brand_confidence": brand_confidence,
        "brand_evidence": brand_evidence,
        "brand_reason": api_analysis.get("brand_reason") if api_analysis else "",
        "issue_types": issue_types,
        "primary_issue_type": issue_primary,
        "issue_reason_summary": issue_reason_summary,
        "issue_confidence": issue_confidence,
        "issue_analysis_source": issue_analysis_source,
        "analysis_note": analysis_note,
        "issue_evidence": issue_evidence,
        "api_analysis": api_analysis,
        "cost": round(cost, 2) if cost is not None else None,
        "fee_amount": round(cost, 2) if cost is not None else None,
        "enterprise_cost_amount": round(cost, 2) if cost is not None else None,
        "enterprise_cost_status": fee.status,
        "enterprise_cost_evidence": fee.evidence,
        "enterprise_cost_reviewed": False,
        "internal_accountability_amount": round(internal_fee.amount, 2)
        if internal_fee.amount is not None
        else None,
        "internal_accountability_status": internal_fee.status,
        "internal_accountability_evidence": internal_fee.evidence,
        "internal_accountability_reviewed": False,
        "external_borne_amount": round(external_fee.amount, 2)
        if external_fee.amount is not None
        else None,
        "external_borne_status": external_fee.status,
        "external_borne_evidence": external_fee.evidence,
        "external_borne_reviewed": False,
        "fee_currency": "CNY",
        "fee_status": fee.status,
        "cost_included": cost_included,
        "cost_status": cost_status,
        "cost_items": cost_items,
        "internal_accountability_items": internal_fee.items,
        "external_borne_items": external_fee.items,
        "cost_rule_amount": round(cost, 2) if cost is not None else None,
        "cost_rule_evidence": fee.evidence,
        "api_cost_candidate": api_cost,
        "api_cost_evidence": api_cost_evidence,
        "cost_evidence": cost_evidence,
        "cost_analysis_source": cost_analysis_source,
        "selected_solution": api_analysis.get("selected_solution")
        if api_analysis
        else fields.get("solution", ""),
        "resolution_status": resolution_status,
        "resolution_reason": resolution_reason,
        "resolution_analysis_source": resolution_analysis_source,
        "closed": closed,
        "closed_reason": resolution_reason,
        "closed_analysis_source": resolution_analysis_source,
        "case_status": case_status,
        "warnings": warnings,
        "needs_review": needs_review,
        "reviewed": False,
        "reviewed_by": None,
        "reviewed_at": None,
        "source_cells": source_cells,
        "feedback": fields.get("feedback", ""),
        "application": fields.get("application", ""),
        "investigation": fields.get("investigation", ""),
        "solution": fields.get("solution", ""),
        "responsibility": fields.get("responsibility", ""),
    }
    return record, anomalies, confirmations


def extract_fields(ws: Any, rules: dict[str, Any]) -> dict[str, Any]:
    fields: dict[str, Any] = {}
    source_cells: dict[str, str] = {}
    for row_idx in range(1, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            raw = ws.cell(row_idx, col_idx).value
            label = clean_cell(raw)
            field = canonical_field(label, rules)
            if not field:
                continue
            value = collect_value(ws, row_idx, col_idx, rules)
            if not value:
                continue
            if field in fields and value not in fields[field]:
                fields[field] = f"{fields[field]}\n{value}"
            else:
                fields[field] = value
            address = value_source_address(ws, row_idx, col_idx, rules)
            if address:
                source_cells[field] = (
                    f"{source_cells[field]},{address}" if field in source_cells else address
                )
    fields["__source_cells"] = source_cells
    return fields


def value_source_address(ws: Any, row_idx: int, col_idx: int, rules: dict[str, Any]) -> str:
    coordinates: list[str] = []
    for c in range(col_idx + 1, ws.max_column + 1):
        if canonical_field(clean_cell(ws.cell(row_idx, c).value), rules):
            break
        if clean_cell(ws.cell(row_idx, c).value):
            coordinates.append(merged_address(ws, ws.cell(row_idx, c).coordinate))
    if col_idx == 1:
        for r in range(row_idx + 1, ws.max_row + 1):
            if canonical_field(clean_cell(ws.cell(r, 1).value), rules):
                break
            for c in range(2, ws.max_column + 1):
                if clean_cell(ws.cell(r, c).value):
                    coordinates.append(merged_address(ws, ws.cell(r, c).coordinate))
    return ",".join(dict.fromkeys(coordinates))


def merged_address(ws: Any, coordinate: str) -> str:
    for merged in ws.merged_cells.ranges:
        if coordinate in merged:
            return str(merged)
    return coordinate


def extract_workbook_text(workbook: Any, max_chars: int = 40_000) -> str:
    """Extract readable text from arbitrary workbook layouts for the AI fallback."""
    sections: list[str] = []
    current_length = 0
    for ws in workbook.worksheets:
        lines = [f"工作表：{ws.title}"]
        row_limit = min(int(ws.max_row or 0), 500)
        column_limit = min(int(ws.max_column or 0), 60)
        for row_idx in range(1, row_limit + 1):
            values = [
                clean_cell(ws.cell(row_idx, col_idx).value)
                for col_idx in range(1, column_limit + 1)
            ]
            while values and not values[-1]:
                values.pop()
            if not any(values):
                continue
            line = f"第{row_idx}行：" + " | ".join(values)
            remaining = max_chars - current_length
            if remaining <= 0:
                break
            lines.append(line[:remaining])
            current_length += len(line)
        sections.append("\n".join(lines))
        if current_length >= max_chars:
            break
    return "\n\n".join(sections).strip()


def collect_value(ws: Any, row_idx: int, col_idx: int, rules: dict[str, Any]) -> str:
    values: list[str] = []
    for c in range(col_idx + 1, ws.max_column + 1):
        if canonical_field(clean_cell(ws.cell(row_idx, c).value), rules):
            break
        text = clean_cell(ws.cell(row_idx, c).value)
        if text:
            values.append(text)

    if col_idx == 1:
        for r in range(row_idx + 1, ws.max_row + 1):
            first = clean_cell(ws.cell(r, 1).value)
            if canonical_field(first, rules):
                break
            row_values = []
            for c in range(2, ws.max_column + 1):
                text = clean_cell(ws.cell(r, c).value)
                if text:
                    row_values.append(text)
            if row_values:
                values.extend(row_values)
    return "\n".join(values).strip()


def canonical_field(label: str, rules: dict[str, Any]) -> str | None:
    if not label:
        return None
    normalized = normalize_label(label)
    if "审批" in normalized or normalized in {"售后方案申请表", "总经理审批"}:
        return None
    # Field names in the source templates are short labels. Long narrative cells
    # often contain words such as "项目" or "产品", and must not be treated as labels.
    if len(normalized) > 18:
        return None
    aliases = rules.get("field_aliases", {})
    for field, words in aliases.items():
        for word in sorted(words, key=len, reverse=True):
            candidate = normalize_label(str(word))
            if normalized == candidate:
                return field
    return None


def infer_brand(text: str, rules: dict[str, Any]) -> tuple[str, float, str]:
    display_order = list(rules.get("brand_order", []))
    priority = list(rules.get("brand_priority", [])) or ["海外", "竹语堂", "桃花江"]
    for brand in display_order:
        if brand not in priority:
            priority.append(brand)
    for brand in priority:
        for keyword in rules.get("brand_keywords", {}).get(brand, []):
            if keyword and keyword in text:
                return brand, 0.95, keyword
    return normalize_brand(rules.get("fallback_brand", "未知")), 0.0, "未找到明确品牌字段或品牌证据"


def normalize_brand(value: Any) -> str:
    text = clean_cell(value)
    if not text or text in {"未识别", "无法识别", "无法判断", "不明确", "无", "null", "None"}:
        return "未知"
    return text


def fallback_reason_summary(fields: dict[str, str], issue_types: list[str]) -> str:
    """从原表中保留最具体的原因描述，避免只有分类没有原因。"""
    keywords = {
        "性能": [
            "开裂",
            "断裂",
            "裂纹",
            "发霉",
            "霉变",
            "腐烂",
            "变形",
            "分层",
            "剥离",
            "膨胀",
            "收缩",
            "破裂",
            "间隙",
        ],
        "外观": ["颜色", "色差", "漆面", "划伤", "脏污", "变色", "表面"],
        "加工": ["尺寸", "规格", "开片", "加工", "图纸", "指令", "安装", "遗漏"],
    }
    wanted = [word for issue_type in issue_types for word in keywords.get(issue_type, [])]
    root_markers = [
        "原因",
        "导致",
        "造成",
        "判定",
        "错误",
        "遗漏",
        "未涂胶",
        "不符",
        "指令",
        "混淆",
        "膨胀",
        "收缩",
        "含水率",
        "未核对",
        "未确认",
        "责任主体",
        "质量问题",
    ]
    investigation = clean_cell(fields.get("investigation", ""))
    if investigation:
        parts = [
            part.strip(" ：:，,。；;|")
            for part in re.split(r"[\n。；;|]", investigation)
            if part.strip()
        ]
        causal = [
            part
            for part in parts
            if any(marker in part for marker in root_markers)
            and not any(word in part for word in ["方案一", "方案二", "费用", "客户意愿"])
        ]
        if causal:
            causal = [re.sub(r"^售后对接人[:：].*?沟通[，,]?", "", part).strip() for part in causal]
            strong_markers = [
                "未涂胶",
                "质量问题",
                "错误造成的原因",
                "导致",
                "膨胀",
                "收缩",
                "遗漏",
                "判定",
            ]
            ranked = sorted(
                enumerate(causal),
                key=lambda item: (
                    -sum(1 for marker in strong_markers if marker in item[1]),
                    item[0],
                ),
            )
            return "；".join(part for _, part in ranked[:3])[:360]

    for field in ["feedback", "application", "solution", "raw_text"]:
        text = clean_cell(fields.get(field, ""))
        if not text:
            continue
        parts = [
            part.strip(" ：:，,。；;|") for part in re.split(r"[\n。；;|]", text) if part.strip()
        ]
        relevant = [part for part in parts if any(word in part for word in wanted)] or parts
        candidate = max(
            relevant,
            key=lambda part: (sum(1 for word in wanted if word in part), min(len(part), 220)),
            default=text,
        )
        if candidate:
            return f"{candidate[:300]}；原表未写明根因"
    return "原表未提供可确认的具体原因"


def infer_primary_issue_type(fields: dict[str, str], issue_types: list[str]) -> str | None:
    if not issue_types:
        return None
    cost_based = infer_primary_from_costed_causes(fields, issue_types)
    if cost_based:
        return cost_based
    cause_text = " ".join(
        str(fields.get(key, "")) for key in ["investigation", "feedback", "application"]
    )
    strong_processing_markers = [
        "未涂胶",
        "胶合",
        "图纸拆解",
        "设计人员",
        "设计遗漏",
        "沟通不到位",
        "存在遗漏",
        "尺寸不符",
        "规格错误",
        "开片错误",
        "指令错误",
        "加工错误",
        "不能安装",
        "无法安装",
        "大于其内径",
        "孔位错误",
        "装配尺寸",
    ]
    if "加工" in issue_types and any(marker in cause_text for marker in strong_processing_markers):
        return "加工"
    for candidate in ["性能", "外观", "加工", "其他"]:
        if candidate in issue_types:
            return candidate
    return issue_types[0]


def infer_primary_from_costed_causes(fields: dict[str, str], issue_types: list[str]) -> str | None:
    """Use the largest attributable fee when the investigation explicitly names mixed causes."""
    if "性能" not in issue_types or "加工" not in issue_types:
        return None
    investigation = clean_cell(fields.get("investigation", ""))
    solution = clean_cell(fields.get("solution", ""))
    if "性能问题" not in investigation or not re.search(
        r"加工.{0,8}(问题|不符)|尺寸不符", investigation
    ):
        return None
    cause_tokens: dict[str, set[str]] = {"性能": set(), "加工": set()}
    for segment in re.split(r"[，,。；;]", investigation):
        category = None
        if "性能问题" in segment or any(word in segment for word in ["开裂", "分层", "霉变"]):
            category = "性能"
        if re.search(r"加工.{0,8}(问题|不符)|尺寸不符", segment):
            category = "加工"
        if category:
            cause_tokens[category].update(
                re.findall(r"(?<!\d)(?:1[0-9]{2}|2[0-9]{2})(?!\d)", segment)
            )
    totals = {"性能": 0.0, "加工": 0.0}
    for segment in re.split(r"[。；;]", solution):
        match = FORMULA_RESULT_RE.search(segment)
        if not match:
            continue
        amount = float(match.group(1))
        tokens = set(re.findall(r"(?<!\d)(?:1[0-9]{2}|2[0-9]{2})(?!\d)", segment))
        scores = {category: len(tokens & values) for category, values in cause_tokens.items()}
        best = max(scores, key=scores.get)
        if scores[best] > 0:
            totals[best] += amount
    if max(totals.values()) > 0 and totals["性能"] != totals["加工"]:
        return max(totals, key=totals.get)
    return None


def infer_issue_types(text: str, rules: dict[str, Any]) -> tuple[list[str], dict[str, list[str]]]:
    found: list[str] = []
    evidence: dict[str, list[str]] = {}
    for issue_type in rules.get("issue_order", []):
        hits = [kw for kw in rules.get("issue_keywords", {}).get(issue_type, []) if kw in text]
        if hits:
            found.append(issue_type)
            evidence[issue_type] = hits[:8]
    return found, evidence


def infer_date(
    fields: dict[str, str], path: Path
) -> tuple[str | None, int | None, int | None, float]:
    complaint = fields.get("complaint_date", "")
    parsed = parse_date_text(complaint)
    if parsed:
        return parsed.isoformat(), parsed.year, parsed.month, 0.95

    month_match = MONTH_DAY_RE.search(str(complaint or ""))
    if month_match:
        month = int(month_match.group(1))
        if not 1 <= month <= 12:
            return None, None, None, 0.0
        day_text = month_match.group(2)
        if day_text:
            day = int(day_text)
            try:
                date(2000, month, day)
            except ValueError:
                return None, None, None, 0.0
            return f"{month}月{day}日", None, month, 0.9
        return f"{month}月", None, month, 0.85

    # 日期只能来自“提出质量问题/报修/投诉日期”字段。
    # 订单号、交货日期、文件名和文件修改时间都不得用于补齐日期。
    return None, None, None, 0.0


def parse_date_text(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    if re.fullmatch(r"\d+(?:\.\d+)?", text):
        serial = float(text)
        if 25000 <= serial <= 60000:
            try:
                return from_excel(serial).date()
            except Exception:
                return None
    match = FULL_DATE_RE.search(text)
    if match:
        year, month, day = map(int, match.groups())
        try:
            return date(year, month, day)
        except ValueError:
            return None
    return None


def parse_order_date(text: str) -> date | None:
    for match in ORDER_DATE_RE.finditer(text):
        year, month, day = map(int, match.groups())
        try:
            return date(year, month, day)
        except ValueError:
            continue
    return None


def infer_year(text: str) -> int | None:
    parsed = parse_order_date(text)
    if parsed:
        return parsed.year
    match = re.search(r"20\d{2}", text)
    return int(match.group(0)) if match else None


def extract_cost(text: str, rules: dict[str, Any]) -> tuple[float, bool, str]:
    clean = text.replace("×", "*").replace("＝", "=")
    if _is_transferred_cost(clean, rules):
        return 0.0, False, "包含转嫁给供应商/合作方等排除关键词"

    total_candidates: list[float] = []
    formula_results: list[float] = []
    fallback_amounts: list[float] = []

    for sentence in re.split(r"[。；;\n]", clean):
        if not sentence.strip():
            continue
        if not any(keyword in sentence for keyword in rules.get("cost_keywords", [])):
            continue
        if "处罚" in sentence or "考核" in sentence:
            continue
        amounts = _money_amounts(sentence)
        if not amounts:
            continue
        if any(keyword in sentence for keyword in rules.get("cost_total_keywords", [])):
            total_candidates.append(amounts[-1][0])
        formula_results.extend(float(v) for v in FORMULA_RESULT_RE.findall(sentence))
        for amount, start, end in amounts:
            suffix = sentence[end : end + 3]
            prefix = sentence[max(0, start - 4) : start]
            if "/" in suffix or "每" in prefix:
                continue
            if amount not in fallback_amounts:
                fallback_amounts.append(amount)

    if total_candidates:
        value = max(total_candidates)
        return value, True, f"识别到汇总费用 {value} 元"

    combined: list[float] = []
    for amount in [*formula_results, *fallback_amounts]:
        if amount not in combined:
            combined.append(amount)
    if combined:
        value = sum(combined) if len(combined) > 1 else combined[0]
        return value, True, f"识别到费用明细 {', '.join(str(round(v, 2)) for v in combined)} 元"
    return 0.0, True, ""


def extract_selected_cost(fields: dict[str, str], rules: dict[str, Any]) -> tuple[float, bool, str]:
    """按最终方案优先级提取费用，避免把调查中的未采用备选方案算入。"""
    for key, label in [
        ("solution", "售后解决方案"),
        ("application", "方案申请明细"),
        ("investigation", "售后负责人调查情况"),
        ("feedback", "反馈内容"),
    ]:
        value = clean_cell(fields.get(key, ""))
        if not value:
            continue
        amount, included, evidence = extract_cost(value, rules)
        if evidence:
            return amount, included, f"{label}：{evidence}"
    return 0.0, True, ""


def infer_resolution_status(fields: dict[str, str], rules: dict[str, Any]) -> tuple[str, str]:
    source_text = " ".join(
        str(fields.get(k, ""))
        for k in ["feedback", "application", "investigation", "solution", "responsibility"]
    )
    if re.search(
        r"(已|已经|后面|最终).{0,12}(解决|完成|处理完|更换完|返工完|修复完|结案|关闭)", source_text
    ):
        return "已解决", "原文明确记录问题已经解决或处理完成"
    if re.search(r"(就此|至此|本次|售后).{0,8}(结案|关闭)|明确.{0,8}(结案|关闭)", source_text):
        return "已解决", "原文明确记录售后已经结案或关闭"
    if "客户自己" in source_text and "解决" in source_text:
        return "已解决", "反馈内容明确记录客户自行处理解决"
    solution = clean_cell(fields.get("solution", ""))
    responsibility = clean_cell(fields.get("responsibility", ""))
    if not solution and not responsibility:
        return "未解决", "原表未提供售后解决方案或完成记录"
    pending_markers = [
        "建议",
        "下售后单",
        "补数生产",
        "重新生产",
        "安排",
        "待",
        "需要",
        "需更换",
        "需维修",
        "提供更换",
        "现场施工",
        "补齐数量",
        "进行更换",
        "进行处理",
    ]
    if any(marker in f"{solution} {responsibility}" for marker in pending_markers):
        return "处理中", "原表已形成方案，但文字显示仍需生产、补数、更换、返修或现场执行"
    return "处理中", "原表已有解决方案，但没有明确的实际完成或结案记录"


def infer_closed(fields: dict[str, str], rules: dict[str, Any]) -> bool:
    """兼容旧调用；仅实际解决完成才返回 True。"""
    return infer_resolution_status(fields, rules)[0] == "已解决"


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return ""
    if text.startswith("=DISPIMG("):
        return ""
    return re.sub(r"\s+", " ", text.replace("\r", "\n")).strip()


def normalize_label(value: str) -> str:
    return re.sub(r"[\s/:：/\\（）()]+", "", value)


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _select_after_sales_sheet(workbook: Any) -> Any | None:
    for ws in workbook.worksheets:
        probe = " ".join(
            clean_cell(ws.cell(r, c).value)
            for r in range(1, min(ws.max_row, 5) + 1)
            for c in range(1, min(ws.max_column, 5) + 1)
        )
        if "售后方案申请表" in probe or "售后" in ws.title:
            return ws
    return workbook.worksheets[0] if workbook.worksheets else None


def _money_amounts(sentence: str) -> list[tuple[float, int, int]]:
    values: list[tuple[float, int, int]] = []
    for match in MONEY_RE.finditer(sentence):
        try:
            values.append((float(match.group(1)), match.start(1), match.end()))
        except ValueError:
            continue
    return values


def _is_transferred_cost(text: str, rules: dict[str, Any]) -> bool:
    excluded = rules.get("excluded_cost_keywords", [])
    return "转嫁给" in text and any(keyword in text for keyword in excluded if keyword != "转嫁给")


def _anomaly(
    path: Path, sheet: str | None, field: str, severity: str, message: str
) -> dict[str, Any]:
    return {
        "source_file": path.name,
        "sheet": sheet,
        "field": field,
        "severity": severity,
        "rule_id": "AFTERSALES-PARSE-001",
        "message": message,
        "suggestion": "请检查源文件是否为售后方案申请表模板",
    }


def _confirmation(
    path: Path,
    field: str,
    reason: str,
    suggested_value: Any,
    options: list[Any],
    confidence: float,
    blocking: bool,
) -> dict[str, Any]:
    return {
        "source_file": path.name,
        "field": field,
        "reason": reason,
        "suggested_value": suggested_value,
        "options": options,
        "confidence": confidence,
        "blocking": blocking,
        "batch_applicable": True,
        "rememberable": True,
    }
