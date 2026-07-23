from __future__ import annotations

from collections.abc import Iterable
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from after_sales_agent.analysis import ISSUE_ORDER, build_analysis_metrics

GREEN = "176B4D"
DARK_GREEN = "254F43"
LIGHT_GREEN = "EAF4EF"
LIGHT_BLUE = "EAF1F8"
LIGHT_RED = "FDE9E7"
LIGHT_GRAY = "F4F6F5"
WHITE = "FFFFFF"
TEXT = "22332B"
MUTED = "63756C"
RED = "C00000"
BLUE = "2F75B5"
CONTENT_WIDTH_DXA = 9360


def write_word_analysis_report(
    records: list[dict[str, Any]],
    result: dict[str, Any],
    output_path: Path,
) -> None:
    """Create a readable business report from reviewed after-sales records."""

    metrics = build_analysis_metrics(records)
    summary = result.get("summary", {})
    pending_count = int(summary.get("confirmation_count") or 0)
    case_count = int(metrics["case_count"])
    totals = metrics["totals"]
    total_fee = float(totals["after_sales_fee_amount"])
    statuses = metrics["status_counts"]
    brand_counts: dict[str, int] = metrics["brand_counts"]
    reason_counts: dict[str, int] = metrics["primary_issue_counts"]
    monthly: list[dict[str, Any]] = metrics["monthly"]
    risks = metrics["risks"]

    document = Document()
    _configure_document(document)
    _add_memo_header(document)

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.add_run("售后分析报告")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run("基于已导入文件、AI 分析结果与人工复核值生成")
    run.font.name = "微软雅黑"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    meta_run = meta.add_run(
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}    "
        f"输入文件：{int(summary.get('input_count') or 0)} 份    "
        f"成功处理：{int(summary.get('processed_count') or 0)} 份"
    )
    meta_run.font.name = "微软雅黑"
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor.from_string(MUTED)

    _add_kpi_strip(
        document,
        (
            ("售后案件", f"{case_count} 件"),
            ("售后费用", f"¥{total_fee:,.2f}"),
            ("已解决", f"{int(statuses['已解决'])} 件"),
            ("待复核项", f"{pending_count} 项"),
        ),
    )

    top_brand = _top_item(brand_counts)
    top_reason = _top_item(reason_counts)
    peak_month = max(
        monthly,
        key=lambda item: float(item["after_sales_fee_amount"]),
        default=None,
    )
    lead_lines = [
        f"本次共汇总 {case_count} 件售后案件，人工复核口径费用合计 "
        f"{total_fee:,.2f} 元。",
        (
            f"案件数最多的品牌为“{top_brand[0]}”（{top_brand[1]} 件）；"
            if top_brand
            else "本次没有可确认的品牌数据；"
        )
        + (
            f"主要售后原因以“{top_reason[0]}”最多（{top_reason[1]} 件）。"
            if top_reason
            else "没有可确认的主要售后原因。"
        ),
    ]
    if peak_month:
        peak_fee = float(peak_month["after_sales_fee_amount"])
        lead_lines.append(
            f"费用最高月份为 {peak_month['period']}，金额 {peak_fee:,.2f} 元，"
            f"{'已达到或超过 10,000 元预警线' if peak_fee >= 10_000 else '未达到预警线'}。"
        )
    _add_callout(document, "执行摘要", lead_lines, LIGHT_GREEN, GREEN)

    document.add_heading("一、数据处理概况", level=1)
    processing_rows = [
        ("输入文件", f"{int(summary.get('input_count') or 0)} 份"),
        ("成功处理", f"{int(summary.get('processed_count') or 0)} 份"),
        ("处理失败", f"{int(summary.get('failed_count') or 0)} 份"),
        ("纳入汇总案件", f"{case_count} 件"),
        ("待人工复核", f"{pending_count} 项"),
    ]
    _add_data_table(
        document,
        ("指标", "结果"),
        processing_rows,
        (3.9, 2.6),
    )

    document.add_page_break()
    document.add_heading("二、品牌方售后占比", level=1)
    _add_body_text(
        document,
        "每份售后文件按人工复核后的品牌计 1 次；无法确认品牌时保留“未知”，"
        "不会用客户名称替代品牌。",
    )
    brand_rows = _distribution_rows(brand_counts, case_count)
    _add_data_table(
        document,
        ("品牌", "案件数量", "占比"),
        brand_rows or [("暂无数据", "0", "0.0%")],
        (3.6, 1.3, 1.6),
    )

    document.add_heading("三、售后原因数量占比", level=1)
    _add_body_text(
        document,
        "每个案件只按人工复核后的一个“主要原因”计数，因此各原因数量合计等于案件总数。",
    )
    reason_rows = [
        (
            reason,
            str(int(reason_counts.get(reason, 0))),
            _percent(int(reason_counts.get(reason, 0)), case_count),
        )
        for reason in ISSUE_ORDER
    ]
    _add_data_table(
        document,
        ("主要原因", "案件数量", "占比"),
        reason_rows,
        (3.6, 1.3, 1.6),
    )

    document.add_page_break()
    document.add_heading("四、售后费用年度逐月明细", level=1)
    _add_body_text(
        document,
        "金额采用人工复核后的最终售后费用。表中以 10,000 元为单月预警线，"
        "达到或超过预警线的月份标记为“预警”。",
    )
    monthly_rows = [
        (
            str(item["period"]),
            str(int(item["case_count"])),
            f"{float(item['after_sales_fee_amount']):,.2f}",
            (
                "预警"
                if float(item["after_sales_fee_amount"]) >= 10_000
                else "正常"
            ),
        )
        for item in monthly
    ]
    _add_data_table(
        document,
        ("报修月份", "案件数", "售后费用（元）", "状态"),
        monthly_rows or [("暂无有效月份", "0", "0.00", "—")],
        (2.1, 1.1, 2.0, 1.3),
        alert_column=3,
    )
    _add_callout(
        document,
        "费用口径",
        [
            f"人工复核口径售后费用合计：{total_fee:,.2f} 元。",
            f"企业承担金额合计：{float(totals['enterprise_cost_amount']):,.2f} 元；"
            f"内部责任扣款合计：{float(totals['internal_accountability_amount']):,.2f} 元；"
            f"外部承担金额合计：{float(totals['external_borne_amount']):,.2f} 元。",
        ],
        LIGHT_BLUE,
        BLUE,
    )

    document.add_heading("五、案件解决状态", level=1)
    status_rows = [
        ("已解决", str(int(statuses["已解决"])), _percent(int(statuses["已解决"]), case_count)),
        ("处理中", str(int(statuses["处理中"])), _percent(int(statuses["处理中"]), case_count)),
        ("未解决", str(int(statuses["未解决"])), _percent(int(statuses["未解决"]), case_count)),
    ]
    _add_data_table(
        document,
        ("处理状态", "案件数量", "占比"),
        status_rows,
        (3.6, 1.3, 1.6),
    )

    document.add_heading("六、数据风险与复核提示", level=1)
    risk_rows = [
        ("售后费用为空或待确认", int(risks["unknown_enterprise_cost_count"])),
        ("品牌未知", int(risks["unknown_brand_count"])),
        ("报修年份或月份缺失", int(risks["unknown_date_count"])),
        ("文件处理失败", int(summary.get("failed_count") or 0)),
        ("尚未完成的人工复核项", pending_count),
    ]
    _add_data_table(
        document,
        ("检查项目", "数量"),
        [(label, f"{value} 项") for label, value in risk_rows],
        (4.9, 1.6),
        alert_column=1,
    )
    suggestions = _build_suggestions(
        pending_count=pending_count,
        failed_count=int(summary.get("failed_count") or 0),
        risks=risks,
        monthly=monthly,
        top_reason=top_reason,
    )
    _add_callout(document, "建议的下一步", suggestions, LIGHT_GRAY, DARK_GREEN)

    disclaimer = document.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(12)
    disclaimer.paragraph_format.space_after = Pt(0)
    run = disclaimer.add_run(
        "说明：本报告用于售后统计、趋势观察和人工复核，不替代责任认定、"
        "财务入账、赔偿审批或正式结案。"
    )
    run.font.name = "微软雅黑"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    _set_east_asia_font(normal.element, "微软雅黑")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = document.styles["Title"]
    title.font.name = "微软雅黑"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK_GREEN)
    _set_east_asia_font(title.element, "微软雅黑")
    title.paragraph_format.space_after = Pt(4)

    for name, size, color in (
        ("Heading 1", 16, DARK_GREEN),
        ("Heading 2", 13, GREEN),
        ("Heading 3", 12, GREEN),
    ):
        style = document.styles[name]
        style.font.name = "微软雅黑"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        _set_east_asia_font(style.element, "微软雅黑")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("竹简工作台 · 售后分析    ")
    run.font.name = "微软雅黑"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    _append_field(footer, "PAGE")
    footer.add_run(" / ")
    _append_field(footer, "NUMPAGES")


def _add_memo_header(document: Document) -> None:
    header = document.sections[0].header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(4.9)
    table.columns[1].width = Inches(1.6)
    left, right = table.rows[0].cells
    left.width = Inches(4.9)
    right.width = Inches(1.6)
    left.text = "竹简工作台  BambooWorkbench"
    right.text = "售后分析"
    _shade_cell(left, DARK_GREEN)
    _shade_cell(right, GREEN)
    for cell in (left, right):
        _set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        for run in cell.paragraphs[0].runs:
            run.font.name = "微软雅黑"
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_table_width(table, CONTENT_WIDTH_DXA)


def _add_kpi_strip(
    document: Document,
    items: Iterable[tuple[str, str]],
) -> None:
    values = list(items)
    table = document.add_table(rows=1, cols=len(values))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(6.5 / len(values))] * len(values)
    for index, ((label, value), cell) in enumerate(
        zip(values, table.rows[0].cells, strict=True)
    ):
        cell.width = widths[index]
        _shade_cell(cell, LIGHT_GREEN if index % 2 == 0 else LIGHT_BLUE)
        _set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_run = paragraph.add_run(f"{label}\n")
        label_run.font.name = "微软雅黑"
        label_run.font.size = Pt(8.5)
        label_run.font.color.rgb = RGBColor.from_string(MUTED)
        value_run = paragraph.add_run(value)
        value_run.font.name = "微软雅黑"
        value_run.font.size = Pt(13)
        value_run.font.bold = True
        value_run.font.color.rgb = RGBColor.from_string(DARK_GREEN)
    _set_table_width(table, CONTENT_WIDTH_DXA)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _add_callout(
    document: Document,
    title: str,
    lines: Iterable[str],
    fill: str,
    accent: str,
) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    _shade_cell(cell, fill)
    _set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    _set_cell_left_border(cell, accent, 20)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    title_run = paragraph.add_run(title)
    title_run.font.name = "微软雅黑"
    title_run.font.size = Pt(11)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor.from_string(accent)
    for line in lines:
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(f"• {line}")
        run.font.name = "微软雅黑"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(TEXT)
    _set_table_width(table, CONTENT_WIDTH_DXA)


def _add_body_text(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True


def _add_data_table(
    document: Document,
    headers: tuple[str, ...],
    rows: Iterable[tuple[Any, ...]],
    widths_inches: tuple[float, ...],
    *,
    alert_column: int | None = None,
) -> None:
    data = list(rows)
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    for index, (header, width) in enumerate(
        zip(headers, widths_inches, strict=True)
    ):
        cell = table.rows[0].cells[index]
        cell.width = Inches(width)
        _shade_cell(cell, DARK_GREEN)
        _set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        run.font.name = "微软雅黑"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
    _repeat_table_header(table.rows[0])

    for row_index, values in enumerate(data):
        cells = table.add_row().cells
        for column_index, (value, width) in enumerate(
            zip(values, widths_inches, strict=True)
        ):
            cell = cells[column_index]
            cell.width = Inches(width)
            _shade_cell(cell, WHITE if row_index % 2 == 0 else LIGHT_GRAY)
            _set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if column_index == 0
                else WD_ALIGN_PARAGRAPH.RIGHT
            )
            run = paragraph.add_run(str(value))
            run.font.name = "微软雅黑"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(TEXT)
            if alert_column == column_index and str(value) not in {
                "0 项",
                "0",
                "—",
                "正常",
            }:
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string(RED)
    _set_table_width(table, CONTENT_WIDTH_DXA)


def _distribution_rows(counts: dict[str, int], total: int) -> list[tuple[str, str, str]]:
    return [
        (name, str(int(count)), _percent(int(count), total))
        for name, count in counts.items()
    ]


def _percent(value: int, total: int) -> str:
    return f"{value / total:.1%}" if total else "0.0%"


def _top_item(counts: dict[str, int]) -> tuple[str, int] | None:
    if not counts or not any(counts.values()):
        return None
    return max(counts.items(), key=lambda item: (item[1], item[0]))


def _build_suggestions(
    *,
    pending_count: int,
    failed_count: int,
    risks: dict[str, int],
    monthly: list[dict[str, Any]],
    top_reason: tuple[str, int] | None,
) -> list[str]:
    suggestions: list[str] = []
    if pending_count:
        suggestions.append("导出最终结果前，逐份完成费用、责任工厂和投诉类型等人工复核项。")
    if failed_count:
        suggestions.append("对处理失败的源文件单独重试，确认文件未损坏且工作表未加密。")
    if int(risks["unknown_brand_count"]):
        suggestions.append("品牌保持“未知”不影响汇总，但建议结合订单或方案申请明细补齐品牌。")
    if int(risks["unknown_date_count"]):
        suggestions.append("优先补齐提出质量问题或报修日期，避免案件无法进入年度月度费用图。")
    if any(float(item["after_sales_fee_amount"]) >= 10_000 for item in monthly):
        suggestions.append("对超过 10,000 元红线的月份复盘责任工厂、重复原因和处理方案。")
    if top_reason:
        suggestions.append(f"围绕数量最高的“{top_reason[0]}”问题开展专项原因复盘与预防措施。")
    if not suggestions:
        suggestions.append("当前数据未发现明显缺口，建议按月持续跟踪品牌、原因和费用变化。")
    return suggestions


def _set_east_asia_font(element: Any, font_name: str) -> None:
    r_pr = element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_name)


def _shade_cell(cell: Any, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def _set_cell_margins(
    cell: Any,
    *,
    top: int,
    bottom: int,
    left: int,
    right: int,
) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (
        ("top", top),
        ("bottom", bottom),
        ("start", left),
        ("end", right),
    ):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_left_border(cell: Any, color: str, size: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    left = borders.find(qn("w:start"))
    if left is None:
        left = OxmlElement("w:start")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:color"), color)


def _repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_table_width(table: Any, width_dxa: int) -> None:
    tbl_pr = table._tbl.tblPr
    width = tbl_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def _append_field(paragraph: Any, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))
    run.font.name = "微软雅黑"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
