from __future__ import annotations

import math
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, PieChart3D, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint, SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from after_sales_agent.manual_review import effective_after_sales_fee

ISSUE_ORDER = ("性能", "外观", "加工", "其他")
GREEN = "176B4D"
DARK_GREEN = "254F43"
BLUE = "2F75B5"
RED = "C00000"
GOLD = "D9A441"
LIGHT_GREEN = "EAF4EF"
LIGHT_BLUE = "EAF1F8"
LIGHT_RED = "FDE9E7"
LIGHT_GRAY = "F4F6F5"
TEXT = "22332B"
MUTED = "63756C"
WHITE = "FFFFFF"
GRID = Border(
    left=Side(style="thin", color="D7E2DC"),
    right=Side(style="thin", color="D7E2DC"),
    top=Side(style="thin", color="D7E2DC"),
    bottom=Side(style="thin", color="D7E2DC"),
)
CHART_COLORS = (
    "176B4D",
    "2F75B5",
    "D9A441",
    "7F8C8D",
    "8E6C9F",
    "D46A6A",
    "4F9D9D",
    "9A7B4F",
    "6B8E4E",
    "5B6C8F",
)


def write_brand_share_workbook(
    records: list[dict[str, Any]], output_path: Path
) -> None:
    included = _included_records(records)
    counts = Counter(_clean_brand(record.get("brand")) for record in included)
    rows = sorted(counts.items(), key=lambda item: (-item[1], item[0]))
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "品牌占比"
    _write_share_dashboard(
        worksheet,
        title="品牌方售后案件占比",
        subtitle=(
            "每个售后案件按人工复核后的品牌计 1 次；"
            "品牌未知时保留“未知”，不使用客户名称替代。"
        ),
        category_label="品牌",
        rows=rows,
        chart_title="品牌方售后案件占比",
        table_name="BrandShareTable",
    )
    _save(workbook, output_path)


def write_issue_share_workbook(
    records: list[dict[str, Any]], output_path: Path
) -> None:
    included = _included_records(records)
    counts: Counter[str] = Counter()
    for record in included:
        issue = str(record.get("primary_issue_type") or "其他").strip()
        counts[issue if issue in ISSUE_ORDER else "其他"] += 1
    rows = [(issue, counts[issue]) for issue in ISSUE_ORDER if counts[issue]]
    _append_zero_categories(rows, counts)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "原因占比"
    _write_share_dashboard(
        worksheet,
        title="售后原因数量占比",
        subtitle=(
            "每个售后案件按照人工复核后的主要原因计 1 次，"
            "性能、外观、加工和其他合计等于案件总数。"
        ),
        category_label="售后原因",
        rows=rows,
        chart_title="售后原因数量占比",
        table_name="IssueShareTable",
    )
    _save(workbook, output_path)


def write_annual_monthly_fee_workbook(
    records: list[dict[str, Any]],
    output_path: Path,
    *,
    threshold: float = 10_000.0,
) -> None:
    included = _included_records(records)
    fees: dict[int, dict[int, float]] = defaultdict(lambda: defaultdict(float))
    cases: dict[int, dict[int, int]] = defaultdict(lambda: defaultdict(int))
    unknown_year_fees: dict[int, float] = defaultdict(float)
    unknown_year_month_cases: dict[int, int] = defaultdict(int)
    unknown_year_fee = 0.0
    for record in included:
        year = _integer(record.get("year"))
        month = _integer(record.get("month"))
        has_month = month is not None and 1 <= month <= 12
        if year is not None and has_month:
            cases[year][month] += 1
        elif has_month:
            unknown_year_month_cases[month] += 1

        amount = effective_after_sales_fee(record)
        if amount is None:
            continue
        if year is None or not has_month:
            unknown_year_fee += amount
            if has_month:
                unknown_year_fees[month] += amount
            continue
        fees[year][month] += amount

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "年度月度费用"
    worksheet.sheet_view.showGridLines = False
    worksheet["A1"] = "售后费用年度逐月明细"
    worksheet["A1"].font = Font(name="微软雅黑", size=18, bold=True, color=WHITE)
    worksheet["A1"].fill = PatternFill("solid", fgColor=GREEN)
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.merge_cells("A1:W1")
    worksheet.row_dimensions[1].height = 34
    worksheet["A2"] = (
        "蓝色柱表示人工复核后的售后费用；红色折线为 10,000 元预警线。"
        "达到或超过红线的月份会在明细表中标红。"
    )
    worksheet["A2"].font = Font(name="微软雅黑", size=10, color=MUTED)
    worksheet["A2"].alignment = Alignment(wrap_text=True)
    worksheet.merge_cells("A2:W2")
    worksheet["G3"] = "费用红线（元）"
    worksheet["H3"] = threshold
    worksheet["G3"].font = Font(name="微软雅黑", bold=True, color=RED)
    worksheet["H3"].font = Font(name="微软雅黑", bold=True, color=RED)
    worksheet["H3"].number_format = '#,##0.00;[Red]-#,##0.00;"-"'

    period_fees: list[tuple[str, dict[int, float], str]] = [
        (f"{year} 年", fees[year], f"MonthlyFee{year}") for year in sorted(fees)
    ]
    if unknown_year_fees:
        period_fees.append(
            ("年份未注明（按原月日）", unknown_year_fees, "MonthlyFeeUnknownYear")
        )
    if not period_fees:
        worksheet["A5"] = "没有可按年份和月份统计的已复核售后费用。"
        worksheet["A5"].font = Font(name="微软雅黑", size=11, color=MUTED)
    for block_index, (period_label, monthly_fees, table_name) in enumerate(period_fees):
        start = 5 + block_index * 27
        _write_year_fee_block(
            worksheet,
            period_label=period_label,
            monthly_fees=monthly_fees,
            threshold=threshold,
            start_row=start,
            table_name=table_name,
        )

    note_row = 5 + len(period_fees) * 27
    unknown_month_fee = sum(unknown_year_fees.values())
    missing_month_fee = max(0.0, unknown_year_fee - unknown_month_fee)
    if missing_month_fee:
        worksheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
        cell = worksheet.cell(
            note_row,
            1,
            f"另有缺少有效月份的已复核费用 {missing_month_fee:,.2f} 元，"
            "因没有可靠月份，不能放入逐月柱状图，但已保留在总表和分析报告中。",
        )
        cell.fill = PatternFill("solid", fgColor=LIGHT_RED)
        cell.font = Font(name="微软雅黑", size=10, color=RED)
        cell.alignment = Alignment(wrap_text=True)

    case_sheet = workbook.create_sheet("逐月案件数量")
    _write_monthly_case_dashboard(
        case_sheet,
        cases=cases,
        unknown_year_cases=unknown_year_month_cases,
    )

    widths = {
        "A": 12,
        "B": 17,
        "C": 17,
        "D": 17,
        "E": 16,
        "F": 3,
        "G": 15,
        "H": 15,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    for column in range(9, 24):
        worksheet.column_dimensions[get_column_letter(column)].width = 10
    worksheet.freeze_panes = "A5"
    worksheet.page_setup.orientation = "landscape"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    _save(workbook, output_path)


def write_yearly_chart_dashboard_workbook(
    records: list[dict[str, Any]],
    output_path: Path,
) -> None:
    """Write one presentation-ready dashboard sheet for every complaint year."""

    included = _included_records(records)
    grouped: dict[int | None, list[dict[str, Any]]] = defaultdict(list)
    for record in included:
        grouped[_integer(record.get("year"))].append(record)

    workbook = Workbook()
    workbook.remove(workbook.active)
    ordered_periods: list[int | None] = sorted(
        period for period in grouped if period is not None
    )
    if None in grouped:
        ordered_periods.append(None)
    if not ordered_periods:
        ordered_periods = [None]

    for period in ordered_periods:
        title = f"{period}年度图表" if period is not None else "年份未注明图表"
        worksheet = workbook.create_sheet(title=title[:31])
        _write_yearly_chart_dashboard(
            worksheet,
            records=grouped.get(period, []),
            period=period,
        )

    _save(workbook, output_path)


def _write_yearly_chart_dashboard(
    worksheet: Any,
    *,
    records: list[dict[str, Any]],
    period: int | None,
) -> None:
    period_label = f"{period}年" if period is not None else "年份未注明"
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 85
    worksheet.merge_cells("A1:N2")
    worksheet["A1"] = f"{period_label}售后统计图表总览"
    worksheet["A1"].font = Font(name="微软雅黑", size=20, bold=True, color=WHITE)
    worksheet["A1"].fill = PatternFill("solid", fgColor=DARK_GREEN)
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 27
    worksheet.row_dimensions[2].height = 27
    worksheet.merge_cells("A3:N3")
    worksheet["A3"] = (
        "品牌和原因按案件数统计；逐月数量依据提出质量问题或报修日期统计。"
        "图上标注分类、数量、百分比及月份坐标，便于直接汇报。"
    )
    worksheet["A3"].font = Font(name="微软雅黑", size=10, color=MUTED)
    worksheet["A3"].alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    brand_counts = Counter(_clean_brand(record.get("brand")) for record in records)
    issue_counts: Counter[str] = Counter()
    monthly_cases: Counter[int] = Counter()
    for record in records:
        issue = str(record.get("primary_issue_type") or "其他").strip()
        issue_counts[issue if issue in ISSUE_ORDER else "其他"] += 1
        month = _integer(record.get("month"))
        if month is not None and 1 <= month <= 12:
            monthly_cases[month] += 1

    brands = sorted(brand_counts.items(), key=lambda item: (-item[1], item[0]))
    issues = [(issue, issue_counts[issue]) for issue in ISSUE_ORDER if issue_counts[issue]]
    total = len(records)
    resolved = sum(
        1
        for record in records
        if str(record.get("resolution_status") or "").strip() == "已解决"
    )
    _write_dashboard_kpi(worksheet, "A4:C5", "售后案件", total, BLUE)
    _write_dashboard_kpi(worksheet, "D4:F5", "品牌数量", len(brands), GREEN)
    _write_dashboard_kpi(
        worksheet,
        "G4:I5",
        "原因分类",
        len(issues),
        GOLD,
    )
    _write_dashboard_kpi(worksheet, "J4:N5", "已解决", resolved, DARK_GREEN)

    brand_header = 52
    issue_header = 52
    month_header = 52
    _write_dashboard_source_table(
        worksheet,
        start_col=1,
        header_row=brand_header,
        headers=("品牌", "案件数量", "占比"),
        rows=[
            (category, count, count / total if total else 0)
            for category, count in (brands or [("暂无数据", 0)])
        ],
    )
    _write_dashboard_source_table(
        worksheet,
        start_col=6,
        header_row=issue_header,
        headers=("售后原因", "案件数量", "占比"),
        rows=[
            (category, count, count / total if total else 0)
            for category, count in (issues or [("暂无数据", 0)])
        ],
    )
    _write_dashboard_source_table(
        worksheet,
        start_col=11,
        header_row=month_header,
        headers=("月份", "售后反馈"),
        rows=[(f"{month}月", int(monthly_cases[month])) for month in range(1, 13)],
    )

    if total:
        brand_chart = _dashboard_pie_chart(
            worksheet,
            title=f"{period_label}品牌售后统计",
            category_col=1,
            value_col=2,
            header_row=brand_header,
            data_count=len(brands),
        )
        worksheet.add_chart(brand_chart, "A7")

        issue_chart = _dashboard_pie_chart(
            worksheet,
            title=f"{period_label}售后问题统计",
            category_col=6,
            value_col=7,
            header_row=issue_header,
            data_count=len(issues),
        )
        worksheet.add_chart(issue_chart, "H7")

    monthly_chart = _dashboard_monthly_case_chart(
        worksheet,
        title=f"{period_label}售后反馈逐月分布图",
        header_row=month_header,
        category_col=11,
        value_col=12,
        maximum=max(monthly_cases.values(), default=0),
    )
    worksheet.add_chart(monthly_chart, "A24")

    for column in range(1, 15):
        worksheet.column_dimensions[get_column_letter(column)].width = 12
    worksheet.column_dimensions["A"].width = 15
    worksheet.column_dimensions["H"].width = 15
    worksheet.column_dimensions["K"].width = 13
    worksheet.freeze_panes = "A6"
    worksheet.print_area = "A1:N49"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1
    worksheet.oddHeader.center.text = f"&B{period_label}售后统计图表总览"
    worksheet.oddFooter.center.text = "第 &P 页，共 &N 页"


def _write_dashboard_kpi(
    worksheet: Any,
    cell_range: str,
    label: str,
    value: int,
    accent: str,
) -> None:
    worksheet.merge_cells(cell_range)
    cell = worksheet[cell_range.split(":")[0]]
    cell.value = f"{label}\n{value:,}"
    cell.font = Font(name="微软雅黑", size=12, bold=True, color=accent)
    cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    cell.border = Border(
        left=Side(style="medium", color=accent),
        right=Side(style="thin", color="D7E2DC"),
        top=Side(style="thin", color="D7E2DC"),
        bottom=Side(style="thin", color="D7E2DC"),
    )
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_dashboard_source_table(
    worksheet: Any,
    *,
    start_col: int,
    header_row: int,
    headers: tuple[str, ...],
    rows: list[tuple[Any, ...]],
) -> None:
    for offset, header in enumerate(headers):
        cell = worksheet.cell(header_row, start_col + offset, header)
        cell.fill = PatternFill("solid", fgColor=DARK_GREEN)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_offset, values in enumerate(rows, start=1):
        for column_offset, value in enumerate(values):
            cell = worksheet.cell(
                header_row + row_offset,
                start_col + column_offset,
                value,
            )
            cell.font = Font(name="微软雅黑", size=10, color=TEXT)
            cell.border = GRID
            cell.alignment = Alignment(
                horizontal="left" if column_offset == 0 else "right",
                vertical="center",
            )
            if column_offset == 2:
                cell.number_format = "0.0%"
            if row_offset % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)


def _dashboard_pie_chart(
    worksheet: Any,
    *,
    title: str,
    category_col: int,
    value_col: int,
    header_row: int,
    data_count: int,
) -> PieChart3D:
    chart = PieChart3D()
    chart.title = title
    chart.height = 10.8
    chart.width = 14.5
    chart.style = 10
    chart.legend.position = "b"
    chart.legend.overlay = False
    chart.firstSliceAng = 270
    chart.add_data(
        Reference(
            worksheet,
            min_col=value_col,
            min_row=header_row,
            max_row=header_row + data_count,
        ),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(
            worksheet,
            min_col=category_col,
            min_row=header_row + 1,
            max_row=header_row + data_count,
        )
    )
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showCatName = True
    chart.dataLabels.showVal = True
    chart.dataLabels.showPercent = True
    chart.dataLabels.showLeaderLines = True
    chart.dataLabels.showLegendKey = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.separator = "\n"
    chart.dataLabels.dLblPos = "bestFit"
    chart.series[0].data_points = [
        DataPoint(
            idx=index,
            spPr=GraphicalProperties(
                solidFill=CHART_COLORS[index % len(CHART_COLORS)]
            ),
        )
        for index in range(data_count)
    ]
    return chart


def _dashboard_monthly_case_chart(
    worksheet: Any,
    *,
    title: str,
    header_row: int,
    category_col: int,
    value_col: int,
    maximum: int,
) -> BarChart:
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = None
    chart.x_axis.title = None
    chart.height = 12.8
    chart.width = 29.5
    chart.gapWidth = 70
    chart.legend.position = "b"
    chart.legend.overlay = False
    chart.varyColors = False
    chart.add_data(
        Reference(
            worksheet,
            min_col=value_col,
            min_row=header_row,
            max_row=header_row + 12,
        ),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(
            worksheet,
            min_col=category_col,
            min_row=header_row + 1,
            max_row=header_row + 12,
        )
    )
    chart.x_axis.delete = False
    chart.x_axis.axPos = "b"
    chart.x_axis.tickLblPos = "nextTo"
    chart.x_axis.tickLblSkip = 1
    chart.x_axis.tickMarkSkip = 1
    chart.y_axis.delete = False
    chart.y_axis.axPos = "l"
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = max(4, maximum + 1)
    chart.y_axis.majorUnit = 1 if maximum <= 10 else max(1, math.ceil(maximum / 5))
    chart.series[0].graphicalProperties.solidFill = BLUE
    chart.series[0].graphicalProperties.line.solidFill = DARK_GREEN
    chart.series[0].tx = SeriesLabel(v="售后反馈")
    chart.y_axis.numFmt = "0"
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showLegendKey = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.showCatName = False
    chart.dataLabels.showPercent = False
    chart.dataLabels.showBubbleSize = False
    chart.dataLabels.dLblPos = "outEnd"
    return chart


def _write_share_dashboard(
    worksheet: Any,
    *,
    title: str,
    subtitle: str,
    category_label: str,
    rows: list[tuple[str, int]],
    chart_title: str,
    table_name: str,
) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:L1")
    worksheet["A1"] = title
    worksheet["A1"].font = Font(name="微软雅黑", size=18, bold=True, color=WHITE)
    worksheet["A1"].fill = PatternFill("solid", fgColor=GREEN)
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 34
    worksheet.merge_cells("A2:L2")
    worksheet["A2"] = subtitle
    worksheet["A2"].font = Font(name="微软雅黑", size=10, color=MUTED)
    worksheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[2].height = 30

    total = sum(value for _, value in rows)
    worksheet["A4"] = "案件总数"
    worksheet["B4"] = total
    worksheet["D4"] = "分类数量"
    worksheet["E4"] = len([value for _, value in rows if value])
    for label in ("A4", "D4"):
        worksheet[label].font = Font(name="微软雅黑", size=10, bold=True, color=MUTED)
    for value in ("B4", "E4"):
        worksheet[value].font = Font(name="微软雅黑", size=15, bold=True, color=GREEN)
        worksheet[value].alignment = Alignment(horizontal="center")

    header_row = 6
    worksheet.cell(header_row, 1, category_label)
    worksheet.cell(header_row, 2, "案件数量")
    worksheet.cell(header_row, 3, "占比")
    data_rows = rows or [("暂无数据", 0)]
    for row_index, (category, count) in enumerate(data_rows, start=header_row + 1):
        worksheet.cell(row_index, 1, category)
        worksheet.cell(row_index, 2, count)
        worksheet.cell(row_index, 3, count / total if total else 0)
        worksheet.cell(row_index, 3).number_format = "0.0%"
        for column in range(1, 4):
            cell = worksheet.cell(row_index, column)
            cell.font = Font(name="微软雅黑", size=10, color=TEXT)
            cell.border = GRID
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "right",
                vertical="center",
            )
        if row_index % 2 == 0:
            for column in range(1, 4):
                worksheet.cell(row_index, column).fill = PatternFill(
                    "solid", fgColor=LIGHT_GRAY
                )

    for column in range(1, 4):
        cell = worksheet.cell(header_row, column)
        cell.fill = PatternFill("solid", fgColor=DARK_GREEN)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center")
    end_row = header_row + len(data_rows)
    table = Table(displayName=table_name, ref=f"A{header_row}:C{end_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    worksheet.add_table(table)

    if total:
        chart = PieChart()
        chart.title = chart_title
        chart.firstSliceAng = 270
        chart.height = 11.5
        chart.width = 18.5
        chart.legend.position = "b"
        chart.add_data(
            Reference(
                worksheet,
                min_col=2,
                min_row=header_row,
                max_row=end_row,
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(
                worksheet,
                min_col=1,
                min_row=header_row + 1,
                max_row=end_row,
            )
        )
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showCatName = True
        chart.dataLabels.showVal = True
        chart.dataLabels.showPercent = True
        chart.dataLabels.showLeaderLines = True
        chart.dataLabels.showLegendKey = False
        chart.dataLabels.dLblPos = "bestFit"
        points = []
        for index in range(len(data_rows)):
            points.append(
                DataPoint(
                    idx=index,
                    spPr=GraphicalProperties(
                        solidFill=CHART_COLORS[index % len(CHART_COLORS)]
                    ),
                )
            )
        chart.series[0].data_points = points
        worksheet.add_chart(chart, "E5")

    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 14
    worksheet.column_dimensions["C"].width = 14
    worksheet.column_dimensions["D"].width = 4
    for column in range(5, 13):
        worksheet.column_dimensions[get_column_letter(column)].width = 11
    worksheet.freeze_panes = "A6"
    worksheet.page_setup.orientation = "landscape"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1


def _write_year_fee_block(
    worksheet: Any,
    *,
    period_label: str,
    monthly_fees: dict[int, float],
    threshold: float,
    start_row: int,
    table_name: str,
) -> None:
    worksheet.merge_cells(
        start_row=start_row,
        start_column=1,
        end_row=start_row,
        end_column=5,
    )
    title = worksheet.cell(start_row, 1, f"{period_label}售后费用逐月明细")
    title.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    title.font = Font(name="微软雅黑", size=13, bold=True, color=GREEN)
    title.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[start_row].height = 26

    header_row = start_row + 2
    headers = ("月份", "售后费用（元）", "1万元预警线", "距离红线（元）", "状态")
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(header_row, column, header)
        cell.fill = PatternFill("solid", fgColor=DARK_GREEN)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.cell(header_row, 6, "售后费用（元）")
    worksheet.cell(header_row, 6).font = Font(name="微软雅黑", size=1, color=WHITE)

    for month in range(1, 13):
        row = header_row + month
        amount = round(float(monthly_fees.get(month, 0.0)), 2)
        worksheet.cell(row, 1, f"{month}月")
        worksheet.cell(row, 2, amount)
        worksheet.cell(row, 3, threshold)
        worksheet.cell(row, 4, round(amount - threshold, 2))
        worksheet.cell(
            row,
            5,
            "达到/超过红线" if amount >= threshold else "正常",
        )
        worksheet.cell(row, 6, amount if amount else None)
        worksheet.cell(row, 6).number_format = "#,##0.00"
        worksheet.cell(row, 6).font = Font(name="微软雅黑", size=1, color=WHITE)
        for column in range(1, 6):
            cell = worksheet.cell(row, column)
            cell.border = GRID
            cell.font = Font(name="微软雅黑", size=10, color=TEXT)
            cell.alignment = Alignment(
                horizontal="left" if column in {1, 5} else "right",
                vertical="center",
            )
        for column in (2, 3, 4):
            worksheet.cell(row, column).number_format = (
                '#,##0.00;[Red]-#,##0.00;"-"'
            )
        if monthly_fees.get(month, 0.0) >= threshold:
            for column in range(1, 6):
                worksheet.cell(row, column).fill = PatternFill(
                    "solid", fgColor=LIGHT_RED
                )
            worksheet.cell(row, 5).font = Font(
                name="微软雅黑", size=10, bold=True, color=RED
            )
        elif month % 2 == 0:
            for column in range(1, 6):
                worksheet.cell(row, column).fill = PatternFill(
                    "solid", fgColor=LIGHT_GRAY
                )

    total_row = header_row + 13
    worksheet.cell(total_row, 1, "年度合计")
    worksheet.cell(total_row, 2, round(sum(monthly_fees.values()), 2))
    worksheet.cell(total_row, 4, round(max(monthly_fees.values(), default=0.0), 2))
    worksheet.cell(total_row, 5, "全年费用 / 单月最高")
    for column in range(1, 6):
        cell = worksheet.cell(total_row, column)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=DARK_GREEN)
        cell.border = GRID
    for column in (2, 4):
        worksheet.cell(total_row, column).number_format = (
            '#,##0.00;[Red]-#,##0.00;"-"'
        )

    table = Table(
        displayName=table_name,
        ref=f"A{header_row}:E{header_row + 12}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    worksheet.add_table(table)

    bar = BarChart()
    bar.type = "col"
    bar.title = f"{period_label}逐月售后费用（红线：10,000 元）"
    bar.y_axis.title = "售后费用（元）"
    bar.x_axis.title = None
    bar.x_axis.tickLblPos = "low"
    bar.x_axis.tickLblSkip = 1
    bar.x_axis.tickMarkSkip = 1
    bar.height = 10
    bar.width = 19
    bar.gapWidth = 45
    bar.legend.position = "b"
    axis_max = _axis_max(max(monthly_fees.values(), default=0.0), threshold)
    bar.y_axis.scaling.min = 0
    bar.y_axis.scaling.max = axis_max
    bar.y_axis.numFmt = "#,##0"
    bar.add_data(
        Reference(
            worksheet,
            min_col=6,
            min_row=header_row,
            max_row=header_row + 12,
        ),
        titles_from_data=True,
    )
    bar.set_categories(
        Reference(
            worksheet,
            min_col=1,
            min_row=header_row + 1,
            max_row=header_row + 12,
        )
    )
    bar.series[0].graphicalProperties.solidFill = BLUE
    bar.series[0].graphicalProperties.line.solidFill = BLUE
    bar.series[0].tx = SeriesLabel(v="售后费用（元）")
    bar.dataLabels = DataLabelList()
    bar.dataLabels.showVal = True
    bar.dataLabels.showLegendKey = False
    bar.dataLabels.showSerName = False
    bar.dataLabels.showCatName = True
    bar.dataLabels.showPercent = False
    bar.dataLabels.showBubbleSize = False
    bar.dataLabels.dLblPos = "outEnd"

    line = LineChart()
    line.add_data(
        Reference(
            worksheet,
            min_col=3,
            min_row=header_row,
            max_row=header_row + 12,
        ),
        titles_from_data=True,
    )
    line.set_categories(
        Reference(
            worksheet,
            min_col=1,
            min_row=header_row + 1,
            max_row=header_row + 12,
        )
    )
    line.y_axis.axId = 200
    line.y_axis.crosses = "max"
    line.y_axis.scaling.min = 0
    line.y_axis.scaling.max = axis_max
    line.y_axis.delete = True
    line.series[0].graphicalProperties.line.solidFill = RED
    line.series[0].graphicalProperties.line.width = 38_100
    line.series[0].graphicalProperties.solidFill = RED
    line.series[0].tx = SeriesLabel(v="1万元预警线")
    bar += line
    worksheet.add_chart(bar, f"G{start_row}")


def _write_monthly_case_dashboard(
    worksheet: Any,
    *,
    cases: dict[int, dict[int, int]],
    unknown_year_cases: dict[int, int],
) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:W1")
    worksheet["A1"] = "售后反馈逐月分布"
    worksheet["A1"].font = Font(name="微软雅黑", size=18, bold=True, color=WHITE)
    worksheet["A1"].fill = PatternFill("solid", fgColor=GREEN)
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 34
    worksheet.merge_cells("A2:W2")
    worksheet["A2"] = (
        "按提出质量问题或报修日期统计案件数量；不同年份分别展示，"
        "年份缺失但月份明确的案件单独列示，不与已知年份混合。"
    )
    worksheet["A2"].font = Font(name="微软雅黑", size=10, color=MUTED)
    worksheet["A2"].alignment = Alignment(wrap_text=True)

    periods: list[tuple[str, dict[int, int], str]] = [
        (f"{year} 年", cases[year], f"MonthlyCases{year}") for year in sorted(cases)
    ]
    if unknown_year_cases:
        periods.append(
            (
                "年份未注明（按原月日）",
                unknown_year_cases,
                "MonthlyCasesUnknownYear",
            )
        )
    if not periods:
        worksheet["A5"] = "没有可按月份统计的售后案件。"
        worksheet["A5"].font = Font(name="微软雅黑", size=11, color=MUTED)
    for block_index, (period_label, monthly_cases, table_name) in enumerate(periods):
        start_row = 5 + block_index * 24
        _write_monthly_case_block(
            worksheet,
            period_label=period_label,
            monthly_cases=monthly_cases,
            start_row=start_row,
            table_name=table_name,
        )

    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 16
    worksheet.column_dimensions["C"].width = 3
    for column in range(4, 18):
        worksheet.column_dimensions[get_column_letter(column)].width = 10
    worksheet.freeze_panes = "A5"
    worksheet.page_setup.orientation = "landscape"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0


def _write_monthly_case_block(
    worksheet: Any,
    *,
    period_label: str,
    monthly_cases: dict[int, int],
    start_row: int,
    table_name: str,
) -> None:
    worksheet.merge_cells(
        start_row=start_row,
        start_column=1,
        end_row=start_row,
        end_column=2,
    )
    title = worksheet.cell(start_row, 1, f"{period_label}售后反馈逐月分布")
    title.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    title.font = Font(name="微软雅黑", size=13, bold=True, color=GREEN)
    title.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[start_row].height = 26

    header_row = start_row + 2
    for column, header in enumerate(("月份", "案件数量"), start=1):
        cell = worksheet.cell(header_row, column, header)
        cell.fill = PatternFill("solid", fgColor=DARK_GREEN)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.cell(header_row, 3, "案件数量")
    worksheet.cell(header_row, 3).font = Font(name="微软雅黑", size=1, color=WHITE)
    for month in range(1, 13):
        row = header_row + month
        worksheet.cell(row, 1, f"{month}月")
        worksheet.cell(row, 2, int(monthly_cases.get(month, 0)))
        worksheet.cell(row, 3, int(monthly_cases.get(month, 0)) or None)
        worksheet.cell(row, 3).font = Font(name="微软雅黑", size=1, color=WHITE)
        for column in range(1, 3):
            cell = worksheet.cell(row, column)
            cell.border = GRID
            cell.font = Font(name="微软雅黑", size=10, color=TEXT)
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "right",
                vertical="center",
            )
            if month % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)

    table = Table(
        displayName=table_name,
        ref=f"A{header_row}:B{header_row + 12}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    worksheet.add_table(table)

    chart = BarChart()
    chart.type = "col"
    chart.title = f"{period_label}售后反馈逐月分布图"
    chart.y_axis.title = "案件数量"
    chart.x_axis.title = None
    chart.height = 10
    chart.width = 19
    chart.gapWidth = 55
    chart.legend = None
    chart.add_data(
        Reference(
            worksheet,
            min_col=3,
            min_row=header_row,
            max_row=header_row + 12,
        ),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(
            worksheet,
            min_col=1,
            min_row=header_row + 1,
            max_row=header_row + 12,
        )
    )
    chart.x_axis.tickLblPos = "low"
    chart.x_axis.tickLblSkip = 1
    chart.x_axis.tickMarkSkip = 1
    maximum = max(monthly_cases.values(), default=0)
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = max(4, maximum + 1)
    chart.y_axis.majorUnit = 1 if maximum <= 10 else max(1, math.ceil(maximum / 5))
    chart.series[0].graphicalProperties.solidFill = GREEN
    chart.series[0].graphicalProperties.line.solidFill = GREEN
    chart.series[0].tx = SeriesLabel(v="案件数量")
    chart.y_axis.numFmt = "0"
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showLegendKey = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.showCatName = True
    chart.dataLabels.showPercent = False
    chart.dataLabels.showBubbleSize = False
    chart.dataLabels.dLblPos = "outEnd"
    worksheet.add_chart(chart, f"D{start_row}")


def _included_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [record for record in records if record.get("included_in_summary", True)]


def _clean_brand(value: Any) -> str:
    text = str(value or "").strip()
    return text if text and text not in {"待确认", "未识别"} else "未知"


def _integer(value: Any) -> int | None:
    try:
        return int(value) if value not in {None, ""} else None
    except (TypeError, ValueError):
        return None


def _append_zero_categories(
    rows: list[tuple[str, int]], counts: Counter[str]
) -> None:
    if rows:
        return
    rows.extend((issue, counts[issue]) for issue in ISSUE_ORDER)


def _axis_max(maximum_fee: float, threshold: float) -> float:
    value = max(maximum_fee, threshold) * 1.25
    step = 5_000.0 if value > 25_000 else 2_000.0
    return max(step, math.ceil(value / step) * step)


def _save(workbook: Workbook, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output_path)
