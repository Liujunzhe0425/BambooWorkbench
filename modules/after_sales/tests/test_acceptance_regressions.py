from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from docx import Document
from openpyxl.utils import get_column_letter

MODULE_ROOT = Path(__file__).resolve().parents[1]
PR16_ACCEPTANCE_FIXTURE = (
    MODULE_ROOT / "tests" / "fixtures" / "acceptance" / "pr16_acceptance_sample.xlsx"
)
sys.path.insert(0, str(MODULE_ROOT / "src"))

from after_sales_agent.analysis import build_analysis_metrics  # noqa: E402
from after_sales_agent.cli import run_job  # noqa: E402
from after_sales_agent.normalizer import determine_fee_breakdown  # noqa: E402
from after_sales_agent.report import write_summary_workbook  # noqa: E402
from after_sales_agent.rules import load_rules  # noqa: E402
from framework_ai_helpers import write_service_response  # noqa: E402

from core.contracts.result import ModuleResult  # noqa: E402


def _word_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join([*paragraphs, *cells])


def _make_workbook(
    path: Path,
    *,
    brand: str = "桃花江竹地板",
    complaint_date: str = "2026-05-10",
    feedback: str = "竹地板出现开裂。",
    investigation: str = "含水率变化导致收缩开裂。",
    solution: str = "已完成更换，企业售后成本1000元。",
) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "售后单"
    rows = [
        ["客户/项目", "脱敏示例客户"],
        ["提出质量问题日期", complaint_date],
        ["产品名称", brand],
        ["反馈内容", feedback],
        ["售后负责人调查情况", investigation],
        ["售后解决方案", solution],
    ]
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _job(input_files: list[Path], output_dir: Path, **overrides: object) -> dict[str, object]:
    payload: dict[str, object] = {
        "schema_version": "1.0",
        "job_id": "acceptance-draft",
        "module_id": "after_sales",
        "phase": "draft",
        "run_mode": "batch" if len(input_files) > 1 else "single",
        "input_files": [str(path) for path in input_files],
        "output_dir": str(output_dir),
        "module_config_file": None,
        "ai_enabled": False,
        "allowed_knowledge_scopes": [],
        "parent_job_id": None,
        "confirmation_file": None,
        "service_response_file": None,
    }
    payload.update(overrides)
    return payload


def _run(tmp_path: Path, payload: dict[str, object], name: str = "job") -> dict[str, object]:
    job_path = tmp_path / f"{name}.json"
    job_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    assert run_job(job_path) == 0
    output_dir = Path(str(payload["output_dir"]))
    result = json.loads((output_dir / "result.json").read_text(encoding="utf-8"))
    ModuleResult.model_validate(result)
    return result


def _run_framework_draft(
    tmp_path: Path,
    payload: dict[str, object],
    name: str = "job",
    *,
    response_status: str = "error",
) -> tuple[dict[str, object], Path]:
    first = _run(tmp_path, payload, f"{name}-first")
    assert first["confirmations"] == []
    assert first["service_requests"]
    response_path = write_service_response(
        tmp_path / f"{name}-service-response.json",
        first["service_requests"],
        task_id=str(payload["job_id"]),
        status=response_status,
    )
    rebound = {
        **payload,
        "ai_enabled": True,
        "service_response_file": str(response_path),
    }
    return _run(tmp_path, rebound, f"{name}-bound"), response_path


def test_offline_complete_file_reuses_recognized_date_and_fee(tmp_path: Path) -> None:
    source = tmp_path / "完整脱敏样表.xlsx"
    _make_workbook(source)
    result, _response = _run_framework_draft(
        tmp_path, _job([source], tmp_path / "output")
    )

    assert result["status"] == "needs_confirmation"
    assert result["summary"]["confirmation_count"] == 2
    prompts = {item["field"]: item for item in result["confirmations"]}
    assert set(prompts) == {"责任工厂", "投诉类型"}
    assert all(item["blocking"] for item in prompts.values())
    assert all(not item["batch_applicable"] for item in prompts.values())
    records = json.loads(
        (tmp_path / "output" / "after_sales_records.json").read_text(encoding="utf-8")
    )
    assert records[0]["analysis_status"] == "AI不可用，离线降级"
    assert records[0]["complaint_date"] == "2026-05-10"
    assert records[0]["after_sales_fee_amount"] == 1000
    assert records[0]["after_sales_fee_status"] == "自动识别"
    assert records[0]["after_sales_fee_reviewed"] is True
    assert records[0]["enterprise_cost_amount"] == 1000


def test_pr16_sanitized_fixture_preserves_second_round_baseline(tmp_path: Path) -> None:
    output = tmp_path / "pr16-fixture-output"
    result, _response = _run_framework_draft(
        tmp_path, _job([PR16_ACCEPTANCE_FIXTURE], output)
    )
    records = json.loads((output / "after_sales_records.json").read_text(encoding="utf-8"))
    record = records[0]

    assert result["status"] == "needs_confirmation"
    assert result["summary"] == {
        "input_count": 1,
        "processed_count": 1,
        "skipped_count": 0,
        "failed_count": 0,
        "anomaly_count": 0,
        "confirmation_count": 2,
    }
    assert {item["type"] for item in result["outputs"]} == {
        "master_summary_excel",
        "analysis_word",
        "brand_share_chart_excel",
        "issue_share_chart_excel",
        "monthly_fee_chart_excel",
        "yearly_chart_dashboard_excel",
    }
    summary = openpyxl.load_workbook(
        output / "after_sales_master_summary.xlsx", data_only=False
    )
    assert summary.sheetnames == ["售后全量总表"]
    assert record["brand"] == "桃花江"
    assert record["complaint_date"] == "2026-05-10"
    assert record["issue_types"] == ["性能", "加工"]
    assert record["enterprise_cost_amount"] == 1000
    assert record["internal_accountability_amount"] == 200
    assert record["external_borne_amount"] == 300
    assert record["resolution_status"] == "已解决"


def test_offline_batch_never_creates_api_config_confirmations(tmp_path: Path) -> None:
    files = [tmp_path / f"完整样表{i}.xlsx" for i in range(3)]
    for index, source in enumerate(files, start=1):
        _make_workbook(
            source,
            complaint_date=f"2026-0{index}-10",
            solution=f"已完成处理，企业售后成本{index * 100}元。",
        )
    result, _response = _run_framework_draft(
        tmp_path, _job(files, tmp_path / "batch-output")
    )

    assert result["status"] == "needs_confirmation"
    assert result["summary"]["confirmation_count"] == 6
    assert not any(item["field"] == "售后费用" for item in result["confirmations"])
    assert not any(item["field"] == "api_config" for item in result["confirmations"])


def test_three_fee_categories_are_extracted_and_never_mixed() -> None:
    breakdown = determine_fee_breakdown(
        {"solution": ("企业实际承担材料及运费，合计1000元；内部责任扣款200元；供应商承担300元。")},
        load_rules(MODULE_ROOT),
    )

    assert breakdown.enterprise_cost.amount == 1000
    assert breakdown.internal_accountability.amount == 200
    assert breakdown.external_borne.amount == 300
    assert (
        sum(
            value.amount or 0
            for value in (
                breakdown.enterprise_cost,
                breakdown.internal_accountability,
                breakdown.external_borne,
            )
        )
        == 1500
    )


def test_external_amount_is_not_enterprise_cost_and_unknown_stays_null() -> None:
    rules = load_rules(MODULE_ROOT)
    external_only = determine_fee_breakdown({"solution": "本次费用300元由供应商承担。"}, rules)
    unknown = determine_fee_breakdown({"solution": "已安排返工，费用尚未确认。"}, rules)

    assert external_only.enterprise_cost.amount is None
    assert external_only.enterprise_cost.status == "transferred"
    assert external_only.external_borne.amount == 300
    assert unknown.enterprise_cost.amount is None
    assert unknown.enterprise_cost.status == "not_found"


def test_master_summary_keeps_one_row_and_all_issue_flags(tmp_path: Path) -> None:
    records = [
        {
            "included_in_summary": True,
            "year": 2026,
            "month": 5,
            "brand": "桃花江",
            "issue_types": ["性能", "加工"],
            "primary_issue_type": "加工",
            "enterprise_cost_amount": 100,
            "resolution_status": "已解决",
        }
    ]
    output = tmp_path / "multilabel.xlsx"
    write_summary_workbook(records, output, load_rules(MODULE_ROOT))
    workbook = openpyxl.load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["售后全量总表"]
    sheet = workbook["售后全量总表"]
    headers = {sheet.cell(3, col).value: col for col in range(1, sheet.max_column + 1)}
    assert sheet.max_row == 4
    assert sheet.cell(4, headers["原因分类（性能/外观/加工）"]).value == "性能、加工"
    assert sheet.cell(4, headers["性能"]).value == "是"
    assert sheet.cell(4, headers["加工"]).value == "是"
    assert sheet.auto_filter.ref == f"A3:{get_column_letter(sheet.max_column)}4"
    assert not sheet.tables
    with ZipFile(output) as archive:
        names = archive.namelist()
        assert not any(name.startswith("xl/tables/") for name in names)
        worksheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert '<autoFilter ref="A3:' in worksheet_xml


def test_visual_workbooks_and_word_report_use_the_same_reviewed_metrics(
    tmp_path: Path,
) -> None:
    first = tmp_path / "一月.xlsx"
    second = tmp_path / "二月.xlsx"
    _make_workbook(first, complaint_date="2026-01-10", solution="已完成处理，企业售后成本100元。")
    _make_workbook(
        second,
        complaint_date="2026-02-10",
        feedback="表面存在色差。",
        investigation="漆面颜色不一致。",
        solution="正在返工，企业售后成本200元。",
    )
    output = tmp_path / "analysis-output"
    result, _response = _run_framework_draft(
        tmp_path, _job([first, second], output)
    )
    records = json.loads((output / "after_sales_records.json").read_text(encoding="utf-8"))
    metrics = build_analysis_metrics(records)
    report = _word_text(output / "after_sales_analysis_report.docx")
    brand_book = openpyxl.load_workbook(
        output / "after_sales_brand_share.xlsx", data_only=False
    )
    reason_book = openpyxl.load_workbook(
        output / "after_sales_reason_share.xlsx", data_only=False
    )
    fee_book = openpyxl.load_workbook(
        output / "after_sales_monthly_fee.xlsx", data_only=False
    )
    dashboard_book = openpyxl.load_workbook(
        output / "after_sales_yearly_dashboard.xlsx", data_only=False
    )

    assert len(brand_book["品牌占比"]._charts) == 1
    assert len(reason_book["原因占比"]._charts) == 1
    assert len(fee_book["年度月度费用"]._charts) == 1
    assert len(fee_book["逐月案件数量"]._charts) == 1
    dashboard = dashboard_book["2026年度图表"]
    assert len(dashboard._charts) == 3
    assert dashboard._charts[0].__class__.__name__ == "PieChart3D"
    assert dashboard._charts[1].__class__.__name__ == "PieChart3D"
    assert dashboard._charts[2].__class__.__name__ == "BarChart"
    assert dashboard._charts[2].x_axis.title is None
    assert dashboard._charts[2].y_axis.title is None
    assert dashboard._charts[2].x_axis.tickLblSkip == 1
    assert dashboard._charts[2].y_axis.majorUnit == 1
    assert dashboard._charts[2].dataLabels.showVal
    brand_labels = brand_book["品牌占比"]._charts[0].dataLabels
    reason_labels = reason_book["原因占比"]._charts[0].dataLabels
    assert brand_labels.showCatName
    assert brand_labels.showVal
    assert brand_labels.showPercent
    assert reason_labels.showCatName
    assert reason_labels.showVal
    assert reason_labels.showPercent
    assert "售后分析报告" in report
    assert f"{result['summary']['input_count']} 份" in report
    assert f"{metrics['totals']['after_sales_fee_amount']:,.2f} 元" in report
    assert f"{metrics['case_count']} 件" in report


def test_user_outputs_never_contain_local_absolute_paths(tmp_path: Path) -> None:
    source = tmp_path / "路径脱敏.xlsx"
    _make_workbook(source)
    output = tmp_path / "path-output"
    _result, _response = _run_framework_draft(
        tmp_path, _job([source], output)
    )
    drive_path = re.compile(r"(?<![A-Za-z])[A-Za-z]:[\\/]")

    records_text = (output / "after_sales_records.json").read_text(encoding="utf-8")
    with ZipFile(output / "after_sales_analysis_report.docx") as archive:
        report_text = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.endswith(".xml")
        )
    assert not drive_path.search(records_text)
    assert not drive_path.search(report_text)
    with ZipFile(output / "after_sales_master_summary.xlsx") as archive:
        workbook_xml = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.endswith(".xml")
        )
    assert not drive_path.search(workbook_xml)


def test_corrupted_workbook_is_isolated_with_stable_chinese_error(tmp_path: Path) -> None:
    valid = tmp_path / "正常.xlsx"
    broken = tmp_path / "损坏.xlsx"
    _make_workbook(valid)
    broken.write_bytes(b"not an xlsx file")
    output = tmp_path / "mixed-output"
    result, _response = _run_framework_draft(
        tmp_path, _job([valid, broken], output)
    )
    records = json.loads((output / "after_sales_records.json").read_text(encoding="utf-8"))

    assert result["status"] == "partial_success"
    assert result["summary"]["processed_count"] == 1
    assert result["summary"]["failed_count"] == 1
    assert records[1]["error_code"] == "WORKBOOK_INVALID_OR_CORRUPTED"
    assert "工作簿损坏或不是有效的 XLSX 文件" in records[1]["analysis_error"]
    assert "File is not a zip file" not in json.dumps(result, ensure_ascii=False)


def test_finalize_writes_all_three_reviewed_amounts_to_every_output(tmp_path: Path) -> None:
    source = tmp_path / "三类金额待复核.xlsx"
    _make_workbook(
        source,
        solution=("企业售后成本待确认；内部责任扣款待确认；供应商承担金额待确认；已完成更换。"),
    )
    draft_output = tmp_path / "draft-output"
    draft, response_path = _run_framework_draft(
        tmp_path, _job([source], draft_output), "draft-job"
    )
    expected_fields = {
        "售后费用",
        "internal_accountability_amount",
        "external_borne_amount",
    }
    prompts = {item["field"]: item for item in draft["confirmations"]}
    assert expected_fields <= prompts.keys()

    values = {
        "售后费用": 1200,
        "internal_accountability_amount": 200,
        "external_borne_amount": 300,
    }
    confirmation_file = tmp_path / "confirmation.json"
    confirmation_file.write_text(
        json.dumps(
            {
                "schema_version": "1.0",
                "job_id": "acceptance-draft",
                "answers": [
                    {
                        "confirmation_id": prompt["confirmation_id"],
                        "value": values.get(field, prompt.get("suggested_value")),
                    }
                    for field, prompt in prompts.items()
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    final_output = tmp_path / "final-output"
    final_payload = _job(
        [source],
        final_output,
        job_id="acceptance-final",
        phase="finalize",
        parent_job_id="acceptance-draft",
        confirmation_file=str(confirmation_file),
        service_response_file=str(response_path),
        ai_enabled=True,
        task_name="售后中央月度验收任务",
        period_year=2026,
        period_month=5,
        learning_enabled=True,
    )
    final = _run(tmp_path, final_payload, "final-job")
    records = json.loads((final_output / "after_sales_records.json").read_text(encoding="utf-8"))
    record = records[0]
    report = _word_text(final_output / "after_sales_analysis_report.docx")
    workbook = openpyxl.load_workbook(
        final_output / "after_sales_master_summary.xlsx", data_only=True
    )
    detail = workbook["售后全量总表"]
    headers = {detail.cell(3, col).value: col for col in range(1, detail.max_column + 1)}

    assert final["status"] == "success"
    assert final["summary"]["confirmation_count"] == 0
    assert record["after_sales_fee_amount"] == 1200
    assert record["enterprise_cost_amount"] == 1200
    assert record["internal_accountability_amount"] == 200
    assert record["external_borne_amount"] == 300
    assert record["fee_amount"] == 1200
    assert detail.cell(4, headers["最终售后费用（企业承担）"]).value == 1200
    assert detail.cell(4, headers["内部责任扣款"]).value == 200
    assert detail.cell(4, headers["外部承担金额"]).value == 300
    assert "企业承担金额合计：1,200.00 元" in report
    assert "内部责任扣款合计：200.00 元" in report
    assert "外部承担金额合计：300.00 元" in report
    snapshot_path = final_output / "after_sales_monthly_snapshot.json"
    assert snapshot_path.is_file()
    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    assert snapshot["payload"]["records"][0]["after_sales_fee_amount"] == 1200
    assert snapshot["payload"]["records"][0]["responsible_factory"]
    assert snapshot["payload"]["records"][0]["primary_issue_type"] in {
        "性能",
        "外观",
        "加工",
        "其他",
    }
    assert prompts["投诉类型"]["learning"]["learning_key"] == (
        "after_sales.complaint_type.v1"
    )
    assert not (final_output / "after_sales_learning_descriptors.json").exists()
    assert final["monthly_snapshot"] == {
        "schema_version": "1.0",
        "path": str(snapshot_path.resolve()),
        "snapshot_format": "module-owned-v1",
        "year": 2026,
        "month": 5,
    }

    mismatched_output = tmp_path / "mismatched-draft-output"
    mismatched_payload = _job(
        [source],
        mismatched_output,
        job_id="acceptance-draft-mismatched-period",
        task_name="售后中央月度验收任务",
        period_year=2026,
        period_month=4,
        learning_enabled=True,
    )
    mismatched = _run(tmp_path, mismatched_payload, "mismatched-draft-job")
    assert mismatched["status"] == "failed"
    assert mismatched.get("monthly_snapshot") is None
    assert mismatched["service_requests"] == []
    assert mismatched["confirmations"] == []
    assert mismatched["outputs"] == []
    assert any(
        "TASK_PERIOD_ZERO_MATCH" in error
        for error in mismatched["errors"]
    )
