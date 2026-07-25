from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
from docx import Document

MODULE_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(MODULE_ROOT / "src"))

from after_sales_agent.cli import _create_sample_workbook, run_job  # noqa: E402
from framework_ai_helpers import write_service_response  # noqa: E402

from core.contracts.result import ModuleResult  # noqa: E402


def _word_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(
        [
            *(paragraph.text for paragraph in document.paragraphs),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )


def _job(input_file: Path, output_dir: Path, **overrides):
    payload = {
        "schema_version": "1.0",
        "job_id": "adapter-draft",
        "module_id": "after_sales",
        "phase": "draft",
        "run_mode": "single",
        "input_files": [str(input_file)],
        "output_dir": str(output_dir),
        "module_config_file": None,
        "ai_enabled": False,
        "allowed_knowledge_scopes": [],
        "parent_job_id": None,
        "confirmation_file": None,
        "service_response_file": None,
    }
    payload.update(overrides)
    return payload


def _bind_framework_response(
    tmp_path: Path,
    job_path: Path,
    payload: dict[str, object],
    output_dir: Path,
) -> tuple[dict[str, object], Path]:
    first = json.loads((output_dir / "result.json").read_text(encoding="utf-8"))
    assert first["confirmations"] == []
    assert first["service_requests"]
    response_path = write_service_response(
        tmp_path / f"{payload['job_id']}-service-response.json",
        first["service_requests"],
        task_id=str(payload["job_id"]),
        status="error",
    )
    rebound = {
        **payload,
        "ai_enabled": True,
        "service_response_file": str(response_path),
    }
    job_path.write_text(json.dumps(rebound, ensure_ascii=False), encoding="utf-8")
    assert run_job(job_path) == 0
    result = json.loads((output_dir / "result.json").read_text(encoding="utf-8"))
    ModuleResult.model_validate(result)
    return result, response_path


def test_draft_and_finalize_write_manual_fee_to_final_excel(tmp_path):
    input_file = tmp_path / "人工复核费用.xlsx"
    _create_sample_workbook(
        input_file,
        customer="示例客户",
        brand_text="桃花江竹地板",
        date_text="2026-05-10",
        feedback="竹地板开裂。",
        solution="已完成更换，费用待人工确认。",
    )

    draft_output = tmp_path / "draft"
    draft_job = _job(input_file, draft_output)
    draft_job_path = tmp_path / "draft-job.json"
    draft_job_path.write_text(json.dumps(draft_job, ensure_ascii=False), encoding="utf-8")
    assert run_job(draft_job_path) == 0
    draft_result, response_path = _bind_framework_response(
        tmp_path, draft_job_path, draft_job, draft_output
    )
    answers = []
    for prompt in draft_result["confirmations"]:
        value = (
            3200 if prompt["field"] == "售后费用" else prompt.get("suggested_value")
        )
        answers.append(
            {
                "confirmation_id": prompt["confirmation_id"],
                "value": value,
                "apply_to_batch": False,
                "remember_mapping": False,
            }
        )
    assert any(item["value"] == 3200 for item in answers)

    confirmation_file = tmp_path / "confirmation.json"
    confirmation_file.write_text(
        json.dumps(
            {"schema_version": "1.0", "job_id": "adapter-draft", "answers": answers},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    final_output = tmp_path / "final"
    final_job = _job(
        input_file,
        final_output,
        job_id="adapter-final",
        phase="finalize",
        parent_job_id="adapter-draft",
        confirmation_file=str(confirmation_file),
        service_response_file=str(response_path),
        ai_enabled=True,
    )
    final_job_path = tmp_path / "final-job.json"
    final_job_path.write_text(json.dumps(final_job, ensure_ascii=False), encoding="utf-8")
    assert run_job(final_job_path) == 0

    final_result = json.loads((final_output / "result.json").read_text(encoding="utf-8"))
    ModuleResult.model_validate(final_result)
    assert final_result["status"] == "success"
    records = json.loads((final_output / "after_sales_records.json").read_text(encoding="utf-8"))
    assert records[0]["fee_amount"] == 3200
    assert records[0]["enterprise_cost_amount"] == 3200
    assert records[0]["enterprise_cost_reviewed"] is True
    assert records[0]["fee_status"] == "exact"
    assert "3,200.00" in _word_text(
        final_output / "after_sales_analysis_report.docx"
    )

    workbook = openpyxl.load_workbook(
        final_output / "after_sales_master_summary.xlsx", data_only=True
    )
    assert any(
        cell.value == 3200
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
    )


def test_required_business_review_updates_fee_factory_and_complaint_type(tmp_path):
    input_file = tmp_path / "三项人工复核.xlsx"
    _create_sample_workbook(
        input_file,
        customer="脱敏复核客户",
        brand_text="桃花江竹地板",
        date_text="2026-07-21",
        feedback="坐凳尺寸错误，现场无法安装。",
        solution="最终方案预计总费用1000元，由二厂承担，已安排补货。",
    )

    draft_output = tmp_path / "review-draft"
    draft_job = _job(input_file, draft_output, job_id="review-draft")
    draft_job_path = tmp_path / "review-draft-job.json"
    draft_job_path.write_text(json.dumps(draft_job, ensure_ascii=False), encoding="utf-8")
    assert run_job(draft_job_path) == 0
    draft, response_path = _bind_framework_response(
        tmp_path, draft_job_path, draft_job, draft_output
    )
    prompts = {item["field"]: item for item in draft["confirmations"]}

    assert {"售后费用", "责任工厂", "投诉类型"} <= prompts.keys()
    assert prompts["售后费用"]["suggested_value"] == 1000
    assert prompts["责任工厂"]["suggested_value"] == "二厂"
    assert prompts["投诉类型"]["options"] == ["性能", "外观", "加工", "其他"]
    assert all(prompts[field]["blocking"] for field in ("售后费用", "责任工厂", "投诉类型"))

    reviewed_values = {"售后费用": "3200", "责任工厂": "成品厂", "投诉类型": "加工"}
    answers = [
        {
            "confirmation_id": prompt["confirmation_id"],
            "value": reviewed_values.get(field, prompt.get("suggested_value")),
            "apply_to_batch": False,
            "remember_mapping": False,
        }
        for field, prompt in prompts.items()
    ]
    confirmation_file = tmp_path / "review-confirmation.json"
    confirmation_file.write_text(
        json.dumps(
            {"schema_version": "1.0", "job_id": "review-draft", "answers": answers},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    final_output = tmp_path / "review-final"
    final_job = _job(
        input_file,
        final_output,
        job_id="review-final",
        phase="finalize",
        parent_job_id="review-draft",
        confirmation_file=str(confirmation_file),
        service_response_file=str(response_path),
        ai_enabled=True,
    )
    final_job_path = tmp_path / "review-final-job.json"
    final_job_path.write_text(json.dumps(final_job, ensure_ascii=False), encoding="utf-8")
    assert run_job(final_job_path) == 0

    final = json.loads((final_output / "result.json").read_text(encoding="utf-8"))
    records = json.loads((final_output / "after_sales_records.json").read_text(encoding="utf-8"))
    record = records[0]
    assert final["status"] == "success"
    assert final["summary"]["confirmation_count"] == 0
    assert record["after_sales_fee_amount"] == 3200
    assert record["responsible_factory"] == "成品厂"
    assert record["primary_issue_type"] == "加工"
    assert record["issue_types"] == ["加工"]
    assert record["manual_review_completed"] is True

    summary = openpyxl.load_workbook(
        final_output / "after_sales_master_summary.xlsx", data_only=True
    )
    detail = summary["售后全量总表"]
    detail_headers = {
        detail.cell(3, col).value: col for col in range(1, detail.max_column + 1)
    }
    assert detail.cell(4, detail_headers["售后费用（人工复核口径）"]).value == 3200
    assert detail.cell(4, detail_headers["责任工厂（人工复核口径）"]).value == "成品厂"
    assert detail.cell(4, detail_headers["主要原因"]).value == "加工"
    assert detail.cell(4, detail_headers["售后费用复核状态"]).value == "已复核"
    assert detail.cell(4, detail_headers["责任工厂复核状态"]).value == "已复核"
    assert detail.cell(4, detail_headers["投诉类型复核状态"]).value == "已复核"
